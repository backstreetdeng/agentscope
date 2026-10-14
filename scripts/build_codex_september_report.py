"""Build the Codex comparison artifact from Codex-authored, validated prose."""
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import sys
from zipfile import ZipFile
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from utils.monthly_report_narrative import (_normalize_batch,_validate_batch_facts,
    _validate_overview_numbers,load_page_card_map,delayed_page_exceptions,
    validate_narrative_package,apply_narratives_to_ppt,narrative_markdown,NARRATIVE_RULE_VERSION)
from utils.monthly_page_narratives import overview_business_packet
from utils.monthly_research_dag import validate_research_ledger

root=Path((ROOT/'.work/september-codex-run.txt').read_text(encoding='utf-8'))
source=Path((ROOT/'.work/september-corrected-package.txt').read_text(encoding='utf-8'))
package=json.loads(source.read_text(encoding='utf-8'))
pages=json.loads((root/'facts.json').read_text(encoding='utf-8'))
edits=json.loads((root/'codex-editorial.json').read_text(encoding='utf-8'))
research=json.loads((root/'research-evidence.json').read_text(encoding='utf-8'))
ledger_errors=validate_research_ledger(root/'research-dag.json')
if ledger_errors:raise ValueError('; '.join(ledger_errors))
approved,evidence,errors=[],{},[]
for edit in edits:
    packet=next(p for p in pages if p['slide']==edit['slide'])
    slide=deepcopy(packet['fact_grounded_draft'])
    slide['headline']=edit['headline']
    refs=[]
    for key in ('business_facts','level_facts','derived_facts','model_facts'):
        for f in packet.get(key,[]):
            if f.get('fact_id') and f.get('period','2026-09')=='2026-09':refs.append(f['fact_id'])
    if not refs:refs=slide['anomaly']['fact_ids']
    slide['body_claims']=[{'role':role,'claim':text,'fact_ids':refs,
        'evidence_ids': [edit.get('evidence',{}).get(str(i))] if edit.get('evidence',{}).get(str(i)) else [],
        'certainty':'qualified' if edit.get('evidence',{}).get(str(i)) else 'unverified'}
        for i,(role,text) in enumerate(zip(('core','structure','models' if edit['slide'] in (9,10,11) else 'support'),edit['claims']))]
    if edit['slide'] in (9,10,11):
        slide['body_claims'][2]['fact_ids']=[f['fact_id'] for f in packet.get('model_facts',[])
                                            if f.get('scope')=='single_month']
    cited={x for c in slide['body_claims'] for x in c['evidence_ids']}
    candidates=[e for e in research['evidence'] if e['id'] in cited]
    result=_normalize_batch({'slides':[slide],'evidence':candidates},'CODEX',(edit['slide'],),'2026-09')
    try:_validate_batch_facts(result,[packet])
    except ValueError as exc:errors.append(str(exc));continue
    result['slides'][0].update(generation_method='codex_authored',execution_origin='codex')
    approved.extend(result['slides']);evidence.update({e['id']:e for e in result['evidence']})
if errors:
    (root/'editorial-errors.json').write_text(json.dumps(errors,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(errors,ensure_ascii=False));raise SystemExit(1)
overview={'slide':3,'headline':'环比修复与同比承压并存，结构向相对抗跌板块倾斜',
    'anomaly':{'type':'structural_divergence','fact_ids':[f for s in approved for f in s['anomaly']['fact_ids']]},
    'headline_evidence_ids':[],'body_claims':[{'claim':
    '9月整体市场零售166.9万辆，同比下降26.6%、环比增长7.3%。SUV较轿车抗跌，新能源渗透率65.0%仍伴随自身同比收缩；BEV较PHEV抗跌，零跑与比亚迪走势分化。环比恢复没有改变累计压力，政策与新品事件须结合时间、对象和口径判断。',
    'fact_ids':[f for s in approved for c in s['body_claims'] for f in c['fact_ids']],
    'evidence_ids':[],'certainty':'unverified'}]}
summary=_normalize_batch({'slides':[overview],'evidence':[]},'OV',(3,),'2026-09')
overview_packet=overview_business_packet(pages,approved,'2026-09')
_validate_overview_numbers(summary,approved,'2026-09',overview_packet)
summary['slides'][0].update(generation_method='codex_authored_summary',execution_origin='codex')
approved=summary['slides']+approved
base=ROOT/'skills/market-analysis-skill/passenger-car-monthly-report'
mapping=next((base/'assets').glob('*.xlsx'));examples=base/'assets/expert-headline-examples.json'
narrative={'run_id':root.name,'report_month':'2026-09','slides':approved,'evidence':list(evidence.values()),
    'research_evidence':research['evidence'],'research':research['research'],'evidence_gaps':research['evidence_gaps'],
    'retained_pages':delayed_page_exceptions(package,load_page_card_map(mapping)),
    'narrative_rule_version':NARRATIVE_RULE_VERSION,'execution_origin':'codex','narrative_author':'Codex',
    'web_model_used_for_narratives':False,'generated_at':datetime.now().isoformat()}
validate_narrative_package(narrative,package,run_id=root.name,mapping_path=mapping,examples_path=examples)
(root/'narrative_evidence.json').write_text(json.dumps(narrative,ensure_ascii=False,indent=2),encoding='utf-8')
output=ROOT/'outputs/ppt/2026年9月中国乘用车市场月度分析报告_Codex对比版.pptx'
from utils.monthly_source_pptx import update_source_template
template=next((base/'assets').glob('*.pptx'))
data_update=update_source_template(template,root/'data-updated.pptx',package,mapping)
(root/'data-update.json').write_text(json.dumps(data_update,ensure_ascii=False,indent=2),encoding='utf-8')
receipt=apply_narratives_to_ppt(root/'data-updated.pptx',output,narrative)
from pptx import Presentation
deck=Presentation(output)
for s in approved:
    notes=[f'执行来源：Codex；报告月份：2026-09；任务：{root.name}',
           f'数据基线 SHA256：{sha256(source.read_bytes()).hexdigest()}',s['headline']]
    for claim in s['body_claims']:
        notes.append(claim['claim'])
        for id in claim['evidence_ids']:
            e=evidence[id]
            notes.append(f"来源：{e['source_name']}；记录/URL：{e['url_or_record_id']}；日期：{e['published_at']}；等级：{e['grade']}；原文：{e.get('quote',e['summary'])}")
    deck.slides[s['slide']-1].notes_slide.notes_text_frame.text='\n'.join(notes)
deck.save(output)
template=next((base/'assets').glob('*.pptx'))
with ZipFile(template) as a,ZipFile(output) as b:
    if a.read('ppt/slides/slide5.xml')!=b.read('ppt/slides/slide5.xml'):raise ValueError('第5页滞后内容被改写')
report=narrative_markdown(narrative)
report+='\n\n## 对比基线\n\n- 执行来源：Codex；主旨句由 Codex 直接撰写，未调用网页端模型代写。\n'
report+='- 数据基线：'+str(source)+'\n- 数据包 SHA256：'+sha256(source.read_bytes()).hexdigest()+'\n'
report+='- 只有第5页库存数据滞后；其余分析页使用9月数据。\n'
report+='- 公网协会报告的数量、增速与 BI 口径不同；本报告数字统一使用 BI，公网材料只作背景核验。\n'
report+='\n## 主旨句与来源\n'
for s in approved:
    report+=f"\n### 第{s['slide']}页：{s['headline']}\n\n"
    for c in s['body_claims']:
        report+=c['claim']+'\n\n'
        for id in c['evidence_ids']:
            e=evidence[id];report+=f"来源：{e['source_name']}；记录/URL：{e['url_or_record_id']}；日期：{e['published_at']}；等级：{e['grade']}。\n\n"
report_path=ROOT/'outputs/ppt/2026年9月月报_Codex主旨句与对比基线.md'
report_path.write_text(report,encoding='utf-8')
manifest={'run_id':root.name,'execution_origin':'codex','narrative_author':'Codex',
    'report_month':'2026-09','source_package':str(source),'source_sha256':sha256(source.read_bytes()).hexdigest(),
    'narrative_package':str(root/'narrative_evidence.json'),'research':research['research'],
    'data_update':json.loads((root/'data-update.json').read_text(encoding='utf-8')),
    'narrative_receipt':receipt,'pptx_sha256':sha256(output.read_bytes()).hexdigest(),
    'retained_pages':[5],'headline_pages':len(approved),'visual_acceptance':'pending'}
typed=[{'slide':p['slide'],**{k:p.get(k,[]) for k in ('business_facts','level_facts','model_facts','derived_facts')}} for p in pages]
manifest['typed_business_facts_sha256']=sha256(json.dumps(typed,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
output.with_suffix('.manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'pptx':str(output),'headlines':len(approved),'retained_pages':[5],'report':str(report_path)},ensure_ascii=False))
