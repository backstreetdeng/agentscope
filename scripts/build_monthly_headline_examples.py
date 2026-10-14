"""Extract page-level headline examples from business-approved monthly decks."""

from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation


SOURCE_ROOT = Path(r"D:\2024年度工作日志和备忘录\数字化转型产品\4.0 同事组\5.4 晓亚\提供邓skill-0915")
OUTPUT = Path(__file__).resolve().parents[1] / "skills" / "market-analysis-skill" / "passenger-car-monthly-report" / "assets" / "expert-headline-examples.json"
SOURCES = [
    "2026年3月中国乘用车市场月度分析报告.pptx",
    "2026年4月中国乘用车市场月度分析报告.pptx",
    "2026年5月中国乘用车市场月度分析报告.pptx",
    "2026年6月中国乘用车市场月度分析报告0715.pptx",
    "2026年7月中国乘用车市场月度分析报告0814.pptx",
    "2026年8月中国乘用车市场月度分析报告0915.pptx",
]


def new_energy_topic(title: str, context: str = "") -> str | None:
    """Route shared 2.8 headings by chart content, not mentioned companies.

    The business decks added a BEV/PHEV level page in May. It has no
    corresponding slide in the 21-page production template and must remain
    a separate local example topic, rather than contaminate the split page.
    ``context`` should contain chart labels, excluding headline/commentary.
    """
    compact = re.sub(r"\s+", "", title)
    labels = re.sub(r"\s+", "", context)
    if not compact.startswith("2.8新能源"):
        return None
    if re.search(r"新能源销量TOP\d+企业", labels, re.I):
        return "new_energy_enterprises"
    if re.search(r"(?:BEV|PHEV|纯电|插混)分级别市场表现", labels, re.I):
        return "new_energy_powertrain_levels"
    if "渗透率" in labels and "销量及增速" in labels:
        return "new_energy_overall"
    if any(token in compact.upper() for token in ("纯电", "插混", "BEV", "PHEV")):
        return "new_energy_powertrain_split"
    if any(token in compact.upper() for token in ("企业", "车企", "TOP10")):
        return "new_energy_enterprises"
    return None


def target_page(title: str, context: str = "") -> int | None:
    compact = re.sub(r"\s+", "", title)
    if "总体概述" in compact:
        return 3
    if compact.startswith("1.2零售量"):
        return 4
    if compact.startswith("1.3库存(行业)") or compact.startswith("1.3库存（行业）"):
        return 5
    if compact.startswith("2.1车型类别"):
        return 7
    if compact.startswith("2.2车型级别"):
        return 8
    if compact.startswith("2.3轿车"):
        return 9
    if compact.startswith("2.3SUV"):
        return 10
    if compact.startswith("2.3MPV"):
        return 11
    if compact.startswith("2.3车型类别×级别"):
        # March used a shared heading. Classify by the slide's chart titles,
        # not by physical page number or a model mentioned in the commentary.
        if "MPV分级别" in context:
            return 11
        if "SUV分级别" in context:
            return 10
        if "轿车分级别" in context:
            return 9
    if compact.startswith("2.4车型系别"):
        return 12
    if compact.startswith("2.5车型价格段"):
        return 13
    if compact.startswith("2.6区域市场"):
        return 15 if any(token in compact for token in ("线级", "城市")) else 14
    if compact.startswith("2.7企业竞争"):
        return 16
    if compact.startswith("2.8新能源"):
        return {
            "new_energy_overall": 17,
            "new_energy_powertrain_split": 18,
            "new_energy_enterprises": 19,
        }.get(new_energy_topic(title, context))
    return None


def slide_texts(slide) -> list[str]:
    return [
        shape.text.strip().replace("\n", " ")
        for shape in slide.shapes
        if hasattr(shape, "text") and shape.text.strip()
    ]


def main() -> None:
    examples: dict[str, list[dict[str, str | int]]] = {}
    unmapped_topics: dict[str, list[dict[str, str | int]]] = {}
    for source_name in SOURCES:
        source = SOURCE_ROOT / source_name
        if not source.is_file():
            raise FileNotFoundError(source)
        deck = Presentation(source)
        for source_slide, slide in enumerate(deck.slides, 1):
            texts = slide_texts(slide)
            if not texts:
                continue
            # Headlines/commentary may mention models, companies or levels
            # which do not define the page. Use actual chart-label text first.
            context = " ".join(texts[2:])
            page = target_page(texts[0], context)
            example = {
                "source": source_name,
                "source_slide": source_slide,
                "headline": texts[0],
                "body": texts[1] if len(texts) > 1 else "",
            }
            if page is None:
                topic = new_energy_topic(texts[0], context)
                if topic == "new_energy_powertrain_levels":
                    unmapped_topics.setdefault(topic, []).append(example)
                continue
            examples.setdefault(str(page), []).append(example)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(
        json.dumps({"version": 2, "pages": examples, "unmapped_topics": unmapped_topics}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    mapped_count = sum(map(len, examples.values()))
    unmapped_count = sum(map(len, unmapped_topics.values()))
    print(json.dumps({"output": str(OUTPUT), "pages": len(examples),
                      "mapped_examples": mapped_count, "unmapped_examples": unmapped_count,
                      "examples": mapped_count + unmapped_count}, ensure_ascii=False))


if __name__ == "__main__":
    main()
