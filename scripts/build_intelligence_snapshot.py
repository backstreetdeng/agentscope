"""Build the homepage policy/rights snapshot from the three reference workbooks.

The source workbooks remain untouched.  The generated JSON files are served from
``/assets/data`` and keep both the compact dashboard summary and normalized
download records.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


DEFAULT_SOURCE_DIR = Path(
    r"E:/AI/migration-backup/5.4 晓亚\乘用车价格与权益监测-0921"
)
PROJECT_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_DIR / "assets" / "data"
SNAPSHOT_DATE = date(2026, 9, 21)

POLICY_HEADERS = [
    "记录编号", "监测日期/发布日期", "省份", "城市", "区县", "政策名称",
    "政策发布时间", "政策开始时间", "政策结束时间",
    "适用客群（首购/增换购；个人/单位消费者；新能源/燃油车）",
    "当前政策补贴内容", "促消费类型（统一表准/不统一标准）", "补贴金额",
    "叠加及限制条件", "较上期变化类型", "较上期变化详情", "信息来源链接", "采集日期",
]

RIGHTS_HEADERS = [
    "记录编号", "监测日期/发布日期", "车企/品牌", "车型", "年款及版本", "动力类型",
    "活动名称", "活动开始日期", "活动结束日期", "适用地区及渠道", "适用客群",
    "原指导价", "现指导价", "当前活动价及口径", "当前活动类型（支持多选）",
    "活动内容", "叠加及限制条件", "较上期变化类型", "较上期变化详情", "信息来源链接", "采集日期",
]

RIGHTS_TYPES = ["现金优惠", "置换及增购补贴", "保险补贴", "购置税补贴", "金融政策", "其他权益"]


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def split_sentences(value: str) -> list[str]:
    return [part.strip(" \n\r\t；。") for part in re.split(r"[。；\n\r]+", value) if part.strip()]


def unique(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            result.append(value)
    return result


def compact_sentences(content: str, pattern: str, limit: int = 4) -> str:
    matched = [sentence for sentence in split_sentences(content) if re.search(pattern, sentence)]
    return "；".join(matched[:limit])


def policy_audience(content: str) -> str:
    rules = [
        (r"个人消费者|个人用户|自然人|居民个人", "个人消费者"),
        (r"单位消费者|企业购车|法人|机关事业单位", "单位消费者"),
        (r"首次购车|首购", "首购"),
        (r"以旧换新|报废更新|置换|换购|增购|转让旧车", "增换购"),
        (r"新能源|纯电|插电|电动车", "新能源汽车"),
        (r"燃油车|汽油车", "燃油车"),
        (r"新车|新购汽车|购置汽车", "新车"),
        (r"二手车", "二手车"),
    ]
    values = [label for pattern, label in rules if re.search(pattern, content)]
    return "；".join(values) if values else "未明确"


def policy_amounts(content: str) -> str:
    pattern = re.compile(
        r"(?:最高(?:补贴)?(?:不超过|可达|至)?\s*)?"
        r"(?:\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?)\s*(?:万元|元|%|折)"
    )
    return "；".join(unique(match.group(0).replace(" ", "") for match in pattern.finditer(content))[:12])


def policy_standard(content: str) -> str:
    tiered = r"分档|第[一二三四五六七八九十]+档|购车.{0,10}(?:价格|金额).{0,18}(?:以上|以下|不含|含)|按.{0,12}%|区间"
    return "不统一标准" if re.search(tiered, content) else "统一标准"


def policy_date(year: Any, month: Any, publish: Any) -> str:
    year_text = text(year)
    month_text = text(month)
    day_match = re.search(r"(\d{1,2})月(\d{1,2})日", text(publish))
    try:
        y = int(float(year_text))
        m = int(float(month_text))
        d = int(day_match.group(2)) if day_match else 1
        return date(y, m, d).isoformat()
    except (TypeError, ValueError):
        return text(publish)


def policy_id(row_number: int, values: tuple[Any, ...]) -> str:
    key = "|".join(text(values[index]) for index in (3, 4, 5, 9, 11, 12))
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:6].upper()
    return f"POL-{row_number:04d}-{digest}"


def normalize_policy(row_number: int, values: tuple[Any, ...]) -> dict[str, Any]:
    content = text(values[13])
    subsidy = compact_sentences(content, r"补贴|消费券|减免|奖励|贴息|直减|优惠", 5)
    restrictions = compact_sentences(
        content,
        r"仅限|不得|不可|不超过|需在|须在|申领|名额|资金|发票|上牌|指定|资格|先到先得|用完即止|发完即止",
        4,
    )
    result = {
        "记录编号": policy_id(row_number, values),
        "监测日期/发布日期": policy_date(values[1], values[2], values[10]),
        "省份": text(values[3]),
        "城市": text(values[4]),
        "区县": text(values[5]),
        "政策名称": text(values[9]),
        "政策发布时间": text(values[10]),
        "政策开始时间": text(values[11]),
        "政策结束时间": text(values[12]),
        "适用客群（首购/增换购；个人/单位消费者；新能源/燃油车）": policy_audience(content),
        "当前政策补贴内容": subsidy or content[:500],
        "促消费类型（统一表准/不统一标准）": policy_standard(content),
        "补贴金额": policy_amounts(content),
        "叠加及限制条件": restrictions,
        "较上期变化类型": "",
        "较上期变化详情": "",
        "信息来源链接": text(values[14]),
        "采集日期": SNAPSHOT_DATE.isoformat(),
    }
    return result


def normalize_rights(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    source = {headers[index]: text(value) for index, value in enumerate(values)}
    active_types = [name for name in RIGHTS_TYPES if source.get(name)]
    activity_content = "；".join(f"{name}：{source[name]}" for name in active_types)
    return {
        "记录编号": source.get("记录编号", ""),
        "监测日期/发布日期": source.get("监测日期/发布日期", ""),
        "车企/品牌": source.get("车企/品牌", ""),
        "车型": source.get("车型", ""),
        "年款及版本": source.get("年款及版本", ""),
        "动力类型": source.get("动力类型", ""),
        "活动名称": source.get("活动名称", ""),
        "活动开始日期": source.get("活动开始日期", ""),
        "活动结束日期": source.get("活动结束日期", ""),
        "适用地区及渠道": source.get("适用地区及渠道", ""),
        "适用客群": source.get("适用客群", ""),
        "原指导价": source.get("原指导价", ""),
        "现指导价": source.get("现指导价", ""),
        "当前活动价及口径": source.get("当前活动价及口径", ""),
        "当前活动类型（支持多选）": active_types,
        "活动内容": activity_content,
        "叠加及限制条件": source.get("叠加及限制条件", ""),
        "较上期变化类型": source.get("较上期变化类型", ""),
        "较上期变化详情": source.get("较上期变化详情", ""),
        "信息来源链接": source.get("信息来源链接", ""),
        "采集日期": source.get("采集日期", ""),
    }


def policy_topics(records: list[dict[str, Any]]) -> Counter[str]:
    rules = [
        ("购车补贴", r"购车补贴|汽车补贴|补助"),
        ("消费券", r"消费券"),
        ("以旧换新", r"以旧换新|报废更新|置换"),
        ("金融贴息", r"贴息|贷款|金融"),
        ("新能源专项", r"新能源|纯电|电动汽车"),
    ]
    counts: Counter[str] = Counter()
    for record in records:
        haystack = " ".join((record["政策名称"], record["当前政策补贴内容"], record["适用客群（首购/增换购；个人/单位消费者；新能源/燃油车）"]))
        for label, pattern in rules:
            if re.search(pattern, haystack):
                counts[label] += 1
    return counts


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def build() -> None:
    source_dir = Path(os.environ.get("INTELLIGENCE_SOURCE_DIR", DEFAULT_SOURCE_DIR))
    policy_path = source_dir / "地方政策数据.xlsx"
    rights_path = source_dir / "权益数据.xlsx"
    header_path = source_dir / "地方政策+权益表头.xlsx"
    for path in (policy_path, rights_path, header_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    template_workbook = load_workbook(header_path, read_only=True, data_only=True)
    policy_template_sheet = template_workbook["地方促消费政策展示"]
    rights_template_sheet = template_workbook["权益展示"]
    policy_template_headers = [text(cell.value) for cell in policy_template_sheet[2] if text(cell.value)]
    rights_template_headers = [text(cell.value) for cell in rights_template_sheet[1] if text(cell.value)]
    if policy_template_headers != POLICY_HEADERS:
        raise ValueError("地方促消费政策展示表头已变化，请先更新字段映射")
    if rights_template_headers != RIGHTS_HEADERS:
        raise ValueError("权益展示表头已变化，请先更新字段映射")

    policy_sheet = load_workbook(policy_path, read_only=True, data_only=True)["地方促消费政策"]
    policy_rows = list(policy_sheet.iter_rows(min_row=2, values_only=True))
    policy_records = [normalize_policy(index, tuple(values)) for index, values in enumerate(policy_rows, start=2)]

    rights_sheet = load_workbook(rights_path, read_only=True, data_only=True)["综合台账"]
    rights_iter = rights_sheet.iter_rows(values_only=True)
    rights_source_headers = [text(value) for value in next(rights_iter)]
    rights_records = [normalize_rights(rights_source_headers, tuple(values)) for values in rights_iter]

    category_counts = Counter()
    multi_type = 0
    for record in rights_records:
        types = record["当前活动类型（支持多选）"]
        category_counts.update(types)
        if len(types) > 1:
            multi_type += 1

    years = Counter(record["监测日期/发布日期"][:4] for record in policy_records if record["监测日期/发布日期"][:4].isdigit())
    topics = policy_topics(policy_records)
    latest_policy = sorted(
        policy_records,
        key=lambda item: (item["监测日期/发布日期"], item["政策名称"]),
        reverse=True,
    )[:5]
    latest_rights = sorted(
        rights_records,
        key=lambda item: (item["监测日期/发布日期"], item["车企/品牌"], item["车型"]),
        reverse=True,
    )[:8]

    generated_at = datetime.now().astimezone().isoformat(timespec="seconds")
    source_files = [path.name for path in (header_path, policy_path, rights_path)]
    snapshot = {
        "generatedAt": generated_at,
        "asOf": SNAPSHOT_DATE.isoformat(),
        "sourceFiles": source_files,
        "policy": {
            "total": len(policy_records),
            "year2026": years.get("2026", 0),
            "provinceCount": len({record["省份"] for record in policy_records if record["省份"]}),
            "cityCount": len({record["城市"] for record in policy_records if record["城市"]}),
            "topics": [{"label": label, "count": topics[label]} for label in ("购车补贴", "消费券", "以旧换新", "金融贴息", "新能源专项")],
            "latest": latest_policy,
            "download": "/assets/data/policy-intelligence.json",
        },
        "rights": {
            "total": len(rights_records),
            "brandCount": len({record["车企/品牌"] for record in rights_records if record["车企/品牌"]}),
            "modelCount": len({(record["车企/品牌"], record["车型"]) for record in rights_records if record["车企/品牌"] or record["车型"]}),
            "multiTypeCount": multi_type,
            "categories": [{"label": label, "count": category_counts[label]} for label in RIGHTS_TYPES],
            "latest": latest_rights,
            "download": "/assets/data/rights-intelligence.json",
        },
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    write_json(OUTPUT_DIR / "intelligence-snapshot.json", snapshot)
    write_json(
        OUTPUT_DIR / "policy-intelligence.json",
        {
            "generatedAt": generated_at,
            "asOf": SNAPSHOT_DATE.isoformat(),
            "source": policy_path.name,
            "template": header_path.name,
            "extractionMethod": "从地方政策数据 N 列“主要内容”进行语义提炼",
            "headers": policy_template_headers,
            "records": policy_records,
        },
    )
    write_json(
        OUTPUT_DIR / "rights-intelligence.json",
        {
            "generatedAt": generated_at,
            "asOf": SNAPSHOT_DATE.isoformat(),
            "source": rights_path.name,
            "template": header_path.name,
            "activityTypeOptions": RIGHTS_TYPES,
            "headers": rights_template_headers,
            "records": rights_records,
        },
    )
    print(json.dumps({"output": str(OUTPUT_DIR), "policy": len(policy_records), "rights": len(rights_records)}, ensure_ascii=False))


if __name__ == "__main__":
    build()
