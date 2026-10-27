"""Read-only comparison of six business decks and the production template."""
from pathlib import Path
import json
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
REFERENCES = Path(r'E:/AI/migration-backup/5.4 晓亚\提供邓skill-0915')
NAMES = [f'2026年{m}月中国乘用车市场月度分析报告{suffix}.pptx'
         for m, suffix in [(3,''),(4,''),(5,''),(6,'0715'),(7,'0814'),(8,'0915')]]
TEMPLATE = ROOT/'skills/market-analysis-skill/passenger-car-monthly-report/assets/乘用车月报PPT兼容模板.pptx'
OUT = ROOT/'outputs/monthly-report-data/business-contract-review-20261008'

def audit(path):
    deck = Presentation(path)
    result = {'source':str(path), 'slide_count':len(deck.slides), 'pages':{}}
    topics = {5:'库存(行业)',7:'分车型类别市场表现',8:'2.2 车型级别',12:'2.4 车型系别',
              13:'2.5 车型价格段',15:'各线',17:'新能源乘用车市场月度走势'}
    for page, topic in topics.items():
        if path == TEMPLATE:
            actual = page-1
        else:
            candidates = []
            for i, candidate in enumerate(deck.slides):
                texts = [s.text.replace('（','(').replace('）',')') for s in candidate.shapes if s.has_text_frame]
                combined = '\n'.join(texts)
                matched = topic in combined
                if page==15:
                    matched = '2.6 区域市场' in combined and any(s.has_chart for s in candidate.shapes) and any('累计份额' in t for t in texts)
                    matched = matched and ('城市' in combined or '线级' in combined or '线城市' in combined)
                if page==17 and matched:
                    matched = not any(term in texts[0] for term in ('纯电','插混','动力','BEV','PHEV'))
                if matched:
                    candidates.append(i)
            if len(candidates)!=1:
                raise ValueError(f'{path.name} topic {page} matches {candidates}')
            actual=candidates[0]
        slide = deck.slides[actual]
        texts = [s.text for s in slide.shapes if s.has_text_frame and s.text.strip()]
        charts = [s for s in slide.shapes if s.has_chart]
        result['pages'][str(page)] = {
            'source_slide':actual+1,
            'texts':texts,
            'chart_geometry':[[s.shape_id,s.left,s.top,s.width,s.height] for s in charts],
            'charts':[{'part':str(s.chart.part.partname),
                       'series':[{'name':series.name,'values':list(series.values)} for series in s.chart.series]}
                      for s in charts],
        }
    return result

if __name__ == '__main__':
    OUT.mkdir(parents=True,exist_ok=True)
    results = [audit(REFERENCES/name) for name in NAMES] + [audit(TEMPLATE)]
    (OUT/'reference-contracts.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
    for result in results:
        print(Path(result['source']).name)
        for page, data in result['pages'].items():
            if page=='5':
                print('  5:', data['texts'][:2], 'charts:',data['chart_geometry'])
            else:
                periods=[t for t in data['texts'] if '2025.1-' in t or '2026.1-' in t]
                print(' ',page,periods)
                if page!='17':
                    print('    prior reference:',[s['values'][2] for s in data['charts'][-1]['series']])
