"""Generic AgentScope tools exposed to every activated Skill package.

The tools are capability primitives, not a hard-coded business workflow.  The
model decides how to combine them from the active Skill's instructions and its
own persisted task graph.
"""

from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys
from typing import Any
from zipfile import ZipFile
from xml.etree import ElementTree

from agentscope.message import TextBlock
from agentscope.permission import PermissionBehavior, PermissionContext, PermissionDecision
from agentscope.tool import ToolBase, ToolChunk
from pptx import Presentation
from pptx.chart.data import CategoryChartData

from tools.custom_tools import record_tool_call

from .core import SkillRunContext, SkillRuntimeError


def _chunk(value: Any) -> ToolChunk:
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2)
    return ToolChunk(content=[TextBlock(text=text)])


class SkillRuntimeTool(ToolBase):
    """Allow tools only inside the already-authorized run workspace."""

    is_concurrency_safe = False

    def __init__(self, context: SkillRunContext) -> None:
        self.context = context

    async def check_permissions(
        self,
        tool_input: dict[str, Any],
        context: PermissionContext,
    ) -> PermissionDecision:
        return PermissionDecision(
            behavior=PermissionBehavior.ALLOW,
            message=f"{self.name} 仅访问已激活 Skill 与隔离运行工作区。",
        )

    def audit(self, status: str, detail: str) -> None:
        record_tool_call(self.name, status, f"run_id={self.context.run_id}; {detail}")


class SkillPackageInspectTool(SkillRuntimeTool):
    name = "skill_package_inspect"
    description = "检查当前已激活的完整 Skill 包、模板清单、资源和依赖；开始执行 Skill 时优先调用。"
    input_schema = {"type": "object", "properties": {}}

    async def call(self) -> ToolChunk:
        package = self.context.package
        result = {
            "run_id": self.context.run_id,
            "skill_id": package.skill_id,
            "name": package.name,
            "description": package.description,
            "package_hash": package.package_hash,
            "artifact_template": package.artifact_template,
            "dependencies": package.dependencies,
            "resources": package.resources(),
            "workspace": {
                "inputs": "inputs/",
                "working": "working/",
                "outputs": "outputs/",
                "state": "run-state.json",
            },
        }
        self.audit("success", f"resources={len(result['resources'])}")
        return _chunk(result)


class SkillReadResourceTool(SkillRuntimeTool):
    name = "skill_read_resource"
    description = "按需读取当前 Skill 包中的文本资源。二进制资产返回元数据，不会把二进制塞入上下文。"
    input_schema = {
        "type": "object",
        "properties": {
            "relative_path": {"type": "string", "description": "Skill 根目录下的相对路径。"},
            "max_chars": {"type": "integer", "default": 20000},
        },
        "required": ["relative_path"],
    }

    async def call(self, relative_path: str, max_chars: int = 20000) -> ToolChunk:
        path = self.context.package.resolve(relative_path)
        if not path.is_file():
            raise SkillRuntimeError(f"不是可读取文件：{relative_path}")
        if path.suffix.lower() in {".md", ".txt", ".json", ".yaml", ".yml", ".csv"}:
            text = path.read_text(encoding="utf-8", errors="replace")[: max(1, min(max_chars, 50000))]
            result: Any = {"path": relative_path, "text": text, "truncated": len(text) >= max_chars}
        else:
            result = {"path": relative_path, "bytes": path.stat().st_size, "binary": True}
        self.audit("success", f"path={relative_path}")
        return _chunk(result)


class SkillPlanTasksTool(SkillRuntimeTool):
    name = "skill_plan_tasks"
    description = "由智能体建立或重建可验证任务图。每个节点必须包含 id、title、success_criteria 和 depends_on。"
    input_schema = {
        "type": "object",
        "properties": {
            "tasks_json": {
                "type": "string",
                "description": "JSON 数组，例如 [{\"id\":\"inspect\",\"title\":\"检查模板\",\"success_criteria\":\"已读取模板结构\",\"depends_on\":[]}]。",
            }
        },
        "required": ["tasks_json"],
    }

    async def call(self, tasks_json: str) -> ToolChunk:
        tasks = json.loads(tasks_json)
        if not isinstance(tasks, list):
            raise SkillRuntimeError("tasks_json 必须是 JSON 数组。")
        state = self.context.set_plan(tasks)
        self.audit("success", f"tasks={len(state['plan'])}")
        return _chunk({"run_id": self.context.run_id, "plan": state["plan"]})


class SkillUpdateTaskTool(SkillRuntimeTool):
    name = "skill_update_task"
    description = "根据真实工具结果更新任务节点，附加产物路径或来源等可核验证据。"
    input_schema = {
        "type": "object",
        "properties": {
            "task_id": {"type": "string"},
            "status": {"type": "string", "enum": ["pending", "running", "completed", "failed", "blocked"]},
            "summary": {"type": "string", "default": ""},
            "evidence_json": {"type": "string", "default": "[]"},
        },
        "required": ["task_id", "status"],
    }

    async def call(
        self,
        task_id: str,
        status: str,
        summary: str = "",
        evidence_json: str = "[]",
    ) -> ToolChunk:
        evidence = json.loads(evidence_json or "[]")
        if not isinstance(evidence, list):
            raise SkillRuntimeError("evidence_json 必须是 JSON 数组。")
        task = self.context.update_task(task_id, status, summary, evidence)
        self.audit("success", f"task={task_id}; status={status}")
        return _chunk(task)


class SkillCheckpointTool(SkillRuntimeTool):
    name = "skill_checkpoint"
    description = "保存可恢复检查点；记录已完成工作、下一步和关键证据，不记录隐藏思维过程。"
    input_schema = {
        "type": "object",
        "properties": {
            "summary": {"type": "string"},
            "next_action": {"type": "string", "default": ""},
            "evidence_json": {"type": "string", "default": "[]"},
        },
        "required": ["summary"],
    }

    async def call(self, summary: str, next_action: str = "", evidence_json: str = "[]") -> ToolChunk:
        evidence = json.loads(evidence_json or "[]")
        self.context.checkpoint(summary, {"next_action": next_action, "evidence": evidence})
        self.audit("success", summary[:120])
        return _chunk({"saved": True, "run_id": self.context.run_id, "next_action": next_action})


class SkillCopyAssetTool(SkillRuntimeTool):
    name = "skill_copy_asset"
    description = "把 Skill 包中的只读资产复制到本次输出目录，后续只修改副本，绝不覆盖原模板。"
    input_schema = {
        "type": "object",
        "properties": {
            "source_relative_path": {"type": "string"},
            "output_name": {"type": "string"},
        },
        "required": ["source_relative_path", "output_name"],
    }

    async def call(self, source_relative_path: str, output_name: str) -> ToolChunk:
        target = self.context.copy_asset(source_relative_path, output_name)
        relative = target.relative_to(self.context.root).as_posix()
        self.audit("success", f"source={source_relative_path}; output={relative}")
        return _chunk({"path": relative, "bytes": target.stat().st_size})


class WorkspaceListTool(SkillRuntimeTool):
    name = "workspace_list"
    description = "列出当前 Skill 运行工作区中的输入、过程文件和输出，便于恢复任务和核对产物。"
    input_schema = {
        "type": "object",
        "properties": {"prefix": {"type": "string", "default": ""}},
    }

    async def call(self, prefix: str = "") -> ToolChunk:
        base = self.context.root if not prefix else self.context.workspace_path(prefix, must_exist=True)
        if not base.is_dir():
            raise SkillRuntimeError(f"不是目录：{prefix}")
        files = [
            {
                "path": path.relative_to(self.context.root).as_posix(),
                "bytes": path.stat().st_size,
            }
            for path in sorted(base.rglob("*"))
            if path.is_file()
        ]
        self.audit("success", f"prefix={prefix}; files={len(files)}")
        return _chunk({"run_id": self.context.run_id, "files": files})


class WorkspaceReadTextTool(SkillRuntimeTool):
    name = "workspace_read_text"
    description = "读取当前运行工作区中的文本、JSON、CSV 或 YAML 文件。"
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "max_chars": {"type": "integer", "default": 50000},
        },
        "required": ["path"],
    }

    async def call(self, path: str, max_chars: int = 50000) -> ToolChunk:
        resolved = self.context.workspace_path(path, must_exist=True)
        if resolved.suffix.casefold() not in {".md", ".txt", ".json", ".csv", ".yaml", ".yml"}:
            raise SkillRuntimeError("该工具只读取文本、JSON、CSV 或 YAML 文件。")
        limit = max(1, min(int(max_chars), 100000))
        raw = resolved.read_text(encoding="utf-8", errors="replace")
        self.audit("success", f"path={path}; chars={min(len(raw), limit)}")
        return _chunk({"path": path, "text": raw[:limit], "truncated": len(raw) > limit})


class WorkspaceWriteTextTool(SkillRuntimeTool):
    name = "workspace_write_text"
    description = "把结构化中间结果或报告写入当前运行工作区；仅允许安全文本格式。"
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "working/ 或 outputs/ 下的相对路径。"},
            "content": {"type": "string"},
            "artifact_kind": {"type": "string", "default": ""},
        },
        "required": ["path", "content"],
    }

    async def call(self, path: str, content: str, artifact_kind: str = "") -> ToolChunk:
        normalized = str(path).replace("\\", "/")
        if not normalized.startswith(("working/", "outputs/")):
            raise SkillRuntimeError("只能写入 working/ 或 outputs/。")
        resolved = self.context.workspace_path(normalized)
        if resolved.suffix.casefold() not in {".md", ".txt", ".json", ".csv", ".yaml", ".yml"}:
            raise SkillRuntimeError("该工具只写入文本、JSON、CSV 或 YAML 文件。")
        resolved.parent.mkdir(parents=True, exist_ok=True)
        resolved.write_text(str(content), encoding="utf-8")
        if artifact_kind:
            self.context.register_artifact(resolved, str(artifact_kind))
        self.context.checkpoint("已写入运行文件", {"path": normalized, "bytes": resolved.stat().st_size})
        self.audit("success", f"path={normalized}; bytes={resolved.stat().st_size}")
        return _chunk({"path": normalized, "bytes": resolved.stat().st_size})


class SkillRunScriptTool(SkillRuntimeTool):
    name = "skill_run_script"
    description = "在隔离运行工作区中执行完整 Skill 包 scripts/ 下的受信脚本；支持 Python、PowerShell 和 Node。"
    input_schema = {
        "type": "object",
        "properties": {
            "relative_path": {"type": "string", "description": "必须位于 Skill 包 scripts/ 下。"},
            "args_json": {"type": "string", "default": "[]", "description": "参数 JSON 数组。"},
            "timeout_seconds": {"type": "integer", "default": 180, "minimum": 1, "maximum": 600},
        },
        "required": ["relative_path"],
    }

    async def call(
        self,
        relative_path: str,
        args_json: str = "[]",
        timeout_seconds: int = 180,
    ) -> ToolChunk:
        normalized = str(relative_path).replace("\\", "/")
        if not normalized.startswith("scripts/"):
            raise SkillRuntimeError("只能执行当前 Skill 包 scripts/ 下的脚本。")
        if self.context.package.policy.get("allow_package_scripts", True) is False:
            raise SkillRuntimeError("该 Skill 策略禁止执行包内脚本。")
        script = self.context.package.resolve(normalized)
        if not script.is_file():
            raise SkillRuntimeError(f"Skill 脚本不存在：{normalized}")
        arguments = json.loads(args_json or "[]")
        if not isinstance(arguments, list) or len(arguments) > 30:
            raise SkillRuntimeError("args_json 必须是最多 30 项的 JSON 数组。")
        arguments = [str(item) for item in arguments]
        extension = script.suffix.casefold()
        if extension == ".py":
            command = [sys.executable, str(script), *arguments]
        elif extension == ".ps1":
            command = [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script),
                *arguments,
            ]
        elif extension in {".js", ".mjs", ".cjs"}:
            command = ["node", str(script), *arguments]
        else:
            raise SkillRuntimeError("仅允许执行 .py、.ps1、.js、.mjs 或 .cjs。")
        runtime_env = os.environ.copy()
        runtime_env.update(
            {
                "SKILL_PACKAGE_ROOT": str(self.context.package.root),
                "SKILL_RUN_ROOT": str(self.context.root),
                "SKILL_INPUTS_DIR": str(self.context.inputs_dir),
                "SKILL_WORKING_DIR": str(self.context.working_dir),
                "SKILL_OUTPUTS_DIR": str(self.context.outputs_dir),
                "PYTHONUTF8": "1",
                "PYTHONDONTWRITEBYTECODE": "1",
            }
        )
        process = subprocess.run(
            command,
            cwd=self.context.working_dir,
            env=runtime_env,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=max(1, min(int(timeout_seconds), 600)),
            check=False,
        )
        result = {
            "script": normalized,
            "returncode": process.returncode,
            "success": process.returncode == 0,
            "stdout": process.stdout[-20000:],
            "stderr": process.stderr[-10000:],
        }
        self.context.checkpoint("Skill 包内脚本执行完成", result)
        self.audit("success" if result["success"] else "warning", f"script={normalized}; rc={process.returncode}")
        return _chunk(result)


def _presentation_inventory(path: Path, include_text: bool = True) -> dict[str, Any]:
    deck = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    chart_count = 0
    table_count = 0
    for slide_number, slide in enumerate(deck.slides, start=1):
        row: dict[str, Any] = {"slide": slide_number, "texts": [], "charts": [], "tables": []}
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if include_text and getattr(shape, "has_text_frame", False):
                text = str(shape.text or "").strip()
                if text:
                    row["texts"].append({"shape": shape_index, "text": text[:1000]})
            if getattr(shape, "has_chart", False):
                chart_count += 1
                chart = shape.chart
                series = []
                for item in chart.series:
                    try:
                        values = list(item.values)
                    except Exception:
                        values = []
                    series.append({"name": str(item.name), "values": values})
                try:
                    categories = [str(item) for item in chart.plots[0].categories]
                except Exception:
                    categories = []
                row["charts"].append(
                    {"shape": shape_index, "chart": len(row["charts"]) + 1, "series": series, "categories": categories}
                )
            if getattr(shape, "has_table", False):
                table_count += 1
                table = shape.table
                row["tables"].append(
                    {
                        "shape": shape_index,
                        "rows": len(table.rows),
                        "columns": len(table.columns),
                        "values": [[cell.text for cell in table_row.cells] for table_row in table.rows],
                    }
                )
        slides.append(row)
    return {
        "path": str(path),
        "slides": len(deck.slides),
        "charts": chart_count,
        "tables": table_count,
        "inventory": slides,
    }


def _external_file_relationships(archive: ZipFile) -> list[str]:
    """Return external file/workbook links while allowing ordinary web hyperlinks."""

    blocked: list[str] = []
    for name in archive.namelist():
        if not name.endswith(".rels"):
            continue
        try:
            root = ElementTree.fromstring(archive.read(name))
        except ElementTree.ParseError:
            continue
        for relationship in root:
            if relationship.attrib.get("TargetMode") != "External":
                continue
            target = relationship.attrib.get("Target", "")
            rel_type = relationship.attrib.get("Type", "")
            target_lower = target.casefold()
            is_web = target_lower.startswith(("http://", "https://", "mailto:"))
            is_hyperlink = rel_type.casefold().endswith("/hyperlink")
            if not (is_web and is_hyperlink):
                blocked.append(f"{name}: {target or rel_type}")
    return blocked


class PresentationInspectTool(SkillRuntimeTool):
    name = "presentation_inspect"
    description = "读取 Skill 模板或运行工作区 PPTX 的页、文本、原生图表、系列、类别和表格结构。"
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "skill:assets/x.pptx 或 outputs/x.pptx。"},
            "include_text": {"type": "boolean", "default": True},
        },
        "required": ["path"],
    }

    def _resolve(self, value: str) -> Path:
        if value.startswith("skill:"):
            return self.context.package.resolve(value[6:])
        return self.context.workspace_path(value, must_exist=True)

    async def call(self, path: str, include_text: bool = True) -> ToolChunk:
        resolved = self._resolve(path)
        result = _presentation_inventory(resolved, bool(include_text))
        result["path"] = path
        self.audit("success", f"path={path}; slides={result['slides']}; charts={result['charts']}")
        return _chunk(result)


class PresentationReplaceTextTool(SkillRuntimeTool):
    name = "presentation_replace_text"
    description = "在运行工作区的 PPTX 副本中按精确字符串批量替换文本，保留模板版式。"
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "replacements_json": {"type": "string", "description": "JSON 对象：旧文本到新文本。"},
        },
        "required": ["path", "replacements_json"],
    }

    async def call(self, path: str, replacements_json: str) -> ToolChunk:
        resolved = self.context.workspace_path(path, must_exist=True)
        replacements = json.loads(replacements_json)
        if not isinstance(replacements, dict) or not replacements:
            raise SkillRuntimeError("replacements_json 必须是非空 JSON 对象。")
        deck = Presentation(str(resolved))
        changes = 0
        for slide in deck.slides:
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        updated = run.text
                        for old, new in replacements.items():
                            updated = updated.replace(str(old), str(new))
                        if updated != run.text:
                            run.text = updated
                            changes += 1
        deck.save(str(resolved))
        self.context.register_artifact(resolved, "presentation")
        self.audit("success", f"path={path}; changed_runs={changes}")
        return _chunk({"path": path, "changed_runs": changes})


class PresentationUpdateChartTool(SkillRuntimeTool):
    name = "presentation_update_chart"
    description = "按页码和图表序号更新 PPTX 副本中的原生可编辑图表数据，保留图表对象而非截图。"
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "slide_number": {"type": "integer", "minimum": 1},
            "chart_number": {"type": "integer", "minimum": 1},
            "categories_json": {"type": "string", "description": "JSON 数组。"},
            "series_json": {"type": "string", "description": "JSON 数组，每项包含 name 和 values。"},
        },
        "required": ["path", "slide_number", "chart_number", "categories_json", "series_json"],
    }

    async def call(
        self,
        path: str,
        slide_number: int,
        chart_number: int,
        categories_json: str,
        series_json: str,
    ) -> ToolChunk:
        resolved = self.context.workspace_path(path, must_exist=True)
        categories = json.loads(categories_json)
        series = json.loads(series_json)
        if not isinstance(categories, list) or not isinstance(series, list) or not series:
            raise SkillRuntimeError("categories_json 和 series_json 必须是有效数组。")
        deck = Presentation(str(resolved))
        if not 1 <= int(slide_number) <= len(deck.slides):
            raise SkillRuntimeError("slide_number 超出范围。")
        charts = [shape.chart for shape in deck.slides[int(slide_number) - 1].shapes if getattr(shape, "has_chart", False)]
        if not 1 <= int(chart_number) <= len(charts):
            raise SkillRuntimeError("chart_number 超出范围。")
        data = CategoryChartData()
        data.categories = categories
        for item in series:
            values = item.get("values") or []
            if len(values) != len(categories):
                raise SkillRuntimeError("每个图表系列的 values 数量必须与 categories 一致。")
            data.add_series(str(item.get("name") or "系列"), tuple(values))
        charts[int(chart_number) - 1].replace_data(data)
        deck.save(str(resolved))
        self.context.register_artifact(resolved, "presentation")
        self.audit("success", f"path={path}; slide={slide_number}; chart={chart_number}")
        return _chunk({"path": path, "slide": slide_number, "chart": chart_number, "series": len(series)})


class ForecastSeriesTool(SkillRuntimeTool):
    name = "forecast_series"
    description = "对已给出的同口径历史序列执行透明的基准预测。支持 linear_trend、seasonal_naive 和 drift；返回预测值及方法说明。"
    input_schema = {
        "type": "object",
        "properties": {
            "periods_json": {"type": "string", "description": "历史期间 JSON 数组。"},
            "values_json": {"type": "string", "description": "历史数值 JSON 数组。"},
            "horizon": {"type": "integer", "minimum": 1, "maximum": 36},
            "method": {"type": "string", "enum": ["linear_trend", "seasonal_naive", "drift"], "default": "linear_trend"},
            "season_length": {"type": "integer", "minimum": 1, "maximum": 24, "default": 12},
        },
        "required": ["periods_json", "values_json", "horizon"],
    }

    async def call(
        self,
        periods_json: str,
        values_json: str,
        horizon: int,
        method: str = "linear_trend",
        season_length: int = 12,
    ) -> ToolChunk:
        periods = json.loads(periods_json)
        raw_values = json.loads(values_json)
        values = [float(value) for value in raw_values]
        if not isinstance(periods, list) or len(periods) != len(values) or len(values) < 2:
            raise SkillRuntimeError("历史期间与数值必须等长，且至少包含两个观测值。")
        horizon = max(1, min(int(horizon), 36))
        if method == "seasonal_naive":
            season = max(1, min(int(season_length), len(values)))
            forecast = [values[(len(values) - season + index) % len(values)] for index in range(horizon)]
        elif method == "drift":
            slope = (values[-1] - values[0]) / max(1, len(values) - 1)
            forecast = [values[-1] + slope * step for step in range(1, horizon + 1)]
        else:
            n = len(values)
            x_mean = (n - 1) / 2
            y_mean = sum(values) / n
            denominator = sum((index - x_mean) ** 2 for index in range(n)) or 1.0
            slope = sum((index - x_mean) * (value - y_mean) for index, value in enumerate(values)) / denominator
            intercept = y_mean - slope * x_mean
            forecast = [intercept + slope * (n - 1 + step) for step in range(1, horizon + 1)]
        result = {
            "method": method,
            "history_periods": periods,
            "history_count": len(values),
            "forecast": [round(value, 6) for value in forecast],
            "warning": "这是透明基准预测；正式结论仍需结合指标定义、外生变量、情景假设与回测。",
        }
        self.audit("success", f"method={method}; history={len(values)}; horizon={horizon}")
        return _chunk(result)


class PresentationValidateTool(SkillRuntimeTool):
    name = "presentation_validate"
    description = "验证运行工作区 PPTX 的 ZIP 完整性、可重新读取性、页数、原生图表/表格和外部文件链接。"
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "expected_slides": {"type": "integer", "minimum": 1, "default": 0},
            "minimum_charts": {"type": "integer", "minimum": 0, "default": 0},
        },
        "required": ["path"],
    }

    async def call(self, path: str, expected_slides: int = 0, minimum_charts: int = 0) -> ToolChunk:
        resolved = self.context.workspace_path(path, must_exist=True)
        checks: list[dict[str, Any]] = []
        try:
            with ZipFile(resolved) as archive:
                bad = archive.testzip()
                external_files = _external_file_relationships(archive)
            checks.append({"id": "package_integrity", "passed": bad is None, "detail": bad or "OK"})
            checks.append(
                {
                    "id": "external_file_links",
                    "passed": not external_files,
                    "detail": "无外部文件/工作簿依赖" if not external_files else external_files[:20],
                }
            )
        except Exception as exc:
            checks.append({"id": "package_integrity", "passed": False, "detail": str(exc)})
        try:
            inventory = _presentation_inventory(resolved, include_text=False)
            checks.append({"id": "pptx_reopen", "passed": True, "detail": "python-pptx reopen OK"})
            if expected_slides:
                checks.append({"id": "slide_count", "passed": inventory["slides"] == expected_slides, "detail": f"{inventory['slides']} / {expected_slides}"})
            checks.append({"id": "editable_charts", "passed": inventory["charts"] >= int(minimum_charts), "detail": f"charts={inventory['charts']}; tables={inventory['tables']}"})
        except Exception as exc:
            inventory = {}
            checks.append({"id": "pptx_reopen", "passed": False, "detail": str(exc)})
        passed = all(item["passed"] for item in checks)
        self.context.register_artifact(resolved, "presentation")
        self.context.record_artifact_validation(resolved, "pptx_structure", passed, checks)
        self.context.checkpoint("PPTX 验收", {"path": path, "passed": passed, "checks": checks})
        self.audit("success" if passed else "warning", f"path={path}; passed={passed}")
        return _chunk({"path": path, "passed": passed, "checks": checks, "inventory": inventory})


class PresentationPowerPointValidateTool(SkillRuntimeTool):
    name = "presentation_powerpoint_validate"
    description = "使用本机 Microsoft PowerPoint 真正打开并逐页渲染 PPTX，检查 Office 兼容性和渲染完整性。"
    input_schema = {
        "type": "object",
        "properties": {"path": {"type": "string"}},
        "required": ["path"],
    }

    async def call(self, path: str) -> ToolChunk:
        resolved = self.context.workspace_path(path, must_exist=True)
        render_dir = self.context.working_dir / "rendered" / resolved.stem
        render_dir.mkdir(parents=True, exist_ok=True)
        script = Path(__file__).resolve().parents[1] / "scripts" / "validate_presentation.ps1"
        if not script.is_file():
            raise SkillRuntimeError(f"PowerPoint 验收脚本不存在：{script}")
        process = subprocess.run(
            [
                "powershell.exe",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script),
                "-InputPath",
                str(resolved),
                "-OutputDir",
                str(render_dir),
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            check=False,
        )
        if process.returncode != 0:
            result = {
                "path": path,
                "passed": False,
                "detail": (process.stderr or process.stdout or "PowerPoint 验收失败")[-3000:],
            }
        else:
            lines = [line.strip() for line in process.stdout.splitlines() if line.strip()]
            try:
                payload = json.loads(lines[-1])
            except (IndexError, json.JSONDecodeError) as exc:
                raise SkillRuntimeError(f"无法解析 PowerPoint 验收结果：{process.stdout[-1000:]}") from exc
            result = {
                "path": path,
                "passed": bool(payload.get("success")),
                "slides": int(payload.get("slides") or 0),
                "rendered": int(payload.get("rendered") or 0),
                "render_dir": render_dir.relative_to(self.context.root).as_posix(),
            }
        self.context.checkpoint("Microsoft PowerPoint 验收", result)
        self.context.record_artifact_validation(
            resolved,
            "powerpoint_render",
            bool(result["passed"]),
            result,
        )
        self.audit("success" if result["passed"] else "warning", f"path={path}; passed={result['passed']}")
        return _chunk(result)


class SkillValidateRunTool(SkillRuntimeTool):
    name = "skill_validate_run"
    description = "独立检查任务图和必需交付物。只有 passed=true 时智能体才可以宣告整个 Skill 完成。"
    input_schema = {
        "type": "object",
        "properties": {
            "required_artifacts_json": {"type": "string", "default": "[]"}
        },
    }

    async def call(self, required_artifacts_json: str = "[]") -> ToolChunk:
        artifacts = json.loads(required_artifacts_json or "[]")
        if not isinstance(artifacts, list):
            raise SkillRuntimeError("required_artifacts_json 必须是 JSON 数组。")
        validation = self.context.validate([str(item) for item in artifacts])
        self.audit("success" if validation["status"] == "passed" else "warning", validation["status"])
        return _chunk(validation)


def build_skill_runtime_tools(context: SkillRunContext) -> list[ToolBase]:
    """Return the same package-aware capability surface for every Skill."""

    return [
        SkillPackageInspectTool(context),
        SkillReadResourceTool(context),
        SkillPlanTasksTool(context),
        SkillUpdateTaskTool(context),
        SkillCheckpointTool(context),
        SkillCopyAssetTool(context),
        WorkspaceListTool(context),
        WorkspaceReadTextTool(context),
        WorkspaceWriteTextTool(context),
        SkillRunScriptTool(context),
        PresentationInspectTool(context),
        PresentationReplaceTextTool(context),
        PresentationUpdateChartTool(context),
        ForecastSeriesTool(context),
        PresentationValidateTool(context),
        PresentationPowerPointValidateTool(context),
        SkillValidateRunTool(context),
    ]
