"""Read-only contracts for intentional annual references and inventory layout."""
from hashlib import sha256
from io import BytesIO
import re
from openpyxl import load_workbook

INVENTORY_TITLE_ID = 1048746
INVENTORY_CHART_TITLE_ID = 1048747
NARRATIVE_BODY_NAME = 'Monthly narrative body'
ANNUAL_REFERENCE_LABELS = {7:1048838, 8:1048889, 12:1049044, 13:1049099, 15:1049158}

def _shape(slide, identifier):
    matches = [s for s in slide.shapes if s.shape_id == identifier]
    if len(matches)!=1:
        raise ValueError(f'Expected one template shape id {identifier}')
    return matches[0]

def prepare_inventory_roles(slide):
    """Return title and no body; preserve chart caption and native geometry."""
    title = _shape(slide, INVENTORY_TITLE_ID)
    caption = _shape(slide, INVENTORY_CHART_TITLE_ID)
    if any(s.name == NARRATIVE_BODY_NAME for s in slide.shapes):
        raise ValueError('库存页不应包含新增正文；请使用未压缩图表的已核验兼容模板。')
    if caption.text != '2025-2026年乘用车库存量月度走势':
        raise ValueError('库存页图表标题不符合已确认模板。')
    geometry = {4194342:(701748,1516912,7046418,4924640),
                4194343:(7851963,1528548,3897125,4924639)}
    for identifier, expected in geometry.items():
        chart = _shape(slide, identifier)
        if not chart.has_chart or (chart.left,chart.top,chart.width,chart.height)!=expected:
            raise ValueError('库存页原生图表位置/尺寸发生变化，不能以压缩图表容纳正文。')
    return title, None

def _verify_share_workbook(chart):
    """Prove all four period values agree without rewriting workbook data."""
    book_part=chart.part.chart_workbook.xlsx_part
    if book_part is None:
        raise ValueError('结构份额图缺少内嵌工作簿。')
    book=load_workbook(BytesIO(book_part.blob),data_only=True)
    try:
        for index, series in enumerate(chart.series):
            nodes=series._ser.xpath('.//c:val/c:numRef/c:f')
            match=re.fullmatch(r"'?([^'!]+)'?!\$?A\$?(\d+):\$?D\$?\2",nodes[0].text if len(nodes)==1 else '')
            if not match or int(match[2])!=index+1 or len(series.values)!=4:
                raise ValueError('结构份额图系列/四期数据引用不符合模板。')
            sheet=match[1]
            if sheet not in book.sheetnames and sheet=='Sheet1' and book.sheetnames==['renameSheet']:
                sheet='renameSheet'
            if sheet not in book.sheetnames:
                raise ValueError('结构份额图引用的工作簿sheet不存在。')
            for col, value in enumerate(series.values,1):
                cell=book[sheet].cell(index+1,col).value
                if not isinstance(cell,(int,float)) or abs(cell-value)>1e-8:
                    raise ValueError('结构份额图cache与内嵌工作簿不一致。')
    finally:
        book.close()

def validate_business_chart_roles(deck, report_month, retained_pages=()):
    """Validate, rather than normalize, page-specific comparison semantics."""
    match=re.fullmatch(r'(\d{4})-(0[1-9]|1[0-2])',report_month)
    if not match:
        raise ValueError('report_month must be YYYY-MM')
    year, month=int(match[1]),int(match[2])
    prepare_inventory_roles(deck.slides[4])
    contracts=[]
    for page,label_id in ANNUAL_REFERENCE_LABELS.items():
        if page in retained_pages:
            contracts.append({'slide':page,'scope':'retained_source_update_delayed'})
            continue
        slide=deck.slides[page-1]
        texts=[s.text for s in slide.shapes if s.has_text_frame]
        if _shape(slide,label_id).text!=f'{year-1}.1-12' or f'{year}.1-{month}' not in texts:
            raise ValueError(f'第{page}页必须保留上年全年结构参照，不得改为同期累计。')
        charts=[s.chart for s in slide.shapes if s.has_chart]
        if len(charts)!=3:
            raise ValueError(f'第{page}页原生图表数量不符合模板。')
        share=charts[2]
        _verify_share_workbook(share)
        if abs(sum(s.values[2] for s in share.series)-100)>0.3:
            raise ValueError(f'第{page}页全年结构份额合计异常。')
        contracts.append({'slide':page,'reference_period':f'{year-1}.1-12',
                          'current_period':f'{year}.1-{month}',
                          'scope':'prior_year_full_year_structure',
                          'reference_values':[s.values[2] for s in share.series],
                          'chart_sha256':sha256(share.part.blob).hexdigest(),
                          'workbook_sha256':sha256(share.part.chart_workbook.xlsx_part.blob).hexdigest()})
    if 17 in retained_pages:
        return contracts
    texts=[s.text for s in deck.slides[16].shapes if s.has_text_frame]
    if f'{year-1}.1-{month}' not in texts or f'{year}.1-{month}' not in texts:
        raise ValueError('第17页累计表现必须使用去年同期与本年同期。')
    contracts.append({'slide':17,'reference_period':f'{year-1}.1-{month}',
                      'current_period':f'{year}.1-{month}','scope':'prior_year_same_period'})
    return contracts
