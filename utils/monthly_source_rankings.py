"""Editable ranking panels populated only from the accepted source package.

The narrative packet contains representative models, not a full ranking. Never
use its selected four models to populate TOP15/TOP20 panels.
"""
from utils.monthly_segment_facts import _model_rows, _level_rows
from utils.monthly_page_facts import _entities, _Packet, _prior, _previous
import re


def _text(cards, code):
    return str(cards.get(code, {}).get('text', ''))


def _required(rows, count, code):
    if len(rows) < count:
        raise ValueError(f'{code}榜单数据不完整：已解析{len(rows)}项，需要{count}项；请补采完整表格，不能沿用模板旧值。')
    return sorted(rows, key=lambda row: -row['sales'])[:count]


def _denominator(value, code):
    if not value or value <= 0:
        raise ValueError(f'{code}缺少目标月份销量分母，不能沿用旧份额。')
    return value


def _rate(value, comparison):
    return (value / comparison - 1) * 100 if comparison else None


def _models(cards, code, month, cumulative=False):
    grid = cards.get(code, {}).get('grid_snapshot')
    if grid:
        if grid.get('month') != month or grid.get('complete') is not True:
            raise ValueError(f'{code}结构化表格月份或完整性未核验。')
        headers = {(int(c['row']), int(c['col'])): str(c.get('text', '')).strip()
                   for c in grid.get('headers', [])}
        byrow = {}
        for cell in grid.get('cells', []):
            byrow.setdefault(int(cell['row']), {})[int(cell['col'])] = str(cell.get('text', '')).strip()
        namecols = [col for (row, col), value in headers.items() if value == '车型']
        levelcols = [col for (row, col), value in headers.items() if value.startswith('车型级别')]
        quantity = {}
        for (row, col), value in headers.items():
            if row == 1 and (re.fullmatch(r'\d{4}', value) if cumulative else re.fullmatch(r'\d{4}-\d{2}', value)):
                if headers.get((0, col)) in ('数量', '销量', '终端销量'):
                    quantity[value] = col
        period = month[:4] if cumulative else month
        prior = str(int(month[:4]) - 1) if cumulative else _prior(month)
        previous = None if cumulative else _previous(month)
        if len(namecols) != 1 or period not in quantity:
            raise ValueError(f'{code}结构化车型/数量/目标期间列不明确。')
        def value(cells, key):
            text = cells.get(quantity.get(key), '')
            return int(text.replace(',', '')) if re.fullmatch(r'\d[\d,]*', text) else None
        rows = []
        for _, cells in sorted(byrow.items()):
            name = cells.get(namecols[0], '')
            if not name or name.startswith(('加载', '合计', '总计')):
                continue
            sales, baseline, previous_sales = value(cells, period), value(cells, prior), value(cells, previous)
            if sales is None:
                # A physically captured empty target cell is unpublished data,
                # not zero sales. Exclude it from the observed ranking; a
                # missing cell boundary is still a collection failure.
                if quantity[period] not in cells:
                    raise ValueError(f'{code}/{name}缺少目标期间销量单元格。')
                continue
            rows.append(dict(model=name, level=cells.get(levelcols[0], '') if len(levelcols) == 1 else '',
                             sales=sales, comparison_sales=baseline, previous_month_sales=previous_sales,
                             yoy_pct=_rate(sales, baseline), mom_pct=_rate(sales, previous_sales),
                             period=month, scope='year_to_date' if cumulative else 'single_month'))
        return rows
    # _model_rows deliberately excludes collapsed rows with ambiguous blanks.
    # Those rows must be reacquired with cell boundaries; they are not zeros.
    return _model_rows(_text(cards, code), month, cumulative)


def _enterprise_rows(cards, code, month, cumulative, warnings):
    grid = cards.get(code, {}).get('grid_snapshot')
    if grid:
        # Same quantity schema, with 企业简称 instead of 车型 and no level.
        proxy = dict(cards[code])
        proxy['grid_snapshot'] = dict(grid, headers=[dict(cell, text='车型' if cell.get('text') == '企业简称' else cell.get('text', '')) for cell in grid.get('headers', [])])
        return [dict(row, entity=row['model']) for row in _models({code: proxy}, code, month, cumulative)]
    parsed = _entities(_text(cards, code), code, month, warnings, cumulative)
    period = month[:4] if cumulative else month
    prior = str(int(month[:4]) - 1) if cumulative else _prior(month)
    previous = None if cumulative else _previous(month)
    return [dict(entity=name, model=name, sales=row['values'][period],
                 comparison_sales=row['values'].get(prior),
                 previous_month_sales=row['values'].get(previous),
                 yoy_pct=_rate(row['values'][period], row['values'].get(prior)),
                 mom_pct=_rate(row['values'][period], row['values'].get(previous)))
            for name, row in parsed.items() if period in row['values']]


def _shapes(slide):
    charts = [s for s in slide.shapes if s.has_chart]
    tables = [s for s in slide.shapes if s.has_table]
    return charts, tables


def _enterprise_labels(slide, count):
    # These template charts deliberately omit c:cat; company names are editable
    # overlay text boxes. Their drawing order is not their left-to-right order.
    groups = {}
    for shape in slide.shapes:
        if shape.has_text_frame and re.fullmatch(r'[\u4e00-\u9fff\-]{3,12}', shape.text):
            groups.setdefault(shape.top, []).append(shape)
    rows = [sorted(shapes, key=lambda s: s.left) for _, shapes in sorted(groups.items())
            if len(shapes) == count]
    if len(rows) != 2:
        raise ValueError('企业图表名称覆盖层数量不符合兼容模板，不能保留旧企业名。')
    return rows


def _neutral_model_colors(slide, chart):
    """Unknown fuel type must not inherit another model's semantic color."""
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.enum.chart import XL_DATA_LABEL_POSITION
    from pptx.dml.color import RGBColor
    for plot in chart.plots:
        plot.has_data_labels = True
        plot.data_labels.position = XL_DATA_LABEL_POSITION.INSIDE_END
        plot.data_labels.font.color.rgb = RGBColor(255,255,255)
    for series in chart.series:
        # Old per-series labels override the plot-level settings and carry
        # manual positions for last month's bar endpoints. Replace that label
        # configuration so current values stay inside the actual bars.
        for old_labels in series._ser.xpath('./c:dLbls'):
            series._ser.remove(old_labels)
        labels=OxmlElement('c:dLbls')
        fmt=OxmlElement('c:numFmt');fmt.set('formatCode','0.0');fmt.set('sourceLinked','0');labels.append(fmt)
        tx=OxmlElement('c:txPr');body=OxmlElement('a:bodyPr');lst=OxmlElement('a:lstStyle');paragraph=OxmlElement('a:p')
        props=OxmlElement('a:pPr');font=OxmlElement('a:defRPr');font.set('sz','900')
        fill=OxmlElement('a:solidFill');white=OxmlElement('a:srgbClr');white.set('val','FFFFFF');fill.append(white);font.append(fill)
        props.append(font);paragraph.append(props);tx.append(body);tx.append(lst);tx.append(paragraph);labels.append(tx)
        position=OxmlElement('c:dLblPos');position.set('val','inEnd');labels.append(position)
        for name,value in [('showLegendKey','0'),('showVal','1'),('showCatName','0'),('showSerName','0'),('showPercent','0')]:
            flag=OxmlElement('c:'+name);flag.set('val',value);labels.append(flag)
        anchor=series._ser.xpath('./c:cat')[0] if series._ser.xpath('./c:cat') else series._ser.xpath('./c:val')[0]
        anchor.addprevious(labels)
        for label in series._ser.xpath('./c:dLbls/c:dLbl'):
            for layout in label.xpath('./c:layout'):
                label.remove(layout)
            position = label.find('{*}dLblPos')
            if position is not None: position.set('val','inEnd')
        for props in series._ser.xpath('./c:dLbls//a:rPr | ./c:dLbls//a:defRPr | ./c:dLbls//a:endParaRPr'):
            for fill in props.xpath('./a:solidFill'):
                props.remove(fill)
            fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val','FFFFFF');fill.append(color);props.append(fill)
        for props in series._ser.xpath('./c:spPr | ./c:dPt/c:spPr'):
            for fill in props.xpath('./a:solidFill | ./a:gradFill | ./a:pattFill | ./a:noFill'):
                props.remove(fill)
            fill = OxmlElement('a:solidFill'); color = OxmlElement('a:srgbClr')
            color.set('val', '6D88A8'); fill.append(color); props.insert(0, fill)
    from utils.monthly_source_pptx import _text as write_text
    for shape in slide.shapes:
        if shape.has_text_frame and '蓝色填充为传统车型' in shape.text:
            value = re.sub(r'[，；,]?蓝色填充为传统车型[，,]绿色填充为新能源车型', '；柱形统一配色，能源类型未作区分', shape.text)
            write_text(shape.text_frame, value)


def _model_value_column(slide, chart_shape, table_shape, data):
    """Keep readable values in a fixed gutter instead of clipping short bars."""
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    chart=chart_shape.chart
    for flag in chart._chartSpace.xpath('.//c:dLbls/c:showVal'):
        flag.set('val','0')
    chart.value_axis.maximum_scale=max(r['sales']/10000 for r in data)*3.5
    y=table_shape.top+table_shape.table.rows[0].height
    for row, item in zip(list(table_shape.table.rows)[1:],data):
        box=slide.shapes.add_textbox(table_shape.left-Inches(.29),y,Inches(.27),row.height)
        box.name='Monthly model value '+item['model']
        box.text_frame.margin_left=box.text_frame.margin_right=0
        box.text_frame.margin_top=box.text_frame.margin_bottom=0
        box.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE
        paragraph=box.text_frame.paragraphs[0];paragraph.alignment=PP_ALIGN.RIGHT
        run=paragraph.add_run();run.text=f"{item['sales']/10000:.1f}";run.font.size=Pt(8)
        run.font.name='Arial';run.font.color.rgb=RGBColor(0,51,112)
        y+=row.height


def update_rankings(slide, page, cards, month, packet=None):
    """Update ranking charts and adjacent tables; raise before any data edits.

    Unknown original comparison ranks are shown as '待核' and explicitly logged,
    never calculated from an observed-only subset of vehicles or enterprises.
    """
    from utils.monthly_source_pptx import set_chart, write_rows, number, set_callouts, _text as write_text
    from utils.monthly_source_pptx import pct as precise_pct
    def pct(value, plus=False):
        return ('>1000%' if value > 1000 else '<-1000%') if value is not None and abs(value)>1000 else precise_pct(value,plus)
    warnings = []
    charts, tables = _shapes(slide)
    period_label = f'{int(month[:4])}年{int(month[5:])}月'
    ytd_label = f'{int(month[:4])}年1—{int(month[5:])}月'
    warnings.append({'page': page, 'reason': 'comparison_rank_unverified',
                     'message': '排名变化显示待核；完整原始比较排名未核验，不以观测子集代替。'})
    if page in (9, 10, 11, 20):
        warnings.append({'page': page, 'reason': 'fuel_type_unverified_neutral_color',
                         'message': '未核验车型能源属性，柱形使用统一中性配色。'})
    if page in (9, 10, 11):
        group = {9: '3', 10: '4', 11: '5'}[page]
        code, cumulative_code = f'P2-{group}.3A', f'P2-{group}.3B'
        current = _required(_models(cards, code, month), 15, code)
        cumulative = {r['model']: r for r in _models(cards, cumulative_code, month, True)}
        missing = [r['model'] for r in current if r['model'] not in cumulative]
        if missing:
            raise ValueError(f'{cumulative_code}缺少当月TOP15车型的累计值：' + '、'.join(missing))
        overall = {r['model']: r for r in _models(cards, 'P2-14.1', month)}
        levels = {}
        for r in current:
            level = r.get('level') or cumulative[r['model']].get('level')
            other = overall.get(r['model'], {})
            if not level and other.get('sales') == r['sales']:
                level = other.get('level')
            if not level:
                raise ValueError(f'{code}/{r["model"]}缺少车型级别。')
            levels[r['model']] = level
        total = _denominator(sum(_level_rows(_text(cards, f'P2-{group}.1'), False).get(month, {}).values()), code)
        ytd_total = _denominator(sum(_level_rows(_text(cards, f'P2-{group}.2'), True).get(month[:4], {}).values()), cumulative_code)
        rows = [['同比', '环比', '级别', '份额', '排名变化', '累计']]
        for row in current:
            rows.append([pct(row['yoy_pct']) or '-', pct(row['mom_pct']) or '-',
                         levels[row['model']], pct(row['sales'] / total * 100), '待核',
                         number(cumulative[row['model']]['sales'])])
        target_chart = max(charts, key=lambda s: s.left)
        target_table_shape = next(s for s in tables if len(s.table.rows) == 16)
        target_table = target_table_shape.table
        share_table = next(s.table for s in tables if len(s.table.rows) == 1 and len(s.table.columns) == 2)
        set_chart(target_chart.chart, [r['model'] for r in current], [(period_label, [r['sales'] / 10000 for r in current])])
        target_chart.chart.value_axis.maximum_scale = max(r['sales']/10000 for r in current)*1.2
        _neutral_model_colors(slide, target_chart.chart)
        write_rows(target_table, rows)
        _model_value_column(slide,target_chart,target_table_shape,current)
        write_rows(share_table, [[share_table.cell(0, 0).text, pct(sum(cumulative[r['model']]['sales'] for r in current) / ytd_total * 100)]])
    elif page in (16, 19):
        group, count = ('10', 15) if page == 16 else ('13', 10)
        code, cumulative_code = f'P2-{group}.1', f'P2-{group}.2'
        current = _required(_enterprise_rows(cards, code, month, False, warnings), count, code)
        cumulative = _required(_enterprise_rows(cards, cumulative_code, month, True, warnings), count, cumulative_code)
        facts = _Packet(page, cards, month)
        denominators = facts.totals('single_month', page == 19)
        total = _denominator(denominators.get(month), code)
        prior_total = _denominator(denominators.get(_prior(month)), code)
        rows = [['环比', *[pct(r['mom_pct']) or '-' for r in current]],
                ['当月份额' if page == 16 else '份额', *[pct(r['sales'] / total * 100) for r in current]],
                ['同期份额Δ', *[f"{r['sales'] / total * 100 - r['comparison_sales'] / prior_total * 100:+.1f}" if r['comparison_sales'] is not None else '-' for r in current]],
                ['环比排名Δ', *['待核' for _ in current]]]
        if len(charts) != 2:
            raise ValueError(f'第{page}页企业图表数量不符合兼容模板。')
        label_rows = _enterprise_labels(slide, count)
        for shape, data in zip(charts, [current, cumulative]):
            set_chart(shape.chart, [r['entity'] for r in data],
                      [(None, [r['comparison_sales'] / 10000 if r['comparison_sales'] is not None else None for r in data]),
                       (None, [r['sales'] / 10000 for r in data])])
        for labels, data in zip(label_rows, [current, cumulative]):
            for shape, row in zip(labels, data):
                write_text(shape.text_frame, row['entity'])
        write_rows(tables[0].table, rows)
        set_callouts(slide, [r['yoy_pct'] for r in current + cumulative])
        warnings.extend(facts.warnings)
    elif page == 20:
        code, cumulative_code = 'P2-14.1', 'P2-14.2'
        current = _required(_models(cards, code, month), 20, code)
        cumulative = _required(_models(cards, cumulative_code, month, True), 20, cumulative_code)
        facts = _Packet(page, cards, month)
        total = _denominator(facts.totals('single_month').get(month), code)
        ytd_total = _denominator(facts.totals('year_to_date').get(month), cumulative_code)
        for row in current + cumulative:
            if not row.get('level'):
                raise ValueError(f'{row["model"]}缺少车型级别。')
        monthly_rows = [['同比', '环比', '份额', '级别', '排名变化']]
        monthly_rows += [[pct(r['yoy_pct']) or '-', pct(r['mom_pct']) or '-', pct(r['sales'] / total * 100), r['level'], '待核'] for r in current]
        cumulative_rows = [['同比', '份额', '级别', '排名变化']]
        cumulative_rows += [[pct(r['yoy_pct']) or '-', pct(r['sales'] / ytd_total * 100), r['level'], '待核'] for r in cumulative]
        set_chart(charts[0].chart, [r['model'] for r in current], [(period_label, [r['sales'] / 10000 for r in current])])
        set_chart(charts[1].chart, [r['model'] for r in cumulative], [(ytd_label, [r['sales'] / 10000 for r in cumulative])])
        charts[0].chart.value_axis.maximum_scale = max(r['sales']/10000 for r in current)*1.2
        charts[1].chart.value_axis.maximum_scale = max(r['sales']/10000 for r in cumulative)*1.2
        _neutral_model_colors(slide, charts[0].chart)
        _neutral_model_colors(slide, charts[1].chart)
        write_rows(next(s.table for s in tables if len(s.table.columns) == 5), monthly_rows)
        write_rows(next(s.table for s in tables if len(s.table.columns) == 4), cumulative_rows)
        _model_value_column(slide,charts[0],next(s for s in tables if len(s.table.columns)==5),current)
        _model_value_column(slide,charts[1],next(s for s in tables if len(s.table.columns)==4),cumulative)
        share_tables = sorted([s for s in tables if len(s.table.rows) == 1], key=lambda s: s.left)
        for shape, data, denominator in zip(share_tables, [current, cumulative], [total, ytd_total]):
            write_rows(shape.table, [[shape.table.cell(0, 0).text, pct(sum(r['sales'] for r in data) / denominator * 100)]])
    else:
        raise ValueError(f'第{page}页没有榜单更新映射。')
    return {'page': page, 'period': month, 'status': 'updated', 'warnings': warnings}
