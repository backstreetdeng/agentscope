"""Page-scoped narrative generation and deterministic PPT text updates."""

from __future__ import annotations

import ast
import asyncio
from copy import deepcopy
import json
import hashlib
from io import BytesIO
import re
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable
import unicodedata
import zipfile

from agentscope.message import UserMsg
from openpyxl import load_workbook
from pptx import Presentation

from agents.market_strategy_agent import build_monthly_narrative_agent
from utils.config import expand_env
from utils.text import strip_thinking
from utils.monthly_segment_facts import SEGMENT_PAGES, segment_page_facts, segment_fact_draft
from utils.monthly_page_facts import page_business_facts
from utils.monthly_page_narratives import BUSINESS_PAGES, business_fact_draft, validate_business_slide, overview_business_packet
from utils.monthly_ppt_layout import prepare_inventory_roles, validate_business_chart_roles


ANOMALY_TYPES = {
    "reversal",
    "streak_change",
    "structural_divergence",
    "share_or_rank_shift",
    "contribution_concentration",
    "ordinary_change",
}
TARGET_PAGES = (4, 5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19)
BATCHES = ((4, 5, 7, 8), (9, 10, 11, 12), (13, 14, 15, 16), (17, 18, 19))
PAGE_LABELS = {
    4: "1.2 零售量", 5: "1.3 库存（行业）", 7: "2.1 车型类别",
    8: "2.2 车型级别", 9: "2.3 轿车级别", 10: "2.3 SUV级别",
    11: "2.3 MPV级别", 12: "2.4 车型系别", 13: "2.5 车型价格段",
    14: "2.6 区域市场", 15: "2.6 区域市场", 16: "2.7 企业竞争",
    17: "2.8 新能源", 18: "2.8 新能源", 19: "2.8 新能源",
}
CAUSAL_MARKERS = ("由于", "受", "推动", "带动", "驱动", "影响下", "受益于", "叠加", "观望", "促销", "政策")
MAX_HEADLINE_DISPLAY_WIDTH = 52
TARGET_BODY_MIN_DISPLAY_WIDTH = 125
MAX_BODY_DISPLAY_WIDTH = 150
BODY_LINE_TARGET_DISPLAY_WIDTH = 50
NARRATIVE_RULE_VERSION = "entity-bound-business-v6"


def _display_width(text: str) -> float:
    """Approximate rendered width using one unit for CJK and half for Latin."""

    width = 0.0
    for char in text:
        width += 1.0 if unicodedata.east_asian_width(char) in {"W", "F", "A"} else 0.55
    return width


def _clip_business_text(text: str, max_width: float) -> str:
    """Keep complete business clauses within the PPT text-box budget."""

    clean = re.sub(r"[\r\n\t]+", "", str(text or ""))
    clean = re.sub(r"[ \u3000]+", " ", clean).strip("，；。 ")
    if _display_width(clean) <= max_width:
        return clean
    candidates = [match.end() for match in re.finditer(r"[。；]", clean)]
    fitting = [end for end in candidates if _display_width(clean[:end]) <= max_width]
    if fitting:
        return clean[: fitting[-1]].rstrip("，；。 ") + "。"
    output: list[str] = []
    for char in clean:
        if _display_width("".join(output) + char) > max_width:
            break
        output.append(char)
    return "".join(output).rstrip("，、；：: ")


def _format_business_text(text: str, report_month: str) -> str:
    """Apply monthly-report wording and unit conventions deterministically."""

    value = str(text or "")
    value = re.sub(r"同比正增(?:长)?", "同比增长", value)
    value = value.replace("双位数", "两位数")
    match = re.fullmatch(r"(\d{4})-(\d{2})", str(report_month or ""))
    if match:
        year, month = match.groups()
        month_label = f"{int(month)}月"
        value = value.replace(f"{year}-{month}", month_label)
        value = value.replace(f"{year}年{int(month)}月", month_label)

    def format_vehicle_count(match_obj: re.Match[str]) -> str:
        number = int(match_obj.group(1).replace(",", ""))
        if number >= 10000:
            return f"{number / 10000:.1f}万辆"
        return f"{number}辆"

    value = re.sub(r"(?<![\d.])(\d[\d,]*)辆", format_vehicle_count, value)
    return value


def _separate_model_counts(text: str, fact_ids: list[str]) -> str:
    """Use exact entity IDs before formatting counts, e.g. GL8 + 5874."""
    for fact_id in fact_ids:
        match = re.fullmatch(r"P2-\d+\.3[AB]:([^:]+):\d{4}-\d{2}", str(fact_id))
        if match:
            text = re.sub(re.escape(match[1]) + r"(?=\d[\d,]*(?:\.\d+)?万?辆)", match[1] + " ", text)
    return text


def _wrap_business_text(text: str, max_lines: int = 3) -> str:
    """Wrap a dense paragraph into stable business clauses for this template."""

    remaining = _clip_business_text(text, MAX_BODY_DISPLAY_WIDTH)
    if _display_width(remaining) <= BODY_LINE_TARGET_DISPLAY_WIDTH * 1.15:
        return remaining
    # Balance the *whole* paragraph. Greedily ending early at a comma left a
    # long final line which PowerPoint wrapped into a fourth physical line.
    count = min(max_lines, max(2, int((_display_width(remaining) - 1) // 50) + 1))
    widths = [0.0]
    for char in remaining:
        widths.append(widths[-1] + _display_width(char))
    target = widths[-1] / count
    limit = BODY_LINE_TARGET_DISPLAY_WIDTH * 1.12

    @lru_cache(None)
    def split(start: int, lines_left: int):
        if lines_left == 1:
            width = widths[-1] - widths[start]
            return ((width - target) ** 2, (len(remaining),)) if width <= limit else None
        best = None
        for end in range(start + 1, len(remaining)):
            width = widths[end] - widths[start]
            if width > limit:
                break
            if width < target * 0.65:
                continue
            rest = split(end, lines_left - 1)
            if rest is None:
                continue
            left, right = remaining[end - 1], remaining[end]
            boundary_cost = 0 if left in "；。，" else 200
            if right in "；。，、%":
                boundary_cost += 1000
            if left.isascii() and right.isascii() and left.isalnum() and right.isalnum():
                boundary_cost += 200  # do not split model names, numbers or units
            score = (width - target) ** 2 + boundary_cost + rest[0]
            candidate = (score, (end, *rest[1]))
            if best is None or candidate[0] < best[0]:
                best = candidate
        return best

    chosen = split(0, count)
    if chosen is None:
        raise ValueError("正文无法排入三行，需凝练重写，不能丢掉末尾内容。")
    lines, start = [], 0
    for end in chosen[1]:
        lines.append(remaining[start:end].strip())
        start = end
    return "\n".join(lines)


def _ppt_narrative_text(item: dict[str, Any], report_month: str) -> tuple[str, str]:
    """Return one headline and one evidence-dense paragraph of about three lines."""

    headline = _format_business_text(str(item.get("headline", "")).strip(), report_month)
    prefix = PAGE_LABELS.get(int(item.get('slide', 0)), '')
    if prefix:
        headline = re.sub(rf'^{re.escape(prefix)}\s*[|：:]\s*(?:\d+\.\d+\s+)?', f'{prefix}：', headline)
    if item.get('coverage_requirement') and _display_width(headline) > MAX_HEADLINE_DISPLAY_WIDTH:
        raise ValueError(f"第{item.get('slide')}页标题超出单行预算，需提炼判断，不能静默截去事实限定。")
    headline = _clip_business_text(headline, MAX_HEADLINE_DISPLAY_WIDTH)
    if int(item.get('slide', 0)) == 5:
        # Inventory is title-only in the approved business template. Local
        # claims support factual validation/overview, not an extra PPT body.
        return headline, ''
    def format_claim(claim):
        text = _separate_model_counts(str(claim.get("claim", "")).strip(), claim.get("fact_ids", []))
        return _format_business_text(text, report_month).rstrip("；。 ")
    claims = [
        format_claim(claim)
        for claim in item.get("body_claims", [])
        if str(claim.get("claim", "")).strip()
    ]
    body_width = _display_width("；".join(claims))
    if item.get("coverage_requirement") and body_width > MAX_BODY_DISPLAY_WIDTH:
        # Return to generation instead of silently clipping away model evidence.
        raise ValueError(f"第{item.get('slide')}页正文超出三行预算：当前{body_width:.1f}，允许{MAX_BODY_DISPLAY_WIDTH}；"
                         "删除重复累计份额/同比或重复判断，保留核心关系及两款车型销量与增速，不能截掉车型段。")
    # A three-line layout is a maximum *and* an information-density target.  Do
    # not discard the structural or model/company evidence merely because the
    # model emitted it as a second claim.  Join all available layers, then only
    # compress the combined paragraph when it actually exceeds the text box.
    body = _wrap_business_text("；".join(claims))
    return headline, body


def normalize_card_code(value: str) -> str:
    value = value.strip().replace("：", ":")
    match = re.search(r"P(?P<section>[12])[.\-](?P<group>\d+)[.\-](?P<item>\d+)(?P<suffix>[AB]?)", value, re.I)
    if not match:
        return ""
    section = match.group("section")
    group = match.group("group")
    item = match.group("item")
    suffix = match.group("suffix").upper()
    return f"P1.{group}-{item}{suffix}" if section == "1" else f"P2-{group}.{item}{suffix}"


def load_page_card_map(mapping_path: Path) -> dict[int, list[str]]:
    workbook = load_workbook(mapping_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        page_map: dict[int, set[str]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            raw_page, raw_cards = row[0], row[8]
            if not isinstance(raw_page, (int, float)) or not raw_cards or raw_cards == "--":
                continue
            page = int(raw_page)
            for raw_code in re.findall(r"P\d+[.\-]\d+[.\-]\d+[AB]?", str(raw_cards), re.I):
                code = normalize_card_code(raw_code)
                if code:
                    page_map.setdefault(page, set()).add(code)
        return {page: sorted(codes) for page, codes in page_map.items()}
    finally:
        workbook.close()


def _card_lookup(package: dict[str, Any]) -> dict[str, dict[str, Any]]:
    cards = package.get("cards", [])
    if isinstance(cards, dict):
        cards = list(cards.values())
    result: dict[str, dict[str, Any]] = {}
    for card in cards if isinstance(cards, list) else []:
        if not isinstance(card, dict):
            continue
        code = normalize_card_code(str(card.get("code", "")))
        if code:
            result[code] = card
    return result


def delayed_page_exceptions(package: dict[str, Any], page_map: dict[int, list[str]]) -> list[dict[str, Any]]:
    """Derive retained pages from source text, never trust caller skip lists."""
    from tools.monthly_report_browser import MonthlyReportBrowserConnector
    cards = _card_lookup(package)
    diagnostics = MonthlyReportBrowserConnector._card_diagnostics(
        {code: str(card.get('text', '')) for code, card in cards.items()},
        str(package.get('report_month', '')),
        context_verified=package.get('month_confirmation', {}).get('context_period_verified', True))
    delayed = {code: item['latest_period'] for code, item in diagnostics.items()
               if item['status'] == 'delayed_allowed'}
    return [{'slide': page, 'reason': 'source_update_delayed',
             'cards': {code: delayed[code] for code in codes if code in delayed},
             'action': 'preserve_existing_content'}
            for page, codes in sorted(page_map.items()) if set(codes).intersection(delayed)]


def _validate_retained_sources(narrative: dict[str, Any], exceptions: list[dict[str, Any]]) -> None:
    delayed_codes = {code for item in exceptions for code in item['cards']}
    for slide in narrative.get('slides', []):
        fact_ids = list(slide.get('anomaly', {}).get('fact_ids', []))
        fact_ids.extend(value for claim in slide.get('body_claims', []) for value in claim.get('fact_ids', []))
        if any(str(value).split(':', 1)[0] in delayed_codes for value in fact_ids):
            raise ValueError('当月主旨句或概述不得引用滞后卡片事实。')


def _previous_month(report_month: str) -> str:
    year, month = (int(value) for value in report_month.split("-"))
    month -= 1
    if month == 0:
        year, month = year - 1, 12
    return f"{year:04d}-{month:02d}"


def _compact_card_facts(code: str, text: str, report_month: str) -> dict[str, Any]:
    """Build the smallest page-fact packet needed by the language model.

    The raw DOM/card text stays in the local snapshot. This function keeps only
    a title, table labels and rows relevant to the report month, comparison
    periods or the visible top-ranked entities.
    """

    clean_lines = [line.strip() for line in text.replace("\r", "").splitlines() if line.strip()]
    title = clean_lines[0] if clean_lines else code
    segments = text.split("---视图分段---")
    current = report_month
    previous = _previous_month(report_month)
    prior_year = f"{int(report_month[:4]) - 1}-{report_month[5:]}"
    month_token = report_month[5:]
    selected: list[str] = []
    added_periods: set[str] = set()
    header_added = False
    rank_added = False
    compact_added = False

    def add(lines: list[str]) -> None:
        for line in lines:
            value = line.strip()
            if value and (not selected or value != selected[-1]):
                selected.append(value)

    for segment in segments:
        lines = [line.strip() for line in segment.replace("\r", "").splitlines() if line.strip()]
        if not lines:
            continue
        # Retain table labels, then only the current/prior comparison row windows.
        header_end = min(len(lines), 32)
        matched = False
        for index, line in enumerate(lines):
            if line in {current, previous, prior_year} or (
                line == month_token and any(value == str(report_month[:4]) for value in lines[:32])
            ):
                period_key = line if line != month_token else f"month-{line}"
                matched = True
                if period_key in added_periods:
                    continue
                if not header_added:
                    add(lines[:header_end])
                    header_added = True
                add(lines[index : min(len(lines), index + 34)])
                added_periods.add(period_key)
        # Rank/detail cards show the business TOP rows in the first viewport.
        if not matched and not rank_added and not compact_added and len(lines) > 10 and any(label in title for label in ("车型", "企业", "省份", "城市")):
            add(lines[: min(len(lines), 90)])
            rank_added = True
        # Cumulative and two-year comparison tables are already compact.
        if not matched and not compact_added and not rank_added and 5 < len(lines) <= 100:
            add(lines)
            compact_added = True
        if len(selected) >= 180:
            break

    selected = selected[:180]
    percentages = [float(value) for value in re.findall(r"(?<!\d)(-?\d+(?:\.\d+)?)%", "\n".join(selected))]
    candidate_signals: list[dict[str, Any]] = []
    if percentages:
        candidate_signals.append({"type": "largest_positive_pct", "value": max(percentages)})
        candidate_signals.append({"type": "largest_negative_pct", "value": min(percentages)})
        if min(percentages) < 0 < max(percentages):
            candidate_signals.append({"type": "direction_divergence", "value": True})
    return {
        "fact_id": code,
        "title": title,
        "period": report_month,
        "facts": selected,
        "candidate_signals": candidate_signals,
        "raw_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
        "raw_char_count": len(text),
        "fact_char_count": len("\n".join(selected)),
    }


def _balanced_json_objects(text: str) -> list[str]:
    """Return balanced top-level JSON object candidates from mixed model text."""

    candidates: list[str] = []
    start = -1
    depth = 0
    in_string = False
    escaped = False
    for index, char in enumerate(text):
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
            continue
        if char == '"':
            in_string = True
        elif char == "{":
            if depth == 0:
                start = index
            depth += 1
        elif char == "}" and depth:
            depth -= 1
            if depth == 0 and start >= 0:
                candidates.append(text[start : index + 1])
                start = -1
    return candidates


def _json_candidate_variants(candidate: str) -> list[str]:
    """Apply conservative repairs to a complete JSON object candidate."""

    variants = [candidate.strip().lstrip("\ufeff")]
    no_trailing_commas = re.sub(r",\s*([}\]])", r"\1", variants[0])
    if no_trailing_commas != variants[0]:
        variants.append(no_trailing_commas)
    return variants


def _extract_json(text: str | None) -> dict[str, Any]:
    """Extract a JSON object while rejecting unfinished reasoning-only replies."""

    raw_text = "" if text is None else str(text)
    cleaned = strip_thinking(raw_text).strip()
    if not cleaned:
        raise ValueError("模型未返回可展示的JSON正文（可能停留在思考阶段）。")

    candidates: list[str] = []
    for fenced in re.finditer(r"```(?:json)?\s*([\s\S]*?)```", cleaned, re.I):
        candidates.extend(_balanced_json_objects(fenced.group(1)))
    candidates.extend(_balanced_json_objects(cleaned))
    if cleaned.startswith("{") and cleaned.endswith("}"):
        candidates.insert(0, cleaned)

    errors: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        for variant in _json_candidate_variants(candidate):
            if variant in seen:
                continue
            seen.add(variant)
            try:
                value = json.loads(variant, strict=False)
            except json.JSONDecodeError as exc:
                errors.append(str(exc))
                try:
                    # Some compatible model endpoints occasionally return a
                    # Python dict literal. literal_eval is local and safe, and
                    # is accepted only when the result is a dictionary.
                    value = ast.literal_eval(variant)
                except (ValueError, SyntaxError):
                    continue
            if isinstance(value, dict):
                return value
    detail = errors[-1] if errors else "未找到完整的JSON对象"
    raise ValueError(f"模型返回内容无法解析为JSON对象：{detail}")


def _structured_retry_config(config: dict[str, Any]) -> dict[str, Any]:
    """Use deterministic, non-thinking output for structured retry attempts."""

    retry_config = deepcopy(config)
    provider_name = str(retry_config.get("active_provider", ""))
    providers = retry_config.get("providers", {})
    provider = providers.get(provider_name) if isinstance(providers, dict) else None
    if not isinstance(provider, dict):
        return retry_config

    provider["temperature"] = min(float(provider.get("temperature", 0.3) or 0.3), 0.1)
    model_name = str(expand_env(provider.get("model", ""))).strip().lower()
    if provider_name.lower() == "minimax" or model_name == "minimax-m3":
        extra_body = dict(provider.get("extra_body") or {})
        extra_body["reasoning_split"] = True
        extra_body["thinking"] = {"type": "disabled"}
        provider["extra_body"] = extra_body
    return retry_config


def _save_reply_diagnostic(
    run_dir: Path,
    batch_id: str,
    attempt: int,
    text: str | None,
    error: Exception | None,
    stage: str = "model",
) -> None:
    """Persist auditable reply metadata without retaining reasoning content."""

    raw_text = "" if text is None else str(text)
    final_text = strip_thinking(raw_text).strip()
    meta = {
        "batch_id": batch_id,
        "attempt": attempt,
        "raw_sha256": hashlib.sha256(raw_text.encode("utf-8")).hexdigest(),
        "raw_char_count": len(raw_text),
        "final_char_count": len(final_text),
        "contained_thinking_tag": bool(re.search(r"</?think\b", raw_text, re.I)),
        "error": str(error) if error else "",
        "saved_at": datetime.now().isoformat(timespec="seconds"),
    }
    safe_stage = "evidence" if stage == "evidence" else "model"
    (run_dir / f"{safe_stage}-reply-{batch_id}-attempt-{attempt}.json").write_text(
        json.dumps(meta, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    # Preserve the exact model-facing final answer when it exists. Deliberately
    # do not persist <think> content, which is not part of the business output.
    if final_text:
        (run_dir / f"{safe_stage}-final-{batch_id}-attempt-{attempt}.txt").write_text(
            final_text,
            encoding="utf-8",
        )


def _page_payload(
    pages: tuple[int, ...],
    page_map: dict[int, list[str]],
    cards: dict[str, dict[str, Any]],
    examples: dict[str, Any],
    report_month: str,
) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for page in pages:
        codes = page_map.get(page, [])
        page_cards = []
        for code in codes:
            card = cards.get(code)
            if not card:
                continue
            page_cards.append(_compact_card_facts(code, str(card.get("text", "")), report_month))
        page_examples = examples.get(str(page), [])
        lengths = [len(str(item.get("headline", ""))) for item in page_examples if item.get("headline")]
        body_lengths = [
            round(_display_width(str(item.get("body", ""))))
            for item in page_examples
            if item.get("body")
        ]
        result.append({
            "slide": page,
            "required_prefix": PAGE_LABELS.get(page, ""),
            "card_codes": codes,
            "cards": page_cards,
            # Only an aggregate style profile leaves the local machine. The
            # business reports themselves and their historical facts are not
            # inserted into a model prompt.
            "expert_style_profile": {
                "sample_count": len(page_examples),
                "median_headline_chars": sorted(lengths)[len(lengths) // 2] if lengths else 32,
                "median_body_display_width": sorted(body_lengths)[len(body_lengths) // 2] if body_lengths else 145,
                "target_body_display_width": f"{TARGET_BODY_MIN_DISPLAY_WIDTH}-{MAX_BODY_DISPLAY_WIDTH}",
                "rules": [
                    "对象与方向明确",
                    "优先拐点和结构分化",
                    "一个主判断",
                    "正文保留核心现象、结构判断、必要的车型企业或原因证据",
                    "约三行，不以缩成一两行为目标",
                ],
            },
        })
        if page in SEGMENT_PAGES:
            structured = segment_page_facts(page, cards, report_month)
            result[-1].update(structured)
            result[-1]["fact_grounded_draft"] = segment_fact_draft(page, structured, report_month)
            # Ambiguous flattened cells must not reach the model alongside the
            # period-labelled facts and invite it to guess column assignments.
            for compact in result[-1]["cards"]:
                compact["facts"] = []
                compact["candidate_signals"] = []
            result[-1]["expert_style_profile"]["rules"].extend([
                "用代表车型对照解释级别分化，不把单月份额和累计份额混用",
                "每页至少两款车型及对应销量或增速，车型支撑必须保留在最终三行正文",
                "历史范文的因果只学习取证方向，不沿用旧月原因或车型数字",
            ])
        elif page in BUSINESS_PAGES:
            structured = page_business_facts(page, cards, report_month)
            draft = business_fact_draft(page, structured, report_month)
            result[-1].update(structured)
            result[-1]['fact_grounded_draft'] = draft
            for compact in result[-1]['cards']:
                compact['facts'] = []
                compact['candidate_signals'] = []
            result[-1]['expert_style_profile']['rules'].extend([
                '仅使用具名business_facts，严格配对对象、同比环比、单月累计与份额分母',
                '优先解释净增减贡献、结构迁移及规模与增速的反差，不用系统操作提示填正文',
                '比较集合不完整时，只比较已配对对象，不能写全榜唯一或领涨',
            ])
    return result


def _normalize_batch(raw: dict[str, Any], batch_id: str, expected_pages: tuple[int, ...], report_month: str) -> dict[str, Any]:
    slides = raw.get("slides", [])
    evidence = raw.get("evidence", [])
    if not isinstance(slides, list) or not isinstance(evidence, list):
        raise ValueError("主旨句JSON缺少slides或evidence数组。")

    id_map: dict[str, str] = {}
    normalized_evidence: list[dict[str, Any]] = []
    for index, item in enumerate(evidence, 1):
        if not isinstance(item, dict):
            continue
        old_id = str(item.get("id", f"EV-{index:03d}"))
        new_id = f"{batch_id}-{old_id}"
        id_map[old_id] = new_id
        grade = str(item.get("grade", "U")).upper()
        if grade not in {"A", "B", "C", "U"}:
            grade = "U"
        entities = item.get("entities", [])
        if not isinstance(entities, list) or not entities:
            entities = ["未标明对象"]
            grade = "U"
        normalized_evidence.append({
            "id": new_id,
            "source_type": str(item.get("source_type", "unknown")) or "unknown",
            "source_name": str(item.get("source_name", "未标明来源")) or "未标明来源",
            "url_or_record_id": str(item.get("url_or_record_id", "")),
            "published_at": str(item.get("published_at", "未标明日期")) or "未标明日期",
            "applicable_month": str(item.get("applicable_month", report_month)) or report_month,
            "entities": [str(value) for value in entities if str(value).strip()],
            "grade": grade,
            "summary": str(item.get("summary", "未提供摘要")) or "未提供摘要",
            "retrieved_at": str(item.get("retrieved_at", datetime.now().isoformat(timespec="seconds"))),
        })
    grade_by_id = {item["id"]: item["grade"] for item in normalized_evidence}

    normalized_slides: list[dict[str, Any]] = []
    for slide in slides:
        if not isinstance(slide, dict) or slide.get("slide") not in expected_pages:
            continue
        anomaly = slide.get("anomaly", {}) if isinstance(slide.get("anomaly"), dict) else {}
        anomaly_type = str(anomaly.get("type", "ordinary_change"))
        if anomaly_type not in ANOMALY_TYPES:
            anomaly_type = "ordinary_change"
        fact_ids = anomaly.get("fact_ids", [])
        if not isinstance(fact_ids, list) or not fact_ids:
            fact_ids = [f"slide:{slide['slide']}"]
        claims = []
        for claim in slide.get("body_claims", []) if isinstance(slide.get("body_claims"), list) else []:
            if not isinstance(claim, dict) or not str(claim.get("claim", "")).strip():
                continue
            ids = [id_map.get(str(value), str(value)) for value in claim.get("evidence_ids", []) if str(value).strip()]
            grades = {grade_by_id.get(value, "U") for value in ids}
            certainty = str(claim.get("certainty", "unverified"))
            if grades.intersection({"A", "B"}):
                certainty = "supported"
            elif grades.intersection({"A", "B", "C"}):
                certainty = "qualified"
            else:
                certainty, ids = "unverified", []
            claim_fact_ids = claim.get("fact_ids", [])
            if not isinstance(claim_fact_ids, list) or not claim_fact_ids:
                claim_fact_ids = list(fact_ids)
            claims.append({
                "claim": _format_business_text(_separate_model_counts(str(claim["claim"]).strip(), claim_fact_ids), report_month),
                "fact_ids": [str(value) for value in claim_fact_ids],
                "evidence_ids": ids,
                "certainty": certainty,
                "role": str(claim.get("role", "")),
            })
        headline = _format_business_text(_separate_model_counts(str(slide.get("headline", "")).strip(),
                         [*fact_ids, *[value for item in claims for value in item['fact_ids']]]), report_month)
        prefix = PAGE_LABELS.get(int(slide["slide"]), "")
        if prefix and not headline.startswith(prefix):
            content = re.sub(r'^\d+\.\d+\s+', '', headline.lstrip('：: '))
            headline = f"{prefix}：{content}"
        headline_ids = [id_map.get(str(value), str(value)) for value in slide.get("headline_evidence_ids", []) if str(value).strip()]
        headline_grades = {grade_by_id.get(value, "U") for value in headline_ids}
        causal_headline = any(marker in headline for marker in CAUSAL_MARKERS if marker != '受') or re.search(
            r'受(?:到)?[^，；。]{1,20}(?:影响|带动|冲击|提振|拖累|作用)', headline)
        if causal_headline and not headline_grades.intersection({"A", "B"}):
            # 原因没有达到标题证据门槛时，优先回退到模型给出的事实型正文观点。
            factual_claim = next((item["claim"] for item in claims if not item["evidence_ids"]), "")
            if factual_claim:
                headline = f"{prefix}：{factual_claim.lstrip('：: ')}" if prefix else factual_claim
            headline_ids = []
        if {c.get('role') for c in claims} != {'core', 'structure', 'support'}:
            headline = _clip_business_text(headline, MAX_HEADLINE_DISPLAY_WIDTH)
        normalized_slides.append({
            "slide": int(slide["slide"]),
            "topic": str(slide.get("topic", "")),
            "anomaly": {
                "type": anomaly_type,
                "priority": int(anomaly.get("priority", 6) or 6),
                "fact_ids": [str(value) for value in fact_ids],
            },
            "headline": headline,
            "headline_evidence_ids": headline_ids,
            "body_claims": claims,
        })
    found = {slide["slide"] for slide in normalized_slides if slide["headline"]}
    missing = sorted(set(expected_pages) - found)
    if missing:
        raise ValueError(f"主旨句批次缺少页面：{missing}")
    return {"slides": normalized_slides, "evidence": normalized_evidence}


def _batch_prompt(report_month: str, payload: list[dict[str, Any]]) -> str:
    month_match = re.fullmatch(r"\d{4}-(\d{2})", report_month)
    month_label = f"{int(month_match.group(1))}月" if month_match else report_month
    return (
        f"报告月份：{report_month}，成稿中统一写作“{month_label}”，不要写“{report_month}”或目标年份。"
        "请为以下页面生成一个主旨标题和一个约三行的高密度正文段落。"
        "先根据卡片识别异常，再按需调用工具核验原因。专家示例只学习写法，严禁复制旧数字或旧原因。\n"
        "输出结构必须是："
        '{"slides":[{"slide":4,"topic":"零售量","anomaly":{"type":"ordinary_change","priority":6,"fact_ids":["P1.1-1"]},'
        '"headline":"...","headline_evidence_ids":[],"body_claims":[{"claim":"...","fact_ids":["P1.1-1"],"evidence_ids":[],"certainty":"unverified"}]}],'
        '"evidence":[{"id":"EV-001","source_type":"new_car_plugin","source_name":"...","url_or_record_id":"...",'
        '"published_at":"YYYY-MM-DD","applicable_month":"YYYY-MM","entities":["..."],"grade":"A","summary":"...","retrieved_at":"..."}]}。\n'
        "异常type只能为reversal、streak_change、structural_divergence、share_or_rank_shift、contribution_concentration、ordinary_change。"
        "标题必须以页面required_prefix开头，只提炼本页最重要的一个结论，控制在52个显示字符内。"
        "正文只允许一条完整段落，按‘核心现象—关键结构—车型/企业证据或经核验原因’组织，"
        f"目标显示宽度{TARGET_BODY_MIN_DISPLAY_WIDTH}-{MAX_BODY_DISPLAY_WIDTH}，尽可能接近三行。"
        "四行以上只做去重和单位压缩；不得为追求简短而删成一两行、不得丢掉结构判断和必要支撑。"
        "数量表达统一：不足1万辆写精确整数且不加千分位，如8294辆；达到1万辆写一位小数的万辆，如71655辆写7.2万辆。"
        "标题如含原因，必须在headline_evidence_ids列出A/B级证据；没有A/B证据就写事实型标题。"
        "事实型观点不需要外部证据，certainty写unverified且evidence_ids为空；原因型观点必须给证据。"
        "第9、10、11页只使用level_facts/model_facts/derived_facts，不猜测空白列位置。scope=single_month是单月，"
        "year_to_date是1月至报告月累计。sales单位辆，yoy_pct是同比、mom_pct是环比。"
        "这三页body_claims分为core、structure、models三条并填写role，每条约40-50显示字符，"
        "合计125-150。models必须写至少两款model_facts中的车型，逐款写对应销量及同比或环比，"
        "引用精确车型fact_id。选择与级别核心分化相关的代表车型对照，不机械追加累计支撑，"
        "不能只列车型名称，不能把甲车型增速写到乙车型，不足事实时明确不足。"
        "models属于正文最后一层但不得被删。此处车型集合并非完整销量榜，未提供排名时不得声称榜首。"
        "标题不能把两款代表车型泛化成主流车型均下滑/全部车型均增长；直接写具体车型名，或写两款代表车型。"
        "这三页每条claim目标40-55显示宽度（中文1、英数0.55），总计125-150。"
        "必须按单位压缩后的成稿计算长度。全文通常需约170-210个实际字符，不能把125误当总字符数。"
        "用增长/下降表述替代+/-速记，用关系句说明同级车型分化或环比回升与同比承压并存；"
        "避免三句都是数字列表，也不写泛泛而谈的展望。只选最关键的级别事实和两款车型；"
        "不要罗列全部级别、全部车型或所有同比环比。models仅用single_month记录，每款只保留销量与一种增速。"
        "所有比例直接引用字段，不自行计算合计份额或差额。level为空的车型不能自行归类；"
        "derived_facts为本地基于原始销量计算并记录operands/method的支撑关系，可直接引用其sales_delta、"
        "group_share_pct、net_decline_share_pct、share_change_pp。负sales_delta表示同比净减，不能写成销量；"
        "有最大净减和贡献比例才可写主要拖累，不从百分比降幅猜贡献。车型名与销量之间留空格，避免GL8与5874拼接。"
        "累计段只能用year_to_date记录，禁止引用单月份额。没有比较份额字段不能写升至/降至/迁移最快。"
        "级别增速逐一配对写清，如A级同比下降X%，B级同比下降Y%，不要用A/B分别X/Y的省略写法。"
        "份额变化须说明较上月或较上年同期；不能把一升一降的两款车型同时描述为与级别方向一致。"
        "fact_grounded_draft是本地从当月事实组织的草案，不是旧报告。以其事实配对和三层分析为骨架，"
        "改善核心判断、删掉重复措辞，保留必要车型支撑；不得把正确的级别、数字或同比方向改错。"
        "范文核心能力是说明谁在支撑、谁在拖累，以及增长是否足以抵补收缩，不是数字列表后加一句表现分化。"
        "保留草案中有测算依据的净增减/贡献关系，不用空泛的市场修复/结构优化替代。对当月有明确事实的车型给出数据对照。"
        "fact_grounded_draft正文已经约三行，不能在它之外再加多组累计或同比环比。结构已有份额变化时，"
        "不重复添加累计销量/同比；每条精炼为一个业务关系，不能越修改越长。"
        "不得仅凭某车型增长声称其带动整个级别，主要拖累/贡献需要实际增减量测算。"
        "不使用未经完整历史检验的创新高/延续/连续/转正，不用销量替代交付、出口或需求。\n"
        "其他分析页只使用business_facts具名记录。value单位见unit，metric明确销量或库存；"
        "scope=stock为库存，single_month为单月，year_to_date为1月至报告月累计。"
        "第5页库存是仅标题页，后台本地处理，不生成正文。第7/8/12/13/15右侧图使用上年全年"
        "结构参照本年累计，不是同比；如正文引用同期累计事实，必须明确去年同期，不能借全年柱解释同比。"
        "所有同比环比和份额必须配对entity与精确fact_id，份额share_scope分母不能互换。"
        "每层正文至少保留一条实际具名数量、同比、环比、份额或贡献，不能只写口径提醒。"
        "指标写成对象+明确指标+数值，例如日系同比下降34.8%，不能省略成日系34.8%。"
        "分号后可以延续上一具名对象，但不能借用别的对象数字；份额必须写清整体市场或新能源市场。"
        "标题不要堆数字，超过单行预算必须重新提炼；缺少完整比较范围不能写领先或领涨。"
        "使用正式业务判断，不用仍是大头等口语；不要重复章节编号。"
        "body_claims必须保留core、structure、support三条role，合计125—150显示宽度。"
        "以本地fact_grounded_draft的业务关系为骨架改善表达，不能把有操作数支撑的贡献段删成套话。"
        "不得新增数字或改写比较方向；净增不能写净减，库存上升不能写同步去化。"
        "带同比/环比/份额/存销比的数值须在同一子句中紧跟具体对象。价格段直接按entity写。"
        "share_change_pp按未取整的销量占比计算，不能自行拿两个一位小数份额相减。"
        "历史时间序列与全榜单没有进入最小模型摘要；仅可保留business_guards中本地验过的极值/连续表述。"
        "business_facts中CALC记录为已保存operands/method的测算值，可直接引用；不得自行创造贡献分母。\n页面输入：\n"
        + json.dumps(payload, ensure_ascii=False)
    )


def _model_payload(payload: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Remove local trace metadata before constructing a model request."""

    result: list[dict[str, Any]] = []
    for page in payload:
        item = {key: value for key, value in page.items() if key != "cards"}
        item["cards"] = [
            {
                "fact_id": card.get("fact_id"),
                "title": card.get("title"),
                "period": card.get("period"),
                "facts": card.get("facts", []),
                "candidate_signals": card.get("candidate_signals", []),
            }
            for card in page.get("cards", [])
        ]
        draft = page.get('fact_grounded_draft')
        if int(page.get('slide', 0)) in SEGMENT_PAGES and isinstance(draft, dict):
            # Give the editor the selected argument, not a menu of unrelated
            # cumulative figures that encourages extending a three-line draft.
            needed = {value for claim in draft.get('body_claims', []) for value in claim.get('fact_ids', [])}
            item['model_facts'] = [row for row in page.get('model_facts', []) if row['fact_id'] in needed]
            item['level_facts'] = [row for row in page.get('level_facts', []) if row['scope'] == 'single_month']
            item['derived_facts'] = [row for row in page.get('derived_facts', []) if row['fact_id'] in needed]
        elif page.get('business_facts') and isinstance(draft, dict):
            needed = {value for claim in draft.get('body_claims', []) for value in claim.get('fact_ids', [])}
            # Only target-period cited facts leave the host. History/raw cells,
            # provenance operands and unavailable ranking rows stay local.
            keys = {'fact_id', 'entity', 'scope', 'period', 'metric', 'unit', 'value',
                    'yoy_pct', 'mom_pct', 'share_pct', 'share_change_pp', 'share_scope',
                    'delta_yoy', 'delta_mom', 'net_change_share_pct', 'yoy_change_pp', 'inventory_sales_ratio'}
            item['business_facts'] = [{key: value for key, value in row.items() if key in keys}
                                      for row in page.get('business_facts', [])
                                      if row['fact_id'] in needed and row['period'] == page['coverage']['period']]
            item.pop('fact_warnings', None)
            item['coverage'] = {'period': page['coverage']['period'], 'ranking_scope': 'only locally verified comparisons'}
        result.append(item)
    return result


def _evidence_payload(payload: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Create a public-only query plan for plugins and Tavily.

    Evidence agents do not receive sales, shares, growth rates or any other
    internal measurement, so those values cannot accidentally enter tool args.
    """

    ignored = {
        "数量", "份额", "同比", "环比", "排名", "销量", "年", "月",
        "日期组合 (月)", "企业简称", "车型", "车型级别", "使用省份", "区域-使用",
    }
    result: list[dict[str, Any]] = []
    for page in payload:
        entities: list[str] = []
        signal_types: list[str] = []
        for card in page.get("cards", []):
            for signal in card.get("candidate_signals", []):
                value = str(signal.get("type", ""))
                if value and value not in signal_types:
                    signal_types.append(value)
            for value in card.get("facts", []):
                text = str(value).strip()
                if (
                    text in ignored
                    or not re.search(r"[\u4e00-\u9fff]", text)
                    or re.search(r"\d|%", text)
                    or len(text) > 18
                    or text.startswith("P")
                ):
                    continue
                if text not in entities:
                    entities.append(text)
                if len(entities) >= 12:
                    break
        for model in page.get("model_facts", []):
            name = str(model.get("model", ""))
            if name and name not in entities:
                entities.append(name)
        draft = page.get('fact_grounded_draft', {})
        selected = {value for claim in draft.get('body_claims', []) for value in claim.get('fact_ids', [])}
        for fact in page.get('business_facts', []):
            name = str(fact.get('entity', ''))
            if fact.get('fact_id') in selected and name and name not in entities:
                entities.append(name)
        if draft.get('anomaly', {}).get('type'):
            signal_types.append(draft['anomaly']['type'])
        result.append({
            "slide": page.get("slide"),
            "topic": page.get("required_prefix", ""),
            "period": next((card.get("period", "") for card in page.get("cards", [])), ""),
            "public_entities": entities,
            "signal_types": signal_types,
        })
    return result


def _evidence_prompt(report_month: str, payload: list[dict[str, Any]]) -> str:
    return (
        f"报告月份：{report_month}。页面输入已经由本地程序从原始卡片压缩为最小事实包，"
        "不得索要或推测原始卡片。请先判断哪些异常确实需要解释原因，再按需调用工具。"
        "政策、权益、新车、口碑工具优先，只有结构化工具不足时才使用公开网络检索。"
        "工具查询只能包含公开的月份、地区、企业、品牌、车型和事件，不得把销量、份额、"
        "内部同比环比或raw_sha256写入查询参数。没有可靠证据时返回空证据，不要凑原因。"
        "最后只输出JSON："
        '{"evidence":[{"id":"EV-001","source_type":"policy_plugin","source_name":"...",'
        '"url_or_record_id":"...","published_at":"YYYY-MM-DD","applicable_month":"YYYY-MM",'
        '"entities":["..."],"grade":"A","summary":"...","retrieved_at":"..."}],'
        '"evidence_gaps":["..."]}。\n页面最小事实包：\n'
        + json.dumps(payload, ensure_ascii=False)
    )


def _validate_batch_facts(batch: dict[str, Any], page_payload: list[dict[str, Any]]) -> None:
    """Reject numbers or fact references that cannot be traced to the local packet."""

    payload_by_slide = {int(item["slide"]): item for item in page_payload}
    for slide in batch.get("slides", []):
        payload = payload_by_slide.get(int(slide["slide"]), {})
        if payload.get('business_facts'):
            validate_business_slide(slide, payload)
            _ppt_narrative_text(slide, payload['coverage']['period'])
            slide.pop('ppt_comparison_facts', None)
            if int(slide['slide']) == 5:
                slide['presentation_role'] = 'title_only'
            elif int(slide['slide']) in (7, 8, 12, 13, 15):
                slide['chart_reference_scope'] = 'prior_year_full_year_structure'
            continue  # Do not fall back to the old bag-of-numbers whitelist.
        if payload.get("coverage_requirement"):
            _validate_segment_support(slide, payload)
        codes = set(payload.get("card_codes", []))
        safe_payload = _model_payload([payload])[0]
        fact_text = json.dumps({key: safe_payload.get(key, []) for key in
                               ("cards", "level_facts", "model_facts", "derived_facts")}, ensure_ascii=False)
        raw_numbers = [float(value.replace(",", "")) for value in re.findall(r"-?\d[\d,]*(?:\.\d+)?", fact_text)]
        allowed: set[float] = set()
        percentages: list[float] = []
        for value in raw_numbers:
            allowed.add(round(value, 2))
            if abs(value) >= 1000:
                allowed.add(round(value / 10000, 0))
                allowed.add(round(value / 10000, 1))
                allowed.add(round(value / 10000, 2))
        for value in re.findall(r"(-?\d+(?:\.\d+)?)%", fact_text):
            number = float(value)
            percentages.append(number)
            allowed.add(round(number, 2))
            allowed.add(round(abs(number), 2))
        # Typed percentage fields are JSON numbers without a '%' suffix.
        # Chinese prose expresses a negative sign as '下降', not always '-'.
        for row in [*payload.get("level_facts", []), *payload.get("model_facts", []),
                    *payload.get("derived_facts", [])]:
            for key in ("yoy_pct", "mom_pct", "share_pct"):
                if row.get(key) is not None:
                    allowed.add(round(abs(float(row[key])), 2))
        # Derived measures can be negative deltas expressed as Chinese '减少'
        # rather than a leading minus sign. Retain their traceable magnitudes.
        derived_text = json.dumps(payload.get("derived_facts", []), ensure_ascii=False)
        for token in re.findall(r"-?\d[\d,]*(?:\.\d+)?", derived_text):
            number = abs(float(token.replace(",", "")))
            allowed.add(round(number, 2))
            if number >= 1000:
                allowed.add(round(number / 10000, 1))
        for left in percentages:
            for right in percentages:
                allowed.add(round(abs(left - right), 1))
        # Section numbers and the report year/month are presentation metadata.
        allowed.update({float(slide["slide"]), float(payload.get("slide", 0)), float(str(payload.get("required_prefix", "")).split(" ")[0].split(".")[0] or 0)})
        allowed.update({
            float(value)
            for value in re.findall(r"\d+(?:\.\d+)?", str(payload.get("required_prefix", "")))
        })
        period = str(payload.get("cards", [{}])[0].get("period", "")) if payload.get("cards") else ""
        allowed.update({float(value) for value in re.findall(r"\d+", period)})

        texts = [str(slide.get("headline", ""))]
        texts.extend(str(item.get("claim", "")) for item in slide.get("body_claims", []))
        for text in texts:
            for token in re.findall(r"(?<![A-Za-z])(?<!P)(-?\d[\d,]*(?:\.\d+)?)", text):
                number = round(float(token.replace(",", "")), 2)
                if number not in allowed:
                    raise ValueError(f"第{slide['slide']}页出现无法回溯的数字：{token}")

        referenced = list(slide.get("anomaly", {}).get("fact_ids", []))
        referenced.extend(value for claim in slide.get("body_claims", []) for value in claim.get("fact_ids", []))
        strong_contribution_markers = ("主要贡献", "最大贡献", "主要拖累", "最大拖累")
        if any(marker in text for text in texts for marker in strong_contribution_markers):
            measured = any(row.get("fact_id") in referenced and row.get("operands") and row.get("method")
                           for row in payload.get("derived_facts", []))
            if slide.get("anomaly", {}).get("type") != "contribution_concentration" or not measured:
                raise ValueError(f"第{slide['slide']}页未完成贡献测算，不得使用主要/最大贡献或拖累表述。")

        if payload.get("coverage_requirement"):
            exact_ids = {row["fact_id"] for key in ("level_facts", "model_facts", "derived_facts")
                         for row in payload.get(key, []) if row.get("fact_id")}
            bad_ids = [value for value in referenced if value not in exact_ids]
        else:
            bad_ids = [value for value in referenced if not any(str(value).startswith(code) for code in codes)]
        if bad_ids:
            raise ValueError(f"第{slide['slide']}页出现非当前页事实ID：{bad_ids}")


def _validate_segment_support(slide: dict[str, Any], payload: dict[str, Any]) -> None:
    """Validate rendered, entity-paired model facts, not just numbers anywhere."""
    slide["coverage_requirement"] = payload["coverage_requirement"]
    if re.search(r"(?:主流|头部|全部|所有)车型.{0,10}(?:均|全线|全部)", str(slide.get('headline', ''))):
        raise ValueError(f"第{slide['slide']}页标题把部分代表车型泛化为全部，需限定为具体车型或两款代表车型。")
    if any("同向背离" in text for text in [str(slide.get("headline", "")),
                                           *(str(claim.get("claim", "")) for claim in slide.get("body_claims", []))]):
        raise ValueError(f"第{slide['slide']}页“同向背离”表述自相矛盾，需明确同向或背离并重写。")
    level_rows = payload.get("level_facts", [])
    level_pattern = r"(?<![A-Za-z0-9])(A00|A0|C\+|A|B|C|D)级"
    growth_pattern = r"\s*(增长|下降|下滑|增|降|\+|−|-)?\s*(\d+(?:\.\d+)?)%"

    def scope_at(text: str, position: int) -> str:
        markers = list(re.finditer(r"累计|单月|当月|本月", text[:position]))
        return "year_to_date" if markers and markers[-1][0] == "累计" else "single_month"

    def level_record(level: str, scope: str, metric: str):
        row = next((row for row in level_rows if row["level"] == level and row["scope"] == scope), None)
        if row is None:
            raise ValueError(f"第{slide['slide']}页{level}级缺少{scope}{metric}事实。")
        return row

    def derived_record(level: str, scope: str, key: str, referenced: list[str]):
        row = next((row for row in payload.get("derived_facts", [])
                    if row.get("fact_id") in referenced and row.get("level") == level
                    and row.get("scope") == scope and row.get(key) is not None
                    and row.get("operands") and row.get("method")), None)
        if row is None:
            raise ValueError(f"第{slide['slide']}页{level}级缺少已引用的{scope}{key}测算事实。")
        return row

    def derived_group_record(labels: list[str], scope: str, key: str, referenced: list[str]):
        row = next((row for row in payload.get("derived_facts", [])
                    if row.get("fact_id") in referenced and sorted(row.get("levels", [])) == sorted(labels)
                    and row.get("scope") == scope and row.get(key) is not None
                    and row.get("operands") and row.get("method")), None)
        if row is None:
            raise ValueError(f"第{slide['slide']}页{'/'.join(labels)}组合缺少已引用的{scope}{key}测算事实。")
        return row

    def share_key(clause: str, start: int, end: int) -> str:
        # Compare the label belonging to this percentage, not the entire
        # sentence. Both '上月77.6%提高...至78.7%' and the reverse order occur.
        earlier = list(re.finditer(r"\d+(?:\.\d+)?%", clause[:start]))
        prefix = clause[earlier[-1].end() if earlier else 0:start]
        suffix_label = re.match(r"\s*[（(]\s*(上月|上年同期|去年同期|去年同月)\s*[）)]", clause[end:])
        if suffix_label:
            marker = suffix_label[1]
        elif re.search(r"(?:提高|提升|上升|增加|下降|降低|回落|减少)\s*[-+]?\d+(?:\.\d+)?\s*(?:个)?百分点.{0,6}(?:至|达|达到|现为)\s*$", prefix):
            return "share_pct"
        else:
            markers = list(re.finditer(r"上月|上年同期|去年同期|去年同月|本月|当月|单月|当前|目前", prefix))
            marker = markers[-1][0] if markers else ""
        if marker == "上月":
            return "previous_month_share_pct"
        if marker in {"上年同期", "去年同期", "去年同月"}:
            return "prior_year_share_pct"
        return "share_pct"

    level_texts = [(str(slide.get("headline", "")), slide.get("anomaly", {}).get("fact_ids", []))]
    level_texts.extend((str(claim.get("claim", "")), claim.get("fact_ids", []))
                       for claim in slide.get("body_claims", []) if claim.get("role") != "models")
    for text, referenced in level_texts:
        matches = list(re.finditer(level_pattern, text))
        for index, match in enumerate(matches):
            clause = text[match.end(): matches[index + 1].start() if index + 1 < len(matches) else len(text)]
            if "同比转正" in clause:
                scope = scope_at(text, match.end() + clause.index("同比转正"))
                row = level_record(match[1], scope, "同比转正")
                if not (row.get("yoy_pct") is not None and row["yoy_pct"] > 0 and
                        row.get("previous_month_yoy_pct") is not None and row["previous_month_yoy_pct"] <= 0):
                    raise ValueError(f"第{slide['slide']}页{match[1]}级缺少同比转正的上月比较事实。")
            for metric, key in (("同比", "yoy_pct"), ("环比", "mom_pct")):
                for growth in re.finditer(metric + growth_pattern, clause):
                    scope = scope_at(text, match.end() + growth.start())
                    row = level_record(match[1], scope, metric)
                    value = float(growth[2]) * (-1 if growth[1] in {"下降", "下滑", "降", "−", "-"} else 1)
                    if row.get(key) is None or abs(value - row[key]) > 0.051:
                        raise ValueError(f"第{slide['slide']}页{match[1]}级{scope}{metric}错配，应为{row.get(key)}%。")
            for count in re.finditer(r"(\d[\d,]*(?:\.\d+)?)(万?)辆", clause):
                scope = scope_at(text, match.end() + count.start())
                delta = re.search(r"(净减|净增|减少|增加|增量|减量)\s*$", clause[:count.start()])
                if delta:
                    row = derived_record(match[1], scope, "sales_delta", referenced)
                    decreasing = delta[1] in {"净减", "减少", "减量"}
                    value = float(count[1].replace(",", "")) * (-1 if decreasing else 1)
                    expected = round(row["sales_delta"] / 10000, 1) if count[2] else row["sales_delta"]
                    if value != expected:
                        raise ValueError(f"第{slide['slide']}页{match[1]}级{scope}净增减错配，应为{row['sales_delta']}辆。")
                    continue
                row = level_record(match[1], scope, "销量")
                expected = round(row["sales"] / 10000, 1) if count[2] else row["sales"]
                if float(count[1].replace(",", "")) != expected:
                    raise ValueError(f"第{slide['slide']}页{match[1]}级{scope}销量错配，应为{row['sales']}辆。")
            for decline_share in re.finditer(r"占(?:轿车|SUV|MPV)?(?:总)?净减(?:量)?\s*(\d+(?:\.\d+)?)%", clause):
                scope = scope_at(text, match.end() + decline_share.start())
                row = derived_record(match[1], scope, "net_decline_share_pct", referenced)
                if abs(float(decline_share[1]) - row["net_decline_share_pct"]) > 0.051:
                    raise ValueError(f"第{slide['slide']}页{match[1]}级净减贡献占比错配，应为{row['net_decline_share_pct']}%。")
            point_changes = list(re.finditer(r"(提高|提升|上升|增加|下降|降低|回落|减少)\s*([-+]?\d+(?:\.\d+)?)\s*(?:个)?百分点", clause))
            for change in point_changes:
                scope = scope_at(text, match.end() + change.start())
                measured = derived_record(match[1], scope, "share_change_pp", referenced)
                intro = re.search(r"(?:较|比)(上月|上年同期|去年同期|去年同月)[，,\s]*$", text[:match.start()])
                comparison_text = clause + (intro[0] if intro else "")
                comparison_scopes = {"previous_month" if marker == "上月" else "prior_year"
                                     for marker in re.findall(r"上月|上年同期|去年同期|去年同月", comparison_text)}
                measured_scope = measured.get("share_comparison_scope", measured.get("comparison_scope"))
                if len(comparison_scopes) != 1 or measured_scope not in comparison_scopes:
                    raise ValueError(f"第{slide['slide']}页{match[1]}级百分点变化比较范围不明或与测算基期不一致。")
                row = level_record(match[1], scope, "份额变化")
                comparison_key = "previous_month_share_pct" if "previous_month" in comparison_scopes else "prior_year_share_pct"
                baseline = row.get(comparison_key)
                value = float(change[2]) * (-1 if change[1] in {"下降", "降低", "回落", "减少"} else 1)
                expected = float(measured["share_change_pp"])
                if baseline is None or row.get("share_pct") is None or abs(expected - round(row["share_pct"] - baseline, 1)) > 1e-8:
                    raise ValueError(f"第{slide['slide']}页{match[1]}级百分点变化缺少一致的当期/基期份额事实。")
                if abs(value - expected) > 1e-8:
                    raise ValueError(f"第{slide['slide']}页{match[1]}级百分点变化错配，应为{expected}个百分点。")
            if point_changes:
                for percentage in re.finditer(r"(-?\d+(?:\.\d+)?)%", clause):
                    earlier = list(re.finditer(r"\d+(?:\.\d+)?%", clause[:percentage.start()]))
                    prefix = clause[earlier[-1].end() if earlier else 0:percentage.start()]
                    if re.search(r"(?:同比|环比)[^，；。%]*$", prefix) or re.search(r"占(?:轿车|SUV|MPV)?(?:总)?净减(?:量)?\s*$", prefix):
                        continue
                    scope = scope_at(text, match.end() + percentage.start())
                    row = level_record(match[1], scope, "比较份额")
                    key = share_key(clause, percentage.start(), percentage.end())
                    if row.get(key) is None or abs(float(percentage[1]) - row[key]) > 1e-8:
                        raise ValueError(f"第{slide['slide']}页{match[1]}级{key}错配，应为{row.get(key)}%。")
            for share in re.finditer(r"(?:份额|占比|占)(?:仍|已)?(?:为|达|占|升至|降至|收缩至)?\s*(\d+(?:\.\d+)?)%", clause):
                scope = scope_at(text, match.end() + share.start())
                row = level_record(match[1], scope, "份额")
                value = float(share[1])
                key = share_key(clause, share.start(1), share.end())
                if row.get(key) is None or abs(value - row[key]) > 0.051:
                    raise ValueError(f"第{slide['slide']}页{match[1]}级{scope}份额错配，应为{row.get(key)}%。")
                if any(marker in share[0] for marker in ("升至", "降至", "收缩至")):
                    comparison_key = ("previous_month_share_pct" if "上月" in clause else
                                      "prior_year_share_pct" if any(marker in clause for marker in ("上年同期", "去年同期", "去年同月")) else None)
                    baseline = row.get(comparison_key) if comparison_key else None
                    rising = "升至" in share[0]
                    if baseline is None or (value <= baseline if rising else value >= baseline):
                        raise ValueError(f"第{slide['slide']}页{match[1]}级份额变化缺少明确基期或方向不匹配。")
            if any(term in clause for term in ("降幅最深", "降幅最大", "领跌")):
                scope = scope_at(text, match.end())
                row = level_record(match[1], scope, "降幅比较")
                comparisons = [other["yoy_pct"] for other in level_rows
                               if other["scope"] == scope and other.get("yoy_pct") is not None]
                if comparisons and row.get("yoy_pct") != min(comparisons):
                    raise ValueError(f"第{slide['slide']}页{match[1]}级不是同比降幅最大的级别。")
        # The explicit aggregate grammar intentionally omits '级': A/B合计占
        # ...同比净减..., followed by A0与C+净增... . Keep each aggregate's
        # metrics attached to its own local derived record, not any allowed
        # number from another group in the same packet.
        group_matches = list(re.finditer(
            r"(?<![A-Za-z0-9])((?:A00|A0|C\+|A|B|C|D)(?:(?:/|与|和|、)(?:A00|A0|C\+|A|B|C|D))+)(?=(?:合计)?(?:占|份额|同比|净增|净减))",
            text,
        ))
        for index, group in enumerate(group_matches):
            labels = re.findall(r"A00|A0|C\+|A|B|C|D", group[1])
            end = group_matches[index + 1].start() if index + 1 < len(group_matches) else len(text)
            clause = text[group.end():end]
            for share in re.finditer(r"(?:份额|占比|占)(?:为|达)?\s*(\d+(?:\.\d+)?)%", clause):
                scope = scope_at(text, group.end() + share.start())
                row = derived_group_record(labels, scope, "group_share_pct", referenced)
                if abs(float(share[1]) - row["group_share_pct"]) > 0.051:
                    raise ValueError(f"第{slide['slide']}页{'/'.join(labels)}组合份额错配，应为{row['group_share_pct']}%。")
            for delta in re.finditer(r"(净减|净增|减少|增加)\s*(\d[\d,]*(?:\.\d+)?)(万?)辆", clause):
                scope = scope_at(text, group.end() + delta.start())
                row = derived_group_record(labels, scope, "sales_delta", referenced)
                value = float(delta[2].replace(",", "")) * (-1 if delta[1] in {"净减", "减少"} else 1)
                expected = round(row["sales_delta"] / 10000, 1) if delta[3] else row["sales_delta"]
                if value != expected:
                    raise ValueError(f"第{slide['slide']}页{'/'.join(labels)}组合净增减错配，应为{row['sales_delta']}辆。")
    _, rendered = _ppt_narrative_text(slide, str(payload["model_facts"][0]["period"]) if payload.get("model_facts") else "")
    rendered_width = _display_width(rendered.replace("\n", ""))
    if rendered_width < 120:
        raise ValueError(f"第{slide['slide']}页压缩单位后只有{rendered_width:.1f}显示宽度，目标125—150；需补充当前事实支持的级别关系及车型分化判断，不能用空话凑字。")
    model_claims = [item for item in slide.get("body_claims", []) if item.get("role") == "models"]
    if not model_claims:
        raise ValueError(f"第{slide['slide']}页缺少role=models车型支撑段。")
    records = {item["fact_id"]: item for item in payload.get("model_facts", [])}
    known_names = {row["model"] for row in records.values()}
    supported: set[str] = set()
    for claim in model_claims:
        text = str(claim.get("claim", ""))
        cited = [records[value] for value in claim.get("fact_ids", []) if value in records]
        cited_names = {row["model"] for row in cited if row.get("scope", "single_month") == "single_month"}
        uncited = [name for name in known_names if name in text and name not in cited_names]
        if uncited:
            raise ValueError(f"第{slide['slide']}页车型未引用对应单月事实ID：{uncited}。")
        # Restrict unknown-entity detection to explicit '...车型...' followed
        # immediately by a volume. Do not treat arbitrary Chinese prose as a
        # model name; known model names above are always checked independently.
        for named in re.finditer(r"(?:^|[，；。])\s*([^\d，；。:：]{0,16}车型[^\d，；。:：]{0,12})(?=\d+(?:\.\d+)?万?辆)", text):
            name = named[1].strip()
            if name not in known_names:
                raise ValueError(f"第{slide['slide']}页出现无对应事实的车型：{name}。")
        if "与级别方向一致" in text:
            directions = {1 if row["yoy_pct"] > 0 else -1 for row in cited if row.get("yoy_pct") is not None}
            if len(directions) > 1:
                raise ValueError(f"第{slide['slide']}页代表车型一升一降，不能同时写与级别方向一致，应说明同级车型分化。")
        for row in cited:
            name = row["model"]
            if name not in text or name not in rendered:
                continue
            if row.get("scope", "single_month") != "single_month":
                raise ValueError(f"第{slide['slide']}页车型段需使用单月记录，累计另列，避免混用{name}。")
            if scope_at(text, text.index(name)) != "single_month":
                raise ValueError(f"第{slide['slide']}页{name}单月记录不能表述为累计。")
            prefix = text[max(0, text.index(name) - 14):text.index(name)]
            assigned = re.search(r"(A00|A0|C\+|A|B|C|D)级(?:代表车型|代表|车型|的)?[：:\s]*$", prefix)
            if not assigned:
                sentence = re.split(r"[；。\n]", text[:text.index(name)])[-1]
                groups = list(re.finditer(r"(A00|A0|C\+|A|B|C|D)级(?:代表车型|代表|车型)?[：:]", sentence))
                assigned = groups[-1] if groups else None
            if assigned and assigned[1] != row.get("level", ""):
                raise ValueError(f"第{slide['slide']}页{name}级别没有对应事实，不能自行归类为{assigned[1]}级。")
            start = text.index(name) + len(name)
            ends = [text.find(other, start) for other in known_names if other != name]
            end = min([pos for pos in ends if pos >= 0] or [len(text)])
            clause = text[start:end]
            count = re.search(r"(\d[\d,]*(?:\.\d+)?)(万?)辆", clause)
            if not count:
                raise ValueError(f"第{slide['slide']}页{name}缺少对应销量，不能只有名称。")
            expected_sales = round(row["sales"] / 10000, 1) if count[2] else row["sales"]
            if float(count[1].replace(",", "")) != expected_sales:
                raise ValueError(f"第{slide['slide']}页{name}销量错配：应为{row['sales']}辆。")
            growth_found = False
            for metric, key in (("同比", "yoy_pct"), ("环比", "mom_pct")):
                for growth in re.finditer(metric + growth_pattern, clause):
                    growth_found = True
                    value = float(growth[2]) * (-1 if growth[1] in {"下降", "下滑", "降", "−", "-"} else 1)
                    if row.get(key) is None or abs(value - row[key]) > 0.051:
                        raise ValueError(f"第{slide['slide']}页{name}{metric}错配，应为{row.get(key)}%。")
            if not growth_found:
                raise ValueError(f"第{slide['slide']}页{name}需保留与级别分化有关的同比或环比，不能只列销量。")
            supported.add(name)
    if len(supported) < payload["coverage_requirement"]["minimum_models"]:
        raise ValueError(f"第{slide['slide']}页最终正文需至少两款可回溯车型，现有{len(supported)}款。")


def _validate_overview_numbers(
    overview: dict[str, Any],
    approved_slide_conclusions: list[dict[str, Any]],
    report_month: str,
    fact_payload: dict[str, Any] | None = None,
) -> None:
    """Ensure the overview only reuses numbers from accepted slide narratives."""

    source_text = "\n".join(
        str(item.get("headline", "")) + "\n" + "\n".join(
            str(claim.get("claim", "")) for claim in item.get("body_claims", [])
        )
        for item in approved_slide_conclusions
    )
    output_text = "\n".join(
        str(item.get("headline", "")) + "\n" + "\n".join(
            str(claim.get("claim", "")) for claim in item.get("body_claims", [])
        )
        for item in overview.get("slides", [])
    )
    source_numbers = set(re.findall(r"-?\d+(?:\.\d+)?", source_text))
    output_numbers = set(re.findall(r"-?\d+(?:\.\d+)?", output_text))
    metadata_numbers = {"3", report_month[:4], str(int(report_month[5:]))}
    invented = output_numbers - source_numbers - metadata_numbers
    if invented:
        raise ValueError(f"概述页出现分页面结论之外的数字：{sorted(invented)}")
    if fact_payload:
        _validate_batch_facts(overview, [fact_payload])


def _deterministic_overview(
    approved_slide_conclusions: list[dict[str, Any]],
    report_month: str,
) -> dict[str, Any]:
    """Build a safe overview from accepted headlines when model retries fail."""

    by_slide = {int(item.get("slide", 0)): item for item in approved_slide_conclusions}
    selected: list[dict[str, Any]] = []
    for slide_number in (4, 17, 7, 16):
        item = by_slide.get(slide_number)
        if item and item not in selected:
            selected.append(item)
    for item in sorted(
        approved_slide_conclusions,
        key=lambda value: (-int(value.get("anomaly", {}).get("priority", 0) or 0), int(value.get("slide", 0))),
    ):
        if item not in selected:
            selected.append(item)
        if len(selected) >= 4:
            break
    if not selected:
        raise ValueError("概述页缺少可复用的已验收分页面结论。")

    def conclusion_text(item: dict[str, Any]) -> str:
        headline = str(item.get("headline", "")).strip()
        return headline.split("：", 1)[-1].strip() or headline

    fact_ids: list[str] = []
    for item in selected:
        for value in item.get("anomaly", {}).get("fact_ids", []):
            value = str(value)
            if value and value not in fact_ids:
                fact_ids.append(value)
    raw = {
        "slides": [{
            "slide": 3,
            "topic": "概述",
            "anomaly": {"type": "ordinary_change", "priority": 6, "fact_ids": fact_ids},
            "headline": f"概述：{conclusion_text(selected[0])}",
            "headline_evidence_ids": [],
            "body_claims": [{
                "claim": "；".join(conclusion_text(item) for item in selected),
                "fact_ids": fact_ids,
                "evidence_ids": [],
                "certainty": "unverified",
            }],
        }],
        "evidence": [],
    }
    return _normalize_batch(raw, "OV", (3,), report_month)


def validate_narrative_package(
    narrative: dict[str, Any],
    package: dict[str, Any],
    *,
    run_id: str,
    mapping_path: Path,
    examples_path: Path,
) -> None:
    """Recheck the exact saved run against current local facts before export.

    This invokes no browser/provider and does not recollect cards. A rule
    upgrade requires restoring the narrative stage, never silently relabels
    old text as newly accepted. Chart comparison proofs are derived locally.
    """
    if narrative.get('run_id') != run_id:
        raise ValueError('主旨句证据包与本轮任务不一致，已停止生成PPT。')
    if narrative.get('narrative_rule_version') != NARRATIVE_RULE_VERSION:
        raise ValueError('逐页核验规则已升级，请新建分析生成新版主旨句；未完成任务可从检查点恢复，已验收的同月卡片无需重新采集。')
    month = str(package.get('report_month', ''))
    if not month or narrative.get('report_month') != month:
        raise ValueError('主旨句证据包与数据包月份不一致，已停止生成PPT。')
    if narrative.get('research'):
        from utils.monthly_research_dag import validate_research_ledger
        research_path = Path(narrative['research'].get('path', ''))
        if research_path.name != 'research-dag.json' or research_path.parent.name != run_id:
            raise ValueError('研究账本不属于本轮任务')
        ledger = json.loads(research_path.read_text(encoding='utf-8'))
        if ledger.get('run_id') != run_id or ledger.get('report_month') != month:
            raise ValueError('研究账本任务或月份不匹配')
        errors = validate_research_ledger(research_path)
        if errors or ledger.get('evidence', []) != narrative.get('research_evidence', []):
            raise ValueError('研究回执或证据不一致：' + '; '.join(errors[:3]))
    slides = narrative.get('slides', [])
    page_map = load_page_card_map(mapping_path)
    exceptions = delayed_page_exceptions(package, page_map)
    retained = {item['slide'] for item in exceptions}
    if narrative.get('retained_pages', []) != exceptions:
        raise ValueError('滞后卡片保留记录与数据包不一致，已停止生成PPT。')
    _validate_retained_sources(narrative, exceptions)
    expected = {3, *TARGET_PAGES} - retained
    actual = [int(slide.get('slide', 0)) for slide in slides]
    if len(actual) != len(expected) or set(actual) != expected:
        raise ValueError('逐页主旨句不完整或包含重复页面，已停止生成PPT。')
    page_map = load_page_card_map(mapping_path)
    examples = json.loads(examples_path.read_text(encoding='utf-8')).get('pages', {})
    cards = _card_lookup(package)
    packets = []
    for pages in BATCHES:
        pages = tuple(page for page in pages if page not in retained)
        if not pages:
            continue
        payload = _page_payload(pages, page_map, cards, examples, month)
        _validate_batch_facts({'slides': [slide for slide in slides if slide['slide'] in pages]}, payload)
        packets.extend(payload)
    overview_packet = overview_business_packet(packets, slides, month)
    if overview_packet is None and not retained:
        raise ValueError('概述页缺少已验收分页面事实，已停止生成PPT。')
    _validate_overview_numbers({'slides': [slide for slide in slides if slide['slide'] == 3]},
                               [slide for slide in slides if slide['slide'] != 3], month, overview_packet)


async def generate_monthly_narratives(
    package: dict[str, Any],
    *,
    run_id: str,
    config: dict[str, Any],
    plugin_tools: list[Any],
    mapping_path: Path,
    examples_path: Path,
    output_dir: Path,
    progress: Callable[[str, int], None] | None = None,
) -> dict[str, Any]:
    report_month = str(package.get("report_month", ""))
    page_map = load_page_card_map(mapping_path)
    exceptions = delayed_page_exceptions(package, page_map)
    retained = {item['slide'] for item in exceptions}
    cards = _card_lookup(package)
    examples_root = json.loads(examples_path.read_text(encoding="utf-8"))
    examples = examples_root.get("pages", {})
    run_dir = output_dir / run_id
    run_dir.mkdir(parents=True, exist_ok=True)

    all_slides: list[dict[str, Any]] = []
    all_evidence: list[dict[str, Any]] = []
    collection_batches: list[dict[str, Any]] = []
    prepared_batches: list[tuple[str, tuple[int, ...], list[dict[str, Any]], Path]] = []
    for batch_index, pages in enumerate(BATCHES, 1):
        pages = tuple(page for page in pages if page not in retained)
        if not pages:
            continue
        batch_id = f"B{batch_index}"
        try:
            page_payload = _page_payload(pages, page_map, cards, examples, report_month)
        except (ValueError, KeyError) as exc:
            raise RuntimeError(f"主旨句最小事实包尚未具备可核验口径，已停止生成：{exc}") from exc
        fact_packet_path = run_dir / f"facts-{batch_id}.json"
        fact_packet_path.write_text(json.dumps(page_payload, ensure_ascii=False, indent=2), encoding="utf-8")
        prepared_batches.append((batch_id, pages, page_payload, fact_packet_path))
    if progress:
        progress("已在本地完成页面最小事实包与异常候选提取；原始卡片未进入模型上下文", 72)

    research_result = None
    if plugin_tools or config.get('providers'):
        from utils.monthly_research_dag import ResearchDAG
        from tools.custom_tools import build_business_toolkit

        async def research_planner(prompt):
            agent = build_monthly_narrative_agent(_structured_retry_config(config), allow_tools=False)
            reply = await agent.reply(UserMsg(name='月报动态研究DAG', content=prompt))
            return _extract_json(reply.get_text_content())

        dag = ResearchDAG(path=run_dir / 'research-dag.json', run_id=run_id,
                          month=report_month, pages=[p for _, _, rows, _ in prepared_batches for p in rows],
                          tools=[*plugin_tools, *build_business_toolkit()], planner=research_planner,
                          progress=progress)
        research_result = await dag.run()
        # The research ledger survives a page falling back to factual prose.
        (run_dir / 'research-evidence.json').write_text(
            json.dumps(research_result, ensure_ascii=False, indent=2), encoding='utf-8')

    for batch_index, (batch_id, pages, page_payload, fact_packet_path) in enumerate(prepared_batches, 1):
        batch_path = run_dir / f"narrative-{batch_id}.json"
        attempts_used = 0
        fact_digest = hashlib.sha256(json.dumps(_model_payload(page_payload), ensure_ascii=False,
                                                sort_keys=True).encode("utf-8")).hexdigest()
        if research_result is not None:
            research_digest = hashlib.sha256(json.dumps(research_result, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
            fact_digest = hashlib.sha256((fact_digest + research_digest).encode()).hexdigest()
        page_digests = {item['slide']: hashlib.sha256(json.dumps(_model_payload([item]), ensure_ascii=False,
                                      sort_keys=True).encode('utf-8')).hexdigest() for item in page_payload}
        if research_result is not None:
            page_digests = {page: hashlib.sha256((value + research_digest).encode()).hexdigest()
                            for page, value in page_digests.items()}
        cached = json.loads(batch_path.read_text(encoding="utf-8")) if batch_path.is_file() else None
        if cached:
            if cached.get("rule_version") != NARRATIVE_RULE_VERSION or cached.get("fact_digest") != fact_digest:
                cached = None
                if progress:
                    progress(f"页面事实分析规则已升级；保留卡片与取证，仅重新生成主旨句批次{batch_id}。", 78)
        if cached:
            try:
                _validate_batch_facts(cached, page_payload)
            except ValueError:
                cached = None
        if cached:
            normalized = cached
        else:
            evidence_path = run_dir / f"evidence-{batch_id}.json"
            if research_result is not None:
                evidence_result = research_result
                evidence_path.write_text(json.dumps(evidence_result, ensure_ascii=False, indent=2), encoding='utf-8')
            elif evidence_path.is_file():
                evidence_result = json.loads(evidence_path.read_text(encoding="utf-8"))
            else:
                if progress:
                    progress(f"正在执行多源取证批次 {batch_index}/{len(BATCHES)}：页面 {','.join(map(str, pages))}", 72 + batch_index * 2)
                evidence_result = {}
                evidence_error: Exception | None = None
                for evidence_attempt in range(1, 4):
                    attempt_config = config if evidence_attempt == 1 else _structured_retry_config(config)
                    evidence_agent = build_monthly_narrative_agent(
                        attempt_config,
                        market_mcp_tools=plugin_tools,
                        allow_tools=True,
                    )
                    evidence_prompt = _evidence_prompt(report_month, _evidence_payload(page_payload))
                    evidence_prompt += (
                        "\n最终回答必须从第一个字符{开始，到最后一个字符}结束；"
                        "不要输出思考标签、Markdown围栏或解释。"
                    )
                    if evidence_error:
                        evidence_prompt += f"\n上一轮取证JSON未通过校验：{evidence_error}。请完整重写JSON。"
                    evidence_text = ""
                    try:
                        evidence_reply = await evidence_agent.reply(UserMsg(
                            name="月报多源取证任务",
                            content=evidence_prompt,
                        ))
                        evidence_text = evidence_reply.get_text_content()
                        parsed_evidence = _extract_json(evidence_text)
                        if not isinstance(parsed_evidence.get("evidence", []), list):
                            raise ValueError(f"取证批次{batch_id}未返回evidence数组。")
                        evidence_result = parsed_evidence
                        _save_reply_diagnostic(
                            run_dir, batch_id, evidence_attempt, evidence_text, None, stage="evidence"
                        )
                        break
                    except Exception as exc:
                        evidence_error = exc
                        _save_reply_diagnostic(
                            run_dir, batch_id, evidence_attempt, evidence_text, exc, stage="evidence"
                        )
                        if evidence_attempt < 3:
                            await asyncio.sleep(1.5)
                if not evidence_result:
                    # External evidence enriches causal interpretation, but a
                    # formatting failure must not discard the verified card
                    # package or completed batches. Continue with internal
                    # facts only; normalization will remove unsupported causes.
                    evidence_result = {
                        "evidence": [],
                        "status": "unavailable",
                        "warning": f"{batch_id}多源取证连续三次未返回有效JSON，仅使用内部事实生成。",
                    }
                    if progress:
                        progress(
                            f"多源取证批次 {batch_index}/{len(BATCHES)} 返回格式异常；已降级为仅使用内部事实，不生成无证据原因。",
                            72 + batch_index * 2,
                        )
                evidence_path.write_text(json.dumps(evidence_result, ensure_ascii=False, indent=2), encoding="utf-8")

            last_error: Exception | None = None
            normalized = {}
            accepted_pages: dict[int, dict[str, Any]] = {}
            accepted_evidence: dict[str, dict[str, Any]] = {}
            for page in pages:
                page_path = run_dir / f"narrative-{batch_id}-page-{page}.json"
                if not page_path.is_file():
                    continue
                saved_page = json.loads(page_path.read_text(encoding="utf-8"))
                if (saved_page.get("rule_version") != NARRATIVE_RULE_VERSION
                        or (saved_page.get('page_fact_digest') and saved_page['page_fact_digest'] != page_digests[page])):
                    continue
                # Old records have a batch-wide digest. A repair to another
                # page must not discard this page: recheck its saved prose
                # against the exact current facts before allowing reuse.
                try:
                    _validate_batch_facts(saved_page, page_payload)
                except ValueError:
                    continue
                if len(saved_page.get("slides", [])) == 1 and int(saved_page["slides"][0]["slide"]) == page:
                    accepted_pages[page] = saved_page["slides"][0]
                    accepted_evidence.update({row["id"]: row for row in saved_page.get("evidence", [])})
            # Do not ask the model to invent a three-line body for a title-only
            # page. Retain local inventory facts for overview/reverse checks.
            if 5 in pages:
                inventory = next(p for p in page_payload if p['slide'] == 5)
                local = _normalize_batch({'slides': [deepcopy(inventory['fact_grounded_draft'])], 'evidence': []},
                                         batch_id, (5,), report_month)
                _validate_batch_facts(local, [inventory])
                local['slides'][0]['generation_method'] = 'fact_grounded_title_only'
                accepted_pages[5] = local['slides'][0]
                local.update(rule_version=NARRATIVE_RULE_VERSION, fact_digest=fact_digest, page_fact_digest=page_digests[5])
                (run_dir / f'narrative-{batch_id}-page-5.json').write_text(
                    json.dumps(local, ensure_ascii=False, indent=2), encoding='utf-8')
            for attempt in range(1, 4):
                remaining_pages = tuple(page for page in pages if page not in accepted_pages)
                if not remaining_pages:
                    normalized = {"slides": [accepted_pages[page] for page in pages],
                                  "evidence": list(accepted_evidence.values())}
                    break
                attempts_used = attempt
                if progress:
                    progress(f"正在完成异常识别、证据分级和主旨句校验：页面 {','.join(map(str, remaining_pages))}", 80 + batch_index * 3)
                attempt_config = _structured_retry_config(config)
                agent = build_monthly_narrative_agent(attempt_config, allow_tools=False)
                prompt = (
                    _batch_prompt(report_month, _model_payload([item for item in page_payload if item["slide"] in remaining_pages]))
                    + "\n以下是上一步取得的候选证据。只能引用其中的证据ID，不得新增来源：\n"
                    + json.dumps(evidence_result, ensure_ascii=False)
                    + "\n本步不得调用任何工具。输出JSON只需包含slides数组。"
                    + "从第一个字符{开始，到最后一个字符}结束；不要输出思考标签、Markdown围栏或解释。"
                )
                if last_error:
                    prompt += f"\n上一轮JSON未通过校验：{last_error}。请完整重写该批次JSON。"
                reply_text = ""
                try:
                    reply = await agent.reply(UserMsg(name="月报主旨句任务", content=prompt))
                    reply_text = reply.get_text_content()
                    raw_narrative = _extract_json(reply_text)
                    raw_narrative["evidence"] = evidence_result.get("evidence", [])
                    # Report all page-specific issues together. A bad first
                    # or omitted page must not discard returned valid pages.
                    raw_slides = raw_narrative.get("slides", [])
                    if not isinstance(raw_slides, list):
                        raise ValueError("主旨句JSON缺少slides数组。")
                    errors = []
                    for page in remaining_pages:
                        try:
                            page_rows = [row for row in raw_slides if isinstance(row, dict) and row.get("slide") == page]
                            candidate = _normalize_batch(
                                {"slides": page_rows, "evidence": raw_narrative["evidence"]},
                                batch_id, (page,), report_month,
                            )
                            if len(candidate["slides"]) != 1:
                                raise ValueError(f"第{page}页返回重复候选，不能选择不确定的版本。")
                            page_candidate = candidate["slides"][0]
                            _validate_batch_facts({"slides": [page_candidate]}, page_payload)
                        except (ValueError, TypeError, KeyError) as error:
                            errors.append(f"第{page}页：{error}")
                        else:
                            page_candidate['generation_method'] = 'model_refined'
                            accepted_pages[page] = page_candidate
                            accepted_evidence.update({row["id"]: row for row in candidate.get("evidence", [])})
                            page_record = {"slides": [page_candidate], "evidence": candidate.get("evidence", []),
                                           "rule_version": NARRATIVE_RULE_VERSION, "fact_digest": fact_digest,
                                           "page_fact_digest": page_digests[page]}
                            (run_dir / f"narrative-{batch_id}-page-{page}.json").write_text(
                                json.dumps(page_record, ensure_ascii=False, indent=2), encoding="utf-8")
                            if progress:
                                progress(f"第{page}页主旨句已通过核验并保存；失败页重试不会覆盖该页。", 80 + batch_index * 3)
                    if errors:
                        raise ValueError("；".join(errors))
                    normalized = {"slides": [accepted_pages[page] for page in pages],
                                  "evidence": list(accepted_evidence.values())}
                    _save_reply_diagnostic(run_dir, batch_id, attempt, reply_text, None)
                    break
                except Exception as exc:
                    last_error = exc
                    normalized = {}
                    _save_reply_diagnostic(run_dir, batch_id, attempt, reply_text, exc)
                    if attempt < 3:
                        await asyncio.sleep(1.5)
            if not normalized:
                # A verified, relation-rich factual draft is a valid business
                # result when prose editing fails. It is not an old-month or
                # fabricated fallback: it must pass the same strict entity,
                # numeric, contribution and three-line gates as model prose.
                for payload in page_payload:
                    page = payload['slide']
                    draft = payload.get('fact_grounded_draft')
                    if page in accepted_pages or page not in (*SEGMENT_PAGES, *BUSINESS_PAGES) or not isinstance(draft, dict):
                        continue
                    try:
                        candidate = _normalize_batch({'slides': [deepcopy(draft)], 'evidence': []},
                                                     batch_id, (page,), report_month)
                        _validate_batch_facts(candidate, page_payload)
                    except (ValueError, TypeError, KeyError):
                        continue
                    fallback_page = candidate['slides'][0]
                    fallback_page['generation_method'] = 'fact_grounded'
                    fallback_page['generation_note'] = '模型润色未通过；使用当月本地事实关系草案，通过相同事实与版式门禁。'
                    accepted_pages[page] = fallback_page
                    page_record = {'slides': [fallback_page], 'evidence': [],
                                   'rule_version': NARRATIVE_RULE_VERSION, 'fact_digest': fact_digest,
                                   'page_fact_digest': page_digests[page],
                                   'model_edit_error': str(last_error)}
                    (run_dir / f'narrative-{batch_id}-page-{page}.json').write_text(
                        json.dumps(page_record, ensure_ascii=False, indent=2), encoding='utf-8')
                    if progress:
                        progress(f'第{page}页模型润色未通过；本地事实分析稿已通过同一门禁并保存。', 80 + batch_index * 3)
                if all(page in accepted_pages for page in pages):
                    normalized = {'slides': [accepted_pages[page] for page in pages],
                                  'evidence': list(accepted_evidence.values())}
            if not normalized:
                raise RuntimeError(f"主旨句批次{batch_id}连续三次未通过生成或验收：{last_error}")
            normalized.update(rule_version=NARRATIVE_RULE_VERSION, fact_digest=fact_digest)
            batch_path.write_text(json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8")
        all_slides.extend(normalized["slides"])
        all_evidence.extend(normalized.get("evidence", []))
        collection_batches.append({
            "batch_id": f"narrative-{batch_id}",
            "status": "completed",
            "evidence_ids": [item["id"] for item in normalized.get("evidence", [])],
            "attempts": attempts_used,
            "fact_packet": str(fact_packet_path),
            "evidence_packet": str(run_dir / f"evidence-{batch_id}.json"),
        })
        if progress:
            progress(f"已完成主旨句批次 {batch_index}/{len(BATCHES)}：页面 {','.join(map(str, pages))}", 72 + batch_index * 4)

    overview_path = run_dir / "narrative-overview.json"
    overview_attempts = 0
    overview_fallback = False
    conclusions_digest = hashlib.sha256(json.dumps(all_slides, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()
    overview_fact_payload = overview_business_packet(
        [page for _, _, pages_payload, _ in prepared_batches for page in pages_payload], all_slides, report_month)
    cached_overview = json.loads(overview_path.read_text(encoding="utf-8")) if overview_path.is_file() else None
    if cached_overview and cached_overview.get("conclusions_digest") == conclusions_digest and cached_overview.get('rule_version') == NARRATIVE_RULE_VERSION:
        overview = cached_overview
    else:
        overview_payload = [{
            "slide": 3,
            "approved_slide_conclusions": [
                {"slide": item["slide"], "headline": item["headline"], "body_claims": item["body_claims"]}
                for item in sorted(all_slides, key=lambda value: value["slide"])
            ],
            "expert_style_profile": {
                "sample_count": len(examples.get("3", [])),
                "rules": ["只汇总已验收分页面结论", "突出总量与结构的核心矛盾", "不新增原因"],
            },
        }]
        approved_conclusions = overview_payload[0]["approved_slide_conclusions"]
        if overview_fact_payload:
            overview_payload[0].update(_model_payload([overview_fact_payload])[0])
        overview = {}
        last_error: Exception | None = None
        for attempt in range(1, 5):
            overview_attempts = attempt
            attempt_config = config if attempt == 1 else _structured_retry_config(config)
            agent = build_monthly_narrative_agent(attempt_config, allow_tools=False)
            prompt = (
                _batch_prompt(report_month, overview_payload)
                + "\n概述页只能归纳已经通过校验的分页面结论，不得调用工具、不得新增原因。"
                + "\n本次唯一待生成页面是slide=3。最终slides数组必须恰好只有一个元素，且slide必须为3。"
                + "approved_slide_conclusions仅为只读输入，禁止复写其中第4、7至19页的JSON或正文；"
                + "只输出第3页的一个标题和三条概述论点，不输出其他页面。"
                + "从第一个字符{开始，到最后一个字符}结束；不要输出思考标签、Markdown围栏或解释。"
            )
            if last_error:
                prompt += f"\n上一轮JSON未通过校验：{last_error}。请完整重写概述页JSON。"
            reply_text = ""
            try:
                reply = await agent.reply(UserMsg(name="月报概述页任务", content=prompt))
                reply_text = reply.get_text_content()
                candidate = _normalize_batch(_extract_json(reply_text), "OV", (3,), report_month)
                _validate_overview_numbers(candidate, approved_conclusions, report_month, overview_fact_payload)
                overview = candidate
                _save_reply_diagnostic(run_dir, "OV", attempt, reply_text, None)
                break
            except Exception as exc:
                last_error = exc
                _save_reply_diagnostic(run_dir, "OV", attempt, reply_text, exc)
                if attempt < 4:
                    await asyncio.sleep(1.5)
        if not overview:
            overview_fallback = True
            if overview_fact_payload:
                overview = _normalize_batch({'slides':[deepcopy(overview_fact_payload['fact_grounded_draft'])], 'evidence':[]}, 'OV', (3,), report_month)
                overview['slides'][0]['generation_method'] = 'fact_grounded_summary'
            else:
                overview = _deterministic_overview(approved_conclusions, report_month)
            _validate_overview_numbers(overview, approved_conclusions, report_month, overview_fact_payload)
            if progress:
                progress(
                    f"概述页模型连续{overview_attempts}次未通过校验；已仅复用已验收分页面结论完成安全汇总。",
                    92,
                )
        overview["conclusions_digest"] = conclusions_digest
        overview['rule_version'] = NARRATIVE_RULE_VERSION
        overview_path.write_text(json.dumps(overview, ensure_ascii=False, indent=2), encoding="utf-8")
    all_slides.extend(overview["slides"])
    all_evidence.extend(overview.get("evidence", []))
    collection_batches.append({
        "batch_id": "narrative-overview",
        "status": "completed",
        "evidence_ids": [],
        "attempts": overview_attempts,
        "fallback": overview_fallback,
    })

    result = {
        "run_id": run_id,
        "report_month": report_month,
        "slides": sorted(all_slides, key=lambda value: value["slide"]),
        "evidence": all_evidence,
        "collection_batches": collection_batches,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "narrative_rule_version": NARRATIVE_RULE_VERSION,
        "retained_pages": exceptions,
    }
    if research_result is not None:
        result['research_evidence'] = research_result['evidence']
        result['research'] = research_result['research']
        result['evidence_gaps'] = research_result['evidence_gaps']
    result_path = run_dir / "narrative_evidence.json"
    expected = {3, *TARGET_PAGES} - retained
    actual = {int(item.get("slide", 0)) for item in result["slides"]}
    if actual != expected:
        raise RuntimeError(f"主旨句页面不完整：缺少{sorted(expected - actual)}，多出{sorted(actual - expected)}")
    _validate_retained_sources(result, exceptions)
    result_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    result["path"] = str(result_path)
    return result


def narrative_markdown(narrative: dict[str, Any]) -> str:
    lines = ["# 乘用车月报主旨句已生成并通过页面级结构校验", ""]
    lines.append(f"- 任务ID：{narrative.get('run_id', '')}")
    lines.append(f"- 报告月份：{narrative.get('report_month', '')}")
    lines.append(f"- 主旨句页面：{len(narrative.get('slides', []))}页")
    if narrative.get('research'):
        used_ids = {i for s in narrative.get('slides', []) for i in
                    [*s.get('headline_evidence_ids', []),
                     *(i for c in s.get('body_claims', []) for i in c.get('evidence_ids', []))]}
        lines.append(f"- 已核验研究材料：{len(narrative.get('research_evidence', []))}条，去重来源{narrative['research'].get('unique_sources', 0)}个；主旨句引用：{len(used_ids)}条")
        names = {'policy': '政策', 'rights': '权益', 'new-car': '新车', 'reputation': '口碑', 'internet': '互联网'}
        for key, stats in narrative['research'].get('domains', {}).items():
            lines.append(f"- {names.get(key, key)}研究：实际查询{stats['queries']}次，成功返回{stats['succeeded']}次")
        gaps = narrative.get('evidence_gaps', [])
        lines.append(f"- 研究闭环：{'仍有待核验项' if gaps else '已完成来源验证'}；未解决缺口{len(gaps)}项")
        lines.extend(f"  - {gap}" for gap in gaps[:5])
    else:
        lines.append(f"- 外部证据：{len(narrative.get('evidence', []))}条（最终主旨句包收录；历史任务未记录完整研究链路）")
    for item in narrative.get('retained_pages', []):
        detail = '、'.join(f'{code}：{month}' for code, month in item['cards'].items())
        lines.append(f"- 第{item['slide']}页数据更新滞后（{detail}），PPT原内容保留；未纳入当月结论。")
    lines.extend(["", "## 逐页主旨句", ""])
    for slide in narrative.get("slides", []):
        lines.append(f"- 第{slide['slide']}页：{slide['headline']}")
    lines.extend(["", "数据、异常、证据等级和主旨句已经按同一 run_id 固化。现在可以点击右侧“生成PPT版”。"])
    return "\n".join(lines)


def _replace_shape_text(shape, text: str) -> None:
    text_frame = shape.text_frame
    first = text_frame.paragraphs[0]
    if first.runs:
        first.runs[0].text = text
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.text = text
    for paragraph in text_frame.paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def apply_narratives_to_ppt(template_path: Path, output_path: Path, narrative: dict[str, Any]) -> dict[str, Any]:
    retained = {int(item['slide']) for item in narrative.get('retained_pages', [])}
    if retained.intersection(int(item['slide']) for item in narrative.get('slides', [])):
        raise ValueError('滞后数据页面不得覆盖原PPT内容。')
    deck = Presentation(template_path)
    report_month = str(narrative.get("report_month", ""))
    chart_roles = validate_business_chart_roles(deck, report_month, retained) if len(deck.slides) == 21 else []
    if not chart_roles and any(int(s.get('slide',0)) in (5,7,8,12,13,15,17) for s in narrative.get('slides',[])):
        raise ValueError('业务页写入必须使用已核验21页兼容模板。')
    written: list[int] = []
    changed_parts: set[str] = set()
    for item in narrative.get("slides", []):
        slide_number = int(item["slide"])
        if slide_number < 1 or slide_number > len(deck.slides):
            continue
        slide = deck.slides[slide_number - 1]
        if slide_number == 5:
            title_shape, body_shape = prepare_inventory_roles(slide)
        else:
            text_shapes = [shape for shape in slide.shapes if hasattr(shape, "text_frame") and shape.text.strip()]
            if len(text_shapes) < 2:
                raise RuntimeError(f"第{slide_number}页未找到标题和正文文本框。")
            title_shape, body_shape = text_shapes[:2]
        headline, body = _ppt_narrative_text(item, report_month)
        _replace_shape_text(title_shape, headline)
        if body:
            _replace_shape_text(body_shape, body)
        written.append(slide_number)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # Text only: retain native charts, their workbooks and business comparison
    # periods byte-for-byte. Full-year structural reference is intentional.
    serialized = BytesIO()
    deck.save(serialized)
    source_bytes = template_path.read_bytes()
    changed_parts.update(f"ppt/slides/slide{number}.xml" for number in written)
    with zipfile.ZipFile(BytesIO(source_bytes)) as original, \
            zipfile.ZipFile(serialized) as edited, \
            zipfile.ZipFile(output_path, "w") as result:
        for info in original.infolist():
            content = edited.read(info.filename) if info.filename in changed_parts else original.read(info.filename)
            result.writestr(info, content)

    check = Presentation(output_path)
    for item in narrative.get("slides", []):
        slide_number = int(item["slide"])
        text_shapes = [shape for shape in check.slides[slide_number - 1].shapes if hasattr(shape, "text_frame") and shape.text.strip()]
        expected_headline, expected_body = _ppt_narrative_text(item, report_month)
        if not text_shapes or text_shapes[0].text.strip() != expected_headline:
            raise RuntimeError(f"第{slide_number}页主旨句写回校验失败。")
        if expected_body and (len(text_shapes) < 2 or
                              text_shapes[1].text.replace("\v", "\n").strip() != expected_body):
            raise RuntimeError(f"第{slide_number}页正文写回校验失败，不能丢失车型支撑。")
    result = {"slides_written": written, "headline_count": len(written)}
    if retained:
        result['retained_pages'] = narrative['retained_pages']
    if chart_roles:
        result.update(title_only_pages=[5], chart_period_contracts=chart_roles)
    return result
