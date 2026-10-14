"""Update the editable monthly template from the exact accepted BI package."""
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZipFile
from hashlib import sha256
import json
from copy import deepcopy
from math import isclose
from math import ceil
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

from utils.monthly_page_facts import page_business_facts
from utils.monthly_segment_facts import segment_page_facts
from utils.monthly_report_narrative import delayed_page_exceptions, load_page_card_map
from utils.monthly_report_pptx_v2 import update_rate_callouts, update_small_share_labels


def pct(value, plus=False):
    return '' if value is None else f'{value:+.1f}%' if plus else f'{value:.1f}%'


def number(value):
    return '' if value is None else f'{value / 10000:.1f}'


def _text(frame, value):
    runs = [r for p in frame.paragraphs for r in p.runs]
    if runs:
        runs[0].text = str(value)
        for r in runs[1:]:
            r.text = ''
    else:
        frame.text = str(value)


def write_rows(table, rows):
    if len(rows) != len(table.rows) or any(len(row) != len(table.columns) for row in rows):
        raise ValueError('数据行列数与原生表格不匹配，不允许残留旧单元格。')
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            _text(table.cell(i, j).text_frame, value)


def set_callouts(slide, values):
    update_rate_callouts(slide, [None if x is None else x / 100 for x in values], plus=True)


def _ref_values(ref, values, workbook):
    formula = ref.find('{*}f')
    match = re.fullmatch(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?", formula.text or '')
    if not match:
        raise ValueError(f'无法核验图表工作簿引用：{formula.text}')
    sheet, col, row, endcol, endrow = match.groups()
    if sheet not in workbook.sheetnames:
        if len(workbook.sheetnames) != 1:
            raise ValueError('图表数据工作表不明确。')
        sheet = workbook.sheetnames[0]
    bounds = range_boundaries(f'{col}{row}:{endcol or col}{endrow or row}')
    horizontal = bounds[1] == bounds[3] and bounds[0] != bounds[2]
    cells = [(bounds[1] + (0 if horizontal else i), bounds[0] + (i if horizontal else 0)) for i in range(len(values))]
    for i in range(max(bounds[2] - bounds[0] + 1, bounds[3] - bounds[1] + 1)):
        workbook[sheet].cell(bounds[1] + (0 if horizontal else i), bounds[0] + (i if horizontal else 0)).value = None
    for (y, x), value in zip(cells, values):
        workbook[sheet].cell(y, x).value = value
    if cells:
        lastrow, lastcol = cells[-1]
        formula.text = f"'{sheet}'!${col}${row}" + (f':${get_column_letter(lastcol)}${lastrow}' if len(values) > 1 else '')
    cache = ref.find('{*}numCache')
    numeric = cache is not None
    if cache is None:
        cache = ref.find('{*}strCache')
    if cache is None:
        raise ValueError('图表缺少可编辑数据缓存。')
    for item in list(cache):
        if etree.QName(item).localname in ('pt', 'ptCount'):
            cache.remove(item)
    count = OxmlElement('c:ptCount'); count.set('val', str(len(values))); cache.append(count)
    for i, value in enumerate(values):
        if value is None:
            continue
        point = OxmlElement('c:pt'); point.set('idx', str(i))
        v = OxmlElement('c:v'); v.text = str(float(value)) if numeric else str(value)
        point.append(v); cache.append(point)


def set_chart(chart, categories, series):
    """Update original cache and referenced cells without replacing chart styles."""
    if len(series) != len(chart.series):
        raise ValueError('原生图表系列数发生变化。')
    part = chart.part.chart_workbook.xlsx_part
    raw = part.blob
    # Some business templates omit names on built-in cell styles; normalize
    # only that invalid metadata in the copied chart workbook.
    with ZipFile(BytesIO(raw)) as source:
        style = etree.fromstring(source.read('xl/styles.xml'))
        missing = style.xpath('//*[local-name()="cellStyle" and not(@name)]')
        if missing:
            for i, node in enumerate(missing): node.set('name', 'Normal' if i == 0 else f'Style{i}')
            normalized = BytesIO()
            with ZipFile(normalized, 'w') as output:
                for info in source.infolist():
                    output.writestr(info, etree.tostring(style) if info.filename == 'xl/styles.xml' else source.read(info.filename))
            raw = normalized.getvalue()
    book = load_workbook(BytesIO(raw))
    try:
        for old, (name, values) in zip(chart.series, series):
            if len(values) != len(categories):
                raise ValueError('图表分类和数据数量不一致。')
            val = old._ser.xpath('./c:val/c:numRef')
            if len(val) != 1:
                raise ValueError('原生图表数值引用不明确。')
            _ref_values(val[0], values, book)
            cat = old._ser.xpath('./c:cat/c:strRef | ./c:cat/c:numRef')
            if cat:
                _ref_values(cat[0], categories, book)
            tx = old._ser.xpath('./c:tx/c:strRef')
            if tx and name is not None:
                _ref_values(tx[0], [name], book)
        buffer = BytesIO(); book.save(buffer); part.blob = buffer.getvalue()
    finally:
        book.close()


def _row(packet, code, entity, scope, period):
    found = [r for r in packet['business_facts'] if r['source_card_code'] == code and r['entity'] == entity
             and r['scope'] == scope and r['period'] == period]
    if len(found) != 1:
        raise ValueError(f'{code}/{entity}/{scope}/{period}缺少唯一完整数据，不能沿用旧值。')
    return found[0]


def _labels(slide, base, month):
    y, m = int(month[:4]), int(month[5:]); by, bm = int(base[:4]), int(base[5:])
    prev = f'{y}.{m - 1}' if m > 1 else f'{y-1}.12'
    replacements = [(f'{by}年{bm}月', f'{y}年{m}月'), (f'{by}年1—{bm}月', f'{y}年1—{m}月'),
                    (f'{by}.1-{bm}', f'{y}.1-{m}'), (f'{by-1}.1-{bm}', f'{y-1}.1-{m}'),
                    (f'{by}.{bm-1}' if bm > 1 else f'{by-1}.12', prev), (f'{by}.{bm}', f'{y}.{m}')]
    for shape in slide.shapes:
        if shape.has_text_frame:
            value = shape.text
            # Simultaneous replacement prevents August -> September -> October cascades.
            mapping = dict(replacements)
            pattern = '|'.join(re.escape(k) for k in sorted(mapping, key=len, reverse=True))
            updated = re.sub(pattern, lambda match: mapping[match[0]], value)
            if updated != value:
                _text(shape.text_frame, updated)


def _overview(slide,cards,month,skip):
    retail=page_business_facts(4,cards,month)
    current=_row(retail,'P1.1-1','整体市场','single_month',month)
    cumulative=_row(retail,'P1.1-3','整体市场','year_to_date',month)
    energy=page_business_facts(17,cards,month) if 17 not in skip else None
    nev=_row(energy,'P2-11.1','新能源','single_month',month) if energy else None
    ytd_nev=_row(energy,'P2-11.2','新能源','year_to_date',month) if energy else None
    metrics={1048640:number(current['value'])+'万辆',1048644:pct(current['yoy_pct']),1048648:pct(current['mom_pct'],True),
             1048652:pct(nev['share_pct']) if nev else '待更新',1048656:number(cumulative['value'])+'万辆',
             1048660:pct(cumulative['yoy_pct']),1048668:'累计NEV渗透率'+(pct(ytd_nev['share_pct']) if ytd_nev else '待更新')}
    def walk(items):
        for shape in items:
            yield shape
            if hasattr(shape,'shapes'):yield from walk(shape.shapes)
    for shape in walk(slide.shapes):
        if shape.shape_id not in metrics:continue
        value=metrics[shape.shape_id]
        # Keep the header/value paragraph distinction and the original red value run.
        if shape.shape_id==1048668:_text(shape.text_frame,value)
        else:
            paragraphs=shape.text_frame.paragraphs
            if len(paragraphs)!=2:raise ValueError('概述指标卡版式不符合模板。')
            runs=paragraphs[1].runs
            if not runs:raise ValueError('概述指标卡缺少数值文本。')
            runs[0].text=value
            for run in runs[1:]:run.text=''
    sections={7:1048627,8:1048630,12:1048634}
    for page,shapeid in sections.items():
        shape=next(s for s in slide.shapes if s.shape_id==shapeid)
        if page in skip:_text(shape.text_frame,'相关卡片更新滞后，保留原分析页；未纳入本月结论。')
        else:
            facts=page_business_facts(page,cards,month)
            current=[r for r in facts['business_facts'] if r['period']==month and r['scope']=='single_month']
            current=sorted(current,key=lambda r:-r['value'])[:2]
            _text(shape.text_frame,'；'.join(f"{r['entity']}零售{number(r['value'])}万辆，同比{r['yoy_pct']:+.1f}%" for r in current)+'。')
    shape=next(s for s in slide.shapes if s.shape_id==1048636)
    if 16 in skip:_text(shape.text_frame,'企业卡片更新滞后，保留原页；未纳入本月结论。')
    else:
        from utils.monthly_source_rankings import _enterprise_rows
        companies=sorted(_enterprise_rows(cards,'P2-10.1',month,False,[]),key=lambda r:-r['sales'])[:2]
        _text(shape.text_frame,'；'.join(f"{r['model']}零售{number(r['sales'])}万辆，同比{r['yoy_pct']:+.1f}%" for r in companies)+'。')


def template_period(path):
    deck = Presentation(path)
    text = '\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    periods = re.findall(r'(20\d{2})年\s*(1[0-2]|0?[1-9])月', text)
    if not periods:
        # Cover can show publication month; analytical chart titles define the period.
        text = '\n'.join(s.text for s in deck.slides[12].shapes if s.has_text_frame)
        periods = re.findall(r'(20\d{2})年\s*(1[0-2]|0?[1-9])月', text)
    if not periods:
        raise ValueError('无法确认模板期间。')
    # Monthly chart captions are the actual data period, even if cover is publication date.
    captions = '\n'.join(s.text for s in deck.slides[12].shapes if s.has_text_frame and '市场表现' in s.text)
    dates = re.findall(r'(20\d{2})年\s*(1[0-2]|0?[1-9])月', captions)
    year, month = (dates or periods)[-1]
    return f'{year}-{int(month):02d}'


def load_ppt_supplement(source_path, supplement_path):
    source = json.loads(Path(source_path).read_text(encoding='utf-8'))
    if not Path(supplement_path).is_file():
        return source
    supplemented = json.loads(Path(supplement_path).read_text(encoding='utf-8'))
    record = supplemented.get('ppt_grid_supplement', {})
    if (record.get('source_sha256') != sha256(Path(source_path).read_bytes()).hexdigest()
            or supplemented.get('report_month') != source.get('report_month')
            or supplemented.get('source_url') != source.get('source_url')):
        raise ValueError('PPT补采数据与本轮原始数据包不匹配。')
    if {c['code'] for c in supplemented['cards']} != {c['code'] for c in source['cards']}:
        raise ValueError('PPT补采数据卡片集合不匹配。')
    return supplemented


def _comparison(slide, page, cards, month):
    charts = [s.chart for s in slide.shapes if s.has_chart]
    tables = [s.table for s in slide.shapes if s.has_table]
    segment = page in (9, 10, 11)
    if segment:
        packet = segment_page_facts(page, cards, month)
        labels = [r['level'] for r in packet['level_facts'] if r['scope'] == 'single_month']
        def row(label, scope):
            return next(r for r in packet['level_facts'] if r['level'] == label and r['scope'] == scope)
        current = [row(label, 'single_month') for label in labels]
        cumulative = [row(label, 'year_to_date') for label in labels]
        values = lambda rows: [r['sales'] / 10000 for r in rows]
        lastvalues = lambda rows: [r['comparison_sales'] / 10000 for r in rows]
    else:
        packet = page_business_facts(page, cards, month)
        groups = packet['coverage']['groups']; code = groups[0]['card_code']; cumcode = groups[1]['card_code']
        labels = groups[0]['expected_entities']
        current = [_row(packet, code, label, 'single_month', month) for label in labels]
        cumulative = [_row(packet, cumcode, label, 'year_to_date', month) for label in labels]
        values = lambda rows: [r['value'] / 10000 for r in rows]
        lastvalues = lambda rows: [r['comparison_value'] / 10000 for r in rows]
    set_chart(charts[0], [''] * len(labels), [(None, lastvalues(current)), (None, values(current))])
    set_chart(charts[1], [''] * len(labels), [(None, lastvalues(cumulative)), (None, values(cumulative))])
    if segment:
        write_rows(tables[0], [['环比', *[pct(r.get('mom_pct')) for r in current]], ['份额', *[pct(r['share_pct']) for r in current]]])
        write_rows(tables[1], [['份额', *[pct(r['share_pct']) for r in cumulative]]])
    else:
        write_rows(tables[0], [['环比', *[pct(r.get('mom_pct')) for r in current]]])
        previous = f'{month[:4]}-{int(month[5:])-1:02d}' if int(month[5:]) > 1 else f'{int(month[:4])-1}-12'
        old = [_row(packet, code, label, 'single_month', previous) for label in labels]
        # The third column is intentionally prior-year FULL-YEAR structure.
        annual = [list(series.values)[2] for series in charts[2].series]
        share = [[r['share_pct'] for r in old], [r['share_pct'] for r in current], list(reversed(annual)),
                 [r['share_pct'] for r in cumulative]]
        set_chart(charts[2], [''] * 4, [(None, [column[i] for column in share]) for i in reversed(range(len(labels)))])
        update_small_share_labels(slide, [[x / 100 for x in column] for column in share])
    set_callouts(slide, [r['yoy_pct'] for r in current + cumulative])
    return packet


def _point_label(series,index,position,number_format,color):
    labels=series._ser.xpath('./c:dLbls')
    if not labels:
        node=OxmlElement('c:dLbls');series._ser.append(node);labels=[node]
    for old in list(labels[0].xpath('./c:dLbl')):
        idx=old.find('{*}idx')
        if idx is not None and int(idx.get('val'))==index:labels[0].remove(old)
    label=OxmlElement('c:dLbl');idx=OxmlElement('c:idx');idx.set('val',str(index));label.append(idx)
    fmt=OxmlElement('c:numFmt');fmt.set('formatCode',number_format);fmt.set('sourceLinked','0');label.append(fmt)
    tx=OxmlElement('c:txPr');tx.append(OxmlElement('a:bodyPr'));tx.append(OxmlElement('a:lstStyle'))
    paragraph=OxmlElement('a:p');props=OxmlElement('a:pPr');font=OxmlElement('a:defRPr');font.set('sz','1000')
    fill=OxmlElement('a:solidFill');rgb=OxmlElement('a:srgbClr');rgb.set('val',color);fill.append(rgb);font.append(fill)
    props.append(font);paragraph.append(props);tx.append(paragraph);label.append(tx)
    pos=OxmlElement('c:dLblPos');pos.set('val',position);label.append(pos)
    for tag,enabled in [('showLegendKey','0'),('showVal','1'),('showCatName','0'),
                        ('showSerName','0'),('showPercent','0'),('showBubbleSize','0')]:
        value=OxmlElement('c:'+tag);value.set('val',enabled);label.append(value)
    labels[0].insert(0,label)


def _retail(slide, cards, month):
    packet = page_business_facts(4, cards, month)
    chart = next(s.chart for s in slide.shapes if s.has_chart)
    table = next(s.table for s in slide.shapes if s.has_table)
    year, end = int(month[:4]), int(month[5:])
    old = [list(s.values) for s in chart.series]
    for i in range(12):
        historical = [r for r in packet['business_facts'] if r['source_card_code']=='P1.1-1'
                      and r['entity']=='整体市场' and r['period']==f'{year-1}-{i+1:02d}']
        if historical:
            old[0][i] = historical[0]['value']/10000
        period = f'{year}-{i+1:02d}'
        if i >= end:
            old[1][i] = old[3][i] = None; continue
        row = _row(packet, 'P1.1-1', '整体市场', 'single_month', period)
        old[1][i] = row['value'] / 10000
        # Earlier reference rows can be absent from virtual fragments; retain
        # historical comparisons but require the newly appended month's pair.
        if row.get('yoy_pct') is not None:
            old[3][i] = row['yoy_pct'] / 100
        elif i == end - 1:
            raise ValueError('目标月份零售同比比较缺失。')
    set_chart(chart, [f'{i}月' for i in range(1, 13)], [(s.name, v) for s, v in zip(chart.series, old)])
    _point_label(chart.series[2],0,'b','0.0%','707070')
    _point_label(chart.series[3],0,'t','0.0%','C00000')
    rows = [[c.text for c in r.cells] for r in table.rows]
    for i in range(12):
        if i >= end:
            rows[2][i+1] = rows[3][i+1] = ''; continue
        current = sum(_row(packet, 'P1.1-1', '整体市场', 'single_month', f'{year}-{m:02d}')['value'] for m in range(1, i+2))
        rows[2][i+1] = number(current)
        if i == end-1:
            actual = _row(packet, 'P1.1-3', '整体市场', 'year_to_date', month)
            if current != actual['value']:
                raise ValueError('分月零售合计与累计卡片不一致。')
            rows[3][i+1] = pct(actual['yoy_pct'])
    write_rows(table, rows)


def _regions(slide, cards, month):
    packet = page_business_facts(14, cards, month)
    chart = next(s.chart for s in slide.shapes if s.has_chart)
    provinces = sorted([r for r in packet['business_facts'] if r['source_card_code'] == 'P2-8.1' and r['period'] == month], key=lambda r: -r['value'])
    if len(provinces) != len(chart.series[0].values):
        raise ValueError('省份卡片覆盖不足，不能保留旧省份数值。')
    national = _row(packet, 'P1.1-1', '整体市场', 'single_month', month)
    set_chart(chart, [r['entity'] for r in provinces], [(f'{month[:4]}年{int(month[5:])}月销量', [r['value']/10000 for r in provinces]),
              ('同比', [r['yoy_pct']/100 for r in provinces]), ('全国同比', [national['yoy_pct']/100]*len(provinces))])
    regions = [_row(packet, 'P2-8.2', n, 'single_month', month) for n in ('东部','中部','西部','东北')]
    table = next(s.table for s in slide.shapes if s.has_table)
    write_rows(table, [['', *[r['entity'] for r in regions]], ['同比', *[pct(r['yoy_pct']) for r in regions]],
                      ['当月份额', *[pct(r['share_pct']) for r in regions]], ['份额变化', *[f"{r['share_change_pp']:+.1f}" for r in regions]]])


def _energy(slide, cards, month):
    packet = page_business_facts(17, cards, month); year, end = int(month[:4]), int(month[5:])
    chart_shapes = [s for s in slide.shapes if s.has_chart]
    charts = [s.chart for s in chart_shapes]; tables = [s.table for s in slide.shapes if s.has_table]
    shares = [list(s.values) for s in charts[0].series]; sales = [list(s.values) for s in charts[1].series]
    mom = []; yoy = []
    for i in range(12):
        if i >= end:
            shares[1][i] = sales[1][i] = None; mom.append(''); continue
        row = _row(packet, 'P2-11.1', '新能源', 'single_month', f'{year}-{i+1:02d}')
        shares[1][i] = row['share_pct']/100; sales[1][i] = row['value']/10000
        mom.append(pct(row.get('mom_pct'))); yoy.append(row.get('yoy_pct'))
    set_chart(charts[0], [f'{i}月' for i in range(1,13)], [(str(year-1), shares[0]), (str(year), shares[1])])
    set_chart(charts[1], ['']*12, [(None,sales[0]), (None,sales[1])])
    charts[1].value_axis.maximum_scale=ceil(max(v for values in sales for v in values if v is not None)*1.35/20)*20
    for series_index,series in enumerate(charts[1].series):
        for label in list(series._ser.xpath('./c:dLbls/c:dLbl')):
            index=label.find('{*}idx')
            if index is not None and int(index.get('val'))==end-1:
                label.getparent().remove(label)
        labels=series._ser.xpath('./c:dLbls')
        if labels:
            label=OxmlElement('c:dLbl');index=OxmlElement('c:idx');index.set('val',str(end-1));label.append(index)
            format_node=OxmlElement('c:numFmt');format_node.set('formatCode','0.0');format_node.set('sourceLinked','0');label.append(format_node)
            tx=OxmlElement('c:txPr');tx.append(OxmlElement('a:bodyPr'));tx.append(OxmlElement('a:lstStyle'))
            paragraph=OxmlElement('a:p');props=OxmlElement('a:pPr');font=OxmlElement('a:defRPr');font.set('sz','900')
            fill=OxmlElement('a:solidFill');color=OxmlElement('a:srgbClr');color.set('val','003370');fill.append(color);font.append(fill)
            props.append(font);paragraph.append(props);tx.append(paragraph);label.append(tx)
            position=OxmlElement('c:dLblPos');position.set('val','inEnd' if series_index==0 else 'outEnd');label.append(position)
            for name,enabled in [('showLegendKey','0'),('showVal','1'),('showCatName','0'),('showSerName','0'),('showPercent','0'),('showBubbleSize','0')]:
                value=OxmlElement('c:'+name);value.set('val',enabled);label.append(value)
            labels[0].insert(0,label)
    current = _row(packet,'P2-11.2','新能源','year_to_date',month)
    set_chart(charts[2], ['',''], [(None,[current['comparison_value']/10000,current['value']/10000])])
    _point_label(charts[2].series[0],1,'ctr','0.0','FFFFFF')
    write_rows(tables[0],[['渗透率',pct(current['comparison_share_pct']),pct(current['share_pct'])]])
    write_rows(tables[1],[['环比',*mom]])
    for row in tables[1].rows:
        row.height=Inches(.23)
        for cell in row.cells:
            cell.margin_left=cell.margin_right=Inches(.02)
            cell.margin_top=cell.margin_bottom=Inches(.01)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size=Pt(9)
                    run.font.color.rgb=RGBColor(0,0,0)
    # Monthly percentage callouts are sparse shapes in the template. Resolve
    # them by x position, leaving future months empty instead of old values.
    percentages = [s for s in slide.shapes if s.has_text_frame and re.fullmatch(r'[+-]?\d+(?:\.\d+)?%',s.text.strip())]
    monthly = sorted([s for s in percentages if s.left < chart_shapes[2].left],key=lambda s:s.left)
    month_labels = {int(s.text.strip()[:-1]):s for s in slide.shapes if s.has_text_frame and re.fullmatch(r'(?:[1-9]|1[0-2])月',s.text.strip())}
    if len(monthly)<end and monthly and len(month_labels)==12:
        offset = monthly[0].left-month_labels[1].left
        for i in range(len(monthly)+1,end+1):
            xml=deepcopy(monthly[0]._element)
            xml.xpath('.//p:cNvPr')[0].set('id',str(max(s.shape_id for s in slide.shapes)+1))
            slide.shapes._spTree.insert_element_before(xml,'p:extLst')
            clone=slide.shapes[-1];clone.left=month_labels[i].left+offset
            monthly.append(clone)
    for i, shape in enumerate(monthly):
        _text(shape.text_frame, pct(yoy[i],True) if i < len(yoy) else '')
    for shape in percentages:
        if shape not in monthly:
            _text(shape.text_frame,pct(current['yoy_pct'],True))


def _inventory(slide,cards,month,base):
    packet=page_business_facts(5,cards,month)
    charts=[s.chart for s in slide.shapes if s.has_chart]
    count=len(charts[0].series[0].values)
    # The inventory template contains prior-year January through the base month.
    start=int(base[:4])-1
    periods=[f'{year}-{m:02d}' for year in (start,start+1) for m in range(1,13)
             if f'{year}-{m:02d}'<=base]
    if len(periods)!=count:raise ValueError('库存历史窗口与模板不一致。')
    values={period:[charts[0].series[0].values[i],charts[0].series[1].values[i]] for i,period in enumerate(periods)}
    for row in packet['business_facts']:
        if row['source_card_code']=='P1.2-1' and row['scope']=='stock' and row['period']<=month:
            sales=_row(packet,'P1.1-1','整体市场','single_month',row['period'])
            values[row['period']]=[row['value'],row['value']/sales['value']]
    if month not in values:raise ValueError('库存目标月份缺少有效数值。')
    periods=sorted(p for p in values if p<=month)
    set_chart(charts[0],[f'{int(p[5:])}月' for p in periods],[(None,[values[p][0] for p in periods]),(None,[values[p][1] for p in periods])])
    groups=[('传统车','P1.2-2A','P1.2-2B'),('新能源','P1.2-2A','P1.2-2B'),('自主','P1.2-2C','P1.2-2D'),('合资','P1.2-2C','P1.2-2D')]
    stock=[_row(packet,a,n,'stock',month)['value'] for n,a,b in groups]
    ratio=[v/_row(packet,b,n,'single_month',month)['value'] for v,(n,a,b) in zip(stock,groups)]
    set_chart(charts[1],[n for n,a,b in groups],[(None,stock),(None,ratio)])


def _powertrains(slide,cards,month,base):
    packet=page_business_facts(18,cards,month)
    charts=[s.chart for s in slide.shapes if s.has_chart];tables=[s.table for s in slide.shapes if s.has_table]
    def rolling(end,count):
        y,m=map(int,end.split('-'));index=y*12+m-1
        return [f'{(i)//12}-{i%12+1:02d}' for i in range(index-count+1,index+1)]
    before=rolling(base,len(charts[0].series[0].values));periods=rolling(month,len(before))
    for shape in slide.shapes:
        if shape.has_text_frame and re.search(r'20\d{2}年\d+—12月同比为空', shape.text):
            _text(shape.text_frame, re.sub(r'20\d{2}年\d+—12月同比为空',
                  f'{periods[0][:4]}年{int(periods[0][5:])}—12月同比为空', shape.text))
    cache={p:[s.values[i] for s in charts[0].series] for i,p in enumerate(before)}
    table_rows=[[c.text for c in r.cells] for r in tables[0].rows]
    old_share={p:[table_rows[r][i+2] for r in range(4)] for i,p in enumerate(before)}
    quantities=[];shares=[]
    for period in periods:
        rows=[r for r in packet['business_facts'] if r['source_card_code']=='P2-12.1' and r['period']==period]
        indexed={r['entity']:r for r in rows}
        if set(indexed)=={'BEV','PHEV'}:
            bev,phev=indexed['BEV'],indexed['PHEV']
            quantities.append([bev['value']/10000,phev['value']/10000,
                               None if bev.get('yoy_pct') is None else bev['yoy_pct']/100,
                               None if phev.get('yoy_pct') is None else phev['yoy_pct']/100])
            # National denominator must be an explicitly dated sales fact.
            national=page_business_facts(4,cards,month)
            total=_row(national,'P1.1-1','整体市场','single_month',period)['value']
            nev=bev['value']+phev['value']
            shares.append([pct(bev['value']/nev*100),pct(phev['value']/nev*100),pct(bev['value']/total*100),pct(phev['value']/total*100)])
        elif period in cache and period!=month:
            quantities.append(cache[period]);shares.append(old_share[period])
        else:raise ValueError(f'新能源动力类型缺少{period}完整数据。')
    set_chart(charts[0],[f'{int(p[5:])}月' for p in periods],[(s.name,[v[i] for v in quantities]) for i,s in enumerate(charts[0].series)])
    # This template uses a two-level year/month axis; generic strRef updates
    # cannot change its labels. Keep cache and both workbook columns aligned.
    from openpyxl.utils.cell import range_boundaries
    chart=charts[0];part=chart.part.chart_workbook.xlsx_part
    book=load_workbook(BytesIO(part.blob))
    try:
        for ref in chart._chartSpace.xpath('//c:multiLvlStrRef'):
            formula=ref.find('{*}f').text
            sheetname,area=formula.rsplit('!',1)
            sheet=book[sheetname.strip("'")]
            left,top,right,bottom=range_boundaries(area)
            if right-left!=1 or bottom-top+1!=len(periods):raise ValueError('动力类型多级日期轴不匹配')
            for i,p in enumerate(periods):
                sheet.cell(top+i,left).value=p[:4]
                sheet.cell(top+i,right).value=f'{int(p[5:])}月'
            levels=ref.findall('./{*}multiLvlStrCache/{*}lvl')
            for level_index,level in enumerate(levels):
                for old in list(level):level.remove(old)
                previous=None
                for i,p in enumerate(periods):
                    value=f'{int(p[5:])}月' if level_index==0 else p[:4]
                    if level_index and value==previous:continue
                    point=OxmlElement('c:pt');point.set('idx',str(i));value_node=OxmlElement('c:v');value_node.text=value;point.append(value_node);level.append(point);previous=value
        stream=BytesIO();book.save(stream);part.blob=stream.getvalue()
    finally:book.close()
    write_rows(tables[0],[[*table_rows[r][:2],*[v[r] for v in shares]] for r in range(4)])
    cumulative=[_row(packet,'P2-12.3',n,'year_to_date',month) for n in ('BEV','PHEV')]
    set_chart(charts[1],['',''],[(None,[r['comparison_value']/10000 for r in cumulative]),(None,[r['value']/10000 for r in cumulative])])
    from pptx.enum.chart import XL_DATA_LABEL_POSITION
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    axis=charts[1]._chartSpace.xpath('//c:valAx/c:scaling')[0]
    for node in axis.xpath('./c:max'):
        node.set('val',str(max(max(r['value'],r['comparison_value'])/10000 for r in cumulative)*1.3))
    for plot in charts[1].plots:
        plot.has_data_labels=True
        plot.data_labels.show_value=True
        plot.data_labels.number_format='0.0'
        plot.data_labels.position=XL_DATA_LABEL_POSITION.OUTSIDE_END
        plot.data_labels.font.size=Pt(10)
        plot.data_labels.font.color.rgb=RGBColor.from_string('003B70')
    for shape in slide.shapes:
        if shape.shape_id in {1049294,1049295,1049296,1049297}:
            _text(shape.text_frame,'')
        if shape.shape_id==1049305:
            _text(shape.text_frame,f'{month[:4]}年1—{int(month[5:])}月累计表现')
    set_callouts(slide,[r['yoy_pct'] for r in cumulative])
    write_rows(tables[1],[['累计占比',*[pct(r['share_pct']) for r in cumulative]],['占比变化',*[f"{r['share_change_pp']:+.1f}" for r in cumulative]]])


def update_source_template(template_path, output_path, package, mapping_path):
    month = str(package['report_month']); base = template_period(template_path)
    if not re.fullmatch(r'20\d{2}-(0[1-9]|1[0-2])',month):
        raise ValueError('目标月份无效。')
    if month[:4] != base[:4]:
        raise ValueError('跨年度更新须具备上一年度完整基准，当前模板年度基准尚未核验。')
    cards = {c['code']:c for c in package['cards']}
    retained = delayed_page_exceptions(package,load_page_card_map(mapping_path)); skip = {r['slide'] for r in retained}
    deck = Presentation(template_path)
    if len(deck.slides) != 21:
        raise ValueError('兼容模板必须为21页。')
    receipts=[]; changed=set()
    _overview(deck.slides[2],cards,month,skip)
    changed.add('ppt/slides/slide3.xml')
    from utils.monthly_source_rankings import update_rankings
    for page in (4,5,7,8,9,10,11,12,13,14,15,16,17,18,19,20):
        if page in skip:
            continue
        slide = deck.slides[page-1]
        ranking_receipt = None
        if page == 4: _retail(slide,cards,month)
        elif page == 5:_inventory(slide,cards,month,base)
        elif page in (7,8,9,10,11,12,13,15):
            packet=_comparison(slide,page,cards,month)
            if page in (9,10,11): ranking_receipt=update_rankings(slide,page,cards,month,packet)
        elif page == 14: _regions(slide,cards,month)
        elif page == 17: _energy(slide,cards,month)
        elif page == 18:_powertrains(slide,cards,month,base)
        elif page in (16,19,20): ranking_receipt=update_rankings(slide,page,cards,month)
        else:
            raise ValueError(f'第{page}页数据更新尚未支持，禁止写入旧图表。')
        _labels(slide,base,month)
        changed.add(f'ppt/slides/slide{page}.xml')
        for s in slide.shapes:
            if s.has_chart:
                changed.add(str(s.chart.part.partname).lstrip('/'))
                changed.add(str(s.chart.part.chart_workbook.xlsx_part.partname).lstrip('/'))
        receipts.append({'slide':page,'report_month':month,'status':'updated',
                         'warnings':ranking_receipt.get('warnings',[]) if ranking_receipt else []})
    _labels(deck.slides[0],base,month)
    by,bm=int(base[:4]),int(base[5:]); y,m=int(month[:4]),int(month[5:])
    old_release=f'{by+(bm==12)}年{1 if bm==12 else bm+1}月'
    new_release=f'{y+(m==12)}年{1 if m==12 else m+1}月'
    for shape in deck.slides[0].shapes:
        if shape.has_text_frame and shape.text==old_release:_text(shape.text_frame,new_release)
    changed.add('ppt/slides/slide1.xml')
    serialized=BytesIO();deck.save(serialized)
    with ZipFile(template_path) as source, ZipFile(serialized) as updated, ZipFile(output_path,'w') as result:
        for info in source.infolist():
            result.writestr(info,updated.read(info.filename) if info.filename in changed else source.read(info.filename))
    with ZipFile(output_path) as result:
        if result.testzip():raise ValueError('PPT ZIP完整性校验失败。')
    check=Presentation(output_path)
    for receipt in receipts:
        for s in check.slides[receipt['slide']-1].shapes:
            if not s.has_chart:continue
            book=load_workbook(BytesIO(s.chart.part.chart_workbook.xlsx_part.blob),data_only=True)
            try:
                for series in s.chart.series:
                    ref=series._ser.xpath('./c:val/c:numRef')[0]
                    match=re.fullmatch(r"'?([^'!]+)'?!(.+)",ref.find('{*}f').text)
                    sheet,span=match.groups();bounds=range_boundaries(span.replace('$',''))
                    values=[book[sheet].cell(y,x).value for y in range(bounds[1],bounds[3]+1) for x in range(bounds[0],bounds[2]+1)]
                    cached=list(series.values)
                    if len(cached)!=len(values) or any(
                            (a is None)!=(b is None) or (a is not None and not isclose(float(a),float(b),rel_tol=1e-10,abs_tol=1e-8))
                            for a,b in zip(cached,values)):
                        raise ValueError(f"第{receipt['slide']}页图表缓存与内嵌工作簿不一致。")
            finally:book.close()
    return {'template_period':base,'report_month':month,'updated_pages':receipts,'retained_pages':retained,
            'changed_parts':sorted(changed),'sha256':sha256(Path(output_path).read_bytes()).hexdigest()}
