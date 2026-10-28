"""Read-only validation of the supplied province/month workbook for Chronos-2.

The user confirmed that absent province/month/energy combinations mean zero.
An entirely absent month is still a data gap. This module never modifies inputs.
"""
from __future__ import annotations

from dataclasses import dataclass
import hashlib
from pathlib import Path

from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / 'forecast-engine' / 'share-files' / '月度销量-省份.xlsx'
ENERGIES = ('传统燃料', '纯电动', '常规混合动力', '插电式混合动力', '燃料电池')
FORECAST_ENERGIES = ('传统燃料', '纯电动', '常规混合动力', '插电式混合动力')
FOUR_TYPES = '四类合计（不含燃料电池）'
ALL_PROVINCES = '31省份合计'


def _sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def _previous(ym):
    year, month = map(int, ym.split('-'))
    return f'{year - 1:04d}-12' if month == 1 else f'{year:04d}-{month - 1:02d}'


@dataclass(frozen=True)
class ProvinceWorkbook:
    source_hash: str
    provinces: tuple
    energies: tuple
    records: dict
    first_month: str
    last_month: str
    blank_rows: int

    def series(self, province, energy, cutoff=None):
        if province != ALL_PROVINCES and province not in self.provinces:
            raise ValueError('未知省份。')
        if energy not in FORECAST_ENERGIES + (FOUR_TYPES,):
            raise ValueError('未知动力类别。')
        if cutoff is not None:
            import re
            if not isinstance(cutoff, str) or not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])', cutoff):
                raise ValueError('截止月份必须为 YYYY-MM。')
            if cutoff > self.last_month:
                raise ValueError('截止月份晚于数据最新月份。')
        end = cutoff or self.last_month
        months = {ym for ym, _, _ in self.records if ym <= end}
        places = self.provinces if province == ALL_PROVINCES else (province,)
        kinds = FORECAST_ENERGIES if energy == FOUR_TYPES else (energy,)
        if province == ALL_PROVINCES:
            if len(self.provinces) != 31:
                raise ValueError('数据中的省份数量不是 31，不能生成 31 省份合计。')
        values = {ym: sum(self.records.get((ym, place, kind), 0.0)
                          for place in places for kind in kinds) for ym in months}
        # An entirely missing month cannot be assumed to have zero sales.
        if end not in values:
            raise ValueError('截止月份在工作簿中完全缺失，不能按零处理。')
        result = []
        while end in values:
            result.append((end, values[end]))
            end = _previous(end)
        result.reverse()
        return result


def read_workbook(path=SOURCE):
    path = Path(path)
    before = _sha256(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(book.worksheets) != 1:
            raise ValueError('需要核对工作簿表数量，未自动选择数据表。')
        rows = book.worksheets[0].iter_rows(values_only=True)
        header = next(rows, ())
        if tuple(header[3:8]) != ('年', '月', '使用省份', '新能源类型', '数量'):
            raise ValueError('工作簿字段与预期不一致。')
        records = {}
        provinces, energies = set(), set()
        empty = 0
        for row_number, row in enumerate(rows, start=2):
            if not any(value is not None for value in row):
                empty += 1
                continue
            year, month, province, energy, value = row[3:8]
            if (isinstance(year, bool) or not isinstance(year, int) or not 1900 <= year <= 2200
                    or isinstance(month, bool) or not isinstance(month, int) or not 1 <= month <= 12
                    or not isinstance(province, str) or not province.strip()
                    or energy not in ENERGIES or isinstance(value, bool)
                    or not isinstance(value, (int, float))):
                raise ValueError(f'工作簿第 {row_number} 行字段不合法。')
            import math
            if not math.isfinite(value) or value < 0:
                raise ValueError(f'工作簿第 {row_number} 行数量不是非负有限数。')
            key = (f'{year:04d}-{month:02d}', province.strip(), energy)
            if key in records:
                raise ValueError(f'工作簿第 {row_number} 行与已有省份、月份、类别重复。')
            records[key] = float(value)
            provinces.add(key[1]); energies.add(energy)
        if not records:
            raise ValueError('工作簿没有可用数据。')
        after = _sha256(path)
        if before != after:
            raise ValueError('读取期间工作簿发生变化，请重新运行。')
        months = [key[0] for key in records]
        return ProvinceWorkbook(before, tuple(sorted(provinces)), tuple(e for e in ENERGIES if e in energies),
                                records, min(months), max(months), empty)
    finally:
        book.close()


def describe(workbook):
    return dict(source=SOURCE.name, source_sha256=workbook.source_hash,
                first_month=workbook.first_month, last_month=workbook.last_month,
                provinces=list(workbook.provinces), energies=list(FORECAST_ENERGIES) + [FOUR_TYPES],
                source_categories=list(workbook.energies), missing_combination='zero_by_user_confirmation',
                rows=len(workbook.records), blank_rows=workbook.blank_rows,
                data_status='real_unvalidated', unit='辆', metric='零售',
                external_citation_allowed=False)
