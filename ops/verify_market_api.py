"""Non-sensitive API acceptance evidence; does not store rows or credentials."""
import argparse
import asyncio
import io
import json
from pathlib import Path
import sys
import time
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import httpx
from dotenv import load_dotenv
from integrations.market_data_api import api_base_url,get_data,export_data,get_all
from integrations.market_mcp_plugins import market_mcp_plugin_catalog
from utils.config import ENV_PATH,load_json_config
from integrations.new_car_live import load_new_car
from integrations.reputation_live import load_reputation
from integrations.rights_live import call_rights,overview
from openpyxl import load_workbook

parser=argparse.ArgumentParser(); parser.add_argument('--phase',choices=['before','after','direct'],default='direct'); args=parser.parse_args()
load_dotenv(ENV_PATH,override=True); config=load_json_config()
out=Path(__file__).resolve().parents[1]/'outputs'/'market-api-migration'; out.mkdir(parents=True,exist_ok=True)
async def main():
    evidence={}
    local_client=httpx.AsyncClient(timeout=100)
    paths={'policy':'/api/intelligence/policy-overview','rights':'/api/intelligence/rights-overview','new-car':'/api/new-car/dashboard','reputation':'/api/reputation/overview'}
    resources={'policy':'policy/overview','rights':'rights/overview','new-car':'new-cars/dashboard','reputation':'reputation/overview'}
    async def check(domain):
        start=time.perf_counter()
        try:
            if args.phase=='direct': data=await get_data(config,resources[domain],plugin_id=domain); status=200
            else:
                response=await local_client.get('http://127.0.0.1:7860'+paths[domain]); status=response.status_code; data=response.json()
            evidence[domain]={'http_status':status,'success':data.get('success'), 'source':data.get('source'), 'elapsed_ms':round((time.perf_counter()-start)*1000)}
        except Exception as exc: evidence[domain]={'error_type':type(exc).__name__,'elapsed_ms':round((time.perf_counter()-start)*1000)}
    await asyncio.gather(*(check(d) for d in paths))
    if args.phase=='after':
        evidence['first_pass']=dict(evidence)
        await asyncio.gather(*(check(d) for d in paths))
        start=time.perf_counter()
        response=await local_client.get('http://127.0.0.1:7860/api/new-car/list'); payload=response.json()
        evidence['new_car_page_list']={'http_status':response.status_code,'rows':len(payload.get('rows',[])),'elapsed_ms':round((time.perf_counter()-start)*1000)}
        response=await local_client.get('http://127.0.0.1:7860/api/new-car/dashboard'); payload=response.json()
        evidence['new_car_latest_date']=payload.get('stats',{}).get('latest_record_date')
    if args.phase=='direct':
        new=await load_new_car(config)
        planned=await load_new_car(config,'planned')
        policy=await get_all(config,'policies',plugin_id='policy')
        evidence['pagination']={'new_car_rows':len(new['rows']),'new_car_total':new['total'],'unique_new_car_ids':len({r['id'] for r in new['rows']}),'planned_rows':len(planned['rows']),'policy_rows':len(policy['items']),'policy_total':policy['total_count']}
        query=await call_rights(config,'rights_search_offers',{'params':{'limit':20,'offset':0,'brand':'蔚来','activity_types':['金融政策'],'exact_names':True}})
        evidence['rights_filter']={'total':query['total_count'],'page_rows':len(query['items'])}
        binary,_=await export_data(config,'rights/export',{'brand':'蔚来','activity_types':['金融政策'],'exact_names':True,'limit':5000,'data_version':query['data_version']},'rights')
        wb=load_workbook(io.BytesIO(binary)); evidence['rights_export']={'rows':wb.active.max_row-1,'columns':wb.active.max_column}
        if policy['items']:
            ids=[policy['items'][0]['记录编号']]
            binary,_=await export_data(config,'policies/export',{'record_ids':ids,'data_version':policy['data_version']},'policy')
            wb=load_workbook(io.BytesIO(binary)); evidence['policy_export']={'rows':wb.active.max_row-1,'columns':wb.active.max_column}
        from mcp import ClientSession
        from mcp.client.streamable_http import streamable_http_client
        checks={}
        for domain in paths:
            base,token=api_base_url(config,domain)
            async with httpx.AsyncClient(timeout=20) as client:
                denied=await client.get(base+'/'+resources[domain]); wrong=await client.get(base+'/'+resources[domain],headers={'Authorization':'Bearer wrong'})
            plugin=next(p for p in market_mcp_plugin_catalog(config) if p.plugin_id==domain)
            async with httpx.AsyncClient(timeout=30,headers={'Authorization':'Bearer '+token}) as client:
                async with streamable_http_client(plugin.url,http_client=client) as (read,write,_):
                    async with ClientSession(read,write) as session:
                        await session.initialize(); tools=await session.list_tools()
                        checks[domain]={'unauthenticated':denied.status_code,'wrong_token':wrong.status_code,'mcp_tools':len(tools.tools)}
        evidence['auth_and_mcp']=checks
    await local_client.aclose()
    from integrations.market_data_api import close_clients
    await close_clients()
    (out/(args.phase+'.json')).write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(evidence,ensure_ascii=False,indent=2))
asyncio.run(main())
