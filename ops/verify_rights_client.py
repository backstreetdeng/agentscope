"""Acceptance probe without logging tokens or full ledger payloads."""
import asyncio
import io
import json
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from dotenv import load_dotenv
from integrations.rights_live import overview,call_rights,export
from integrations.market_mcp_plugins import load_market_mcp_plugins
from utils.config import load_json_config,ENV_PATH
from openpyxl import load_workbook
from urllib.request import Request, urlopen
from urllib.parse import urlencode

async def main():
    load_dotenv(ENV_PATH,override=True); config=load_json_config()
    data=await overview(config)
    assert data['status']['storage_backend']=='mysql'
    page=await call_rights(config,'rights_search_offers',{'params':{'limit':20}})
    content,filename=await export(config,{'record_ids':[r['记录编号'] for r in page['items']], 'limit':20,'data_version':page['data_version']})
    book=load_workbook(io.BytesIO(content))
    assert book.active.max_column==21 and book.active.max_row==21
    bundle=await load_market_mcp_plugins(config)
    filters={'brand':'蔚来','activity_types':['金融政策'],'exact_names':True,'limit':20,'offset':0}
    with urlopen('http://127.0.0.1:7860/api/intelligence/rights-query?'+urlencode({'params':json.dumps(filters)}),timeout=60) as response:
        live=json.load(response)
    assert live['total_count']==44
    export_filters={**filters,'limit':5000,'data_version':live['data_version']}
    request=Request('http://127.0.0.1:7860/api/intelligence/export',data=json.dumps({'domain':'rights','rightsParams':export_filters}).encode(),headers={'Content-Type':'application/json'},method='POST')
    with urlopen(request,timeout=60) as response:
        filtered_book=load_workbook(io.BytesIO(response.read()))
    assert filtered_book.active.max_row-1==live['total_count'] and filtered_book.active.max_column==21
    print(json.dumps({'source':data['status']['storage_table'],'count':data['rights']['total'],
                      'plugins':[{k:p.as_dict()[k] for k in ('id','status','tool_count')} for p in bundle.statuses],
                      'export_rows':book.active.max_row-1,'headers':book.active.max_column,
                      'filtered_full_export_rows':filtered_book.active.max_row-1},ensure_ascii=False))
asyncio.run(main())
