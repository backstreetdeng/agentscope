"""Complete the user's 12-column baseline without changing either source file."""
from copy import copy
from collections import Counter
from datetime import datetime
from pathlib import Path
import hashlib
import json
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

SOURCE = Path(r'D:/2024年度工作日志和备忘录/数字化转型产品/4.0 同事组/5.0 邓/2026/5.0 平日日志/1012')
BASE = SOURCE/'2026.1-8新车数据TO晓亚冰雨_202609.xlsx'
REF = SOURCE/'2026年W1-W40（20260101-20261003）新车上市数据汇总-202609月.xlsx'
DEST = Path(__file__).resolve().parents[1]/'outputs'/'new-car-baseline'

def number(value):
    m = re.fullmatch(r'\s*(\d+)(?:年|月)?\s*', str(value or ''))
    return int(m.group(1)) if m else None

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def text(value):
    return str(value or '').strip()

def add_sheet(book, name, headers, rows):
    sheet=book.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes='A2'
    sheet.auto_filter.ref=sheet.dimensions
    sheet.sheet_view.showGridLines=False
    for cell in sheet[1]:
        cell.font=Font(name='等线',bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='24466B')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.column_dimensions[cell.column_letter].width=24
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='等线',size=11)
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,str):
                cell.data_type='s'
    return sheet

def main():
    hashes={str(p):digest(p) for p in (BASE,REF)}
    base=load_workbook(BASE)
    original=base['Sheet1']
    headers=[c.value for c in original[1]]
    assert len(headers)==12
    baseline=[]
    source_errors=[]
    for row in original.iter_rows(min_row=2):
        values=[c.value for c in row]
        if values[0] is not None:
            assert number(values[0]) is not None
            baseline.append((row[0].row,values,[copy(c._style) for c in row]))
            for cell in row:
                if cell.data_type=='e' or cell.value=='#N/A':
                    source_errors.append([number(values[0]),headers[cell.column-1],cell.value,cell.coordinate,'原表缺失标记作为文本保留，业务值待确认'])
    ref=load_workbook(REF,read_only=True,data_only=True)
    simple=ref['简版']
    assert list(next(simple.iter_rows(min_row=3,max_row=3,values_only=True)))==headers
    source_rows=[]
    for index,values in enumerate(simple.iter_rows(min_row=4,values_only=True),4):
        if number(values[0]) is not None:
            source_rows.append((index,list(values)))
    september=[(index,row) for index,row in source_rows if number(row[1])==2026 and number(row[2])==9]
    assert len(september)==302 and len({number(row[0]) for _,row in september})==302
    detailed=ref['数据表']
    detail_headers=list(next(detailed.iter_rows(min_row=1,max_row=1,values_only=True)))
    detail={number(row[0]):(index,dict(zip(detail_headers,row))) for index,row in enumerate(detailed.iter_rows(min_row=2,values_only=True),2) if number(row[0]) is not None}
    for _,row in september:
        assert all(text(row[i])==text(detail[number(row[0])][1][field]) for i,field in enumerate(headers))
    prior_ids={number(row[0]) for _,row,_ in baseline}
    old={number(row[0]):row for _,row,_ in baseline}
    retained=[]
    removed=[]
    for index,row,styles in baseline:
        year,month=number(row[1]),number(row[2])
        if number(row[0])==76796 and month==5:
            removed.append((index,row))
            continue
        if year==2026 and 1<=month<=8:
            retained.append((index,row,styles))
        else:
            assert year==2026 and month==9
    assert len(removed)==1 and len(retained)==1926
    assert len({number(row[0]) for _,row,_ in retained})==len(retained)
    assert next(row for _,row,_ in retained if number(row[0])==76796)[2]=='6月'
    modifications=[]
    for source_row,row in september:
        previous=old.get(number(row[0]))
        if previous:
            for i,field in enumerate(headers[3:],3):
                if text(previous[i])!=text(row[i]):
                    modifications.append([number(row[0]),field,previous[i],row[i],'9月字段按汇总文件简版对齐',source_row])
    new_rows=[row for _,row in september if number(row[0]) not in prior_ids]
    template_style=[copy(c._style) for c in original[2]]
    original.delete_rows(2,original.max_row-1)
    final_rows=[]
    for _,row,styles in retained:
        original.append(row)
        for cell,style in zip(original[original.max_row],styles):
            cell._style=copy(style)
            if cell.data_type=='e':
                cell.data_type='s'
                cell.comment=Comment('原始Excel中的错误值，现作为文本保留。业务分类待确认，不推测补填。','核验说明')
        final_rows.append(row)
    for _,row in september:
        row=list(row)
        row[1],row[2]='2026年','9月'
        original.append(row)
        for cell,style in zip(original[original.max_row],template_style):
            cell._style=copy(style)
        final_rows.append(row)
    original.freeze_panes='A2'
    original.auto_filter.ref=original.dimensions
    add_sheet(base,'9月完整款型',headers,[row for _,row in september])
    add_sheet(base,'9月新增款型',headers,new_rows)
    add_sheet(base,'9月字段对账',['款型ID','字段','原表值','汇总文件值','处理规则','简版来源行'],modifications)
    add_sheet(base,'原表异常待确认',['款型ID','字段','原始错误文本','原表单元格','处理说明'],source_errors)
    add_sheet(base,'9月日期来源',['款型ID','车型','上市时间原文','来源工作表','来源行'],[
        [number(row[0]),row[5],str(detail[number(row[0])][1]['上市时间']),'数据表',detail[number(row[0])][0]] for _,row in september])
    notes=[
        ['范围','原表1—8月＋汇总文件9月；10月8款不进入主表。'],
        ['主表结构','Sheet1保持原12列，不加入配置字段；1—8月原值保留，唯一例外为用户确认的76796去重。'],
        ['76796处理','用户明确确认以6月为准；删除补齐版中的5月重复行，原文件不修改。'],
        ['9月重复处理','原表已含36个9月ID；9月整批以汇总文件简版302条为准，不重复追加36条。'],
        ['9月字段变更','11个已有9月款型的品牌、企业或级别按汇总文件对齐，详情见9月字段对账。'],
        ['日期精度','主表仍保留上市年/月；9月准确日期原文另附，后续入库不要把全部9月日期补成9月1日。'],
        ['来源互检','参考文件简版与数据表的12个基础字段逐ID一致。空白、总计行均不计入。'],
        ['源文件保护','两份输入文件SHA256读取前后不变。'],
        ['数据库','本次未导入本地或146数据库，不修改MCP或口碑。'],
        ['原表错误值','3个车型系别字段的#N/A转为文本保留；业务值未改，见原表异常待确认。'],
    ]
    note=add_sheet(base,'补齐说明',['项目','说明'],notes)
    note.column_dimensions['B'].width=105
    assert len(final_rows)==2228 and len({number(row[0]) for row in final_rows})==2228
    assert Counter(number(row[2]) for row in final_rows)[9]==302
    DEST.mkdir(parents=True,exist_ok=True)
    path=DEST/('2026.1-9新车基础款型补齐版_'+datetime.now().strftime('%Y%m%d-%H%M%S')+'.xlsx')
    base.save(path)
    ref.close();base.close()
    check=load_workbook(path,read_only=True,data_only=False)
    actual=[list(r) for r in check['Sheet1'].iter_rows(min_row=2,values_only=True)]
    assert actual==final_rows
    assert list(next(check['Sheet1'].values))==headers
    assert all(c.data_type not in ('e','f') for s in check for row in s for c in row)
    check.close()
    assert hashes=={str(p):digest(p) for p in (BASE,REF)}
    report={'output':str(path),'final_unique_styles':2228,'baseline_jan_aug':1926,'september_total':302,'new_september_ids':len(new_rows),'existing_september_ids':36,'changed_september_ids':len({r[0] for r in modifications}),'changed_cells':len(modifications),'dedup_confirmed_id':'76796','confirmed_month':6,'excluded_october':sum(number(row[2])==10 for _,row in source_rows),'source_errors_preserved_as_text':source_errors,'source_hashes':hashes,'month_counts':dict(sorted(Counter(number(row[2]) for row in final_rows).items())),'validation':'passed: headers, source mapping, row values, unique IDs, date scope, unchanged sources, no formulas/errors','database_modified':False}
    path.with_suffix('.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()
