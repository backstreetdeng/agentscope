"""Regression tests for entity/scope facts and real page checkpoint recovery.

The model boundary is mocked; normalization, validation, rendering budgets and
checkpoint files use the production implementations. No external calls occur.
"""
from __future__ import annotations

import json
import asyncio
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

from utils.monthly_report_narrative import (
    PAGE_LABELS,
    _format_business_text,
    _normalize_batch,
    _model_payload,
    _ppt_narrative_text,
    _validate_batch_facts,
    apply_narratives_to_ppt,
    delayed_page_exceptions,
    generate_monthly_narratives,
)
from utils.monthly_segment_facts import SEGMENT_PAGES, _derived_segment_facts, segment_fact_draft


def supported_page(page: int = 10, month: str = "2026-08"):
    group = SEGMENT_PAGES[page]
    level_code = f"P2-{group}.1"
    ytd_code = f"P2-{group}.2"
    model_code = f"P2-{group}.3A"
    levels = []
    for scope, code, entries in (
        ("single_month", level_code, (("B", 60000, 60.0, -10.0), ("A", 40000, 40.0, 20.0))),
        ("year_to_date", ytd_code, (("B", 550000, 55.0, 5.0), ("A", 450000, 45.0, -5.0))),
    ):
        for level, sales, share, yoy in entries:
            levels.append({"fact_id": f"{code}:{level}:{month}", "level": level,
                           "scope": scope, "period": month, "sales": sales,
                           "share_pct": share, "yoy_pct": yoy, "mom_pct": 1.0})
    models = [
        {"fact_id": f"{model_code}:山海甲:{month}", "model": "山海甲", "level": "B",
         "scope": "single_month", "period": month, "sales": 24419,
         "yoy_pct": -37.6, "mom_pct": 14.8},
        {"fact_id": f"{model_code}:山海乙:{month}", "model": "山海乙", "level": "B",
         "scope": "single_month", "period": month, "sales": 21984,
         "yoy_pct": 14.6, "mom_pct": 10.6},
    ]
    payload = {"slide": page, "required_prefix": PAGE_LABELS[page],
               "card_codes": [level_code, ytd_code, model_code],
               "cards": [{"fact_id": level_code, "period": month, "facts": []}],
               "level_facts": levels, "model_facts": models,
               "coverage_requirement": {"minimum_models": 2}}
    single_ids = [row["fact_id"] for row in levels if row["scope"] == "single_month"]
    ytd_ids = [row["fact_id"] for row in levels if row["scope"] == "year_to_date"]
    slide = {"slide": page, "headline": f"{PAGE_LABELS[page]}：级别与同级车型走势分化",
             "anomaly": {"type": "structural_divergence", "priority": 7,
                         "fact_ids": single_ids},
             "body_claims": [
                 {"role": "core", "claim": "8月B级单月同比下降10.0%、份额60.0%，A级同比增长20.0%、份额40.0%，级别走势差异，增减方向应分别观察。", "fact_ids": single_ids},
                 {"role": "structure", "claim": "累计B级份额55.0%、同比增长5.0%，A级份额45.0%、同比下降5.0%，单月与累计结构不同，不能将口径合并判断。", "fact_ids": ytd_ids},
                 {"role": "models", "claim": "山海甲2.4万辆同比下降37.6%，山海乙2.2万辆同比增长14.6%，同级车型一升一降，恢复并不均衡。", "fact_ids": [row["fact_id"] for row in models]},
             ]}
    return slide, payload


def grounded_page(page: int):
    """A dense monthly draft backed by calculated contribution operands."""
    slide, payload = supported_page(page)
    pairs = [("A0", 10000, 14000), ("A", 40000, 33333), ("B", 60000, 66667), ("C+", 30000, 34000)]
    if page == 9:
        pairs.insert(0, ("A00", 10000, 10000))
    total = sum(current for _, current, _ in pairs)
    levels = [{"fact_id": f"P2-{SEGMENT_PAGES[page]}.1:{level}:2026-08", "level": level,
               "scope": "single_month", "sales": current, "comparison_sales": baseline,
               "share_pct": round(current / total * 100, 1),
               "yoy_pct": round((current / baseline - 1) * 100, 1), "mom_pct": 1.0}
              for level, current, baseline in pairs]
    derived = _derived_segment_facts(page, levels, "2026-08")
    source = {row["level"]: row for row in levels}
    drag = next(row for row in derived if row.get("level") == "B")
    category = {9: "轿车", 10: "SUV", 11: "MPV"}[page]
    slide["headline"] = f"{PAGE_LABELS[page]}：B级收缩为主要拖累，同级车型一升一降"
    slide["anomaly"].update(type="contribution_concentration", fact_ids=[drag["fact_id"]])
    slide["body_claims"][0].update(
        claim=f"8月B级单月同比下降10.0%、份额{source['B']['share_pct']:.1f}%，A级同比增长20.0%、份额{source['A']['share_pct']:.1f}%，级别走势差异，增减方向应分别观察。",
        fact_ids=[source["B"]["fact_id"], source["A"]["fact_id"]],
    )
    slide["body_claims"][1].update(
        claim=f"B级同比净减6667辆，占{category}总净减{drag['net_decline_share_pct']:.1f}%，是主要拖累，A级增长仍难抵补其他级别收缩，需区分观察。",
        fact_ids=[drag["fact_id"], source["A"]["fact_id"]],
    )
    payload.update(level_facts=levels, derived_facts=derived, fact_grounded_draft=deepcopy(slide))
    return slide, payload


class MonthlySegmentFactsRegressionTest(unittest.TestCase):
    def assert_valid_layout(self, slide, payload):
        # Rejection tests must not pass accidentally because a fixture is long.
        layout = dict(slide, coverage_requirement=payload["coverage_requirement"])
        _ppt_narrative_text(layout, "2026-08")

    def test_baseline_pairs_entity_scope_and_growth(self):
        for page in (9, 10, 11):
            with self.subTest(page=page):
                slide, payload = supported_page(page)
                _validate_batch_facts({"slides": [slide]}, [payload])

    def test_segment_model_payload_keeps_only_the_drafts_models_and_monthly_facts(self):
        for page in (9, 10, 11):
            with self.subTest(page=page):
                draft, payload = grounded_page(page)
                extra = dict(payload["model_facts"][0], model="未选车型", fact_id=f"P2-{SEGMENT_PAGES[page]}.3A:未选车型:2026-08")
                yearly = dict(extra, scope="year_to_date", fact_id=f"P2-{SEGMENT_PAGES[page]}.3B:未选车型:2026-08")
                payload["model_facts"].extend([extra, yearly])
                payload["level_facts"].append(dict(payload["level_facts"][0], scope="year_to_date", fact_id=f"P2-{SEGMENT_PAGES[page]}.2:A0:2026-08"))
                original = deepcopy(payload)
                packet = _model_payload([payload])[0]
                needed = {value for claim in draft["body_claims"] for value in claim["fact_ids"]}
                self.assertEqual({row["model"] for row in packet["model_facts"]}, {"山海甲", "山海乙"})
                self.assertEqual(len(packet["model_facts"]), 2)
                self.assertTrue(all(row["scope"] == "single_month" for row in packet["level_facts"]))
                self.assertEqual({row["fact_id"] for row in packet["derived_facts"]},
                                 {row["fact_id"] for row in payload["derived_facts"] if row["fact_id"] in needed})
                self.assertEqual(payload, original)

    def test_mixed_single_month_and_cumulative_clause_uses_each_scope(self):
        slide, payload = supported_page()
        slide["body_claims"][1] = {
            "role": "structure",
            "claim": "B级单月份额60.0%、同比下降10.0%，累计份额55.0%、同比增长5.0%，双口径结构不同，应区分观察，不能合并判断。",
            "fact_ids": [row["fact_id"] for row in payload["level_facts"] if row["level"] == "B"],
        }
        self.assert_valid_layout(slide, payload)
        _validate_batch_facts({"slides": [slide]}, [payload])

    def test_headline_growth_direction_is_not_just_a_known_number(self):
        slide, payload = supported_page()
        slide["headline"] = f"{PAGE_LABELS[10]}：B级同比增长10.0%，A级同比下降20.0%"
        self.assert_valid_layout(slide, payload)
        with self.assertRaises(ValueError):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_two_declining_representatives_do_not_prove_all_mainstream_models_decline(self):
        slide, payload = supported_page()
        payload["model_facts"][1]["yoy_pct"] = -14.6
        slide["body_claims"][2]["claim"] = slide["body_claims"][2]["claim"].replace(
            "同比增长14.6%", "同比下降14.6%"
        ).replace("一升一降", "降幅不同")
        slide["headline"] = f"{PAGE_LABELS[10]}：主流车型同比均下滑"
        with self.assertRaisesRegex(ValueError, "泛化为全部"):
            _validate_batch_facts({"slides": [slide]}, [payload])
        slide["headline"] = f"{PAGE_LABELS[10]}：两款代表车型同比均下滑"
        _validate_batch_facts({"slides": [slide]}, [payload])

    def test_level_volume_must_belong_to_the_named_level(self):
        slide, payload = supported_page()
        slide["body_claims"][0]["claim"] = (
            "8月B级销量4.0万辆、同比下降10.0%，A级同比增长20.0%、份额40.0%，级别走势出现差异，增减方向应分别观察。"
        )
        self.assert_valid_layout(slide, payload)
        with self.assertRaises(ValueError):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_uncited_third_model_cannot_borrow_a_cited_models_metrics(self):
        slide, payload = supported_page()
        slide["body_claims"][2]["claim"] = (
            "山海甲2.4万辆同比下降37.6%，山海乙2.2万辆同比增长14.6%，无据车型2.4万辆同比下降37.6%。"
        )
        self.assert_valid_layout(slide, payload)
        with self.assertRaises(ValueError):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_known_third_model_still_requires_its_own_exact_fact_id(self):
        slide, payload = supported_page()
        third = dict(payload["model_facts"][0], model="山海丙", fact_id="P2-4.3A:山海丙:2026-08")
        payload["model_facts"].append(third)
        slide["body_claims"][2]["claim"] = (
            "山海甲2.4万辆同比下降37.6%，山海乙2.2万辆同比增长14.6%，山海丙2.4万辆同比下降37.6%。"
        )
        self.assert_valid_layout(slide, payload)
        with self.assertRaisesRegex(ValueError, "未引用"):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_colon_group_level_assignment_cannot_override_model_level(self):
        slide, payload = supported_page()
        slide["body_claims"][2]["claim"] = (
            "A级代表车型：山海甲2.4万辆同比下降37.6%，山海乙2.2万辆同比增长14.6%，车型一升一降。"
        )
        self.assert_valid_layout(slide, payload)
        with self.assertRaises(ValueError):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_validation_does_not_append_claim_ids_to_the_anomaly(self):
        slide, payload = supported_page()
        original = deepcopy(slide["anomaly"])
        _validate_batch_facts({"slides": [slide]}, [payload])
        _validate_batch_facts({"slides": [slide]}, [payload])
        self.assertEqual(slide["anomaly"], original)

    def test_typed_fact_id_requires_exact_current_period_record(self):
        slide, payload = supported_page()
        slide["body_claims"][0]["fact_ids"][0] = "P2-4.1:B:2031-01"
        with self.assertRaises(ValueError):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_share_increase_requires_a_comparison_share_not_only_current_share(self):
        slide, payload = supported_page()
        slide["body_claims"][0]["claim"] = slide["body_claims"][0]["claim"].replace("份额60.0%", "份额升至60.0%")
        self.assert_valid_layout(slide, payload)
        with self.assertRaises(ValueError):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def mpv_share_change(self):
        slide, payload = supported_page(11)
        row = next(row for row in payload["level_facts"] if row["level"] == "B" and row["scope"] == "single_month")
        row.update(share_pct=78.7, previous_month_share_pct=77.6, mom_pct=0.9)
        derived = {"fact_id": "P2-5.1:derived:B:2026-08", "level": "B", "scope": "single_month",
                   "share_change_pp": 1.1, "sales_delta": -6667,
                   "sales_comparison_scope": "prior_year_same_month", "share_comparison_scope": "previous_month",
                   "operands": {"current_share": 78.7, "previous_share": 77.6,
                                "current_sales": 60000, "comparison_sales": 66667},
                   "method": "current_sales-comparison_sales;current_share-previous_share"}
        payload["derived_facts"] = [derived]
        slide["body_claims"][0]["fact_ids"] = [row["fact_id"], derived["fact_id"]]
        slide["body_claims"][0]["claim"] = (
            "8月B级份额78.7%较上月77.6%提高1.1个百分点，A级同比增长20.0%，结构与同比方向不同，不能合并判断。"
        )
        return slide, payload

    def test_mpv_share_point_change_cannot_borrow_the_pages_0_9_metric(self):
        slide, payload = self.mpv_share_change()
        slide["body_claims"][0]["claim"] = slide["body_claims"][0]["claim"].replace("提高1.1", "提高0.9")
        self.assert_valid_layout(slide, payload)
        with self.assertRaisesRegex(ValueError, "百分点变化错配"):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_mpv_share_point_change_validates_both_percentage_word_orders(self):
        prefixes = (
            "B级份额78.7%较上月77.6%提高1.1个百分点",
            "B级较上月77.6%的份额提高1.1个百分点至78.7%",
            "B级较上月提高1.1个百分点至78.7%（上月77.6%）",
            "B级上月份额77.6%，本月份额78.7%，提高1.1个百分点",
            "B级份额78.7%较77.6%（上月）提高1.1个百分点",
        )
        for prefix in prefixes:
            with self.subTest(prefix=prefix):
                slide, payload = self.mpv_share_change()
                slide["body_claims"][0]["claim"] = prefix + "，A级同比增长20.0%，结构与同比方向不同，需分别观察。"
                self.assert_valid_layout(slide, payload)
                _validate_batch_facts({"slides": [slide]}, [payload])
        # A legacy share-only metadata field remains a read-compatible fallback.
        slide, payload = self.mpv_share_change()
        payload["derived_facts"][0]["comparison_scope"] = payload["derived_facts"][0].pop("share_comparison_scope")
        _validate_batch_facts({"slides": [slide]}, [payload])
        for direction in ("下降", "降低", "回落", "减少"):
            with self.subTest(direction=direction):
                slide, payload = self.mpv_share_change()
                row = next(row for row in payload["level_facts"] if row["level"] == "B" and row["scope"] == "single_month")
                row.update(share_pct=77.6, previous_month_share_pct=78.7)
                payload["derived_facts"][0]["share_change_pp"] = -1.1
                payload["derived_facts"][0]["operands"].update(current_share=77.6, previous_share=78.7)
                slide["body_claims"][0]["claim"] = (
                    f"8月B级份额77.6%较上月78.7%{direction}1.1个百分点，A级同比增长20.0%，结构与同比方向不同，不能合并判断。"
                )
                _validate_batch_facts({"slides": [slide]}, [payload])

    def test_mpv_share_point_change_rejects_wrong_percentages_and_unclear_basis(self):
        variants = (
            ("份额78.7%", "份额77.6%", "share_pct错配"),
            ("上月77.6%", "上月78.7%", "previous_month_share_pct错配"),
            ("较上月77.6%", "较77.6%", "比较范围不明"),
            ("提高1.1", "提高1.0", "百分点变化错配"),
            ("提高1.1", "降低1.1", "百分点变化错配"),
        )
        for original, replacement, error in variants:
            with self.subTest(replacement=replacement):
                slide, payload = self.mpv_share_change()
                claim = slide["body_claims"][0]
                claim["claim"] = claim["claim"].replace(original, replacement)
                self.assert_valid_layout(slide, payload)
                with self.assertRaisesRegex(ValueError, error):
                    _validate_batch_facts({"slides": [slide]}, [payload])

    def test_point_change_requires_this_levels_precise_derived_fact_and_share_scope(self):
        slide, payload = self.mpv_share_change()
        other = dict(payload["derived_facts"][0], level="A", fact_id="P2-5.1:derived:A:2026-08")
        payload["derived_facts"].append(other)
        slide["body_claims"][0]["fact_ids"][-1] = other["fact_id"]
        with self.assertRaisesRegex(ValueError, "B级缺少已引用"):
            _validate_batch_facts({"slides": [slide]}, [payload])
        slide, payload = self.mpv_share_change()
        payload["derived_facts"][0]["share_comparison_scope"] = "prior_year"
        payload["derived_facts"][0]["comparison_scope"] = "previous_month"
        with self.assertRaisesRegex(ValueError, "比较范围不明"):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_body_rejects_self_contradictory_same_direction_divergence(self):
        slide, payload = supported_page(11)
        slide["body_claims"][2]["claim"] = slide["body_claims"][2]["claim"].replace("一升一降", "同向背离")
        with self.assertRaisesRegex(ValueError, "自相矛盾"):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_adjacent_tenth_is_not_an_allowed_rounding_for_typed_percentages(self):
        for index, old, new in ((0, "下降10.0%", "下降10.1%"), (0, "份额60.0%", "份额60.1%"),
                                (2, "下降37.6%", "下降37.7%")):
            with self.subTest(replacement=new):
                slide, payload = supported_page()
                slide["body_claims"][index]["claim"] = slide["body_claims"][index]["claim"].replace(old, new)
                with self.assertRaises(ValueError):
                    _validate_batch_facts({"slides": [slide]}, [payload])

    def test_contribution_type_without_calculation_is_not_evidence(self):
        slide, payload = supported_page()
        slide["anomaly"]["type"] = "contribution_concentration"
        slide["headline"] = f"{PAGE_LABELS[10]}：B级为主要拖累"
        with self.assertRaisesRegex(ValueError, "贡献测算"):
            _validate_batch_facts({"slides": [slide]}, [payload])

    def test_derived_contribution_number_and_exact_reference_are_accepted(self):
        slide, payload = supported_page()
        derived = {"fact_id": "P2-4.1:contribution:B:2026-08", "level": "B",
                   "scope": "single_month", "contribution_pct": 12.3,
                   "operands": {"level_sales_delta": -6000, "total_sales_delta": -48780},
                   "method": "round(level_sales_delta / total_sales_delta * 100, 1)"}
        payload["derived_facts"] = [derived]
        slide["anomaly"]["type"] = "contribution_concentration"
        slide["anomaly"]["fact_ids"].append(derived["fact_id"])
        slide["headline"] = f"{PAGE_LABELS[10]}：B级为主要拖累，贡献率12.3%"
        _validate_batch_facts({"slides": [slide]}, [payload])
        for missing_field in ("operands", "method"):
            with self.subTest(missing_field=missing_field):
                broken = deepcopy(payload)
                broken["derived_facts"][0].pop(missing_field)
                with self.assertRaisesRegex(ValueError, "贡献测算"):
                    _validate_batch_facts({"slides": [deepcopy(slide)]}, [broken])

    def test_level_net_decline_is_paired_with_delta_not_current_volume(self):
        slide, payload = supported_page()
        derived = {"fact_id": "P2-4.1:derived:B:2026-08", "level": "B", "scope": "single_month",
                   "sales_delta": -6667, "net_decline_share_pct": 46.4,
                   "operands": {"current_sales": 60000, "comparison_sales": 66667, "total_delta": -14368},
                   "method": "current_sales-comparison_sales; negative_delta/total_negative_delta*100",
                   "source_fact_ids": [payload["level_facts"][0]["fact_id"]]}
        payload["derived_facts"] = [derived]
        slide["body_claims"][0]["claim"] = (
            "8月B级同比净减0.7万辆，占SUV总净减46.4%，A级同比增长20.0%、份额40.0%，增减方向有差异，需分别观察。"
        )
        slide["body_claims"][0]["fact_ids"] = [*slide["body_claims"][0]["fact_ids"], derived["fact_id"]]
        self.assert_valid_layout(slide, payload)
        _validate_batch_facts({"slides": [slide]}, [payload])

        # The numeric value also exists for A; citing A must not authorize B.
        other = dict(derived, level="A", fact_id="P2-4.1:derived:A:2026-08")
        payload["derived_facts"].append(other)
        wrong = deepcopy(slide)
        wrong["body_claims"][0]["fact_ids"] = [value if value != derived["fact_id"] else other["fact_id"]
                                               for value in wrong["body_claims"][0]["fact_ids"]]
        with self.assertRaisesRegex(ValueError, "B级缺少已引用"):
            _validate_batch_facts({"slides": [wrong]}, [payload])

    def test_model_trailing_digit_is_not_consumed_as_vehicle_volume(self):
        slide, payload = supported_page(11)
        model = payload["model_facts"][0]
        old_id = model["fact_id"]
        model.update(model="别克GL8", sales=5874, fact_id="P2-5.3A:别克GL8:2026-08")
        claim = slide["body_claims"][2]
        claim["claim"] = claim["claim"].replace("山海甲2.4万辆", "别克GL85874辆")
        claim["fact_ids"] = [model["fact_id"] if value == old_id else value for value in claim["fact_ids"]]

        _, body = _ppt_narrative_text(slide, "2026-08")
        self.assertIn("别克GL8 5874辆", body.replace("\n", ""))
        self.assertNotIn("8.6万辆", body)
        _validate_batch_facts({"slides": [slide]}, [payload])

        # The production generator normalizes before validating and rendering.
        # Protect that path too, not only direct rendering of a raw claim.
        normalized = _normalize_batch({"slides": [deepcopy(slide)]}, "B2", (11,), "2026-08")
        _validate_batch_facts(normalized, [payload])
        _, body = _ppt_narrative_text(normalized["slides"][0], "2026-08")
        self.assertIn("别克GL8 5874辆", body.replace("\n", ""))
        self.assertNotIn("8.6万辆", body)

    def test_aggregate_net_decline_and_net_gain_cannot_borrow_each_others_numbers(self):
        slide, payload = supported_page()
        for row in payload["level_facts"]:
            if row["scope"] == "single_month":
                if row["level"] == "B":
                    row.update(share_pct=37.5, comparison_sales=66667)
                else:
                    row.update(share_pct=25.0, comparison_sales=43333, yoy_pct=-7.7)
        for level in ("A0", "C+"):
            payload["level_facts"].append({"fact_id": f"P2-4.1:{level}:2026-08", "level": level,
                                           "scope": "single_month", "period": "2026-08", "sales": 30000,
                                           "comparison_sales": 27500, "share_pct": 18.8, "yoy_pct": 9.1})
        main = {"fact_id": "P2-4.1:derived:A+B:2026-08", "scope": "single_month", "levels": ["A", "B"],
                "group_share_pct": 62.5, "sales_delta": -10000,
                "operands": {"current_sales": {"A": 40000, "B": 60000},
                             "comparison_sales": {"A": 43333, "B": 66667}, "total_sales": 160000},
                "method": "sum(current_sales)/total_sales*100;sum(current_sales-comparison_sales)"}
        ends = {"fact_id": "P2-4.1:derived:A0+C+:2026-08", "scope": "single_month", "levels": ["A0", "C+"],
                "group_share_pct": 37.5, "sales_delta": 5000,
                "operands": {"current_sales": {"A0": 30000, "C+": 30000},
                             "comparison_sales": {"A0": 27500, "C+": 27500}, "total_sales": 160000},
                "method": main["method"]}
        payload["derived_facts"] = [main, ends]
        slide["body_claims"][0]["claim"] = (
            "8月B级单月同比下降10.0%、份额37.5%，A级同比下降7.7%、份额25.0%，级别降幅有差异，需分别观察。"
        )
        claim = slide["body_claims"][1]
        claim["claim"] = "A/B合计占62.5%，同比净减1.0万辆，A0与C+净增0.5万辆仍不足抵补主流收缩，应分别观察结构变化。"
        claim["fact_ids"] = [main["fact_id"], ends["fact_id"]]
        self.assert_valid_layout(slide, payload)
        _validate_batch_facts({"slides": [slide]}, [payload])
        for original, replacement in (("净减1.0", "净减0.5"), ("净增0.5", "净增1.0"), ("占62.5", "占37.5")):
            with self.subTest(replacement=replacement):
                wrong = deepcopy(slide)
                wrong["body_claims"][1]["claim"] = wrong["body_claims"][1]["claim"].replace(original, replacement)
                with self.assertRaisesRegex(ValueError, "组合.*错配"):
                    _validate_batch_facts({"slides": [wrong]}, [payload])

    def test_future_month_fact_draft_does_not_invent_declines_or_opposite_signs(self):
        for page in (9, 10, 11):
            with self.subTest(page=page):
                _, payload = supported_page(page, "2030-03")
                rows = payload["level_facts"]
                for row in rows:
                    row["yoy_pct"] = 10.0
                template = deepcopy(rows[0])
                for level in ("A0", "C+"):
                    added = dict(template, level=level, sales=20000, share_pct=20.0)
                    added["fact_id"] = f"P2-{SEGMENT_PAGES[page]}.1:{level}:2030-03"
                    rows.append(added)
                for row in payload["model_facts"]:
                    row["yoy_pct"] = 10.0
                draft = segment_fact_draft(page, payload, "2030-03")
                self.assertIsNotNone(draft)
                text = draft["headline"] + "".join(item["claim"] for item in draft["body_claims"])
                for false_relationship in ("同比收缩", "销量收缩", "方向相反", "同比承压", "一升一降"):
                    self.assertNotIn(false_relationship, text)

    def test_future_mpv_same_negative_sign_is_not_opposite_directions(self):
        _, payload = supported_page(11, "2030-12")
        for row in [*payload["level_facts"], *payload["model_facts"]]:
            row["yoy_pct"] = -10.0
        added = dict(payload["level_facts"][0], level="A0", sales=20000, share_pct=20.0,
                     fact_id="P2-5.1:A0:2030-12")
        payload["level_facts"].append(added)
        draft = segment_fact_draft(11, payload, "2030-12")
        self.assertIsNotNone(draft)
        text = draft["headline"] + "".join(item["claim"] for item in draft["body_claims"])
        self.assertNotIn("方向相反", text)
        self.assertNotIn("一升一降", text)


def response(slides):
    return SimpleNamespace(get_text_content=lambda: json.dumps({"slides": slides}, ensure_ascii=False))


def prompt_pages(message):
    content = message.get_text_content()
    raw = content.split("页面输入：\n", 1)[1]
    payload, _ = json.JSONDecoder().raw_decode(raw)
    return [item["slide"] for item in payload]


class MonthlyBusinessTextAndPptxRegressionTest(unittest.TestCase):
    def test_positive_growth_wording_is_normalized_and_idempotent(self):
        text = "同比正增、同比正增长、同比增长"
        expected = "同比增长、同比增长、同比增长"
        formatted = _format_business_text(text, "2026-08")
        self.assertEqual(formatted, expected)
        self.assertNotIn("增长长", formatted)
        for _ in range(3):
            formatted = _format_business_text(formatted, "2026-08")
            self.assertEqual(formatted, expected)

    def test_ppt_write_changes_only_requested_slide_xml_bytes(self):
        # This minimal package is a test fixture, not a newly authored deck.
        from pptx import Presentation
        from pptx.util import Inches

        deck = Presentation()
        for number in range(1, 4):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for index, text in enumerate((f"原始标题{number}", f"原始正文{number}", "换行标记")):
                box = slide.shapes.add_textbox(Inches(1), Inches(1 + index), Inches(6), Inches(1))
                box.text = text
        serialized = BytesIO()
        deck.save(serialized)
        with ZipFile(serialized) as original:
            parts = {info.filename: original.read(info.filename) for info in original.infolist()}

        untouched = "ppt/slides/slide1.xml"
        source_xml = parts[untouched]
        self.assertIn("<a:t>换行标记</a:t>".encode(), source_xml)
        self.assertIn(b"<a:lstStyle/>", source_xml)
        source_xml = source_xml.replace("<a:t>换行标记</a:t>".encode(), b"<a:t>\n</a:t>", 1)
        source_xml = source_xml.replace(b"<a:lstStyle/>", b"<a:lstStyle></a:lstStyle>", 1)
        parts[untouched] = source_xml
        parts["customXml/preserved.bin"] = b"opaque original bytes\x00\xff\r\n"

        item = {"slide": 2, "headline": "8月级别走势分化",
                "body_claims": [{"role": "core", "claim": "单月同比增长与累计走势不同，应区分口径。", "fact_ids": []}]}
        expected_headline, expected_body = _ppt_narrative_text(item, "2026-08")
        with tempfile.TemporaryDirectory() as directory:
            template = Path(directory) / "source.pptx"
            output = Path(directory) / "updated.pptx"
            with ZipFile(template, "w") as package:
                for name, content in parts.items():
                    package.writestr(name, content)
            before = template.read_bytes()

            # Prove the fixture catches the former full-package save behavior.
            roundtrip = BytesIO()
            Presentation(template).save(roundtrip)
            with ZipFile(roundtrip) as reserialized:
                self.assertNotEqual(reserialized.read(untouched), source_xml)

            retained = [{'slide': 1, 'reason': 'source_update_delayed',
                         'cards': {'P1.2-1': '2026-07'}, 'action': 'preserve_existing_content'}]
            result = apply_narratives_to_ppt(template, output, {"report_month": "2026-08", "slides": [item], 'retained_pages': retained})
            self.assertEqual(result, {"slides_written": [2], "headline_count": 1, 'retained_pages': retained})
            self.assertEqual(template.read_bytes(), before)
            with ZipFile(output) as updated:
                self.assertEqual(set(updated.namelist()), set(parts))
                changed = {name for name, content in parts.items() if updated.read(name) != content}
                self.assertEqual(changed, {"ppt/slides/slide2.xml"})
                for name, content in parts.items():
                    if name != "ppt/slides/slide2.xml":
                        with self.subTest(part=name):
                            self.assertEqual(updated.read(name), content)
                self.assertIn(b"<a:t>\n</a:t>", updated.read(untouched))
            reread = Presentation(output)
            self.assertEqual(reread.slides[1].shapes[0].text, expected_headline)
            self.assertEqual(reread.slides[1].shapes[1].text.replace("\v", "\n"), expected_body)


class MonthlyDerivedArithmeticRegressionTest(unittest.TestCase):
    @staticmethod
    def levels(page, pairs):
        total = sum(current for _, current, _ in pairs)
        return [{"fact_id": f"P2-{SEGMENT_PAGES[page]}.1:{level}:2026-08", "level": level,
                 "scope": "single_month", "sales": current, "comparison_sales": baseline,
                 "share_pct": round(current / total * 100, 1)}
                for level, current, baseline in pairs]

    def sedan_levels(self):
        # A uses the observed 161310/295850 pair. The other fixture levels
        # supply a complete comparison universe with total delta -289957.
        return self.levels(9, [("A00", 100000, 200000), ("A0", 60000, 50000),
                               ("A", 161310, 295850), ("B", 250000, 315417), ("C+", 40000, 40000)])

    def test_suv_group_share_uses_raw_sales_not_sum_of_rounded_level_shares(self):
        rows = self.levels(10, [("A0", 71655, 71655), ("A", 325648, 325648),
                                ("B", 309619, 309619), ("C+", 142185, 142185)])
        derived = _derived_segment_facts(10, rows, "2026-08")
        group = next(row for row in derived if row.get("levels") == ["A", "B"])
        rounded_sum = round(sum(row["share_pct"] for row in rows if row["level"] in {"A", "B"}), 1)
        self.assertEqual(rounded_sum, 74.9)
        self.assertEqual(group["group_share_pct"], 74.8)
        self.assertEqual(group["operands"]["current_sales"], {"A": 325648, "B": 309619})
        self.assertEqual(group["operands"]["total_sales"], 849107)
        self.assertEqual(set(group["source_fact_ids"]), {row["fact_id"] for row in rows})

    def test_sedan_a_net_decline_and_contribution_retain_their_raw_operands(self):
        rows = self.sedan_levels()
        derived = _derived_segment_facts(9, rows, "2026-08")
        record = next(row for row in derived if row.get("level") == "A")
        self.assertEqual(record["sales_delta"], -134540)
        self.assertEqual(record["net_decline_share_pct"], 46.4)
        self.assertEqual(record["operands"]["current_sales"], 161310)
        self.assertEqual(record["operands"]["comparison_sales"], 295850)
        self.assertEqual(record["operands"]["total_delta"], -289957)
        self.assertTrue(record["largest_decline"])
        self.assertEqual(set(record["source_fact_ids"]), {row["fact_id"] for row in rows})

    def test_missing_level_or_comparison_cell_disables_category_contribution(self):
        rows = self.sedan_levels()
        missing_level = deepcopy(rows[:-1])
        missing_cell = deepcopy(rows)
        missing_cell[-1]["comparison_sales"] = None
        for label, incomplete in (("level", missing_level), ("comparison", missing_cell)):
            with self.subTest(missing=label):
                derived = _derived_segment_facts(9, incomplete, "2026-08")
                self.assertTrue(any(row.get("sales_delta", 0) < 0 for row in derived))
                self.assertTrue(all("total_delta" not in row["operands"] for row in derived))
                self.assertTrue(all("net_decline_share_pct" not in row for row in derived))
                self.assertTrue(all("largest_decline" not in row for row in derived))

    def test_month_on_month_share_gain_is_independent_of_year_on_year_volume_decline(self):
        rows = self.sedan_levels()
        source = next(row for row in rows if row["level"] == "A")
        source["previous_month_share_pct"] = 25.3
        self.assertEqual(source["share_pct"], 26.4)
        record = next(row for row in _derived_segment_facts(9, rows, "2026-08") if row.get("level") == "A")
        self.assertEqual(record["sales_delta"], -134540)
        self.assertEqual(record["share_change_pp"], 1.1)
        self.assertEqual(record["share_comparison_scope"], "previous_month")
        self.assertEqual(record["sales_comparison_scope"], "prior_year_same_month")
        self.assertEqual(record["operands"]["current_share"], 26.4)
        self.assertEqual(record["operands"]["previous_share"], 25.3)


class MonthlyPageCheckpointRegressionTest(unittest.IsolatedAsyncioTestCase):
    async def test_retained_page_is_excluded_from_facts_model_and_overview(self):
        with tempfile.TemporaryDirectory() as directory:
            root, examples, _, slides, payloads = self.setup_grounded_run(directory)
            retained = [{'slide': 11, 'reason': 'source_update_delayed',
                         'cards': {'P2-5.1': '2026-07'}, 'action': 'preserve_existing_content'}]
            requests = []
            async def reply(message):
                pages = prompt_pages(message)
                requests.append(pages)
                if pages == [3] or pages == (3,):
                    return self.overview_response()
                return response([slides[page] for page in pages])
            agent = SimpleNamespace(reply=AsyncMock(side_effect=reply))
            with patch('utils.monthly_report_narrative.delayed_page_exceptions', return_value=retained):
                result = await self.run_pipeline(root, examples, payloads, agent)
            self.assertEqual({item['slide'] for item in result['slides']}, {3, 9, 10})
            self.assertEqual(result['retained_pages'], retained)
            self.assertTrue(all(11 not in pages for pages in requests))
            facts = json.loads((root/'partial-run'/'facts-B1.json').read_text(encoding='utf-8'))
            self.assertEqual({item['slide'] for item in facts}, {9, 10})

    def test_delayed_inventory_retains_entire_mapped_page_from_actual_source(self):
        package = {'report_month': '2026-09', 'delayed_cards': ['P1.1-1'], 'cards': [
            {'code': 'P1.2-1', 'text': 'P1.2-1：行业库存完整月份与数量数据\n日期\n2026-08\n4119751\n同比\n-22.5%\n'},
            {'code': 'P1.1-1', 'text': 'P1.1-1：行业零售完整月份与数量数据\n日期\n2026-09\n1668934\n同比\n-22.5%\n'}]}
        exceptions = delayed_page_exceptions(package, {4: ['P1.1-1'], 5: ['P1.2-1', 'P1.2-2A']})
        self.assertEqual(exceptions, [{'slide': 5, 'reason': 'source_update_delayed',
                                     'cards': {'P1.2-1': '2026-08'}, 'action': 'preserve_existing_content'}])

    def test_retained_page_cannot_be_written_even_when_caller_provides_narrative(self):
        with self.assertRaisesRegex(ValueError, '不得覆盖'):
            apply_narratives_to_ppt(Path('unused.pptx'), Path('unused-output.pptx'), {
                'retained_pages': [{'slide': 5}], 'slides': [{'slide': 5}]})

    def setup_run(self, directory):
        root = Path(directory)
        examples = root / "examples.json"
        examples.write_text('{"pages": {}}', encoding="utf-8")
        run_dir = root / "partial-run"
        run_dir.mkdir()
        (run_dir / "evidence-B1.json").write_text('{"evidence": []}', encoding="utf-8")
        pairs = [supported_page(page) for page in (9, 10, 11)]
        return root, examples, run_dir, {slide["slide"]: slide for slide, _ in pairs}, [payload for _, payload in pairs]

    def setup_grounded_run(self, directory):
        root, examples, run_dir, _, _ = self.setup_run(directory)
        pairs = [grounded_page(page) for page in (9, 10, 11)]
        return root, examples, run_dir, {slide["slide"]: slide for slide, _ in pairs}, [payload for _, payload in pairs]

    @staticmethod
    def rejected_slides(slides, pages):
        rejected = [deepcopy(slides[page]) for page in pages]
        for slide in rejected:
            slide["body_claims"][2]["claim"] = slide["body_claims"][2]["claim"].replace("下降37.6%", "下降14.6%")
        return rejected

    @staticmethod
    def overview_response():
        return response([{"slide": 3, "headline": "概述：级别与车型走势分化",
                          "body_claims": [{"claim": "仅归纳已经核验的级别及车型关系。"}]}])

    async def run_pipeline(self, root, examples, payloads, agent):
        def page_payload(pages, *args):
            return deepcopy([item for item in payloads if item["slide"] in pages])
        with patch("utils.monthly_report_narrative.BATCHES", ((9, 10, 11),)), \
             patch("utils.monthly_report_narrative.TARGET_PAGES", (9, 10, 11)), \
             patch("utils.monthly_report_narrative.load_page_card_map", return_value={}), \
             patch("utils.monthly_report_narrative._page_payload", side_effect=page_payload), \
             patch("utils.monthly_report_narrative.build_monthly_narrative_agent", return_value=agent) as build, \
             patch("utils.monthly_report_narrative.asyncio.sleep", new=AsyncMock()):
            try:
                return await generate_monthly_narratives(
                    {"report_month": "2026-08", "cards": []}, run_id="partial-run", config={},
                    plugin_tools=[], mapping_path=root / "unused.xlsx", examples_path=examples, output_dir=root,
                )
            finally:
                self.assertTrue(all(call.kwargs.get("allow_tools") is False for call in build.call_args_list))

    async def test_passed_page_is_saved_immediately_then_reused_on_new_invocation(self):
        with tempfile.TemporaryDirectory() as directory:
            root, examples, run_dir, slides, payloads = self.setup_run(directory)
            requested = []

            async def failing_reply(message):
                pages = prompt_pages(message)
                requested.append(pages)
                if len(requested) > 1:
                    self.assertTrue((run_dir / "narrative-B1-page-9.json").is_file())
                batch = [deepcopy(slides[page]) for page in pages]
                for slide in batch:
                    if slide["slide"] != 9:
                        slide["body_claims"][2]["claim"] = slide["body_claims"][2]["claim"].replace("下降37.6%", "下降14.6%")
                return response(batch)

            agent = SimpleNamespace(reply=AsyncMock(side_effect=failing_reply))
            with self.assertRaisesRegex(RuntimeError, "连续三次"):
                await self.run_pipeline(root, examples, payloads, agent)
            self.assertEqual(requested, [[9, 10, 11], [10, 11], [10, 11]])
            saved = (run_dir / "narrative-B1-page-9.json").read_bytes()
            self.assertFalse((run_dir / "narrative-B1.json").exists())
            self.assertFalse((run_dir / "narrative-B1-page-10.json").exists())
            self.assertFalse((run_dir / "narrative-B1-page-11.json").exists())
            resumed = []

            async def succeeding_reply(message):
                pages = prompt_pages(message)
                resumed.append(pages)
                if pages == [3]:
                    return response([{"slide": 3, "headline": "概述：级别与车型走势分化",
                                      "body_claims": [{"claim": "仅归纳已经核验的级别及车型关系。"}]}])
                return response([deepcopy(slides[page]) for page in pages])

            result = await self.run_pipeline(root, examples, payloads, SimpleNamespace(reply=AsyncMock(side_effect=succeeding_reply)))
            self.assertEqual(resumed, [[10, 11], [3]])
            self.assertEqual((run_dir / "narrative-B1-page-9.json").read_bytes(), saved)
            self.assertEqual([item["slide"] for item in result["slides"]], [3, 9, 10, 11])

    async def test_omitted_page_does_not_discard_valid_returned_page(self):
        with tempfile.TemporaryDirectory() as directory:
            root, examples, run_dir, slides, payloads = self.setup_run(directory)
            requested = []

            async def partial_reply(message):
                requested.append(prompt_pages(message))
                return response([deepcopy(slides[9])])

            with self.assertRaisesRegex(RuntimeError, "连续三次"):
                await self.run_pipeline(root, examples, payloads, SimpleNamespace(reply=AsyncMock(side_effect=partial_reply)))
            self.assertTrue((run_dir / "narrative-B1-page-9.json").is_file())
            self.assertEqual(requested, [[9, 10, 11], [10, 11], [10, 11]])
            self.assertFalse((run_dir / "narrative-B1.json").exists())

    async def test_three_rejected_model_edits_publish_only_a_strictly_valid_fact_draft(self):
        with tempfile.TemporaryDirectory() as directory:
            root, examples, run_dir, slides, payloads = self.setup_grounded_run(directory)
            requested = []

            async def rejected_editor(message):
                pages = prompt_pages(message)
                requested.append(pages)
                if pages == [3]:
                    return self.overview_response()
                return response(self.rejected_slides(slides, pages))

            result = await self.run_pipeline(root, examples, payloads, SimpleNamespace(reply=AsyncMock(side_effect=rejected_editor)))
            self.assertEqual(requested, [[9, 10, 11]] * 3 + [[3]])
            self.assertTrue((run_dir / "narrative_evidence.json").is_file())
            self.assertEqual(result["collection_batches"][0]["attempts"], 3)
            for page in (9, 10, 11):
                saved = json.loads((run_dir / f"narrative-B1-page-{page}.json").read_text(encoding="utf-8"))
                slide = next(row for row in result["slides"] if row["slide"] == page)
                self.assertEqual(slide["generation_method"], "fact_grounded")
                self.assertIn("模型润色未通过", slide["generation_note"])
                self.assertIn("同比错配", saved["model_edit_error"])
                self.assertEqual(saved["slides"][0]["generation_method"], "fact_grounded")
                _validate_batch_facts({"slides": [slide]}, payloads)

    async def test_invalid_fact_draft_still_cannot_publish_after_three_model_failures(self):
        for failure in ("wrong_number", "too_short", "missing_model_fact"):
            with self.subTest(failure=failure), tempfile.TemporaryDirectory() as directory:
                root, examples, run_dir, slides, payloads = self.setup_grounded_run(directory)
                for payload in payloads:
                    draft = payload["fact_grounded_draft"]
                    if failure == "wrong_number":
                        draft["body_claims"][2]["claim"] = draft["body_claims"][2]["claim"].replace("下降37.6%", "下降14.6%")
                    elif failure == "too_short":
                        draft["body_claims"][0]["claim"] = "8月B级同比下降10.0%。"
                        draft["body_claims"][1]["claim"] = "A级同比增长20.0%。"
                        draft["body_claims"][2]["claim"] = "山海甲2.4万辆同比下降37.6%，山海乙2.2万辆同比增长14.6%。"
                    else:
                        draft["body_claims"][2]["fact_ids"] = draft["body_claims"][2]["fact_ids"][:1]

                async def rejected_editor(message):
                    return response(self.rejected_slides(slides, prompt_pages(message)))

                agent = SimpleNamespace(reply=AsyncMock(side_effect=rejected_editor))
                with self.assertRaisesRegex(RuntimeError, "连续三次"):
                    await self.run_pipeline(root, examples, payloads, agent)
                self.assertEqual(agent.reply.await_count, 3)
                self.assertFalse((run_dir / "narrative_evidence.json").exists())
                self.assertFalse((run_dir / "narrative-B1.json").exists())
                self.assertTrue(all(not (run_dir / f"narrative-B1-page-{page}.json").exists() for page in (9, 10, 11)))

    async def test_fact_draft_completion_after_restart_does_not_rewrite_approved_9_10(self):
        with tempfile.TemporaryDirectory() as directory:
            root, examples, run_dir, slides, payloads = self.setup_grounded_run(directory)
            first_requests = []

            async def interrupted_editor(message):
                pages = prompt_pages(message)
                first_requests.append(pages)
                if len(first_requests) > 1:
                    raise asyncio.CancelledError("simulated process interruption after page checkpoint")
                rows = [deepcopy(slides[page]) for page in pages if page in {9, 10}]
                return response(rows + self.rejected_slides(slides, [11]))

            with self.assertRaises(asyncio.CancelledError):
                await self.run_pipeline(root, examples, payloads, SimpleNamespace(reply=AsyncMock(side_effect=interrupted_editor)))
            self.assertEqual(first_requests, [[9, 10, 11], [11]])
            snapshots = {page: (run_dir / f"narrative-B1-page-{page}.json").read_bytes() for page in (9, 10)}
            self.assertFalse((run_dir / "narrative-B1-page-11.json").exists())
            resumed_requests = []

            async def rejected_editor(message):
                pages = prompt_pages(message)
                resumed_requests.append(pages)
                if pages == [3]:
                    return self.overview_response()
                return response(self.rejected_slides(slides, pages))

            result = await self.run_pipeline(root, examples, payloads, SimpleNamespace(reply=AsyncMock(side_effect=rejected_editor)))
            self.assertEqual(resumed_requests, [[11], [11], [11], [3]])
            methods = {row["slide"]: row.get("generation_method") for row in result["slides"]}
            self.assertEqual([methods[9], methods[10], methods[11]], ["model_refined", "model_refined", "fact_grounded"])
            for page in (9, 10):
                self.assertEqual((run_dir / f"narrative-B1-page-{page}.json").read_bytes(), snapshots[page])


if __name__ == "__main__":
    unittest.main()
