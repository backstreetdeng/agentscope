"""Real September migration, editable-data and retained-page acceptance."""
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from hashlib import sha256
import json
import unittest
from pptx import Presentation
from utils.monthly_source_pptx import update_source_template, load_ppt_supplement

ROOT=Path(__file__).resolve().parents[1]
ASSETS=ROOT/'skills/market-analysis-skill/passenger-car-monthly-report/assets'
SOURCE=ROOT/'outputs/monthly-report-data/data-source-package-20261013-011831.json'
SUPPLEMENT=ROOT/'outputs/monthly-report-data/month-20261013-014046-cd5ae8ff/ppt-data-supplement.json'


@unittest.skipUnless(SUPPLEMENT.is_file(),'captured cell-boundary supplement unavailable')
class MonthlySourcePptxAcceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=TemporaryDirectory();cls.output=Path(cls.temp.name)/'september.pptx'
        cls.template=ASSETS/'乘用车月报PPT兼容模板.pptx'
        cls.original_hash=sha256(cls.template.read_bytes()).hexdigest()
        cls.package=load_ppt_supplement(SOURCE,SUPPLEMENT)
        cls.receipt=update_source_template(cls.template,cls.output,cls.package,ASSETS/'乘用车月报-数据源映射指标.xlsx')
        cls.deck=Presentation(cls.output)

    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()

    def test_september_cover_data_and_future_month_blanks(self):
        self.assertIn('2026年9月中国乘用车市场分析报告','\n'.join(s.text for s in self.deck.slides[0].shapes if s.has_text_frame))
        self.assertIn('2026年10月',[s.text for s in self.deck.slides[0].shapes if s.has_text_frame])
        chart=next(s.chart for s in self.deck.slides[3].shapes if s.has_chart)
        self.assertAlmostEqual(chart.series[1].values[8],166.8934)
        self.assertTrue(all(v is None for v in chart.series[1].values[9:]))
        self.assertEqual(sha256(self.template.read_bytes()).hexdigest(),self.original_hash)

    def test_retained_pages_and_their_workbooks_are_identical(self):
        original=Presentation(self.template)
        with ZipFile(self.template) as before,ZipFile(self.output) as after:
            for entry in self.receipt['retained_pages']:
                page=entry['slide']
                self.assertEqual(before.read(f'ppt/slides/slide{page}.xml'),after.read(f'ppt/slides/slide{page}.xml'))
                for shape in original.slides[page-1].shapes:
                    if shape.has_chart:
                        for part in (shape.chart.part,shape.chart.part.chart_workbook.xlsx_part):
                            name=str(part.partname).lstrip('/')
                            self.assertEqual(before.read(name),after.read(name))

    def test_annual_reference_remains_full_year_while_current_ytd_advances(self):
        original=Presentation(self.template)
        for page in (13,15):
            old=[s.chart for s in original.slides[page-1].shapes if s.has_chart][2]
            new=[s.chart for s in self.deck.slides[page-1].shapes if s.has_chart][2]
            self.assertEqual([s.values[2] for s in old.series],[s.values[2] for s in new.series])
            texts=[s.text for s in self.deck.slides[page-1].shapes if s.has_text_frame]
            self.assertIn('2025.1-12',texts);self.assertIn('2026.1-9',texts)

    def test_native_chart_and_embedded_workbook_counts_preserved(self):
        self.assertEqual(len(self.deck.slides),21)
        self.assertEqual(sum(s.has_chart for slide in self.deck.slides for s in slide.shapes),39)
        with ZipFile(self.template) as before,ZipFile(self.output) as after:
            self.assertEqual(set(before.namelist()),set(after.namelist()))
            self.assertIsNone(after.testzip())

    def test_supplement_cannot_be_reused_for_another_source_hash(self):
        modified=Path(self.temp.name)/'modified.json'
        modified.write_bytes(SOURCE.read_bytes()+b'\n')
        with self.assertRaisesRegex(ValueError,'不匹配'):
            load_ppt_supplement(modified,SUPPLEMENT)

    def test_future_target_without_current_data_is_rejected(self):
        data=json.loads(json.dumps(self.package));data['report_month']='2026-10'
        with self.assertRaises((ValueError,KeyError,StopIteration)):
            update_source_template(self.template,Path(self.temp.name)/'bad.pptx',data,ASSETS/'乘用车月报-数据源映射指标.xlsx')
