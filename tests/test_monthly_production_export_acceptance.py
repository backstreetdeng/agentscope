"""Offline acceptance of the exact-run webpage PPT export contract.

The fixtures are already collected August data and accepted page narratives.
Only the completed-job store and filesystem root are substituted. The actual
app export function, release validator and deterministic PPT writer execute.
This is NOT a claim that a new live webpage/background/model task has run.
"""
from __future__ import annotations

from contextlib import ExitStack
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
import shutil
import tempfile
import unittest
from unittest.mock import Mock, patch
from urllib.parse import unquote
from zipfile import ZipFile

from openpyxl import load_workbook
from pptx import Presentation

import app
from utils.monthly_ppt_layout import NARRATIVE_BODY_NAME
from utils.monthly_report_narrative import NARRATIVE_RULE_VERSION, _ppt_narrative_text


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "skills/market-analysis-skill/passenger-car-monthly-report/assets"
PACKAGE = ROOT / "outputs/monthly-report-data/data-source-package-20261008-081855.json"
NARRATIVE = ROOT / "outputs/monthly-report-data/business-contract-review-20261008/narrative_evidence.json"
TEMPLATE = ASSETS / "乘用车月报PPT兼容模板.pptx"
MAPPING = ASSETS / "乘用车月报-数据源映射指标.xlsx"
EXAMPLES = ASSETS / "expert-headline-examples.json"
PAGES = {3, 4, 5, *range(7, 20)}


def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


class CompletedJobStore:
    """No disk writes, worker creation, recovery or production job mutation."""

    def __init__(self, run_id: str, package: Path, narrative: Path):
        self.job = {"run_id": run_id, "status": "completed", "checkpoints": {
            "data_collect": {"snapshot_path": str(package)},
            "narrative": {"path": str(narrative)},
        }}
        self.saved = []

    def get(self, run_id):
        return deepcopy(self.job) if run_id == self.job["run_id"] else None

    def save_output_checkpoint(self, run_id, stage, output):
        if run_id != self.job["run_id"]:
            raise AssertionError("Unexpected foreign run checkpoint")
        self.saved.append((run_id, stage, deepcopy(output)))
        self.job["checkpoints"][stage] = deepcopy(output)


@unittest.skipUnless(all(path.is_file() for path in
                        (PACKAGE, NARRATIVE, TEMPLATE, MAPPING, EXAMPLES)),
                     "Complete private August acceptance fixtures not installed")
class ProductionExportAcceptanceTest(unittest.TestCase):
    def setUp(self):
        stack = self.stack = ExitStack()
        self.addCleanup(stack.close)
        self.root = Path(stack.enter_context(tempfile.TemporaryDirectory(
            prefix="monthly-export-acceptance-"))).resolve()
        self.package_dir = self.root / "outputs/monthly-report-data"
        self.package_dir.mkdir(parents=True)
        self.package_path = self.package_dir / "bound-data.json"
        self.narrative_path = self.package_dir / "bound-narrative.json"
        shutil.copy2(PACKAGE, self.package_path)
        shutil.copy2(NARRATIVE, self.narrative_path)
        self.narrative = json.loads(self.narrative_path.read_text(encoding="utf-8"))
        self.run_id = self.narrative["run_id"]
        # A stale private artifact must fail visibly, never be relabelled v5.
        self.assertEqual(self.narrative["narrative_rule_version"], NARRATIVE_RULE_VERSION)
        asset_dir = self.root / "assets"
        asset_dir.mkdir()
        copied = {}
        for source in (TEMPLATE, MAPPING, EXAMPLES):
            copied[source] = asset_dir / source.name
            shutil.copy2(source, copied[source])
        self.template = copied[TEMPLATE]
        # The production function fingerprints these real renderer sources.
        (self.root / "utils").mkdir()
        for filename in ("monthly_report_narrative.py", "monthly_ppt_layout.py",
                         "monthly_source_pptx.py", "monthly_source_rankings.py",
                         "monthly_page_facts.py", "monthly_page_narratives.py",
                         "monthly_segment_facts.py"):
            shutil.copy2(ROOT / "utils" / filename, self.root / "utils" / filename)
        self.output_dir = self.root / "outputs/ppt"
        self.jobs = CompletedJobStore(self.run_id, self.package_path, self.narrative_path)
        overrides = {
            "PROJECT_DIR": self.root,
            "PPT_OUTPUT_DIR": self.output_dir,
            "PASSENGER_MONTHLY_PACKAGE_DIR": self.package_dir,
            "PASSENGER_MONTHLY_TEMPLATE_PATH": self.template,
            "PASSENGER_MONTHLY_MAPPING_PATH": copied[MAPPING],
            "PASSENGER_MONTHLY_EXAMPLES_PATH": copied[EXAMPLES],
            "MONTHLY_REPORT_JOBS": self.jobs,
            # Deliberately unusable foreign state: export must ignore it.
            "LATEST_AGENT_RUN": {"run_id": "unbound-history",
                                 "skill": {"source_package": {
                                     "snapshot_path": "MUST-NOT-READ-HISTORY.json"}}},
        }
        for name, value in overrides.items():
            stack.enter_context(patch.object(app, name, value))
        self.forbidden = {}
        for name in ("MONTHLY_REPORT_BROWSER", "generate_monthly_narratives",
                     "_latest_passenger_monthly_package_path", "store_project_run",
                     "store_latest_agent_run"):
            guard = Mock(name=name, side_effect=AssertionError(
                f"Export must not invoke {name}"))
            self.forbidden[name] = stack.enter_context(patch.object(app, name, guard))
        stack.enter_context(patch.object(app, "record_tool_call"))
        # Spies preserve actual validator/writer execution, not fake success.
        self.validator = stack.enter_context(patch.object(
            app, "validate_narrative_package", wraps=app.validate_narrative_package))
        self.writer = stack.enter_context(patch.object(
            app, "apply_narratives_to_ppt", wraps=app.apply_narratives_to_ppt))

    def tearDown(self):
        for guard in self.forbidden.values():
            guard.assert_not_called()
            self.assertEqual(guard.mock_calls, [], "Export touched a forbidden method/property chain")

    def export(self, run_id=None):
        return app.generate_passenger_monthly_pptx(
            run_id or self.run_id, "离线验收：同一运行编号月报导出")

    def resolve_download(self, url):
        self.assertTrue(url.startswith("/outputs/ppt/"), url)
        target = (self.root / unquote(url).lstrip("/")).resolve()
        target.relative_to(self.output_dir)
        self.assertTrue(target.is_file(), url)
        return target

    def write_narrative(self, mutate):
        narrative = deepcopy(self.narrative)
        mutate(narrative)
        self.narrative_path.write_text(json.dumps(narrative, ensure_ascii=False), encoding="utf-8")

    def assert_rejected_before_writing(self, expression=None):
        if expression:
            with self.assertRaisesRegex((RuntimeError, ValueError), expression):
                self.export()
        else:
            with self.assertRaises((RuntimeError, ValueError)):
                self.export()
        self.writer.assert_not_called()
        self.assertEqual(self.jobs.saved, [])
        self.assertFalse(self.output_dir.exists())

    def test_real_export_is_downloadable_same_run_native_and_idempotent(self):
        result = self.export()
        self.assertTrue(result["success"])
        self.assertEqual(result["run_id"], self.run_id)
        self.assertEqual((result["cards_usable"], result["cards_expected"]), (45, 45))
        self.assertEqual(result["headline_count"], 16)
        self.validator.assert_called_once()
        self.writer.assert_called_once()
        output = self.resolve_download(result["pptx_url"])
        manifest = json.loads(self.resolve_download(result["manifest_url"]).read_text(encoding="utf-8"))
        self.assertEqual(manifest["run_id"], self.run_id)
        self.assertEqual(set(manifest["slides_written"]), PAGES)
        self.assertEqual(manifest["source_package"], str(self.package_path))
        self.assertEqual(manifest["narrative_package"], str(self.narrative_path))
        self.assertEqual(manifest["source_package_sha256"], digest(self.package_path))
        self.assertEqual(manifest["narrative_package_sha256"], digest(self.narrative_path))
        self.assertEqual(manifest["template_sha256"], digest(self.template))
        self.assertEqual(manifest["output_sha256"], digest(output))
        self.assertTrue(manifest["legacy_excel_disabled"])
        self.assertEqual(manifest["release_mode"], "run_scoped_monthly_data_and_narrative_update")
        with ZipFile(output) as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual(sum(bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", p))
                                 for p in archive.namelist()), 21)
            self.assertEqual(sum(bool(re.fullmatch(r"ppt/charts/chart\d+\.xml", p))
                                 for p in archive.namelist()), 39)
        deck = Presentation(output)
        self.assertEqual(len(deck.slides), 21)
        self.assertEqual(sum(s.has_chart for slide in deck.slides for s in slide.shapes), 39)
        for item in self.narrative["slides"]:
            slide = deck.slides[item["slide"] - 1]
            expected_headline, expected_body = _ppt_narrative_text(item, "2026-08")
            texts = [shape.text for shape in slide.shapes if shape.has_text_frame]
            self.assertIn(expected_headline, texts)
            if item["slide"] == 5:
                self.assertEqual(expected_body, '')
                self.assertFalse(any(s.name == NARRATIVE_BODY_NAME for s in slide.shapes))
            else:
                self.assertIn(expected_body, texts)
        self.assert_annual_reference_and_inventory_preserved(deck)
        cached = self.export()
        self.assertEqual(cached["run_id"], self.run_id)
        self.assertEqual(cached["pptx_url"], result["pptx_url"])
        self.assertEqual(cached["manifest_url"], result["manifest_url"])
        self.writer.assert_called_once()
        self.assertEqual(len(self.jobs.saved), 1)
        self.assertEqual(self.jobs.saved[0][:2], (self.run_id, "ppt_update"))

    def assert_annual_reference_and_inventory_preserved(self, deck):
        original=Presentation(self.template)
        for page in (7,8,12,13,15):
            slide=deck.slides[page-1]
            self.assertIn('2025.1-12',[s.text for s in slide.shapes if s.has_text_frame])
            before=[s.chart for s in original.slides[page-1].shapes if s.has_chart][2]
            after=[s.chart for s in slide.shapes if s.has_chart][2]
            self.assertEqual(before.part.blob,after.part.blob)
            self.assertEqual(before.part.chart_workbook.xlsx_part.blob,after.part.chart_workbook.xlsx_part.blob)
        before=[s for s in original.slides[4].shapes if s.has_chart]
        after=[s for s in deck.slides[4].shapes if s.has_chart]
        self.assertEqual([(s.left,s.top,s.width,s.height) for s in before],
                         [(s.left,s.top,s.width,s.height) for s in after])
        self.assertIn('2025-2026年乘用车库存量月度走势',
                      [s.text for s in deck.slides[4].shapes if s.has_text_frame])
        self.assertIn('2025.1-8',[s.text for s in deck.slides[16].shapes if s.has_text_frame])

    def configure_missing_reused_path(self):
        self.jobs.job["report_month"] = "2026-08"
        self.jobs.job["checkpoints"]["data_collect"].update(snapshot_path="", report_month="2026-08")
        self.jobs.job["checkpoints"]["preflight"] = {
            "snapshot_reused": True, "report_month": "2026-08",
            "data_package_path": str(self.package_path),
        }

    def test_missing_reuse_path_recovers_exact_run_and_exports_idempotently(self):
        self.configure_missing_reused_path()
        original_hashes = (digest(self.package_path), digest(self.narrative_path))
        result = self.export()
        self.resolve_download(result["pptx_url"])
        recovered = self.jobs.job["checkpoints"]["data_collect"]
        self.assertEqual(recovered["snapshot_path"], str(self.package_path))
        self.assertEqual(recovered["source_package_sha256"], original_hashes[0])
        self.assertEqual(recovered["path_recovery"]["from"], "preflight.data_package_path")
        self.assertEqual(recovered["path_recovery"]["run_id"], self.run_id)
        self.assertEqual((digest(self.package_path), digest(self.narrative_path)), original_hashes)
        self.assertEqual(self.export()["pptx_url"], result["pptx_url"])
        self.writer.assert_called_once()
        self.validator.assert_called_once()
        self.assertEqual([item[1] for item in self.jobs.saved], ["data_collect", "ppt_update"])

    def test_explicit_bad_paths_never_fall_back_to_preflight(self):
        self.configure_missing_reused_path()
        outside = self.root / "outside.json"
        shutil.copy2(self.package_path, outside)
        for value in (str(outside), str(self.package_dir / "missing.json"), 123):
            with self.subTest(value=value):
                self.jobs.job["checkpoints"]["data_collect"]["snapshot_path"] = value
                self.assert_rejected_before_writing()

    def test_missing_path_without_exact_run_reuse_record_is_rejected(self):
        self.jobs.job["checkpoints"]["data_collect"]["snapshot_path"] = None
        self.assert_rejected_before_writing("同轮复用记录")

    def test_recovery_rejects_wrong_locked_month(self):
        self.configure_missing_reused_path()
        self.jobs.job["checkpoints"]["preflight"]["report_month"] = "2026-09"
        self.assert_rejected_before_writing("锁定月份不一致")

    def test_recovery_rejects_duplicate_cards_despite_45_claim(self):
        self.configure_missing_reused_path()
        package = json.loads(self.package_path.read_text(encoding="utf-8"))
        package["cards"][-1] = deepcopy(package["cards"][0])
        self.package_path.write_text(json.dumps(package), encoding="utf-8")
        self.assert_rejected_before_writing("缺失、重复")

    def test_recovery_rejects_changed_previously_fingerprinted_package(self):
        self.configure_missing_reused_path()
        self.jobs.job["checkpoints"]["data_collect"]["source_package_sha256"] = "0" * 64
        self.assert_rejected_before_writing("内容与本轮检查点记录不一致")

    def test_recovery_rejects_foreign_narrative_before_saving_any_checkpoint(self):
        self.configure_missing_reused_path()
        self.write_narrative(lambda n: n.update(run_id="another-run"))
        self.assert_rejected_before_writing("本轮任务不一致")

    def test_recovery_rejects_fact_gate_failure_before_saving_any_checkpoint(self):
        self.configure_missing_reused_path()
        self.write_narrative(lambda n: n.update(slides=n["slides"][:-1]))
        self.assert_rejected_before_writing("不完整")
        self.validator.assert_called_once()

    def test_unknown_run_does_not_fall_back_to_history(self):
        with self.assertRaisesRegex(RuntimeError, "未找到对应"):
            self.export("unbound-history")
        self.writer.assert_not_called()
        self.assertEqual(self.jobs.saved, [])

    def test_uncompleted_job_cannot_export(self):
        self.jobs.job["status"] = "running"
        self.assert_rejected_before_writing("尚未完成")

    def test_foreign_narrative_run_cannot_export(self):
        self.write_narrative(lambda n: n.update(run_id="unbound-history"))
        self.assert_rejected_before_writing("本轮任务不一致")

    def test_stale_completed_rule_requires_new_narrative_not_browser_collection(self):
        self.write_narrative(lambda n: n.update(narrative_rule_version="entity-bound-business-v4"))
        self.assert_rejected_before_writing("请新建分析生成新版主旨句")
        self.validator.assert_called_once()

    def test_wrong_new_energy_share_is_rejected_by_real_entity_gate(self):
        def corrupt(narrative):
            page = next(item for item in narrative["slides"] if item["slide"] == 17)
            claim = next(item for item in page["body_claims"] if "63.2%" in item["claim"])
            claim["claim"] = claim["claim"].replace("63.2%", "36.8%")
        self.write_narrative(corrupt)
        self.assert_rejected_before_writing()
        self.validator.assert_called_once()

    def test_missing_page_is_rejected_by_real_completeness_gate(self):
        self.write_narrative(lambda n: n.update(slides=[s for s in n["slides"] if s["slide"] != 15]))
        self.assert_rejected_before_writing("不完整")
        self.validator.assert_called_once()


if __name__ == "__main__":
    unittest.main()
