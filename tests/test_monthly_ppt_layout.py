"""Business contracts, including negative tests for the previous wrong rules."""
from pathlib import Path
import unittest
import gc
from pptx import Presentation
from utils.monthly_ppt_layout import (prepare_inventory_roles,validate_business_chart_roles,
    ANNUAL_REFERENCE_LABELS,NARRATIVE_BODY_NAME)

TEMPLATE=Path(__file__).resolve().parents[1]/'skills/market-analysis-skill/passenger-car-monthly-report/assets/乘用车月报PPT兼容模板.pptx'

@unittest.skipUnless(TEMPLATE.exists(),'Template unavailable')
class BusinessContractTest(unittest.TestCase):
    def setUp(self):
        gc.collect()
    def test_inventory_is_title_only_readonly_and_full_height(self):
        slide=Presentation(TEMPLATE).slides[4]
        before=slide.part.blob
        title,body=prepare_inventory_roles(slide)
        self.assertEqual(title.shape_id,1048746)
        self.assertIsNone(body)
        self.assertEqual(before,slide.part.blob)
        self.assertFalse(any(s.name==NARRATIVE_BODY_NAME for s in slide.shapes))

    def test_annual_and_same_period_are_distinct_and_readonly(self):
        deck=Presentation(TEMPLATE)
        before=[s.part.blob for s in deck.slides]
        result=validate_business_chart_roles(deck,'2026-08')
        self.assertEqual([r['slide'] for r in result],[7,8,12,13,15,17])
        self.assertEqual([r['reference_period'] for r in result],['2025.1-12']*5+['2025.1-8'])
        self.assertEqual(before,[s.part.blob for s in deck.slides])

    def test_changing_annual_label_to_ytd_is_rejected_for_each_page(self):
        deck=Presentation(TEMPLATE)
        for page,identifier in ANNUAL_REFERENCE_LABELS.items():
            with self.subTest(page=page):
                shape=next(s for s in deck.slides[page-1].shapes if s.shape_id==identifier)
                shape._element.xpath('.//a:t')[0].text='2025.1-8'
                with self.assertRaisesRegex(ValueError,'全年结构'):
                    validate_business_chart_roles(deck,'2026-08')
                shape._element.xpath('.//a:t')[0].text='2025.1-12'

    def test_same_period_page17_must_not_be_changed_to_annual(self):
        deck=Presentation(TEMPLATE)
        shape=next(s for s in deck.slides[16].shapes if s.has_text_frame and s.text=='2025.1-8')
        shape._element.xpath('.//a:t')[0].text='2025.1-12'
        with self.assertRaisesRegex(ValueError,'第17页'):
            validate_business_chart_roles(deck,'2026-08')

    def test_inventory_added_body_or_shrunk_charts_fail_closed(self):
        for corruption in ('body','geometry','caption'):
            with self.subTest(corruption=corruption):
                slide=Presentation(TEMPLATE).slides[4]
                if corruption=='body':
                    slide.shapes.add_textbox(1,1,1,1).name=NARRATIVE_BODY_NAME
                elif corruption=='geometry':
                    next(s for s in slide.shapes if s.has_chart).height-=100
                else:
                    next(s for s in slide.shapes if s.shape_id==1048747).text='乘用车库存量月度走势'
                with self.assertRaises(ValueError):prepare_inventory_roles(slide)

    def test_cache_workbook_disagreement_is_rejected(self):
        deck=Presentation(TEMPLATE)
        chart=[s.chart for s in deck.slides[6].shapes if s.has_chart][2]
        chart.series[0]._ser.xpath('.//c:val/c:numRef/c:numCache/c:pt[@idx="2"]/c:v')[0].text='0.8'
        with self.assertRaisesRegex(ValueError,'cache'):
            validate_business_chart_roles(deck,'2026-08')

if __name__=='__main__':unittest.main()
