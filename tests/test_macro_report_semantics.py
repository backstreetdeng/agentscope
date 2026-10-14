# -*- coding: utf-8 -*-
"""Independent semantic regression tests; fixtures here are SYNTHETIC, not data.

ZIP or Office-open success is not business acceptance.  These checks read the
original/native XML and embedded spreadsheets rather than believing model prose
or a generator's own 'passed' flag.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from datetime import date, datetime, timedelta
import importlib.util
import json
import posixpath
import sys
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.datetime import to_excel

PROJECT = Path(__file__).resolve().parents[1]
SKILL = PROJECT / "skills/macro-monthly-report"
REFERENCE = SKILL / "assets/reference.pptx"
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/package/2006/relationships"}


def load_engine():
    path = SKILL / "scripts/macro_engine.py"
    if str(path.parent) not in sys.path:
        sys.path.insert(0, str(path.parent))
    spec = importlib.util.spec_from_file_location("tested_macro_engine", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def load_contract():
    path = SKILL / "scripts/macro_contract.py"
    spec = importlib.util.spec_from_file_location("tested_macro_contract", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def chart_xml(deck: Path, chart_number: int):
    with ZipFile(deck) as archive:
        return etree.fromstring(archive.read(f"ppt/charts/chart{chart_number}.xml"))


def named_series(chart):
    return {"".join(ser.xpath("c:tx//c:v/text() | c:tx/c:v/text()", namespaces=NS)): ser
            for ser in chart.findall(".//c:ser", NS)}


def values(series):
    node = series.find("c:val/c:numRef/c:numCache", NS)
    if node is None:
        node = series.find("c:val/c:numLit", NS)
    if node is None:
        return []
    count = int(node.find("c:ptCount", NS).get("val"))
    result = [None] * count
    for point in node.findall("c:pt", NS):
        result[int(point.get("idx"))] = float(point.findtext("c:v", namespaces=NS))
    return result


def categories(series):
    multi = series.find("c:cat/c:multiLvlStrRef/c:multiLvlStrCache",NS)
    if multi is not None:
        return multi.xpath("c:lvl[1]/c:pt/c:v/text()",namespaces=NS)
    return series.xpath("c:cat//c:pt/c:v/text()", namespaces=NS)


def assert_cache_workbook(test, deck: Path):
    """Cross-check formulas, point lengths, true blanks, numbers and series name."""
    with ZipFile(deck) as archive:
        parts = sorted(name for name in archive.namelist()
                       if name.startswith("ppt/charts/chart") and name.endswith(".xml"))
        for part in parts:
            chart = etree.fromstring(archive.read(part))
            relname = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part)+".rels")
            rels = etree.fromstring(archive.read(relname))
            targets = [rel.get("Target") for rel in rels if rel.get("Type").endswith("/package")]
            test.assertEqual(len(targets), 1, part + " requires one embedded workbook")
            workbook_path = posixpath.normpath(posixpath.join(posixpath.dirname(part), targets[0]))
            workbook = load_workbook(BytesIO(archive.read(workbook_path)), data_only=True)
            for series in chart.findall(".//c:ser", NS):
                for ref in series.findall("c:cat/c:numRef",NS)+series.findall("c:cat/c:strRef",NS):
                    cache=ref.find("c:numCache",NS)
                    numeric=cache is not None
                    if cache is None:cache=ref.find("c:strCache",NS)
                    if cache is None:continue
                    count=int(cache.find("c:ptCount",NS).get("val"))
                    formula=ref.findtext("c:f",namespaces=NS)
                    sheet,cells=formula.rsplit("!",1)
                    mincol,minrow,maxcol,maxrow=range_boundaries(cells.replace("$",""))
                    test.assertEqual((maxcol-mincol+1)*(maxrow-minrow+1),count,part+" category formula/cache length")
                    raw=[workbook[sheet.strip("'")].cell(row,col).value for row in range(minrow,maxrow+1) for col in range(mincol,maxcol+1)]
                    byidx={int(p.get("idx")):p.findtext("c:v",namespaces=NS) for p in cache.findall("c:pt",NS)}
                    for i,actual in enumerate(raw):
                        cached=byidx.get(i)
                        if actual is None:test.assertIsNone(cached)
                        elif numeric:
                            actual=to_excel(actual) if isinstance(actual,(date,datetime)) else float(actual)
                            test.assertAlmostEqual(actual,float(cached),places=7,msg=part+" calendar category mismatch")
                        else:test.assertEqual(str(actual),cached,part+" fixed category mismatch")
                for ref in series.findall("c:val/c:numRef", NS):
                    formula = ref.findtext("c:f", namespaces=NS)
                    test.assertTrue(formula, part + " has missing formula")
                    sheet, cells = formula.rsplit("!", 1)
                    mincol, minrow, maxcol, maxrow = range_boundaries(cells.replace("$", ""))
                    raw = [workbook[sheet.strip("'")].cell(row, col).value
                           for row in range(minrow,maxrow+1) for col in range(mincol,maxcol+1)]
                    cached = values(series)
                    test.assertEqual(len(raw), len(cached), part + " range/cache length mismatch")
                    for index, (actual, cached_value) in enumerate(zip(raw, cached)):
                        if actual is None:
                            test.assertIsNone(cached_value, f"{part} point{index}: actual blank became zero/forecast")
                        else:
                            test.assertIsNotNone(cached_value, f"{part} point{index}: workbook has value missing from cache")
                            test.assertAlmostEqual(float(actual), cached_value, places=7, msg=part)


def synthetic_observations(contract):
    """SYNTHETIC structural fixture; must never be accepted as publisher proof."""
    result = {}
    for index, request in enumerate(contract["source_requests"]):
        frequency = request["frequency"]
        if frequency == "quarterly":
            periods = [date(2025,9,30),date(2025,12,31),date(2026,3,31),date(2026,6,30)]
        elif frequency == "weekly":
            periods = [date(2026,6,12)+timedelta(days=7*n) for n in range(8)]
        elif frequency == "daily":
            periods = []
            cursor = date(2026,7,20)
            while len(periods) < 10:
                if cursor.weekday() < 5: periods.append(cursor)
                cursor += timedelta(days=1)
        elif frequency == "announcement_monthly":
            periods = [date(2026,5,20),date(2026,6,22),date(2026,7,20)]
        else:
            periods = [date(2026,5,31),date(2026,6,30),date(2026,7,31)]
        base = (40+index%9 if request["unit"] == "指数" else
                20000+index*100 if request["unit"] in {"亿元","辆"} else
                70+index%10 if request["unit"] == "美元/桶" else 5+index%7)
        result[request["id"]] = [{"request_id":request["id"],"receipt_id":"SYNTHETIC-NOT-PUBLISHER",
            "definition":request["definition"],"unit":request["unit"],
            "source_url":"https://synthetic.invalid/NOT-BUSINESS-EVIDENCE",
            "period":(f"{d.year}-Q{(d.month-1)//3+1}" if frequency=="quarterly" else
                      d.isoformat() if frequency in {"daily","weekly","announcement_monthly"} else d.strftime("%Y-%m")),
            "published_at":"2026-07-31","date":d,"value":float(base+n),"synthetic":True}
            for n,d in enumerate(periods)]
    return result


class OriginalMacroAssetContractTest(unittest.TestCase):
    """Audit the original immutable asset's intentional semantic structure."""

    def test_native_chart_workbooks_are_internally_consistent(self):
        assert_cache_workbook(self, REFERENCE)

    def test_fixed_categories_are_not_ambiguous_numeric_dates(self):
        for number, first, last, count in [(10,"生产","生产经营活动预期",12),
                                          (18,"零售总额","石油及制品类",13),
                                          (25,"食品","其他用品服务",7),
                                          (20,"第1周","第52周",52),
                                          (28,"1月","12月",12),
                                          (30,"1月","12月",12),
                                          (34,"1月","12月",12)]:
            ser = next(iter(named_series(chart_xml(REFERENCE,number)).values()))
            labels = categories(ser)
            self.assertEqual((labels[0],labels[-1],len(labels)),(first,last,count))

    def test_quarterly_and_constant_series_are_explicit(self):
        self.assertEqual(len(values(named_series(chart_xml(REFERENCE,7))["GDP:不变价:当季同比"])),20)
        self.assertEqual(len(values(named_series(chart_xml(REFERENCE,8))["GDP:累计同比"])),17)
        for number, name in [(2,"扩张收缩分界线"),(9,"荣枯线")]:
            self.assertTrue(all(value == 50 for value in values(named_series(chart_xml(REFERENCE,number))[name])))

    def test_original_combo_workbook_column_order_is_not_series_index(self):
        chart = chart_xml(REFERENCE,15)
        bar = chart.find(".//c:barChart",NS)
        line = chart.find(".//c:lineChart",NS)
        self.assertEqual(bar.xpath("c:ser/c:idx/@val",namespaces=NS),["0","2"])
        self.assertEqual(line.xpath("c:ser/c:idx/@val",namespaces=NS),["1","3"])
        self.assertEqual(named_series(chart)["民间固定资产投资累计值"].findtext("c:val/c:numRef/c:f",namespaces=NS),"Sheet1!$D$2:$D$32")

    def test_overview_is_visible_ole_not_native_table(self):
        with ZipFile(REFERENCE) as archive:
            slide = etree.fromstring(archive.read("ppt/slides/slide7.xml"))
            frames = slide.findall(".//p:graphicFrame",NS)
            visible = [frame for frame in frames if frame.find("p:nvGraphicFramePr/p:cNvPr",NS).get("id")=="24"]
            self.assertEqual(len(visible),1)
            self.assertTrue(visible[0].xpath(".//p:oleObj[@progId='Excel.Sheet.12']",namespaces=NS))
            self.assertFalse(visible[0].xpath(".//a:tbl",namespaces=NS))


class ReviewedMacroTypingTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.contract = load_contract()

    def test_all_37_native_chart_ids_have_explicit_rules(self):
        expected = {
            "s4-shape5","s4-shape10","s8-shape20","s8-shape21","s8-shape25","s8-shape11",
            "s10-shape7","s10-shape28","s11-shape12","s11-shape15","s12-shape13","s12-shape16",
            "s13-shape10","s13-shape17","s14-shape15","s14-shape11","s15-shape9","s16-shape8",
            "s17-shape15","s17-shape20","s18-shape17","s18-shape21","s19-shape13","s21-shape11",
            "s21-shape12","s22-shape9","s22-shape10","s24-shape9","s24-shape17","s25-shape9",
            "s25-shape10","s26-shape11","s26-shape12","s28-shape11","s28-shape17","s30-shape13","s30-shape18",
        }
        self.assertEqual(set(self.contract.RULES),expected)

    def test_nonmonthly_chart_semantics_are_not_guessed(self):
        contract = self.contract.RULES
        self.assertEqual(contract["s10-shape7"][1],"quarterly")
        self.assertEqual(contract["s10-shape28"][1],"quarterly")
        for key in ["s11-shape15","s16-shape8","s24-shape17"]:
            self.assertEqual(contract[key][0],"snapshot")
        self.assertEqual(contract["s21-shape12"][0],"category_metrics")
        self.assertEqual(contract["s17-shape20"][:2],("seasonal","weekly"))
        self.assertEqual(contract["s24-shape9"][:2],("seasonal","monthly"))
        self.assertEqual(contract["s25-shape9"][:2],("seasonal","monthly"))
        self.assertEqual(contract["s28-shape11"][0],"seasonal_current")
        self.assertEqual(contract["s8-shape11"][1],"daily")
        self.assertEqual(contract["s30-shape18"][1],"announcement_monthly")

    def test_raw_units_are_not_reconverted_display_units(self):
        for key in ["s14-shape15","s15-shape9","s24-shape9","s25-shape9","s25-shape10","s30-shape13"]:
            self.assertEqual(self.contract.RULES[key][2],"亿元")
        self.assertEqual(self.contract.RULES["s28-shape11"][2],"辆")
        self.assertEqual(self.contract.RULES["s17-shape20"][2],"万平方米")
        self.assertEqual(self.contract.series_unit("s10-shape28","GDP:累计同比","百分点"),"%")
        self.assertEqual(self.contract.series_unit("s12-shape16","企业景气指数:预期指数","%"),"指数")

    def test_period_normalization_retains_typed_frequency(self):
        parse = self.contract.parse_period
        self.assertEqual(parse("2026-Q3","quarterly").isoformat(),"2026-09-30")
        self.assertEqual(parse("2022-12-01","quarterly").isoformat(),"2022-12-31")
        self.assertEqual(parse("2026-08","monthly").isoformat(),"2026-08-31")
        self.assertEqual(parse("46234","monthly").isoformat(),"2026-07-31")
        self.assertEqual(parse("2026-06-22","announcement_monthly").isoformat(),"2026-06-22")
        self.assertEqual(parse("2026-07-24","weekly").isoformat(),"2026-07-24")
        for label in ("生产","第52周","8月"):
            with self.assertRaises(ValueError):parse(label)

    def test_loan_repeated_categories_have_department_specific_request_ids(self):
        original = named_series(chart_xml(REFERENCE,29))
        request = self.contract.request_contract({"id":"s24-shape17","categories":categories(next(iter(original.values()))),"series":[]})
        self.assertEqual([r["definition"] for r in request],["住户短期贷款","住户中长期贷款","企事业单位短期贷款","企事业单位中长期贷款","票据融资"])
        self.assertEqual(len(set(r["id"] for r in request)),5)

    def test_constant_pmi_lines_have_no_external_data_request(self):
        ser = [{"name":"中国","values":[49.2,50.1]}, {"name":"扩张收缩分界线","values":[50,50]}]
        requests = self.contract.request_contract({"id":"s4-shape10","categories":["2026-07","2026-08"],"series":ser})
        self.assertEqual([r["series_name"] for r in requests],["中国"])

    def test_actual_pmi_equal_to_50_is_not_a_static_reference_line(self):
        self.assertFalse(self.contract.static_series("中国",[50,50]))
        self.assertFalse(self.contract.static_series("制造业PMI",[50,50,50]))
        self.assertTrue(self.contract.static_series("扩张收缩分界线",[50,50]))

    def test_publisher_country_and_measure_scopes_are_specific(self):
        official = self.contract.official_labels
        us = official("s4-shape10","美国","制造业PMI：美国")
        cn = official("s4-shape10","中国","制造业PMI：中国")
        self.assertEqual(us[2],["ismworld.org"])
        self.assertEqual(cn[2],["stats.gov.cn"])
        self.assertIn(["制造业","Manufacturing"],us[1])
        self.assertIn(["ISM","Institute for Supply Management"],us[1])
        short = official("s24-shape17","snapshot","住户短期贷款",0)
        enterprise = official("s24-shape17","snapshot","企事业单位短期贷款",2)
        self.assertNotEqual(short[1],enterprise[1])
        self.assertNotIn("PMI",official("s13-shape17","当月同比","工业增加值同比：当月同比")[0])

    def test_overview_requests_have_precise_history_scope_not_future_query_placeholders(self):
        engine=load_engine();contract=engine.inventory(REFERENCE)
        for request in contract["source_requests"]:
            self.assertNotIn("{year}",request["query"])
            self.assertNotIn("{month}",request["query"])
            self.assertTrue(request.get("history_count",0)>0)
            self.assertTrue(request.get("definition_aliases"))
        for request in contract["source_requests"][-3:]:
            self.assertEqual(request["frequency"],"monthly")
            self.assertEqual(request["history_count"],3)
            self.assertIn(["同比"],request["context_groups"])
            self.assertEqual(request["template_latest_period"],"2026-07")


class TypedEngineBehaviorTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine=load_engine()
        cls.engine_sha256=cls.engine.digest(Path(cls.engine.__file__))
        cls.contract=cls.engine.inventory(REFERENCE)
        cls.charts={c["id"]:c for s in cls.contract["slides"] for c in s["charts"]}
        cls.observations=synthetic_observations(cls.contract)

    def test_normalization_sorts_quarters_and_never_zero_fills_missing_cells(self):
        normal=self.engine.normalize_series(["2026-06","2026-05","2026-03","2025-Q4"],[6,99,3,4],"quarterly")
        self.assertEqual(normal,[(date(2025,12,31),4),(date(2026,3,31),3),(date(2026,6,30),6)])
        normal=self.engine.normalize_series(["2026-07","2026-06","2026-05"],[None,0,float("nan")],"monthly")
        self.assertEqual(normal,[(date(2026,6,30),0)])
        with self.assertRaisesRegex(ValueError,"duplicate_period_conflict"):
            self.engine.normalize_series(["2026-07","2026-07"],[1,2],"monthly")

    def test_forecast_uses_verified_training_only_and_retains_negative_net_flows(self):
        points=[(date(2026,5,31),-120),(date(2026,6,30),-130),(date(2026,7,31),-125)]
        result=self.engine.forecast_series(points,date(2026,9,30),unit="亿元")
        self.assertEqual(result["training_values"],[-120,-130,-125])
        self.assertTrue(all(p["value"]<0 and p["status"]=="forecast" for p in result["values"]))
        self.assertEqual([p["period"] for p in result["values"]],["2026-08-31","2026-09-30"])
        with self.assertRaisesRegex(ValueError,"forecast_verified_history_insufficient"):
            self.engine.forecast_series(points[:2],date(2026,9,30))
        with self.assertRaisesRegex(ValueError,"forecast_horizon_too_long"):
            self.engine.forecast_series(points,date(2027,2,28))

    def test_quarterly_axes_and_mixed_frequency_do_not_become_monthly_gdp(self):
        for cid in ("s10-shape7","s10-shape28"):
            c=self.engine.rebuild_chart(self.charts[cid],self.observations,"2026-09",date(2026,10,3))
            dates=[self.engine.parse_period(v,"daily") for v in c["categories"]]
            self.assertTrue(all(d.month in {3,6,9,12} for d in dates))
            self.assertIn(date(2026,9,30),dates)
            self.assertEqual(len(c["categories"]),len(self.charts[cid]["categories"]))
            for s in c["series"]:
                ix=dates.index(date(2026,9,30))
                self.assertIsNone(s["actual_values"][ix])
                self.assertIsNotNone(s["forecast_values"][ix])
        mixed=self.engine.rebuild_chart(self.charts["s12-shape16"],self.observations,"2026-09",date(2026,10,3))
        dates=[self.engine.parse_period(v,"daily") for v in mixed["categories"]]
        enterprise=next(s for s in mixed["series"] if "景气" in s["name"])
        for i,d in enumerate(dates):
            if d.month not in {3,6,9,12}:
                self.assertIsNone(enterprise["actual_values"][i]);self.assertIsNone(enterprise["forecast_values"][i])

    def test_crosssections_seasonal_calendars_and_policy_dates_are_preserved(self):
        for cid in ("s11-shape15","s16-shape8","s21-shape12","s24-shape17","s17-shape20","s24-shape9","s25-shape9","s28-shape11"):
            c=self.engine.rebuild_chart(self.charts[cid],self.observations,"2026-09",date(2026,10,3))
            self.assertEqual(c["categories"],self.charts[cid]["categories"],cid)
        loans=self.engine.rebuild_chart(self.charts["s24-shape9"],self.observations,"2026-09",date(2026,10,3))
        current=next(s for s in loans["series"] if str(s["name"])=="2026")
        self.assertTrue(all(v is None for v in current["actual_values"][7:]))
        self.assertTrue(all(v is not None for v in current["forecast_values"][7:9]))
        for old in loans["series"]:
            if old["name"]!="2026": self.assertTrue(all(v is None for v in old["forecast_values"]))
        lpr=self.engine.rebuild_chart(self.charts["s30-shape18"],self.observations,"2026-09",date(2026,10,3))
        self.assertIn(date(2026,6,22),[self.engine.parse_period(v,"daily") for v in lpr["categories"]])
        self.assertTrue(all(p["assumption"].startswith("维持现行政策") for s in lpr["series"] for p in s["forecasts"]))

    def test_constant_reference_lines_move_with_axis_without_source_request(self):
        c=self.engine.rebuild_chart(self.charts["s4-shape10"],self.observations,"2026-09",date(2026,10,3))
        line=next(s for s in c["series"] if "分界" in s["name"])
        self.assertEqual(line["actual_values"],[50]*len(c["categories"]))
        self.assertFalse(line["request_ids"])
        self.assertTrue(all(v is None for v in line["forecast_values"]))

    def test_synthetic_source_package_cannot_pass_formal_business_gate(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError,"synthetic_data_not_business_evidence"):
                self.engine.verified_observations({"synthetic":True,"report_month":"2026-09","cutoff_date":"2026-10-03"},self.contract,tmp)
            with self.assertRaisesRegex(ValueError,"synthetic_data_not_business_evidence"):
                self.engine.build(REFERENCE,"2026-09",{"synthetic":True,"report_month":"2026-09","cutoff_date":"2026-10-03"},Path(tmp)/"not-a-report.pptx",Path(tmp)/"ledger.json",tmp)

    def test_shared_loan_overview_uses_current_year_not_first_archival_series(self):
        charts=[self.engine.rebuild_chart(c,self.observations,"2026-09",date(2026,10,3)) for c in self.charts.values()]
        rows=self.engine.collect_table_facts(charts,self.observations,"2026-09")
        row=next(r for r in rows if r["label"]=="新增人民币贷款")
        self.assertIsNotNone(row["facts"][-1],"current 2026 loan projection must reach shared overview")
        self.assertEqual(row["facts"][-1]["status"],"forecast")
        self.assertEqual(row["facts"][-1]["unit"],"万亿元")
        expected=next(c for c in charts if c["chart_id"]=="s24-shape9")
        series=next(s for s in expected["series"] if s["name"]=="2026")
        self.assertAlmostEqual(row["facts"][-1]["value"],series["forecast_values"][8]/10000)


class FullNativeChartUpdateTest(unittest.TestCase):
    """Build the entire layout with isolated synthetic observations, not sources.

    The source gate is patched ONLY to test transformations. Independent final
    validation replays the true gate and MUST refuse these synthetic artifacts.
    """
    @classmethod
    def setUpClass(cls):
        cls.engine=load_engine()
        cls.engine_sha256=cls.engine.digest(Path(cls.engine.__file__))
        cls.contract=cls.engine.inventory(REFERENCE)
        cls.observations=synthetic_observations(cls.contract)
        cls.directory=PROJECT/".work/macro-repair/audit/semantic-fixture"
        cls.directory.mkdir(parents=True,exist_ok=True)
        cls.output=PROJECT/".work/macro-repair/audit/typed-fixture.pptx"
        cls.ledger=cls.directory/"ledger.json"
        cls.package={"synthetic":True,"data_kind":"synthetic","mode":"forecast",
            "report_month":"2026-09","cutoff_date":"2026-10-03","observations":[],
            "fixture_notice":"SYNTHETIC STRUCTURE TEST — NEVER A BUSINESS REPORT"}
        cls.engine.dump(cls.directory/"working/data-package.json",cls.package)
        with patch.object(cls.engine,"verified_observations",return_value=cls.observations):
            cls.result=cls.engine.build(REFERENCE,"2026-09",cls.package,cls.output,cls.ledger,cls.directory)
        cls.engine.dump(cls.directory/"fixture-manifest.json",{
            "synthetic":True,"purpose":"STRUCTURAL TEST ONLY — NOT BUSINESS EVIDENCE",
            "engine_sha256":cls.engine_sha256,"pptx_sha256":cls.engine.digest(cls.output),
            "snapshot_stable":cls.engine_sha256==cls.engine.digest(Path(cls.engine.__file__)),
            "source_gate":"patched for structure only; true independent source replay must reject"})
        if cls.engine_sha256!=cls.engine.digest(Path(cls.engine.__file__)):
            raise RuntimeError("engine_changed_during_fixture_build; rebuild from one frozen snapshot")
        cls.data=json.loads(cls.ledger.read_text(encoding="utf-8"))

    def test_all_37_native_charts_have_one_workbook_and_consistent_caches(self):
        self.assertEqual(self.result["summary"]["charts"],37)
        assert_cache_workbook(self,self.output)

    def test_combo_plot_roles_and_original_value_columns_are_preserved(self):
        for number in (4,8,15,17,24,27,34,36):
            before=chart_xml(REFERENCE,number);after=chart_xml(self.output,number)
            oldplots=[p for p in before.findall(".//c:plotArea/*",NS) if p.tag.endswith("Chart")]
            newplots=[p for p in after.findall(".//c:plotArea/*",NS) if p.tag.endswith("Chart")]
            self.assertEqual(len(oldplots),len(newplots))
            for plot,newplot in zip(oldplots,newplots):
                self.assertIsNotNone(newplot)
                self.assertEqual(plot.tag,newplot.tag)
                for name,old in named_series(plot).items():
                    new=named_series(newplot)[name]
                    oldformula=old.findtext("c:val/c:numRef/c:f",namespaces=NS)
                    newformula=new.findtext("c:val/c:numRef/c:f",namespaces=NS)
                    oldcol=oldformula.split("!")[-1].split("$")[1]
                    newcol=newformula.split("!")[-1].split("$")[1]
                    self.assertEqual(oldcol,newcol,(number,name))
                    self.assertEqual(old.find("c:idx",NS).get("val"),new.find("c:idx",NS).get("val"))

    def test_every_plot_series_precedes_axis_and_plot_metadata(self):
        for number in range(1,38):
            chart=chart_xml(self.output,number)
            for plot in chart.findall(".//c:plotArea/*",NS):
                if not plot.tag.endswith("Chart"):continue
                past_series=False
                for node in plot:
                    local=etree.QName(node).localname
                    if local in {"dLbls","dropLines","hiLowLines","upDownBars","gapWidth","overlap","serLines","axId","extLst"}:
                        past_series=True
                    if local=="ser":self.assertFalse(past_series,f"chart{number}: forecast c:ser appended after plot metadata")

    def test_multilevel_loan_categories_keep_both_labels_and_sheet_columns(self):
        before=chart_xml(REFERENCE,29);after=chart_xml(self.output,29)
        b=next(iter(named_series(before).values())).find("c:cat/c:multiLvlStrRef",NS)
        for series in named_series(after).values():
            a=series.find("c:cat/c:multiLvlStrRef",NS)
            self.assertEqual(etree.tostring(a),etree.tostring(b))
        self.assertEqual(categories(next(iter(named_series(after).values()))),["短期","中长期","短期","中长期","票据"])

    def test_native_date_axes_have_updated_bounds_and_unambiguous_calendar_formats(self):
        bynumber={int(Path(c["part"]).stem.replace("chart","")):c for slide in self.contract["slides"] for c in slide["charts"]}
        ledgers={c["chart_id"]:c for c in self.data["charts"]}
        for number,contract in bynumber.items():
            if contract["semantics"] not in {"time","mixed"}:continue
            expected=ledgers[contract["id"]]["categories"]
            before=chart_xml(REFERENCE,number);after=chart_xml(self.output,number)
            for axis in after.findall(".//c:dateAx",NS):
                fmt=axis.find("c:numFmt",NS).get("formatCode")
                self.assertTrue(any(token in fmt for token in ("y","m","d")),"calendar axis cannot show raw Excel serials")
                span=max(expected)-min(expected)
                if span>366:self.assertIn("y",fmt,"multi-year monthly/weekly/daily timeline needs a year label")
                if contract["frequency"] in {"monthly","quarterly","announcement_monthly"}:
                    self.assertEqual(axis.find("c:majorTimeUnit",NS).get("val"),"months")
                else:self.assertEqual(axis.find("c:majorTimeUnit",NS).get("val"),"days")
                self.assertEqual(float(axis.find("c:scaling/c:min",NS).get("val")),min(expected))
                self.assertEqual(float(axis.find("c:scaling/c:max",NS).get("val")),max(expected))

    def test_visible_overview_is_native_and_shared_forecast_values_agree(self):
        from pptx import Presentation
        deck=Presentation(self.output)
        self.assertEqual(len(deck.slides),33)
        self.assertEqual(sum(bool(getattr(s,"has_table",False)) for slide in deck.slides for s in self.engine.all_shapes(slide.shapes)),2)
        table=self.engine._shape(deck.slides[31],9).table
        rows=self.data["table_updates"][0]["facts"]
        for index,row in enumerate(list(table.rows)[1:],1):
            shared=next(r for r in rows if r["label"]==row.cells[0].text)
            self.assertEqual(row.cells[2].text,self.engine._formatted(shared["facts"][-1]))
        with ZipFile(self.output) as archive:
            page=etree.fromstring(archive.read("ppt/slides/slide7.xml"))
            self.assertFalse(page.xpath(".//p:oleObj[@progId='Excel.Sheet.12']",namespaces=NS))

    def test_true_source_replay_refuses_structurally_valid_synthetic_deck(self):
        result=self.engine.validate(self.output,self.ledger)
        self.assertFalse(result["passed"])
        self.assertTrue(any(i["code"]=="source_replay_failed" and "synthetic_data_not_business_evidence" in i.get("message","") for i in result["issues"]))
        # Deliberately independent placeholder GDP values cannot satisfy the
        # economic identity. They must fail business acceptance, not be balanced
        # to make an artifact fixture look like a real monthly report.
        permitted={"source_replay_failed","business_incomplete","gdp_decomposition_mismatch",
                   "cumulative_forecast_decrease","cumulative_forecast_cross_year_unmodeled"}
        codes={i["code"] for i in result["issues"]}
        self.assertIn("gdp_decomposition_mismatch",codes)
        self.assertIn("business_incomplete",codes)
        unexpected=[i for i in result["issues"] if i["code"] not in permitted]
        self.assertFalse(unexpected,"structure-only fixture may fail source/economic gates, never silent native/chart/table/text/axis errors: "+str(unexpected[:4]))

    def test_independent_validator_rejects_stale_axis_body_table_and_numeric_cache(self):
        from pptx import Presentation
        deck=Presentation(self.output)
        self.engine._shape(deck.slides[9],6).text="旧模板7月结论，不得作为9月分析"
        self.engine._shape(deck.slides[31],9).table.cell(1,2).text="STALE_999"
        gdp=self.engine._shape(deck.slides[9],7).chart._chartSpace
        for axis in gdp.findall(".//c:dateAx",NS):
            axis.find("c:scaling/c:max",NS).set("val","46234")
        chart=self.engine._shape(deck.slides[13],15).chart._chartSpace
        point=chart.find(".//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v",NS)
        point.text=str(float(point.text)+99)
        target=self.directory/"adversarial-stale.pptx";deck.save(target)
        ledger={**self.data,"pptx_sha256":self.engine.digest(target)}
        path=self.directory/"adversarial-stale-ledger.json";self.engine.dump(path,ledger)
        result=self.engine.validate(target,path);codes={i["code"] for i in result["issues"]}
        self.assertFalse(result["passed"])
        self.assertTrue({"stale_date_axis","text_ledger_mismatch","table_ledger_mismatch","cache_workbook_mismatch","native_values_ledger_mismatch"}.issubset(codes),codes)
        self.assertNotIn("pptx_changed_after_build",codes,"test must reach independent semantics, not hash-only refusal")

    def test_category_swap_is_rejected_even_when_cache_and_workbook_agree(self):
        from pptx import Presentation
        deck=Presentation(self.output)
        chart=self.engine._shape(deck.slides[15],8).chart
        xml=chart._chartSpace
        external=xml.find("c:externalData",NS)
        part=chart.part.related_part(external.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
        workbook=load_workbook(BytesIO(part.blob))
        for ser in xml.xpath(".//c:ser"):
            ref=ser.find("c:cat/c:strRef",NS)
            cache=ref.find("c:strCache",NS)
            cache.find("c:pt/c:v",NS).text="FALSE_CATEGORY_IN_SAME_WORKBOOK"
            sheet,col,start,_,_=self.engine._range(ref.findtext("c:f",namespaces=NS),int(cache.find("c:ptCount",NS).get("val")))
            workbook[sheet].cell(start,col).value="FALSE_CATEGORY_IN_SAME_WORKBOOK"
        stream=BytesIO();workbook.save(stream);part._blob=stream.getvalue()
        target=self.directory/"adversarial-category.pptx";deck.save(target)
        path=self.directory/"adversarial-category-ledger.json"
        self.engine.dump(path,{**self.data,"pptx_sha256":self.engine.digest(target)})
        result=self.engine.validate(target,path)
        codes={i["code"] for i in result["issues"]}
        self.assertTrue(any("categor" in code for code in codes),"cache↔workbook agreement is insufficient when category identity changed: "+str(codes))


if __name__ == "__main__":
    unittest.main()
