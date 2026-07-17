"""月报可编辑 PPTX 的离线回归测试，不调用 MiniMax API。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from utils.monthly_report_pptx import build_monthly_report_pptx


PROJECT_DIR = Path(__file__).resolve().parents[1]
DATASET_PATH = PROJECT_DIR / "skills" / "market-analysis-skill" / "1-5月报 AI验数.xlsx"
TEMPLATE_PATH = (
    PROJECT_DIR
    / "skills"
    / "market-analysis-skill"
    / "月报skill"
    / "2026年5月月报部分.pptx"
)

SAMPLE_REPORT = """# 2026年5月乘用车市场月报

### 2.1 车型类别
轿车和SUV仍是主体，SUV降幅相对较小，建议继续跟踪结构份额变化。

### 2.2 车型级别
B级市场体量最大，A00级承压明显，需观察新品补位节奏。

### 2.3 轿车级别分析
轿车B级和A级承压，TOP车型分化加快。

### 2.4 SUV级别分析
SUV A0与C+级表现韧性较强，新能源产品贡献增量。

### 2.5 MPV级别分析
MPV市场规模较小，A级同比正增长但需持续验证。
"""


@unittest.skipUnless(DATASET_PATH.is_file() and TEMPLATE_PATH.is_file(), "缺少业务验数表或 PPTX 模板")
class MonthlyReportPptxTest(unittest.TestCase):
    def test_builds_five_editable_slides_from_business_template(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "月报测试.pptx"
            result = build_monthly_report_pptx(SAMPLE_REPORT, output_path, DATASET_PATH)

            self.assertTrue(result["success"])
            self.assertEqual(result["slides_count"], 5)
            self.assertEqual(result["period"], "2026年5月")
            self.assertTrue(output_path.is_file())

            presentation = Presentation(output_path)
            self.assertEqual(len(presentation.slides), 5)
            self.assertEqual(
                len([shape for shape in presentation.slides[0].shapes if shape.has_chart]),
                3,
            )
            self.assertEqual(
                len([shape for shape in presentation.slides[3].shapes if shape.has_table]),
                4,
            )

            first_chart = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart][0]
            first_series = list(first_chart.series[0].values)
            self.assertEqual([round(value, 4) for value in first_series], [85.0076, 92.7674, 9.6054, 1.4636])

            share_chart = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart][2]
            self.assertEqual(
                [[round(value, 1) for value in series.values] for series in share_chart.series],
                [
                    [0.8, 0.7, 0.7, 0.6],
                    [5.0, 4.4, 4.8, 4.9],
                    [52.6, 54.4, 48.9, 53.7],
                    [41.6, 40.5, 45.6, 40.8],
                ],
            )

            self.assertEqual(
                [category.label for category in share_chart.plots[0].categories],
                ["2026.4", "2026.5", "2025.1-12", "2026.1-5"],
            )

            slide_text = " ".join(
                shape.text
                for shape in presentation.slides[3].shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("2.4 SUV级别", slide_text)
            self.assertIn("SUV A0与C+级表现韧性较强", slide_text)
