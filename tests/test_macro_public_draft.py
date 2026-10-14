"""Draft mechanics only; synthetic observations here are not source acceptance."""
from datetime import date
from io import BytesIO
from pathlib import Path
import sys
import unittest

from openpyxl import load_workbook
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
SCRIPTS=ROOT/'skills/macro-monthly-report/scripts'
sys.path.insert(0,str(SCRIPTS))
from inspect_macro_template import inventory, all_shapes
from macro_draft import rebuild_draft_chart, clear_chart_archive, draft_table_facts, update_draft_tables, update_draft_narrative
import macro_engine as engine


class PublicDraftTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.reference=SCRIPTS.parent/'assets/reference.pptx'
        cls.contract=inventory(cls.reference)
        cls.charts={c['id']:c for slide in cls.contract['slides'] for c in slide['charts']}

    def test_missing_charts_do_not_retain_template_values(self):
        for chart in self.charts.values():
            rebuilt=rebuild_draft_chart(chart,{},'2026-09',date(2026,10,3))
            for series in rebuilt['series']:
                if series['request_ids']:
                    self.assertTrue(all(v is None for v in series['actual_values']),chart['id'])
                self.assertTrue(all(v is None for v in series['forecast_values']))
                self.assertEqual(series['forecasts'],[])

    def test_actuals_keep_their_own_month_not_report_month(self):
        chart=self.charts['s11-shape12'];rid=chart['requests'][0]['id']
        obs={rid:[{'request_id':rid,'date':date(2026,8,31),'value':49.4,'source_url':'fixture-not-business-proof'}]}
        rebuilt=rebuild_draft_chart(chart,obs,'2026-09',date(2026,10,3))
        item=next(s for s in rebuilt['series'] if rid in s['request_ids'])
        dates=[parse_date(v) for v in rebuilt['categories']]
        self.assertEqual(item['actual_values'][dates.index(date(2026,8,31))],49.4)
        self.assertIsNone(item['actual_values'][dates.index(date(2026,9,30))])

    def test_static_pmi_boundary_is_constant_not_an_actual(self):
        chart=self.charts['s4-shape10']
        rebuilt=rebuild_draft_chart(chart,{},'2026-09',date(2026,10,3))
        constants=[s for s in rebuilt['series'] if not s['request_ids']]
        self.assertTrue(constants)
        self.assertTrue(all(v==50 for s in constants for v in s['actual_values']))
        self.assertTrue(all(not s['verified_points'] for s in constants))

    def test_previous_seasonal_years_are_not_archive_proof(self):
        rebuilt=rebuild_draft_chart(self.charts['s24-shape9'],{},'2026-09',date(2026,10,3))
        self.assertTrue(all(v is None for s in rebuilt['series'] for v in s['actual_values']))

    def test_snapshot_uses_common_actual_period(self):
        chart=self.charts['s11-shape15'];obs={}
        for i,r in enumerate(chart['requests']):
            obs[r['id']]=[{'request_id':r['id'],'date':date(2026,8,31),'value':50+i}]
            if i: obs[r['id']].append({'request_id':r['id'],'date':date(2026,9,30),'value':70+i})
        result=rebuild_draft_chart(chart,obs,'2026-09',date(2026,10,3))
        self.assertEqual(result['target'],'2026-08-31')
        self.assertTrue(any(s['name']=='2026-08' for s in result['series']))

    def test_overview_never_forecasts_supplemental_rows(self):
        charts=[rebuild_draft_chart(c,{},'2026-09',date(2026,10,3)) for c in self.charts.values()]
        rows=draft_table_facts(charts,{},'2026-09')
        self.assertTrue(all(f is None for r in rows for f in r['facts']))

    def test_workbook_is_compact_and_extensions_removed(self):
        deck=Presentation(str(self.reference))
        native=next(s.chart for s in all_shapes(deck.slides[27].shapes) if getattr(s,'has_chart',False) and s.shape_id==11)
        clear_chart_archive(native)
        self.assertFalse(native._chartSpace.xpath('.//c:dLbl | .//c:extLst'))
        external=native._chartSpace.find(engine.C('externalData'))
        part=native.part.related_part(external.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
        workbook=load_workbook(BytesIO(part.blob))
        self.assertLess(sum(ws.max_row*ws.max_column for ws in workbook),2000)
        self.assertTrue(all(ws.sheet_state=='visible' for ws in workbook))

    def test_draft_summary_and_missing_indicators_have_clear_labels(self):
        deck=Presentation(str(self.reference))
        charts=[rebuild_draft_chart(c,{},'2026-09',date(2026,10,3)) for c in self.charts.values()]
        rows=draft_table_facts(charts,{},'2026-09')
        tables=update_draft_tables(deck,rows,'2026-09')
        texts=update_draft_narrative(deck,charts,rows,'2026-09',date(2026,10,3))
        for sid,label in ((12,'GDP'),(10,'工业'),(9,'CPI'),(40,'PPI'),(41,'消费'),(33,'PMI'),(34,'投资')):
            self.assertEqual(engine._shape(deck.slides[1],sid).text,label)
        self.assertIn('万亿元',engine._shape(deck.slides[13],7).text)
        self.assertEqual(tables[1]['cells'][0][1],'9月方向')
        self.assertTrue(all(row[1]=='未发布/暂无数据' for row in tables[1]['cells'][1:]))
        self.assertTrue(any(t['slide']==24 and t['new']=='暂无可核验数据' and t.get('added') for t in texts))


def parse_date(value):
    from datetime import timedelta
    return date(1899,12,30)+timedelta(days=int(value))


if __name__=='__main__': unittest.main()
