import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from pptx import Presentation

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'skills/macro-monthly-report/scripts'))
import macro_history as history
import macro_historical_engine as engine
import macro_historical_derivations as derivations
from utils.macro_report_jobs import MacroReportJobManager


class HistoricalReportTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.contract=history.inventory()
        cls.history=history.build_history(cls.contract)

    def test_semantics_map_reordered_template_and_literal_series(self):
        charts={c['id']:c for s in self.contract['slides'] for c in s['charts']}
        self.assertEqual(len(charts),36)
        self.assertEqual(charts['s13-shape10']['slide'],12)
        self.assertEqual(charts['s28-shape11']['shape_id'],26)
        self.assertEqual(charts['s8-shape25']['series'][0]['name'],'汽车轮胎:全钢胎')
        self.assertNotIn('s12-shape13',charts)

    def test_conflicting_june_car_sales_excluded_from_comparison(self):
        rid='auto-sales::current';period='2026-06-30'
        self.assertTrue(any(c['request_id']==rid and c['period']==period for c in self.history['conflicts']))
        self.assertFalse(any(c['request_id']==rid and c['period']==period for c in self.history['observations']))
        self.assertTrue(any(c['request_id']=='overview::钢材产量当月同比' for c in self.history['observations']))

    def test_missing_actuals_do_not_become_forecasts_or_old_month_actuals(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);(root/'working').mkdir()
            package={'report_month':'2026-08','cutoff_date':'2026-09-30','observations':[],
                     'evidence_observations':[],'status':'blocked'}
            (root/'working/data-package.json').write_text(json.dumps(package),encoding='utf-8')
            result=engine.build('2026-08',package,root/'report.pptx',root/'ledger.json',root)
            self.assertFalse(result['business_complete'])
            ledger=json.loads((root/'ledger.json').read_text(encoding='utf-8'))
            self.assertEqual(ledger['outlook_month'],'2026-09')
            self.assertEqual(ledger['issue_month'],'2026-09')
            industrial=next(c for c in ledger['charts'] if c['chart_id']=='s13-shape17')
            self.assertIsNone(industrial['series'][1]['values'][-1])
            self.assertTrue(all('forecast' not in s['name'] and '预测' not in s['name'] for c in ledger['charts'] for s in c['series']))
            deck=Presentation(str(root/'report.pptx'))
            self.assertEqual(sum(s.has_chart for slide in deck.slides for s in slide.shapes),36)
            table=next(s.table for s in deck.slides[31].shapes if s.has_table)
            self.assertIn('09月预测',table.cell(0,2).text)
            self.assertEqual(table.cell(1,2).text,'证据不足')
            v=engine.validate(root/'report.pptx',root/'ledger.json')
            self.assertFalse(v['passed'])
            self.assertFalse(any(x['code']=='cache_workbook_mismatch' for x in v['issues']),v['issues'][:5])
            # Tamper with a cache and rehash the ledger: independent workbook/recipe
            # checks must still catch the corruption, not just the file hash.
            chart=next(s.chart for s in deck.slides[10].shapes if s.has_chart)
            chart._chartSpace.xpath('.//c:ser/c:val/c:numRef/c:numCache/c:pt/c:v')[0].text='999'
            deck.save(str(root/'report.pptx'));ledger['pptx_sha256']=engine.digest(root/'report.pptx')
            (root/'ledger.json').write_text(json.dumps(ledger),encoding='utf-8')
            bad=engine.validate(root/'report.pptx',root/'ledger.json')
            self.assertTrue(any(x['code']=='native_chart_data_changed' for x in bad['issues']))

    def test_future_month_names_and_partial_scope_are_rejected_before_job(self):
        with tempfile.TemporaryDirectory() as tmp:
            jobs=MacroReportJobManager(Path(tmp),lambda *args:{})
            with patch.object(jobs,'_spawn'):
                job=jobs.start({'report_month':'2026-08','mode':'historical_update'})
                self.assertEqual(job['payload']['mode'],'historical_update')
                with self.assertRaises(ValueError):jobs.start({'report_month':'2026-08','mode':'historical_update','slides':[4]})
                with self.assertRaises(ValueError):jobs.start({'report_month':'2026-08','mode':'historical_update','delivery_mode':'public_draft','public_draft_confirmed':True})

    def test_cumulative_flow_difference_preserves_rounding_and_refuses_stock(self):
        def page(month,amount,published):
            cn='七' if month==7 else '八'
            return {'receipt_id':str(month),'url':'https://www.pbc.gov.cn/report/'+str(month),
                    'published_at':published,'text':f'2026年{month}月金融统计数据报告\n五、前{cn}个月人民币贷款增加{amount}万亿元。\n六、{month}月份银行间人民币市场'}
        pages={'7':page(7,'10.38','2026-08-15'),'8':page(8,'10.44','2026-09-14')}
        values=derivations.replay(pages,'2026-08','2026-09-30')
        self.assertEqual(len(values),1)
        self.assertEqual(values[0]['value'],600)
        self.assertEqual(values[0]['uncertainty_half_width'],100)
        self.assertTrue(values[0]['approximate'])
        self.assertEqual(len(values[0]['derivation']['inputs']),2)
        self.assertEqual(derivations.replay(pages,'2026-08','2026-08-31'),[])
        pages['8']['text']=pages['8']['text'].replace('人民币贷款增加','人民币贷款余额为')
        self.assertEqual(derivations.replay(pages,'2026-08','2026-09-30'),[])


if __name__=='__main__':unittest.main()
