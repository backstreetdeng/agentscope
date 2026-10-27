"""Independent weekly news/content/template validator; no model or network."""
from __future__ import annotations
import argparse, hashlib, html, json, re
from datetime import date
from decimal import Decimal
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import urlsplit
from zipfile import ZipFile
from docx import Document
from lxml import etree
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
SECTION_ROWS = json.loads((Path(__file__).resolve().parents[1] / 'references/sections.json').read_text(encoding='utf-8'))
NUM_RE = re.compile(r'\d+(?:\.\d+)?')
ACTION_RE = r'开展|阐明|发布|印发|公布|签署|签订|签约|达成|成立|投产|量产|下线|交付|定点|送样|上市|预售|开工|启用|任命|出任|履新|上任|接任|离任|卸任|辞任|辞职|当选|生效|通过|召开|举行|举办|开幕|启动|获批|实施|完成|推出'

def normalize(value):
    return re.sub(r'[^\w\u4e00-\u9fff]', '', html.unescape(str(value))).lower()

def coverage_entries(text):
    result=[]
    for line in text.splitlines():
        columns=[s.strip() for s in line.strip().strip('|').split('|')]
        if len(columns) >= 5 and re.fullmatch(r'\d{4}\.\d{4}-(?:\d{4}\.)?\d{4}', columns[0]):
            result.append(dict(period=columns[0], section=columns[1], title=columns[2], date=columns[3], source=columns[4]))
    return result

def coverage_for_period(text,start,end):
    period=start[:4]+'.'+start[5:].replace('-','')+'-'+(end[:4]+'.' if end[:4]!=start[:4] else '')+end[5:].replace('-','')
    return [row for row in coverage_entries(text) if row['period']!=period]

def duplicate(item, coverage):
    title=normalize(item.get('title','')); url=item.get('source','')
    for row in coverage:
        previous=normalize(row.get('title',''))
        if (url and url==row.get('source')) or (title and previous and
                (title==previous or SequenceMatcher(None,title,previous).ratio()>=0.83)):
            return row
    return None

def date_in_quote(event_date, quote, source):
    day=date.fromisoformat(event_date)
    quote=re.sub(r'\s+', '', quote)
    patterns=[rf'{day.year}[年./-]0?{day.month}[月./-]0?{day.day}(?:日)?',
              rf'(?<!\d)0?{day.month}[月./-]0?{day.day}(?:日)?(?!\d)']
    if not any(re.search(pattern, quote) for pattern in patterns): return False
    years=set(re.findall(r'20\d{2}',quote))
    published=str(source.get('published',''))
    if years: return str(day.year) in years
    return published.startswith(str(day.year)) or str(day.year) in source.get('text','')

def occurrence_supported(event_date,quote,source):
    """An article timestamp beside '近日' is not an actual event date."""
    lines=re.split(r'[。！？\n]',str(quote))
    action=ACTION_RE
    for line in lines:
        if re.search(r'URL:|Content-Type:|Title:|Published|发布时间|发表时间|发布日期|发布于|更新时间|来源[:：]',line,re.I):continue
        if line.strip()==str(source.get('title','')).strip():continue
        if not date_in_quote(event_date,line,source):continue
        prose=re.sub(r'新闻发布会|发布会上|发布会|发布于','',line)
        if re.search(action,prose) and not re.search(r'近日|日前|将于|将开展|即将|计划|拟于',prose):return True
    # Official circulars carry their issue date in the sign-off, separately.
    compact=re.sub(r'\s+','',str(quote))
    day=date.fromisoformat(event_date)
    signed=rf'{day.year}年0?{day.month}月0?{day.day}日'
    return bool(re.search(signed,compact) and re.search(r'现将|印发给你们|予以公布|特此通知',compact) and not re.search(r'近日|日前',compact))

def validate_item(item, spec, sources, start, end, coverage):
    errors=[]
    if not isinstance(item,dict): return ['资讯条目必须为对象']
    title=str(item.get('title','')).strip(); body=str(item.get('body','')).strip()
    source=sources.get(str(item.get('source_id','')))
    if not title or not body: errors.append('缺少标题或正文')
    if not 40<=len(body)<=320: errors.append('正文需40–320字，禁止空泛填充或长篇扩写')
    if spec['id'] in {'macro','joint'} and len(body)>150: errors.append('宏观/合资正文不得超过150字')
    if '\ufffd' in title+body or re.search(r'<[^>]+>',title+body): errors.append('正文含乱码或HTML标签')
    try:
        event=date.fromisoformat(str(item.get('event_date','')))
        if not date.fromisoformat(start)<=event<=date.fromisoformat(end): errors.append('事件真实发生日超出本轮窗口')
        if item.get('date')!=event.strftime('%m.%d'): errors.append('显示日期与事件日不一致')
    except (ValueError,TypeError): errors.append('事件日缺失或不是完整ISO日期')
    if not source or not source.get('fetched') or not source.get('text'): errors.append('没有实际抓取成功的来源证据')
    else:
        if item.get('source')!=source.get('url'): errors.append('来源URL与抓取记录不一致')
        quote=str(item.get('event_quote','')).strip()
        if len(quote)<15 or normalize(quote) not in normalize(source['text']): errors.append('事件证据不是来源正文逐字摘录')
        else:
            try:
                if not date_in_quote(item['event_date'],quote,source): errors.append('原文证据不能确认事件日与年度')
                elif not occurrence_supported(item['event_date'],quote,source):errors.append('事件日仅来自文章时间或预告，未锚定实际发生动作')
            except ValueError: pass
        if not re.search(ACTION_RE, quote):
            errors.append('事件日期证据缺少已发生的动作')
        if spec.get('preserve_title') and title!=source.get('title'): errors.append('需沿用实际来源原标题')
        numeric={Decimal(n) for n in NUM_RE.findall(body)} - {Decimal(n) for n in NUM_RE.findall(source['text'])}
        if numeric: errors.append('正文数字不在来源内：'+','.join(str(n) for n in sorted(numeric)))
    if item.get('status')!='occurred': errors.append('只收已发生事件')
    if item.get('scope')!='china' and spec['id']!='macro': errors.append('只收中国本土动作')
    if spec['id']!='macro' and re.search(r'(?:在|赴)(?:日本|德国|美国|法国|英国|泰国|马来西亚|印度|印尼|印度尼西亚|巴西|墨西哥|匈牙利|西班牙|南非|意大利|阿联酋|沙特|越南|韩国|新加坡|菲律宾|澳大利亚|土耳其|俄罗斯)',body+str(item.get('event_quote',''))):errors.append('原文显示事件在海外发生，不能标为中国本土动作')
    if not item.get('event_key'): errors.append('缺少事件去重键')
    if duplicate(item,coverage): errors.append('历史台账已报道该事件或来源')
    if re.search(r'财报|业绩快报|交付量|销量|销量突破|产销数据|累计售出|充电量|客流|成交额|以旧换新.*(?:万辆|亿元)|订单量|渗透率|市场份额',title):
        errors.append('属于禁收的运营或政策成效数据资讯')
    if re.search(r'定档|申报|招标|采购公告|将于.*(?:上市|发布|亮相)|即将上市|即将发布|预告|金辑奖',title):
        errors.append('属于预告、申报或招标资讯')
    if spec['id']=='luxury' and re.search(r'召回|危机|暴跌|辟谣|质量问题',title): errors.append('豪华品牌仅收正面动作')
    if spec['section']=='企业要闻' and re.search(r'上市|预售|新车亮相|车型首发',title): errors.append('车型事件应归新车，不得回流企业板块')
    if spec['id'] in {'parts','energy','intelligent','othertech'}:
        if not re.search(r'量产|投产|定点|送样|样品|上车|装车|中试|性能|参数|芯片|架构|技术',body): errors.append('技术/零部件缺少产品技术及落地证据')
    if spec['id']=='newcar':
        proof=item.get('series_history') or {}; record=sources.get(str(proof.get('source_id','')))
        host=urlsplit(record.get('url','')).hostname if record else ''
        if (proof.get('is_new_series') is not True or not record or not record.get('fetched')
            or not any((host or '')==d or (host or '').endswith('.'+d) for d in ('autohome.com.cn','bitauto.com','yiche.com'))
            or len(str(proof.get('quote','')))<15 or normalize(proof.get('quote','')) not in normalize(record.get('text',''))):
            errors.append('全新车系缺少汽车之家/易车独立历史核验')
        name=str(proof.get('series_name','')).strip()
        if not name or not record or normalize(name) not in normalize(title+body) or normalize(name) not in normalize(record.get('title','')) or normalize(name) not in normalize(proof.get('quote','')):
            errors.append('车系历史证明没有精确绑定本条车型名称')
        if record and record.get('series_name') and normalize(name)!=normalize(record['series_name']):errors.append('证明页面属于另一准确车系，禁止相似名称混配')
        if record and record.get('fetch_method')=='autohome/public-suggest':errors.append('搜索建议只证明车系身份，须引用返回的年款目录source_id核对历史')
        if record and record.get('catalog_years'):
            if not record.get('catalog_year_selector'):errors.append('车系年款目录未取得有效筛选记录')
            if min(record['catalog_years'])<date.fromisoformat(start).year:errors.append('车系已有早于本期年度的车型，不是全新车系')
        if re.search(r'改款|年款|换代|新增配置|衍生',title): errors.append('新车板块不收改款/换代/衍生')
    if re.search(r'全球首发|首次亮相|首秀',title+body):
        proof=item.get('first_occurrence_proof') or {}
        evidence=sources.get(str(proof.get('source_id',''))) if isinstance(proof,dict) else None
        if not evidence or not evidence.get('fetched') or len(str(proof.get('quote','')))<15 or normalize(proof.get('quote','')) not in normalize(evidence.get('text','')):
            errors.append('首次/首秀表述没有历史核验证据')
    return errors

def section_issues(spec,items):
    if spec['id']=='joint' and len({str(item.get('company','')).strip() for item in items})<3:
        return ['合资板块需至少三个不同主体；通过条目保留，只补另一家企业的真实消息']
    return []

def same_story(left,right):
    if left.get('source')==right.get('source') or left.get('event_key')==right.get('event_key'):return True
    if left.get('event_date')!=right.get('event_date'):return False
    def actor(value):
        value=normalize(value)
        for specific in ('上汽大众','上汽通用','一汽大众','一汽丰田','广汽丰田','广汽本田','东风本田','北京现代','东风日产'):
            if specific in value:return specific
        if '丰田' in value or '雷克萨斯' in value:return '丰田'
        return value
    if actor(left.get('company',''))!=actor(right.get('company','')):return False
    def action(body):
        hit=re.search(r'签约|签署|签订|发布|公布|印发|投产|量产|下线|任命|出任|辞任|成立|开工|上市',body)
        if not hit:return ''
        word=hit.group(0)
        return '签约' if word in ('签约','签署','签订') else word
    if not action(left.get('body','')) or action(left.get('body',''))!=action(right.get('body','')):return False
    a=normalize(left.get('body',''));b=normalize(right.get('body',''))
    for match in SequenceMatcher(None,a,b,autojunk=False).get_matching_blocks():
        shared=a[match.a:match.a+match.size]
        if match.size>=8 and ('有限公司' in shared or '有限责任公司' in shared):return True
    return False

def validate_content(data,sources,coverage,start,end):
    errors=[]; warnings=[]
    if data.get('period_start')!=start or data.get('period_end')!=end: errors.append('报告日期与本轮输入不一致')
    blocks=data.get('body_sections',[])
    if not isinstance(blocks,list) or [(b.get('section'),b.get('subsection','')) for b in blocks]!=[(s['section'],s['subsection']) for s in SECTION_ROWS]:
        return {'passed':False,'errors':errors+['九章十四板块缺失、重复或顺序不正确'],'warnings':warnings}
    titles=set(); events=set(); all_items=[]
    for block,spec in zip(blocks,SECTION_ROWS):
        items=block.get('items',[])
        if not isinstance(items,list) or not items: errors.append(spec['section']+'/'+spec['subsection']+'：没有通过核验的资讯'); continue
        if len(items)>spec['target_max']: errors.append(spec['id']+'超出条数上限')
        for item in items:
            issues=validate_item(item,spec,sources,start,end,coverage)
            errors.extend([str(item.get('title',''))+'：'+v for v in issues])
            key=normalize(item.get('title','')); event=item.get('event_key')
            if key in titles or event in events: errors.append(str(item.get('title',''))+'：本期重复故事线')
            titles.add(key); events.add(event);all_items.append((item,spec))
    merged=(date.fromisoformat(end)-date.fromisoformat(start)).days>6
    lower,upper=(12,13) if merged else (9,10)
    summary=data.get('summary',[])
    if not isinstance(summary,list) or not lower<=len(summary)<=upper: errors.append(f'摘要需{lower}–{upper}条')
    ranking={'宏观动态':0,'政策法规':1,'企业要闻':2,'新车上市':3}; order=[]; chosen=set()
    for item in summary:
        matches=[(v,spec) for v,spec in all_items if (v.get('title'),v.get('body'),v.get('source'))==(item.get('title'),item.get('body'),item.get('source'))]
        if not matches: errors.append('摘要没有逐字复制正文：'+str(item.get('title','')));continue
        spec=matches[0][1]
        if spec['section'] not in ranking: errors.append('摘要包含技术、零部件或其他禁选板块')
        else: order.append(ranking[spec['section']])
        if item['title'] in chosen: errors.append('摘要条目重复')
        chosen.add(item['title'])
    if order!=sorted(order): errors.append('摘要顺序需宏观→政策→企业→新车')
    joint=blocks[4]['items']
    if len(set(str(i.get('company','')) for i in joint))<3: errors.append('合资板块不足三个不同主体')
    if sum('上汽' in str(i.get('company','')) for i in joint)>2: errors.append('合资板块上汽系超过两条')
    for index,(item,_) in enumerate(all_items):
        for previous,_ in all_items[:index]:
            if same_story(item,previous):errors.append(item['title']+'：与 '+previous['title']+' 为同一已收录故事')
    return {'passed':not errors,'errors':errors,'warnings':warnings,'items':len(all_items),'summary_items':len(summary),'sections':9,'subsections':14}

def validate_artifacts(data,folder,template=None):
    errors=[]; metadata=[]
    for kind in ('html','docx','sources'):
        matches=[p for p in Path(folder).glob('*') if p.is_file() and ((kind=='html' and p.suffix=='.html') or (kind=='docx' and p.suffix=='.docx' and not p.stem.endswith('_来源')) or (kind=='sources' and p.stem.endswith('_来源') and p.suffix=='.docx'))]
        if len(matches)!=1: errors.append(kind+'交付物缺失或有多个版本');continue
        path=matches[0];metadata.append({'kind':kind,'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'bytes':path.stat().st_size})
        if kind=='html':
            text=path.read_text(encoding='utf-8')
            if '\ufffd' in text or re.search(r'<(?:script|link)[^>]+(?:src|href)=["\']https?://',text,re.I): errors.append('HTML有乱码或外链依赖')
            for item in [*data['summary'],*[i for b in data['body_sections'] for i in b['items']]]:
                if html.escape(item['title'],quote=False) not in text or html.escape(item['body'],quote=False) not in text: errors.append('HTML与数据标题/正文不一致：'+item['title'])
            if len(re.findall(r'<section id="s\d+"',text))!=9: errors.append('HTML九章不完整')
            continue
        try:
            with ZipFile(path) as z:
                if z.testzip() is not None: errors.append(kind+'ZIP损坏')
            doc=Document(path);texts=[p.text.strip() for p in doc.paragraphs]
            if kind=='sources':
                for b in data['body_sections']:
                    for item in b['items']:
                        if not any(item['title'] in t and item['source'] in t for t in texts): errors.append('来源文档漏项：'+item['title'])
                continue
            if len(doc.sections)!=2: errors.append('Word分节不是封面/正文两节')
            for paragraph in doc.paragraphs:
                if paragraph.style.name == 'Heading 3' and paragraph.text.strip() != '重要周度资讯盘点':
                    properties=paragraph._element.find(f'{{{W}}}pPr')
                    numbering=properties.find(f'{{{W}}}numPr') if properties is not None else None
                    if numbering is None:errors.append('Word标题未保留自动编号：'+paragraph.text[:50])
            toc_title=next((p for p in doc.paragraphs if p.text.strip()=='目录'),None)
            if toc_title is None or not toc_title.paragraph_format.page_break_before:errors.append('目录没有另起一页')
            if not any(node.get(f'{{{W}}}type')=='page' for node in doc.element.iter(f'{{{W}}}br')):
                errors.append('目录与正文之间缺少分页符')

            cover=''.join(t.text or '' for t in doc.element.iter(f'{{{W}}}t'))
            expected=data['period_start'].replace('-','.')+'-'+data['period_end'].replace('-','.')
            if cover.count(expected)<2: errors.append('封面两个日期区没有正确更新')
            for para in doc.element.iter(f'{{{W}}}p'):
                if not any(parent.tag==f'{{{W}}}txbxContent' for parent in para.iterancestors()):continue
                direct=[]
                for node in para.iter(f'{{{W}}}t'):
                    nearest=next((parent for parent in node.iterancestors() if parent.tag==f'{{{W}}}p'),None)
                    if nearest is para:direct.append(node.text or '')
                title_year=''.join(direct).strip()
                if re.fullmatch(r'20\d{2}',title_year) and title_year!=data['period_start'][:4]:errors.append('封面艺术字年份与本期不一致')

            chapter_order=[p.text.strip() for p in doc.paragraphs if p.style.name == 'Heading 3' and p.text.strip() in {r['section'] for r in SECTION_ROWS}]
            # TOC entries include tabs; only exact heading paragraphs count.
            expected_sections=list(dict.fromkeys(r['section'] for r in SECTION_ROWS))
            if chapter_order!=expected_sections: errors.append('Word章节顺序不一致')
            for item in [*data['summary'],*[i for b in data['body_sections'] for i in b['items']]]:
                if item['title'] not in texts or item['body'] not in texts: errors.append('Word与正文数据不一致：'+item['title'])
            if not any('TOC ' in (t.text or '') for t in doc.element.iter(f'{{{W}}}instrText')): errors.append('缺少真实Word目录域')
            if template:
                with ZipFile(template) as original,ZipFile(path) as output:
                    for name in original.namelist():
                        if name.startswith('word/media/') and (name not in output.namelist() or original.read(name)!=output.read(name)):
                            errors.append('模板封面图片未完整保留')
        except Exception as exc: errors.append(kind+'不能重新读取：'+str(exc))
    return {'passed':not errors,'errors':errors,'artifacts':metadata}

def validate_source_files(sources, root):
    errors=[];root=Path(root).resolve()
    for record in sources.values():
        if not record.get('fetched'):continue
        try:
            path=(root/record['text_path']).resolve();path.relative_to(root)
            actual=hashlib.sha256(path.read_bytes()).hexdigest()
            inline=hashlib.sha256(record.get('text','').encode('utf-8')).hexdigest()
            if actual!=record.get('text_sha256') or inline!=actual:errors.append(record['id']+'来源正文指纹不一致')
        except (KeyError,ValueError,OSError):errors.append(record.get('id','未知来源')+'原文检查点缺失或路径无效')
    return errors

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--data',required=True);parser.add_argument('--sources',required=True);parser.add_argument('--coverage',required=True);parser.add_argument('--start-date',required=True);parser.add_argument('--end-date',required=True);parser.add_argument('--artifacts');parser.add_argument('--template');parser.add_argument('--report',required=True)
    args=parser.parse_args(); data=json.loads(Path(args.data).read_text(encoding='utf-8')); sources=json.loads(Path(args.sources).read_text(encoding='utf-8'))
    if isinstance(sources,list):sources={row['id']:row for row in sources}
    coverage=coverage_for_period(Path(args.coverage).read_text(encoding='utf-8'),args.start_date,args.end_date)
    content=validate_content(data,sources,coverage,args.start_date,args.end_date)
    content['errors'].extend(validate_source_files(sources,Path(args.data).resolve().parent.parent))
    content['passed']=not content['errors']
    artifacts=validate_artifacts(data,args.artifacts,args.template) if args.artifacts else {'passed':True,'artifacts':[]}
    result={'passed':content['passed'] and artifacts['passed'],'content':content,'artifacts':artifacts}
    Path(args.report).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(result,ensure_ascii=False));raise SystemExit(0 if result['passed'] else 2)
if __name__=='__main__':main()
