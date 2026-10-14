"""Explicit, incomplete public-data draft. Never a formal-report fallback.

The reference supplies geometry and names, not numeric history. Every nonconstant
chart cell comes from independently replayed source receipts. No forecasting is
performed. The web adapter must persist the user's explicit draft choice.
"""
from __future__ import annotations

from datetime import date
import hashlib
import io
import json
from pathlib import Path
import re
import posixpath
import zipfile

from openpyxl import Workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

import macro_engine as engine
from macro_contract import CONTRACT_VERSION, TEXT_SLOTS, OVERVIEW_ROWS, parse_period, month_end, static_series
from inspect_macro_template import all_shapes, inventory

MODE = "public_draft"
NOTICE = "公开数据草稿 · 非完整正式报告"


def rebuild_draft_chart(contract, observations, report_month, cutoff):
    semantic, freq = contract["semantics"], contract["frequency"]
    target = engine._target(report_month, freq, cutoff)
    requests = {r["id"]: r for r in contract["requests"]}
    proven = {rid: [p for p in observations.get(rid, []) if p["date"] <= target] for rid in requests}
    series = []
    if semantic in {"time", "mixed"}:
        original = [parse_period(p, freq) for p in contract["categories"]]
        ascending = original[0] <= original[-1]
        dates = sorted({p for p in original if p <= target} |
                       {p["date"] for rid, points in proven.items() for p in points if requests[rid]["frequency"] == freq})
        if freq == "announcement_monthly":
            actual_dates = {p["date"] for points in proven.values() for p in points}
            months = {p.strftime("%Y-%m") for p in actual_dates}
            dates = sorted({p for p in dates if p.strftime("%Y-%m") not in months} | actual_dates)
        while dates and engine.advance(dates[-1], freq) <= target:
            dates.append(engine.advance(dates[-1], freq))
        if not dates:
            dates = [target]
        dates = dates[-len(original):]
        display_dates = dates if ascending else list(reversed(dates))
        categories = [(p - date(1899, 12, 30)).days for p in display_dates]
    else:
        categories = list(contract["categories"])
        display_dates = []
        if semantic in {"snapshot", "category_metrics"}:
            latest = [max(p["date"] for p in points) for points in proven.values() if points]
            # A snapshot comparison must share one observation period, not one
            # category from August alongside another from September.
            if latest:
                target = min(latest)

    for old_index, old in enumerate(contract["series"]):
        name = old["name"]
        actual = [None] * len(categories)
        rids, points = [], []
        display_name = name
        if static_series(name, old["values"]):
            actual = [next(v for v in old["values"] if v is not None)] * len(categories)
            unit = contract["raw_unit"]
        elif semantic in {"time", "mixed"}:
            rid = contract["id"] + "::" + name
            rids = [rid]; points = proven[rid]; unit = requests[rid]["unit"]
            values = {p["date"]: p["value"] for p in points}
            actual = [values.get(p) if requests[rid]["frequency"] != "quarterly" or p.month in {3, 6, 9, 12} else None
                      for p in display_dates]
        elif semantic in {"seasonal", "seasonal_current"}:
            rid = contract["id"] + ("::current" if semantic == "seasonal" else "::" + name)
            rids = [rid]; points = proven[rid]; unit = requests[rid]["unit"]
            years = re.findall(r"20\d{2}", str(name))
            year = int(years[-1]) if semantic == "seasonal" and years else int(report_month[:4])
            for point in points:
                if point["date"].year == year:
                    index = point["date"].isocalendar()[1] - 1 if freq == "weekly" else point["date"].month - 1
                    if 0 <= index < len(actual):
                        actual[index] = point["value"]
        else:
            display_target = target
            if semantic == "snapshot" and len(contract["series"]) > 1:
                old_dates = [parse_period(s["name"]) for s in contract["series"]]
                rank = sorted(old_dates).index(old_dates[old_index])
                key = target.year * 12 + target.month - 1 - (len(old_dates) - 1 - rank)
                display_target = month_end(key // 12, key % 12 + 1)
            if semantic == "snapshot":
                display_name = display_target.strftime("%Y-%m")
            for index in range(len(categories)):
                rid = contract["id"] + (f"::category{index}" if semantic == "snapshot" else "::" + name + f"::category{index}")
                rids.append(rid); own = proven[rid]; points.extend(own)
                actual[index] = next((p["value"] for p in own if p["date"] == display_target), None)
            unit = contract["raw_unit"]
        series.append({"original_name": name, "name": display_name, "actual_values": actual,
                       "forecast_values": [None] * len(categories), "unit": unit, "request_ids": rids,
                       "frequency": requests[rids[0]]["frequency"] if rids else freq, "forecasts": [],
                       "verified_points": [{**p, "date": p["date"].isoformat()} for p in points]})
    return {"chart_id": contract["id"], "semantics": semantic, "frequency": freq, "categories": categories,
            "target": target.isoformat(), "definition": contract["definition"], "series": series, "issues": [],
            "archival_history": "未使用模板数值；仅呈现原始来源已核验实际；未取证单元格为空"}


def clear_chart_archive(chart):
    """Fresh compact workbook; copy only immutable categories and series names.

    One retained workbook has a million formatted rows. Never materialize that
    workbook in openpyxl merely to clear data. No unreferenced/hidden cells survive.
    """
    xml = chart._chartSpace
    external = xml.find(engine.C("externalData"))
    part = chart.part.related_part(external.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"))
    keep = set()
    for formula in xml.xpath(".//c:ser/c:cat//c:f | .//c:ser/c:tx//c:f"):
        sheet, area = formula.text.rsplit("!", 1)
        sheet = sheet.strip("'")
        c1, r1, c2, r2 = range_boundaries(area.replace("$", ""))
        keep.update((sheet, row, col) for row in range(r1, r2 + 1) for col in range(c1, c2 + 1))
    workbook = Workbook()
    names = list(dict.fromkeys(sheet for sheet, _, _ in sorted(keep)))
    workbook.active.title = names[0]
    for name in names[1:]: workbook.create_sheet(name)
    sns={"s":"http://schemas.openxmlformats.org/spreadsheetml/2006/main","r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with zipfile.ZipFile(io.BytesIO(part.blob)) as archive:
        strings=[]
        if "xl/sharedStrings.xml" in archive.namelist():
            shared=etree.fromstring(archive.read("xl/sharedStrings.xml"))
            strings=["".join(item.itertext()) for item in shared]
        wbxml=etree.fromstring(archive.read("xl/workbook.xml"))
        rels=etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets={rel.get("Id"): rel.get("Target") for rel in rels}
        for sheet in wbxml.findall(".//s:sheet",sns):
            name=sheet.get("name")
            cells={(row,col) for own,row,col in keep if own==name}
            if not cells: continue
            target=targets[sheet.get("{"+sns['r']+"}id")]
            path=target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/"+target)
            maximum=max(row for row,_ in cells)
            with archive.open(path) as content:
                for _, node in etree.iterparse(content,events=("end",),tag="{"+sns['s']+"}row"):
                    row=int(node.get("r"))
                    if row>maximum: break
                    for cell in node:
                        coordinate=coordinate_to_tuple(cell.get("r"))
                        if coordinate not in cells: continue
                        kind=cell.get("t");raw=cell.findtext("{"+sns['s']+"}v")
                        if kind=="s": value=strings[int(raw)] if raw is not None else None
                        elif kind=="inlineStr": value="".join(cell.xpath(".//s:t/text()",namespaces=sns))
                        elif kind=="str": value=raw
                        elif kind=="b": value=raw=="1"
                        else: value=float(raw) if raw is not None else None
                        workbook[name].cell(*coordinate).value=value
                    node.clear()
    stream = io.BytesIO(); workbook.save(stream); part._blob = stream.getvalue()
    # Literal manual chart labels/callouts would otherwise retain old numbers.
    for label in xml.xpath(".//c:dLbl | .//c:extLst | .//c:chart/c:title"):
        label.getparent().remove(label)


def draft_table_facts(charts, observations, report_month):
    # The shared overview implementation forecasts the three supplemental
    # rows. Remove that capability in the draft: actuals only for all rows.
    byid = {c["chart_id"]: c for c in charts}
    year, month = map(int, report_month.split("-"))
    periods = []
    for offset in (2, 1, 0):
        key = year * 12 + month - 1 - offset
        periods.append(f"{key // 12:04d}-{key % 12 + 1:02d}")
    rows = []
    for label, cid, key, unit in OVERVIEW_ROWS:
        facts = []
        for period in periods:
            if cid.startswith("overview::"):
                point = next((p for p in observations.get(cid, []) if p["date"].strftime("%Y-%m") == period), None)
                fact = {"value": point["value"], "status": "actual", "unit": unit, "source_url": point["source_url"]} if point else None
            else:
                fact = engine._fact(byid[cid], key, period)
            if fact and unit == "万亿元" and fact["unit"] == "亿元":
                fact = {**fact, "value": fact["value"] / 10000, "unit": unit}
            facts.append(fact)
        rows.append({"label": label, "unit": unit, "periods": periods, "facts": facts, "chart_id": cid})
    return rows


def update_draft_tables(deck, rows, report_month):
    receipts = engine.update_tables(deck, rows, report_month)
    # Missing current-month data says nothing about the validity of prior actuals.
    receipt = receipts[1]
    receipt["cells"][0][1] = str(int(report_month.split("-")[1])) + "月方向"
    for row in receipt["cells"][1:]:
        if row[1] == "待核验":
            row[1] = "未发布/暂无数据"
    table = engine._shape(deck.slides[31], 9).table
    for ri, row in enumerate(receipt["cells"]):
        for ci, text in enumerate(row):
            engine.set_text(table.cell(ri, ci), text)
    return receipts


def update_draft_narrative(deck, charts, rows, report_month, cutoff):
    updates = []
    def record(slide, sid, text):
        shape = engine._shape(deck.slides[slide - 1], sid)
        before = shape.text
        engine.set_text(shape, text)
        updates.append({"slide": slide, "shape_id": sid, "old": before, "new": text})
    # Preserve non-data chapter headings, units, logo, indicator names and page
    # numbers. Discard unrecomputed numeric prose, dated captions and old sources.
    roles = {(slide, sid) for slide, (title, bodies, source) in TEXT_SLOTS.items() for sid in (title, *bodies, source)}
    for number, slide in enumerate(deck.slides, 1):
        for shape in all_shapes(slide.shapes):
            if not getattr(shape, "has_text_frame", False) or (number, shape.shape_id) in roles:
                continue
            text = shape.text.strip()
            if "政策分析及下月经济前瞻" in text or "下月经济前瞻" in text:
                record(number, shape.shape_id, text.replace("政策分析及下月经济前瞻", "政策与公开数据状态").replace("下月经济前瞻", "公开数据状态"))
                continue
            dynamic = (bool(re.search(r"20\d{2}年|20\d{2}[-/]\d|\d+月|[+-]?\d[\d,.]*\s*(?:%|％|亿|万|百分点)|上半年|下半年|来源|wind|mysteel", text, re.I))
                       or (number in {2, 31} and len(text) > 20))
            if dynamic and not re.fullmatch(r"\d+(?:\.\d+)?", text):
                record(number, shape.shape_id, "")
        slide.notes_slide.notes_text_frame.text = NOTICE + "\n信息截止日：" + cutoff.isoformat()
    y, m = map(int, report_month.split("-"))
    record(1, 18, f"{y}年{m}月宏观经济月报\n公开数据草稿")
    record(1, 8, "信息截止日：" + cutoff.isoformat())
    byslide = {}
    for chart in charts:
        byslide.setdefault(int(chart["chart_id"].split("-")[0][1:]), []).append(chart)
    for slide, (title, bodies, source) in TEXT_SLOTS.items():
        if slide == 32:
            continue
        pagecharts = byslide.get(slide, [])
        points, urls, gaps = [], [], []
        for chart in pagecharts:
            for series in chart["series"]:
                for rid in series["request_ids"]:
                    own = [p for p in series["verified_points"] if p["request_id"] == rid]
                    if not own:
                        gaps.append(rid); continue
                    point = max(own, key=lambda p: p["date"])
                    label = series["name"]
                    if "::category" in rid:
                        label = str(chart["categories"][int(rid.rsplit("category", 1)[1])])
                        if chart["semantics"] == "category_metrics": label += series["name"]
                    if chart["semantics"] in {"seasonal", "seasonal_current"}: label = chart["definition"]
                    label = re.sub(r"[（(](?:万亿(?:元)?|亿元|%|％|指数|百分点)[）)]", "", label)
                    value = f"{point['value']:.2f}".rstrip("0").rstrip(".")
                    points.append(f"{label}：{point['date']}实际{value}{series['unit']}")
                    urls.append(point["source_url"])
        # Basic factual narrative, not a claim about causes or unobserved month.
        facts = list(dict.fromkeys(points))
        description = "；".join(facts[:2]) if facts else "暂无可核验数据，图表未填入旧月数字"
        record(slide, title, "公开数据草稿：" + (pagecharts[0]["definition"] if pagecharts else "指标观测"))
        record(slide, bodies[0], description + ("。本页存在缺项，详见备注。" if gaps else "。"))
        for sid in bodies[1:]: record(slide, sid, "")
        footer = engine._shape(deck.slides[slide - 1], source)
        footer.width = min(Pt(792), deck.slide_width - footer.left - Pt(54))
        footer.text_frame.margin_top = footer.text_frame.margin_bottom = Pt(0)
        record(slide, source, f"公开数据草稿；截止{cutoff.isoformat()}；空白=暂无可核验数据；实际期见正文/备注。")
        deck.slides[slide - 1].notes_slide.notes_text_frame.text = NOTICE + "\n" + "\n".join(facts) + "\n来源：\n" + "\n".join(dict.fromkeys(urls)) + "\n缺项：\n" + "\n".join(gaps)
    record(2, 2, f"{m}月宏观经济：公开数据草稿，按实际观测期披露")
    for sid, label in ((12,"GDP"),(10,"工业"),(9,"CPI"),(40,"PPI"),(41,"消费"),(33,"PMI"),(34,"投资"),
                       (25,"价格与消费"),(29,"景气与投资")):
        record(2, sid, label)
    for sid, row in zip((23, 54, 55, 45, 43, 36, 37), rows[:7]):
        available = [(period, fact) for period, fact in zip(row["periods"], row["facts"]) if fact]
        text = row["label"] + "：暂无可核验数据"
        if available:
            period, fact = available[-1]; text = row["label"] + "：" + period + "实际" + engine._formatted(fact) + row["unit"]
        elif sid == 23:
            gdp = next(c for c in charts if c["chart_id"] == "s10-shape7")
            points = [p for s in gdp["series"] if "GDP:不变价:当季同比" in s["original_name"] for p in s["verified_points"]]
            if points:
                point = max(points, key=lambda p:p["date"])
                quarter = (date.fromisoformat(point["date"]).month - 1) // 3 + 1
                text = f"GDP：{quarter}季度实际{point['value']:g}%；本季度未发布"
        record(2, sid, text)
    record(7, 9, "公开数据草稿；—表示暂无可核验数据或该月未发布；不预测、不沿用模板数值。")
    record(8, 9, "1.5 高频指标")
    for sid, cid in ((22, "s8-shape20"), (23, "s8-shape21"), (24, "s8-shape11"), (8, "s8-shape25")):
        chart = next(c for c in charts if c["chart_id"] == cid)
        dates = [p["date"] for s in chart["series"] for p in s["verified_points"]]
        record(8, sid, chart["definition"] + ("：实际至" + max(dates) if dates else "：暂无可核验数据"))
    record(8, 26, "公开数据草稿；高频指标需对应原口径；无数据不填入旧值。")
    for number, sid, cid in ((28, 9, "s28-shape11"), (28, 10, "s28-shape17"), (30, 12, "s30-shape13"), (30, 17, "s30-shape18")):
        record(number, sid, next(c for c in charts if c["chart_id"]==cid)["definition"])
    for sid in (6, 7): record(4, sid, "指数")
    for sid in (6,7):
        unit_shape=engine._shape(deck.slides[3],sid)
        unit_shape.width=Pt(40); unit_shape.height=Pt(18)
        unit_shape.text_frame.margin_top=unit_shape.text_frame.margin_bottom=Pt(0)
        unit_shape.text_frame.margin_left=unit_shape.text_frame.margin_right=Pt(0)
    for number, sid, cid in ((4,14,"s4-shape5"),(4,32,"s4-shape10"),(10,8,"s10-shape7"),(10,10,"s10-shape28"),
                             (12,7,"s12-shape13"),(12,8,"s12-shape16"),(14,7,"s14-shape15"),(14,8,"s14-shape11"),
                             (17,7,"s17-shape15"),(17,8,"s17-shape20"),(19,12,"s19-shape13")):
        chart=next(c for c in charts if c["chart_id"]==cid)
        text=chart["definition"]
        if cid=="s12-shape13": text+="（指数）"
        if cid=="s14-shape15": text+="（万亿元）"
        if cid=="s17-shape20": text+="（万平方米）"
        record(number, sid, text)
    for sid in (8, 11):
        record(32, sid, "已核验实际" if sid==8 else "未执行预测")
    record(31, 2, "财政与利率：公开数据观测，不推断政策效果")
    for sid, text in {4: "仅记录已核验实际；不形成政策因果结论", 5: "公开数据草稿。财政、利率实际数据见前页；未经取证的政策观点不保留。",
                      11: "需求指标", 13: "实际期和缺项见经济指标表。", 15: "财政收支", 30: "同口径已核验数据见前页。", 32: "LPR公告", 34: "采用原公告实际期，不给出利率预测。",
                      35: "01 · 需求观测", 36: "02 · 财政观测", 38: "03 · 利率观测"}.items():
        record(31, sid, text)
    for sid in (35, 36, 38): engine._shape(deck.slides[30], sid).height = Pt(26)
    record(32, 10, "经济指标实况：公开数据草稿（未执行预测）")
    record(32, 6, "仅列原始来源已核验实际。—表示暂无可核验数据或未发布，不以模板历史或预测值补齐。")
    record(32, 7, "公开数据草稿；信息截止日：" + cutoff.isoformat())
    # Neutral, compact content avoids long old annotations colliding with plots.
    for slide in deck.slides:
        for shape in all_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    if len(paragraph.text) > 50:
                        for run in paragraph.runs: run.font.size = Pt(14)
    for number,(_, bodies, source) in TEXT_SLOTS.items():
        for sid in bodies:
            frame=engine._shape(deck.slides[number-1],sid).text_frame
            frame.margin_left=frame.margin_right=Pt(0)
        frame=engine._shape(deck.slides[number-1],source).text_frame
        for paragraph in frame.paragraphs:
            for run in paragraph.runs: run.font.size=Pt(8)
    for number,sid in ((7,9),(8,26)):
        shape=engine._shape(deck.slides[number-1],sid)
        shape.width=min(Pt(792),deck.slide_width-shape.left-Pt(54)); shape.height=Pt(16)
        frame=shape.text_frame
        frame.margin_top=frame.margin_bottom=frame.margin_left=frame.margin_right=Pt(0)
        for paragraph in frame.paragraphs:
            for run in paragraph.runs: run.font.size=Pt(8)
    for chart in charts:
        actuals = [v for series in chart["series"] if series["request_ids"] for v in series["actual_values"] if v is not None]
        if actuals:
            continue
        number, sid = map(int, re.fullmatch(r"s(\d+)-shape(\d+)", chart["chart_id"]).groups())
        native = engine._shape(deck.slides[number-1], sid)
        notice = deck.slides[number-1].shapes.add_textbox(native.left + native.width//8, native.top + native.height//2,
                                                       native.width*3//4, Pt(24))
        notice.text_frame.margin_top=notice.text_frame.margin_bottom=Pt(0)
        notice.text = "暂无可核验数据"
        notice.fill.solid(); notice.fill.fore_color.rgb=RGBColor(255,255,255)
        for run in notice.text_frame.paragraphs[0].runs:
            run.font.size=Pt(15); run.font.color.rgb=RGBColor(100,100,100)
        updates.append({"slide":number,"shape_id":notice.shape_id,"old":"","new":notice.text,"added":True})
    for number, slide in enumerate(deck.slides, 1):
        stamp=slide.shapes.add_textbox(Pt(4),Pt(0),Pt(270),Pt(13))
        stamp.text_frame.margin_top=stamp.text_frame.margin_bottom=Pt(0)
        stamp.text_frame.margin_left=stamp.text_frame.margin_right=Pt(0)
        stamp.text=NOTICE
        for run in stamp.text_frame.paragraphs[0].runs:
            run.font.size=Pt(8); run.font.color.rgb=RGBColor(128,88,36)
        updates.append({"slide":number,"shape_id":stamp.shape_id,"old":"","new":NOTICE,"added":True})
    # Record the final state once per textbox, including cleared old content.
    final = {}
    for update in updates:
        key = (update["slide"], update["shape_id"])
        if key not in final: final[key] = update
        else: final[key]["new"] = update["new"]
    return list(final.values())


def build_draft(reference, report_month, package, output, ledger_path, source_root=None):
    reference, output = Path(reference), Path(output)
    if package.get("report_month") != report_month: raise ValueError("package_report_month_mismatch")
    cutoff = date.fromisoformat(package["cutoff_date"])
    if cutoff > date.today(): raise ValueError("cutoff_in_future")
    source_root = Path(source_root or Path(ledger_path).parent.parent)
    contract = inventory(reference)
    observations = engine.verified_observations(package, contract, source_root)
    usable = sum(bool(observations.get(r["id"])) for r in contract["source_requests"])
    if not usable: raise ValueError("public_draft_has_no_verified_data")
    deck = Presentation(str(reference)); charts = []
    for slide in contract["slides"]:
        for chart in slide["charts"]:
            rebuilt = rebuild_draft_chart(chart, observations, report_month, cutoff)
            native = engine._shape(deck.slides[slide["slide"] - 1], chart["shape_id"]).chart
            clear_chart_archive(native); engine.update_native_chart(native, rebuilt); charts.append(rebuilt)
            if chart["id"] == "s16-shape8":
                for labels in native._chartSpace.xpath(".//c:dLbls"):
                    labels.getparent().remove(labels)
            if chart["id"] in {"s18-shape17", "s18-shape21", "s21-shape11"}:
                for axis in native._chartSpace.xpath(".//c:valAx"):
                    number_format = axis.find(engine.C("numFmt"))
                    if number_format is not None:
                        number_format.set("formatCode", "0.00")
                        number_format.set("sourceLinked", "0")
    rows = draft_table_facts(charts, observations, report_month)
    removed_ole=engine._shape(deck.slides[6],24)._element
    rns="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ole_relations={value for node in removed_ole.iter() for attr,value in node.attrib.items() if attr.startswith("{"+rns+"}")}
    tables = update_draft_tables(deck, rows, report_month)
    remaining={value for node in deck.slides[6]._element.iter() for attr,value in node.attrib.items() if attr.startswith("{"+rns+"}")}
    for relid in ole_relations-remaining:
        deck.slides[6].part.drop_rel(relid)
    texts = update_draft_narrative(deck, charts, rows, report_month, cutoff)
    missing = [r["id"] for r in contract["source_requests"] if not observations.get(r["id"])]
    output.parent.mkdir(parents=True, exist_ok=True); deck.save(str(output))
    ledger = {"schema_version": 2, "contract_version": CONTRACT_VERSION, "delivery_mode": MODE,
              "formal_delivery_allowed": False, "business_complete": False, "draft_content_complete": True,
              "report_month": report_month, "cutoff_date": cutoff.isoformat(), "run_id": source_root.name,
              "reference_sha256": engine.digest(reference), "pptx_sha256": engine.digest(output),
              "source_package_sha256": hashlib.sha256(json.dumps(package, ensure_ascii=False, sort_keys=True).encode()).hexdigest(),
              "source_root": str(source_root), "summary": {"slides": 33, "charts": 37, "tables": 2,
              "required_requests": len(contract["source_requests"]), "verified_requests": usable},
              "coverage": {"required": len(contract["source_requests"]), "verified": usable},
              "charts": charts, "table_updates": tables, "slides": texts, "issues": [], "missing_requests": missing,
              "forecast_fit_policy": "未执行预测；不使用模板历史；仅展示已独立核验公开实际"}
    engine.dump(ledger_path, ledger)
    return {"pptx": str(output), "ledger": str(ledger_path), "delivery_mode": MODE, "formal_delivery_allowed": False,
            "draft_content_complete": True, "business_complete": False, "summary": ledger["summary"],
            "coverage": ledger["coverage"], "missing_requests": missing, "issues": []}
