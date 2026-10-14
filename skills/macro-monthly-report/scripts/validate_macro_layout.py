"""Read-only, macro-specific Office text layout acceptance gate.

Run from the Skill run root:
    validate_macro_layout.py outputs/report.pptx working/macro-layout-validation.json [working/ledger.json]

No static/font-size estimate can grant a pass. Actual Office text bounds are
measured from the same SHA-bound artifact, then independently checked here.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import uuid
from datetime import datetime, timezone

from pptx import Presentation
from macro_contract import TEXT_SLOTS


SCRIPT = Path(__file__).with_suffix(".ps1")
TOLERANCE_PT = 2.0


def workspace_path(path, root, *, must_exist=False):
    root = Path(root).resolve()
    candidate = Path(path)
    candidate = (root / candidate if not candidate.is_absolute() else candidate).resolve()
    if not candidate.is_relative_to(root) or candidate == root:
        raise ValueError("macro_layout_path_outside_run_root")
    if must_exist and not candidate.is_file():
        raise FileNotFoundError(str(candidate))
    return candidate


def _all_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from _all_shapes(shape.shapes)


def build_scope(deck, digest, ledger=None):
    """Select reviewed macro data slots, not template decorations/slide IDs."""
    roles = {}
    if ledger and ledger.get("mode") == "historical_update":
        for update in ledger.get("slides", []):
            if update.get("new", "").strip():
                roles[(int(update["slide"]), int(update["shape_id"]))] = update.get("role", "body")
        # Inspect all remaining data text, including historical policy boxes.
        for n, slide in enumerate(deck.slides, 1):
            if n in {2, 30, 31, 32}:
                for shape in _all_shapes(slide.shapes):
                    if getattr(shape, "has_text_frame", False) and len(shape.text.strip()) > 20:
                        roles.setdefault((n, shape.shape_id), "body")
    else:
        return _legacy_build_scope(deck, digest, ledger)
    return _scope_from_roles(deck, digest, roles)


def _legacy_build_scope(deck, digest, ledger=None):
    roles = {}
    for slide, (title, bodies, source) in TEXT_SLOTS.items():
        roles[(slide, title)] = "title"
        roles[(slide, source)] = "source"
        for body in bodies:
            roles[(slide, body)] = "body"
    roles.update({(7, 12): "title", (7, 9): "source", (8, 9): "title", (8, 26): "source"})
    for slide, shapes in ((4, (6, 7)), (28, (9, 10)), (30, (12, 17))):
        for shape in shapes:
            roles[(slide, shape)] = "caption"
    for shape in (8, 22, 23, 24):
        roles[(8, shape)] = "caption"
    for shape in (2, 4, 11, 15, 32, 35, 36, 38):
        roles[(31, shape)] = "title"
    for shape in (5, 13, 30, 34):
        roles[(31, shape)] = "body"
    for shape in (8, 11):
        roles[(32, shape)] = "caption"
    if ledger:
        for update in ledger.get("slides", []):
            if ledger.get("delivery_mode")=="public_draft" and not str(update.get("new", "")).strip():
                continue  # intentionally removed stale literal annotations
            key = (int(update["slide"]), int(update["shape_id"]))
            if key[0] in {slot[0] for slot in roles}:
                roles.setdefault(key, "caption")
    return _scope_from_roles(deck, digest, roles)


def _scope_from_roles(deck, digest, roles):
    targets = [{"slide": slide, "shape_id": shape, "role": role, "kind": "text"}
               for (slide, shape), role in sorted(roles.items())]
    tables_found = set()
    for slide in (7, 32):
        if slide > len(deck.slides):
            continue
        for shape in _all_shapes(deck.slides[slide - 1].shapes):
            if not getattr(shape, "has_table", False):
                continue
            tables_found.add(slide)
            table = shape.table
            for row in range(len(table.rows)):
                for column in range(len(table.columns)):
                    cell=table.cell(row,column)
                    if cell.is_spanned:
                        continue
                    targets.append({"slide": slide, "shape_id": shape.shape_id, "role": "table_cell", "kind": "table_cell",
                                    "cell_row": row + 1, "cell_column": column + 1,
                                    "row_span":cell.span_height if cell.is_merge_origin else 1,
                                    "column_span":cell.span_width if cell.is_merge_origin else 1})
    scope_issues = []
    if len(deck.slides) != 33:
        scope_issues.append({"code": "macro_slide_count_unexpected", "slides": len(deck.slides)})
    for slide in ({7, 32} - tables_found):
        scope_issues.append({"code": "required_macro_table_missing", "slide": slide})
    return {"schema_version": 1, "artifact_sha256": digest, "scope": "macro_data_slots_only",
            "targets": targets, "issues": scope_issues}


def _rect(value):
    rect = {key: float(value[key]) for key in ("x", "y", "width", "height")}
    if not all(math.isfinite(number) for number in rect.values()) or rect["width"] < 0 or rect["height"] < 0:
        raise ValueError("invalid_text_bounds")
    return rect


def _intersection(first, second):
    return (max(0.0, min(first["x"] + first["width"], second["x"] + second["width"]) - max(first["x"], second["x"])),
            max(0.0, min(first["y"] + first["height"], second["y"] + second["height"]) - max(first["y"], second["y"])))


def audit_measurements(raw, scope):
    issues = [{**item, "severity": "error"} for item in scope.get("issues", [])]
    issues.extend({**item, "severity": "error"} for item in raw.get("issues", []))
    records = []
    if not raw.get("success"):
        issues.append({"code": "office_measurement_not_completed", "severity": "error"})
    if not raw.get("office_passed") or any(raw.get(key) != 33 for key in ("slides", "rendered", "roundtrip_slides")):
        issues.append({"code": "macro_office_render_or_roundtrip_not_passed", "severity": "error"})
    if raw.get("artifact_sha256") != scope["artifact_sha256"]:
        issues.append({"code": "layout_artifact_hash_mismatch", "severity": "error"})
    if raw.get("measurement_id") != scope.get("measurement_id"):
        issues.append({"code": "layout_measurement_invocation_mismatch", "severity": "error"})
    if len(raw.get("measurements", [])) != len(scope["targets"]):
        issues.append({"code": "layout_scope_measurements_missing", "severity": "error"})
    expected = {(item["slide"], item["shape_id"], item.get("cell_row"), item.get("cell_column")): item for item in scope["targets"]}
    seen = set()
    for entry in raw.get("measurements", []):
        record = {**entry, "artifact_sha256": scope["artifact_sha256"], "issues": []}
        key = (entry["slide"], entry["shape_id"], entry.get("cell_row"), entry.get("cell_column"))
        reference = expected.get(key)
        if reference is None or key in seen:
            record["issues"].append({"code": "unexpected_or_duplicate_measurement", "severity": "error"})
        seen.add(key)
        if not entry.get("measured"):
            record["issues"].append({"code": "text_measurement_unavailable", "severity": "error"})
        else:
            try:
                if not entry.get("visible") or abs(float(entry.get("rotation", 0))) > 0.01:
                    raise ValueError("text_visibility_or_rotation_unreviewed")
                box = _rect(entry["text_box"])
                bounds = _rect(entry["rendered_bounds"])
                margins = entry["text_box"]
                inner = {"x": box["x"] + float(margins["margin_left"]), "y": box["y"] + float(margins["margin_top"]),
                         "width": box["width"] - float(margins["margin_left"]) - float(margins["margin_right"]),
                         "height": box["height"] - float(margins["margin_top"]) - float(margins["margin_bottom"])}
                record["inner_text_box"] = inner
                if reference and reference["kind"] == "text" and not str(entry.get("text", "")).strip():
                    record["issues"].append({"code": "required_macro_text_empty", "severity": "error"})
                if str(entry.get("text", "")).strip():
                    if bounds["width"] <= 0 or bounds["height"] <= 0:
                        raise ValueError("nonempty_text_bounds_empty")
                    # Office includes a small glyph/paragraph padding. Ignore
                    # only two points, not font shrinking or multi-line spill.
                    spills = {"left": inner["x"] - bounds["x"], "top": inner["y"] - bounds["y"],
                              "right": bounds["x"] + bounds["width"] - inner["x"] - inner["width"],
                              "bottom": bounds["y"] + bounds["height"] - inner["y"] - inner["height"]}
                    if max(spills.values()) > TOLERANCE_PT:
                        record["issues"].append({"code": "text_box_overflow", "severity": "error", "spill_points": spills})
                    width = float(raw["slide_width"]); height = float(raw["slide_height"])
                    if bounds["x"] < -TOLERANCE_PT or bounds["y"] < -TOLERANCE_PT or bounds["x"] + bounds["width"] > width + TOLERANCE_PT or bounds["y"] + bounds["height"] > height + TOLERANCE_PT:
                        record["issues"].append({"code": "text_outside_slide", "severity": "error"})
                    if reference and reference["role"] in {"body", "title", "source"}:
                        for obstacle in raw.get("obstacles", []):
                            if obstacle["slide"] != entry["slide"] or obstacle["shape_id"] == entry["shape_id"]:
                                continue
                            overlap_width, overlap_height = _intersection(bounds, _rect(obstacle["box"]))
                            if overlap_width > 4 and overlap_height > TOLERANCE_PT:
                                record["issues"].append({"code": "text_overlaps_" + obstacle["kind"], "severity": "error",
                                                         "obstacle_shape_id": obstacle["shape_id"], "overlap_points": [overlap_width, overlap_height]})
            except (ValueError, KeyError, TypeError) as exc:
                record["issues"].append({"code": "text_measurement_unusable", "severity": "error", "detail": str(exc)})
        records.append(record)
    for missing in expected.keys() - seen:
        issues.append({"code": "required_text_measurement_missing", "severity": "error", "slide": missing[0], "shape_id": missing[1], "cell_row": missing[2], "cell_column": missing[3]})
    # Catch a long body hitting its nearby source/caption/title even when no
    # native chart rectangle occupies that location.
    visible = [item for item in records if item.get("measured") and item.get("text", "").strip()]
    for index, first in enumerate(visible):
        for second in visible[index + 1:]:
            if first["slide"] != second["slide"] or first["kind"] == "table_cell" or second["kind"] == "table_cell":
                continue
            try:
                overlap_width, overlap_height = _intersection(_rect(first["rendered_bounds"]), _rect(second["rendered_bounds"]))
            except (ValueError, KeyError, TypeError):
                continue
            if overlap_width > 4 and overlap_height > TOLERANCE_PT:
                first["issues"].append({"code": "visible_text_overlap", "severity": "error", "other_shape_id": second["shape_id"], "overlap_points": [overlap_width, overlap_height]})
    for record in records:
        issues.extend({**issue, "slide": record["slide"], "shape_id": record["shape_id"],
                       "cell_row": record.get("cell_row"), "cell_column": record.get("cell_column")} for issue in record["issues"])
    return {"passed": not issues and bool(records), "measurements": records, "issues": issues}


def validate_layout(input_path, report_path, ledger_path=None, *, run_root=None, timeout=180, runner=subprocess.run):
    root = Path(run_root or os.environ.get("SKILL_RUN_ROOT") or Path.cwd()).resolve()
    report = workspace_path(report_path, root)
    report.parent.mkdir(parents=True, exist_ok=True)
    result = {"schema_version": 1, "gate": "macro_layout", "passed": False, "status": "unavailable",
              "office_passed": False, "slides": 0, "rendered": 0, "roundtrip_slides": 0,
              "validated_at": datetime.now(timezone.utc).isoformat(), "issues": [], "measurements": []}
    try:
        pptx = workspace_path(input_path, root, must_exist=True)
        digest = hashlib.sha256(pptx.read_bytes()).hexdigest()
        result.update({"artifact_path": str(pptx), "artifact_sha256": digest})
        ledger = None
        if ledger_path:
            ledger = json.loads(workspace_path(ledger_path, root, must_exist=True).read_text(encoding="utf-8-sig"))
            if ledger.get("pptx_sha256") != digest:
                raise ValueError("macro_layout_ledger_artifact_hash_mismatch")
        scope = build_scope(Presentation(str(pptx)), digest, ledger)
        scope["measurement_id"] = str(uuid.uuid4())
        measurements_dir = workspace_path(report.parent / "macro-layout", root)
        measurements_dir.mkdir(parents=True, exist_ok=True)
        scope_path = measurements_dir / "scope.json"
        scope_path.write_text(json.dumps(scope, ensure_ascii=False, indent=2), encoding="utf-8")
        raw_path = measurements_dir / "powerpoint-measurements.json"
        process = runner(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(SCRIPT),
                          "-InputPath", str(pptx), "-ScopePath", str(scope_path), "-OutputDir", str(measurements_dir)],
                         capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout, check=False)
        process_receipt = measurements_dir / "office-process.json"
        process_receipt.write_text(json.dumps({"returncode": process.returncode, "stdout": process.stdout[-5000:],
                                              "stderr": process.stderr[-5000:], "measurement_id": scope["measurement_id"]},
                                             ensure_ascii=False, indent=2), encoding="utf-8")
        # A previous audit cannot be used if this invocation failed before
        # writing its measurements; the PS payload also binds the input SHA.
        lines = [line.strip() for line in process.stdout.splitlines() if line.strip()]
        payload = json.loads(lines[-1]) if lines else {}
        if Path(payload.get("full_report", "")).resolve() != raw_path.resolve() or not raw_path.is_file():
            raise ValueError("office_measurement_report_not_returned: " + (process.stderr or process.stdout or "empty Office process output")[-1800:])
        raw = json.loads(raw_path.read_text(encoding="utf-8-sig"))
        result.update(audit_measurements(raw, scope))
        result.update({key: raw.get(key) for key in ("office_passed", "slides", "rendered", "roundtrip_slides")})
        result["office_passed"] = bool(raw.get("office_passed") and raw.get("artifact_sha256") == digest
                                      and raw.get("measurement_id") == scope["measurement_id"]
                                      and all(raw.get(key) == 33 for key in ("slides", "rendered", "roundtrip_slides")))
        result["status"] = "passed" if result["passed"] else "blocked" if raw.get("success") else "unavailable"
        result["measurement_report"] = str(raw_path)
        if process.returncode != 0 and result["passed"]:
            result["passed"] = False; result["status"] = "unavailable"
            result["issues"].append({"code": "office_process_failed", "severity": "error", "detail": (process.stderr or process.stdout)[-2000:]})
        if hashlib.sha256(pptx.read_bytes()).hexdigest() != digest:
            result["passed"] = False; result["office_passed"] = False; result["status"] = "blocked"
            result["issues"].append({"code": "artifact_changed_during_layout_validation", "severity": "error"})
    except (OSError, ValueError, KeyError, subprocess.TimeoutExpired) as exc:
        result["issues"].append({"code": "macro_layout_validation_unavailable", "severity": "error", "detail": str(exc)})
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return {**result, "full_report": str(report)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input_path")
    parser.add_argument("report_path")
    parser.add_argument("ledger_path", nargs="?")
    args = parser.parse_args()
    try:
        result = validate_layout(args.input_path, args.report_path, args.ledger_path)
    except (ValueError, OSError) as exc:
        result = {"gate": "macro_layout", "passed": False, "status": "unavailable", "issues": [{"code": "run_root_path_guard", "detail": str(exc)}]}
    print(json.dumps({key: value for key, value in {**result, "issues": result.get("issues", [])[:25]}.items() if key != "measurements"}, ensure_ascii=False))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
