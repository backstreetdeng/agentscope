"""Independent content gate for the macro forecast baseline artifact."""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

from pptx import Presentation


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: validate_forecast_baseline.py <pptx> <forecast-ledger.json>")
    run_root = Path(os.environ.get("SKILL_RUN_ROOT", Path.cwd())).resolve()
    pptx_path, ledger_path = map(Path, sys.argv[1:])
    pptx_path = (run_root / pptx_path).resolve() if not pptx_path.is_absolute() else pptx_path.resolve()
    ledger_path = (run_root / ledger_path).resolve() if not ledger_path.is_absolute() else ledger_path.resolve()
    if run_root not in pptx_path.parents or run_root not in ledger_path.parents:
        raise ValueError("PPTX 和台账必须位于本次 Skill 运行工作区内")
    ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
    report_month = str(ledger.get("report_month", ""))
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", report_month):
        raise SystemExit("台账缺少有效 report_month")
    presentation = Presentation(str(pptx_path))
    actual_charts = sum(1 for slide in presentation.slides for shape in slide.shapes if getattr(shape, "has_chart", False))
    expected_charts = int((ledger.get("summary") or {}).get("charts", -1))
    slides = ledger.get("slides") or []
    footer_pages = 0
    for slide in presentation.slides:
        if any("预测版｜目标月：" in getattr(shape, "text", "") for shape in slide.shapes):
            footer_pages += 1
    year, month = report_month.split("-")
    cover_target = f"{year}年{int(month)}月宏观经济分析报告"
    cover_text = "\n".join(
        getattr(shape, "text", "") for shape in presentation.slides[0].shapes
    )
    cover_dates = re.findall(r"20\d{2}年\d{1,2}月", cover_text)
    checks = {
        "slide_count": len(presentation.slides) == 33,
        "chart_count": actual_charts == expected_charts == 37,
        "ledger_slide_count": len(slides) == len(presentation.slides),
        "all_chart_bindings_recorded": sum(len(item.get("charts") or []) for item in slides) == actual_charts,
        "forecast_boundary_on_every_page": footer_pages == len(presentation.slides),
        "cover_month_updated": cover_target in cover_text,
        "cover_metadata_month_consistent": bool(cover_dates) and all(
            item == f"{year}年{int(month)}月" for item in cover_dates
        ),
    }
    result = {"passed": all(checks.values()), "checks": checks, "slides": len(presentation.slides), "charts": actual_charts}
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["passed"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
