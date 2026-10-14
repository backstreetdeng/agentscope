"""Build/replay historical reports with archived provenance and actual-only charts."""
from __future__ import annotations
from copy import deepcopy
from datetime import date, datetime, timedelta
import io
import json
import math
import re
from pathlib import Path
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel
from pptx import Presentation
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Pt
from macro_history import inventory, build_history, archived_points, VERSION, REFERENCE
from macro_engine import C, A, digest, dump, _shape, _range, _series_name, set_text, forecast_series
from macro_contract import parse_period, month_end, static_series, OVERVIEW_ROWS
from inspect_macro_template import all_shapes


def period_end(month):
    return month_end(*map(int,month.split('-')))


def replay_sources(package, contract, root):
    """Keep archived and publisher evidence separate; replay every official proof."""
    from macro_sources import validate_observation, verify_source_corroboration
    import hashlib
    if package.get('synthetic') or package.get('data_kind') in {'fixture','synthetic','test'}:
        raise ValueError('synthetic_data_not_business_evidence')
    pages={}
    for p in (Path(root)/'working/macro-sources/pages').glob('*.json'):
        page=json.loads(p.read_text(encoding='utf-8'))
        if page.get('text_sha256')!=hashlib.sha256(page.get('text','').encode()).hexdigest():
            raise ValueError('source_receipt_hash_mismatch')
        if page.get('fetch_provider') not in {'direct-primary-publisher','direct-public-publisher'} or not page.get('fetched_at'):
            raise ValueError('source_receipt_not_live')
        pages[page['receipt_id']]=page
    requests={r['id']:r for r in contract['source_requests']}
    checked=verify_source_corroboration(package,contract,pages)
    if checked['issues']:
        raise ValueError('source_corroboration_failed:'+str(checked['issues'][:3]))
    values=[]
    for o in checked['observations']:
        req=requests[o['request_id']]
        v=validate_observation(o,req,pages[o['receipt_id']],date.fromisoformat(package['cutoff_date']),package['report_month'])
        values.append({**v,'period':parse_period(v['period'],req['frequency']).isoformat(),
                       'provenance':'official_actual'})
    from macro_historical_derivations import replay
    direct={(v['request_id'],v['period']) for v in values}
    values.extend(v for v in replay(pages,package['report_month'],package['cutoff_date'])
                  if (v['request_id'],v['period']) not in direct)
    return values


def merged_facts(contract, history, actuals, month):
    target=period_end(month)
    facts={(o['request_id'],o['period']):deepcopy(o) for o in history['observations'] if o['period']<=target.isoformat()}
    # New-period actuals always win over archived history, with a retained revision trail.
    revisions=[]
    for o in actuals:
        key=(o['request_id'],o['period'])
        if key in facts and not math.isclose(facts[key]['value'],o['value'],abs_tol=1e-8):
            revisions.append({'old':facts[key],'new':o})
        facts[key]=deepcopy(o)
    return facts,revisions


def chart_values(chart, facts, month):
    sem=chart['semantics'];frequency=chart['frequency'];target=period_end(month)
    if frequency=='quarterly':target=month_end(target.year,(target.month//3)*3) if target.month>=3 else month_end(target.year-1,12)
    requests=chart['requests'];out=[]
    def lookup(rid,day):return facts.get((rid,day.isoformat()))
    def value(rid,day):
        f=lookup(rid,day);return f['value'] if f else None
    if sem in {'time','mixed'}:
        original=[parse_period(c,frequency) for c in chart['categories']]
        days=sorted({*original,*[date.fromisoformat(p) for rid,p in facts if rid in {r['id'] for r in requests} and p<=target.isoformat()]})
        days=[d for d in days if d<=target][-len(original):]
        if frequency=='monthly' and target not in days:days=(days+[target])[-len(original):]
        categories=[(d-date(1899,12,30)).days for d in days]
        for s in chart['series']:
            if static_series(s['name'],s['values']):
                vals=[50.0]*len(days);rids=[]
            else:
                req=next(r for r in requests if r.get('display_series_name',r['series_name'])==s['name'])
                vals=[value(req['id'],d) for d in days];rids=[req['id']]
            out.append({'original_name':s['name'],'name':s['name'],'values':vals,'request_ids':rids})
    elif sem in {'snapshot','category_metrics'} and any(lookup(r['id'],target) is None for r in requests):
        # A partially evidenced cross-section cannot borrow an old value under
        # a new month label. Keep the complete dated template chart instead.
        categories=chart['categories']
        out=[{'original_name':s['name'],'name':s['name'],'values':s['values'],
              'request_ids':[r['id'] for r in requests]} for s in chart['series']]
    elif sem in {'snapshot','category_metrics'}:
        categories=chart['categories']
        if chart['id']=='s24-shape17':categories=['居民短期','居民中长期','企业短期','企业中长期','票据']
        dates=sorted(parse_period(s['name']) for s in chart['series']) if sem=='snapshot' else []
        for s in chart['series']:
            day=target
            if sem=='snapshot':
                rank=dates.index(parse_period(s['name']));ordinal=target.year*12+target.month-1-(len(dates)-1-rank)
                day=month_end(ordinal//12,ordinal%12+1)
            rids=[next(r['id'] for r in requests if r.get('category_index')==i and
                  (sem=='snapshot' or r['series_name']==s['name'])) for i in range(len(categories))]
            out.append({'original_name':s['name'],'name':day.strftime('%Y-%m') if sem=='snapshot' else s['name'],
                        'values':[value(r,day) for r in rids],'request_ids':rids})
    else:
        categories=chart['categories']
        for s in chart['series']:
            if chart['id']=='s28-shape11':
                year=int(s['name']) if s['name'].isdigit() else target.year
                rid='auto-sales::'+('current' if s['name'].isdigit() else 'yoy')
            else:
                year=int(s['name']);rid=requests[0]['id']
            vals=[]
            for i in range(len(categories)):
                try:day=date.fromisocalendar(year,i+1,7) if frequency=='weekly' else month_end(year,i+1)
                except ValueError:vals.append(None);continue
                vals.append(value(rid,day) if day<=target else None)
            out.append({'original_name':s['name'],'name':s['name'],'values':vals,'request_ids':[rid]})
        if chart['id']=='s28-shape11':
            current=next((s for s in out if s['name']==month[:4]),None)
            yoy=next((s for s in out if s['name']=='销量当月同比'),None)
            if current and yoy:
                yoy['values']=[v if current['values'][i] is not None else None for i,v in enumerate(yoy['values'])]
    return {'chart_id':chart['id'],'physical_id':chart['physical_id'],'slide':chart['slide'],
            'shape_id':chart['shape_id'],'semantics':sem,'frequency':frequency,'categories':categories,
            'series':out,'definition':chart['definition']}


def write_chart(chart, values):
    """Replace caches and embedded data together, retaining plots and series axes.

    Literal-only charts receive an embedded workbook; oversized old workbooks
    are never materialized. No forecast series is added to an actual-data chart.
    """
    xml=chart._chartSpace;wb=Workbook();ws=wb.active;ws.title='Data'
    ext=xml.find(C('externalData'))
    if ext is None:
        ext=etree.SubElement(xml,C('externalData'))
        part=EmbeddedXlsxPart.new(b'',chart.part.package)
        ext.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id',chart.part.relate_to(part,RT.PACKAGE))
    else:
        part=chart.part.related_part(ext.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
    for index,(node,item) in enumerate(zip(xml.xpath('.//c:ser'),values['series'])):
        col=index+2;letter=get_column_letter(col)
        tx=node.find(C('tx'))
        if tx is None:tx=etree.Element(C('tx'));node.insert(2,tx)
        for ch in list(tx):tx.remove(ch)
        etree.SubElement(tx,C('v')).text=item['name'];ws.cell(1,col).value=item['name']
        cat=node.find(C('cat'))
        if cat is None:
            cat=etree.Element(C('cat'));node.insert(list(node).index(node.find(C('val'))),cat)
        for ch in list(cat):cat.remove(ch)
        numeric=values['semantics'] in {'time','mixed'}
        ref=etree.SubElement(cat,C('numRef' if numeric else 'strRef'))
        etree.SubElement(ref,C('f')).text=f"Data!$A$2:$A${len(values['categories'])+1}"
        cache=etree.SubElement(ref,C('numCache' if numeric else 'strCache'))
        if numeric:etree.SubElement(cache,C('formatCode')).text='yyyy-mm'
        etree.SubElement(cache,C('ptCount'),val=str(len(values['categories'])))
        for i,c in enumerate(values['categories']):
            ws.cell(i+2,1).value=c
            pt=etree.SubElement(cache,C('pt'),idx=str(i));etree.SubElement(pt,C('v')).text=str(c)
        val=node.find(C('val'))
        for ch in list(val):val.remove(ch)
        ref=etree.SubElement(val,C('numRef'));etree.SubElement(ref,C('f')).text=f'Data!${letter}$2:${letter}${len(item["values"])+1}'
        cache=etree.SubElement(ref,C('numCache'));etree.SubElement(cache,C('formatCode')).text='General'
        etree.SubElement(cache,C('ptCount'),val=str(len(item['values'])))
        for i,v in enumerate(item['values']):
            ws.cell(i+2,col).value=v
            if v is not None:
                pt=etree.SubElement(cache,C('pt'),idx=str(i));etree.SubElement(pt,C('v')).text=str(v)
    # Removing stale manual endpoint labels does not remove chart styles/axes.
    for node in xml.xpath('.//c:dLbl | .//c:ser/c:extLst'):
        node.getparent().remove(node)
    for labels in xml.xpath('.//c:dLbls'):
        show=labels.find(C('showVal'))
        if show is None:show=etree.SubElement(labels,C('showVal'))
        show.set('val','0')
    blank=xml.find('./'+C('chart')+'/'+C('dispBlanksAs'))
    if blank is None:blank=etree.SubElement(xml.find(C('chart')),C('dispBlanksAs'))
    blank.set('val','gap')
    epoch=xml.find(C('date1904'))
    if epoch is not None:epoch.set('val','0')
    for axis in xml.xpath('.//c:dateAx | .//c:valAx'):
        for bound in axis.xpath('./c:scaling/c:min | ./c:scaling/c:max | ./c:majorUnit | ./c:minorUnit'):
            bound.getparent().remove(bound)
    if values['semantics'] in {'time','mixed'}:
        for axis in xml.xpath('.//c:catAx | .//c:dateAx'):
            fmt=axis.find(C('numFmt'))
            if fmt is None:fmt=etree.SubElement(axis,C('numFmt'))
            fmt.set('formatCode','yyyy-mm');fmt.set('sourceLinked','0')
            tick=axis.find(C('tickLblPos'))
            if tick is not None:tick.set('val','nextTo')
            orient=axis.find('./'+C('scaling')+'/'+C('orientation'))
            if orient is not None:orient.set('val','minMax')
            skip=axis.find(C('tickLblSkip'))
            if axis.tag==C('catAx'):
                if skip is None:skip=etree.SubElement(axis,C('tickLblSkip'))
                skip.set('val',str(max(1,math.ceil(len(values['categories'])/8))))
    if values['chart_id']=='hf-land':
        for axis in xml.xpath('.//c:valAx/c:axPos'):axis.set('val','l')
    if values['chart_id'] in {'hf-land','s8-shape21'}:
        for axis in xml.xpath('.//c:catAx | .//c:dateAx'):
            tick=axis.find(C('tickLblPos'))
            if tick is not None:tick.set('val','low')
        for labels in xml.xpath('.//c:legend//c:tx//a:t'):
            if labels.text:labels.text=labels.text.replace('（右轴）','').replace('(右轴)','')
    if values['chart_id']=='s11-shape15':
        layout=xml.find('./'+C('chart')+'/'+C('plotArea')+'/'+C('layout'))
        if layout is None:layout=etree.SubElement(xml.find('./'+C('chart')+'/'+C('plotArea')),C('layout'))
        for child in list(layout):layout.remove(child)
        manual=etree.SubElement(layout,C('manualLayout'))
        etree.SubElement(manual,C('layoutTarget'),val='inner')
        for key in ('xMode','yMode','wMode','hMode'):etree.SubElement(manual,C(key),val='factor')
        for key,value in [('x','.25'),('y','.15'),('w','.50'),('h','.70')]:etree.SubElement(manual,C(key),val=value)
    stream=io.BytesIO();wb.save(stream);part._blob=stream.getvalue()


def facts_for_page(charts, contract, facts, month):
    result=[];target=period_end(month)
    previous=month_end(target.year if target.month>1 else target.year-1,target.month-1 if target.month>1 else 12)
    for c in charts:
        for r in c['requests']:
            points=sorted([v for (rid,p),v in facts.items() if rid==r['id'] and p<=target.isoformat()],key=lambda x:x['period'])
            if not points:continue
            current=points[-1];prior=next((v for v in points if v['period']==previous.isoformat()),None)
            if r['frequency'] in {'weekly','daily','quarterly'}:prior=points[-2] if len(points)>1 else None
            label=r['definition'].split('：')[-1]
            if r.get('category_index') is not None:label=str(c['categories'][r['category_index']])+(r['series_name'] if c['semantics']=='category_metrics' else '')
            result.append({'id':r['id'],'label':label,'value':current['value'],'period':current['period'],
                'unit':r['unit'],'provenance':current['provenance'],'prior':prior,
                'approximate':current.get('approximate',False),'derivation':current.get('derivation'),
                'statistical_definition':r['definition'],
                'delta':current['value']-prior['value'] if prior else None,'source':current.get('source_url',current.get('source_file'))})
    return result


def fmt(v):return f'{v:.2f}'.rstrip('0').rstrip('.')


def theme(pagefacts, topic, month):
    """Compositional thematic claims from identified facts, without causal invention."""
    m=int(month[-2:]);current=[a for a in pagefacts if a['period'][:7]==month]
    f=current if current else pagefacts
    def pick(word):return next((x for x in f if word in x['label']),None)
    def trend(x):
        d=x.get('delta');return '保持平稳' if d is None or abs(d)<1e-8 else '增速回升' if d>0 else '增速回落'
    x=pick('当月同比') or (f[0] if f else None)
    if not x:return {'headline':f'{m}月{topic}：本期数据待核验','body':'本期尚未取得精确口径的新增数据，保留历史观测期，不能形成当月结论。','fact_ids':[]}
    if x['period'][:7]!=month and topic!='经济增长':
        return {'headline':f'{m}月{topic}：本期取证未完成，历史更新至{x["period"][:7]}',
                'body':f'本期暂无已核验的新增数据。历史最新观测为{x["period"]}，不能将其改写为{m}月实际。',
                'fact_ids':[x['id']],'method':'period_boundary','causal_claims':[]}
    if topic=='PMI':
        x=next((a for a in f if a['label']=='PMI'),x)
        zone='扩张区间' if x['value']>50 else '收缩区间' if x['value']<50 else '荣枯线'
        title=f'{m}月制造业景气{trend(x).replace("增速","")}，处于{zone}'
    elif topic=='工业生产':
        title=f'{m}月工业{trend(x)}'
        hi=pick('高技术')
        total_cum=next((a for a in f if a['label']=='累计同比'),None)
        if hi:title+='，高技术制造业累计'+('增速高于整体' if total_cum and hi['value']>total_cum['value'] else '保持增长' if hi['value']>0 else trend(hi))
    elif topic=='投资':
        x=pick('固定资产投资累计同比') or x
        title=f'1—{m}月投资'+('降幅扩大' if x['value']<0 and x.get('delta') is not None and x['delta']<0 else trend(x))
        sectors=[p for p in f if p['label'] in {'房地产','制造业','基建'}]
        if sectors:title+='，分领域走势'+('均偏弱' if all(a['value']<0 for a in sectors) else '仍有分化')
    elif topic=='消费':title=f'{m}月消费{trend(x)}，'+('社零仍保持增长' if x['value']>0 else '社零同比下降')
    elif topic=='消费结构':
        up=[a for a in f if a['value']>0];down=[a for a in f if a['value']<0]
        title=f'{m}月消费结构'+('分化延续' if up and down else '整体改善' if not down else '偏弱')
        car=pick('汽车')
        if car:title+='，汽车零售'+('同比下降' if car['value']<0 else '保持增长')
    elif topic=='房地产':title=f'{m}月房地产'+('销售仍在调整' if (pick('销售面积') or x)['value']<0 else '销售边际改善')
    elif topic=='房价':title=f'{m}月房价走势仍有城市层级差异'
    elif topic=='进出口':
        x=pick('出口') or x;title=f'{m}月出口{trend(x)}，'+('仍保持增长' if x['value']>0 else '同比下降')
    elif topic in {'CPI','PPI'}:
        x=pick('核心') if topic=='CPI' and not pick('CPI:当月同比') else pick('CPI:当月同比') if topic=='CPI' else pick('同比')
        x=x or f[0];title=f'{m}月{topic}同比{trend(x).replace("增速","")}，价格走势需结合分项观察'
        food=pick('食品:');mom=pick('环比')
        if topic=='CPI' and food:title=f'{m}月CPI同比{trend(x).replace("增速","")}，食品价格'+('仍同比下降' if food['value']<0 else '同比上涨')
        if topic=='PPI' and mom:title=f'{m}月PPI同比{trend(x).replace("增速","")}，环比'+('上涨' if mom['value']>0 else '下降' if mom['value']<0 else '持平')
    elif topic=='新增人民币贷款':title=f'{m}月信贷'+('净减少' if x['value']<0 else '实现净增加')+'，居民与企业融资分项需分别观察'
    elif topic=='社融':
        total=next((a for a in f if a['id']=='s25-shape9::current'),x);gov=pick('政府债券')
        title=f'{m}月社融'+trend(total)
        if gov and total['value']>0:title+='，政府债券融资占比'+fmt(gov['value']/total['value']*100)+'%'
    elif topic=='货币':
        a=pick('M1');b=pick('M2');title=f'{m}月货币增速分化'
        if a and b:
            gap=b['value']-a['value'];priorgap=b['prior']['value']-a['prior']['value'] if a['prior'] and b['prior'] else None
            word='收窄' if priorgap is not None and gap<priorgap else '扩大' if priorgap is not None and gap>priorgap else '保持平稳'
            title=f'{m}月M2—M1增速差{word}至{fmt(gap)}个百分点'
    elif topic=='汽车市场':
        sales=next((a for a in f if a['id']=='auto-sales::yoy'),None)
        if sales:
            x=sales;title=f'{m}月汽车销量'+('同比增长' if x['value']>0 else '同比下降' if x['value']<0 else '同比持平')+'，产销与利润走势分别观察'
        else:title=f'{m}月汽车销量本期取证未完成，产量与利润分列观察'
    elif topic=='财政货币':title=f'{m}月财政收支与利率：关注已发布数据变化'
    elif topic=='海外经济':
        title=f'{m}月全球制造业与主要经济体景气仍有差异' if any(a['id'].startswith('s4-shape5') for a in current) else f'{m}月海外景气取证未完成，各经济体按实际观测期列示'
    else:title=f'{m}月经济运行：供需与结构变化需结合观察'
    chosen=[x];seen={(x['label'],x['period'],x['unit'],x['value'])}
    for a in f:
        key=(a['label'],a['period'],a['unit'],a['value'])
        if key not in seen:chosen.append(a);seen.add(key)
        if len(chosen)==3:break
    sentences=[]
    for a in chosen:
        unit=a['unit'];v=a['value'];label=a['label'].replace('(%,右轴)','').replace('(右轴)','')
        if unit=='亿元' and abs(v)>=10000:v/=10000;unit='万亿元'
        if unit=='辆':v/=10000;unit='万辆'
        period=a['period'][:7].replace('-','年')+'月'
        if '累计' in a.get('statistical_definition',''):period=f'1—{int(a["period"][5:7])}月累计'
        line=f'{period}{label}为'+('约' if a.get('approximate') else '')+f'{fmt(v)}{unit}'
        if a['delta'] is not None:
            changeunit='个百分点' if a['unit']=='%' else '点' if a['unit']=='指数' else a['unit']
            line+='，较上期'+('上升' if a['delta']>0 else '下降' if a['delta']<0 else '持平')+(('约' if a.get('approximate') else '')+fmt(abs(a['delta']))+changeunit if a['delta'] else '')
        sentences.append(line)
    return {'headline':title,'body':'；'.join(sentences)+'。','fact_ids':[a['id'] for a in f],
            'method':'deterministic_fact_composition','causal_claims':[]}


TOPICS={'s4-shape5':'海外经济','s10-shape7':'经济增长','s11-shape12':'PMI','s13-shape10':'工业生产',
 's14-shape15':'投资','s15-shape9':'消费','s16-shape8':'消费结构','s17-shape15':'房地产',
 's18-shape17':'房价','s19-shape13':'进出口','s21-shape11':'CPI','s22-shape9':'PPI',
 's24-shape9':'新增人民币贷款','s25-shape9':'社融','s26-shape11':'货币','s28-shape11':'汽车市场','s30-shape13':'财政货币'}


def text_slots(contract):
    slots={}
    for s in contract['slides']:
        if not s['charts'] or s['slide']==8:continue
        texts=s['texts'];source=[t for t in texts if '数据来源' in t['text']]
        bodies=[t for t in texts if len(t['text'])>65 and '数据来源' not in t['text']]
        titles=[t for t in texts if 10<len(t['text'])<=65 and '数据来源' not in t['text'] and '|' not in t['text'] and not any(w in t['text'] for w in ['（%）','(%)','累计值','指数：','变动(万辆)','利润(%)'])]
        if bodies and titles:
            title=min(titles,key=lambda t:t['y']);body=min(bodies,key=lambda t:t['y'])
            slots[s['slide']]={'title':title['shape_id'],'body':body['shape_id'],'source':source[0]['shape_id'] if source else None}
    return slots


def render_contents(deck,contract,facts,month,cutoff):
    updates=[];narratives=[];slots=text_slots(contract)
    def put(n,sid,text,role='body'):
        shape=_shape(deck.slides[n-1],sid);old=shape.text;set_text(shape,text)
        updates.append({'slide':n,'shape_id':sid,'old':old,'new':text,'role':role})
    y,m=map(int,month.split('-'));nextday=month_end(y if m<12 else y+1,m+1 if m<12 else 1)
    put(1,18,f'{y}年{m}月宏观经济分析报告','title');put(1,8,f'{nextday.year}年{nextday.month}月 · 信息截止{cutoff}','source')
    _shape(deck.slides[0],8).width=Pt(620)
    put(7,9,'缺项：暂无可核验数据；历史列保留；--为非发布期。','source')
    for slide in contract['slides']:
        n=slide['slide']
        if n not in slots:continue
        f=facts_for_page(slide['charts'],contract,facts,month)
        topic=next((TOPICS[c['id']] for c in slide['charts'] if c['id'] in TOPICS),'经济运行')
        narrative=theme(f,topic,month);narratives.append({'slide':n,'topic':topic,'facts':f,**narrative})
        slot=slots[n];put(n,slot['title'],narrative['headline'],'title');put(n,slot['body'],narrative['body'])
        if slot['source']:
            latest=max((v['period'] for v in f),default='待核验')
            put(n,slot['source'],f'来源见备注；数据至{latest}','source')
            footer=_shape(deck.slides[n-1],slot['source']);footer.width=min(Pt(760),deck.slide_width-footer.left-Pt(60))
            footer.text_frame.margin_left=footer.text_frame.margin_right=Pt(0)
        deck.slides[n-1].notes_slide.notes_text_frame.text=json.dumps({'cutoff':cutoff,'facts':f,'narrative':narrative},ensure_ascii=False)
    # Data-bearing manual captions and stale axis tick boxes must not carry last month.
    for t in contract['slides'][7]['texts']:
        if re.fullmatch(r'20\d\d-\d\d',t['text']):put(8,t['shape_id'],'','removed')
    high=facts_for_page(contract['slides'][7]['charts'],contract,facts,month)
    high_period=max((x['period'] for x in high),default='待核验')
    put(8,54,f'高频分别截至原始观测日；最新{high_period}；来源与单位见备注','source')
    for sid,text in [(10,'全钢胎开工率（%）'),(12,'唐山高炉开工率（%）'),(14,'地产销售（万平方米）')]:put(8,sid,text,'caption')
    for sid in (6,7):put(4,sid,'指数','caption')
    for sid in (6,7):
        box=_shape(deck.slides[3],sid);box.height=Pt(27)
        box.text_frame.margin_top=box.text_frame.margin_bottom=Pt(0)
        for p in box.text_frame.paragraphs:
            for run in p.runs:run.font.size=Pt(10)
    for sid in (10,12,14):
        box=_shape(deck.slides[7],sid);box.height=Pt(30)
        for p in box.text_frame.paragraphs:
            for run in p.runs:run.font.size=Pt(11)
    # The supplied July template overlays two independent native charts in one
    # panel. Separate land and housing so their axes and units can be read.
    for sid,left in [(76,75),(77,356.7),(78,638.4)]:
        box=_shape(deck.slides[7],sid);box.left=Pt(left);box.width=Pt(252)
        box.chart.has_title=False
    for sid,left in [(18,69),(19,350.7)]:
        box=_shape(deck.slides[7],sid);box.left=Pt(left);box.width=Pt(264)
    frame=deck.slides[7].shapes.add_shape(1,Pt(632.4),Pt(335),Pt(264),Pt(176))
    frame.fill.background();frame.line.color.rgb=__import__('pptx.dml.color',fromlist=['RGBColor']).RGBColor(176,176,176)
    # Put the added border behind the charts, avoiding new opaque plot covers.
    tree=deck.slides[7].shapes._spTree;tree.remove(frame._element);tree.insert(2,frame._element)
    _shape(deck.slides[7],12).width=Pt(264)
    put(8,26,'','removed');put(8,28,'','removed')
    for sid in (25,27):
        _shape(deck.slides[7],sid).width=0
    put(8,14,'土地与商品房成交（万平方米）','caption')
    _shape(deck.slides[7],14).left=Pt(356.7);_shape(deck.slides[7],14).width=Pt(534)
    for sid in (77,78):
        chart=_shape(deck.slides[7],sid).chart;chart.has_legend=True
        from pptx.enum.chart import XL_LEGEND_POSITION
        chart.legend.position=XL_LEGEND_POSITION.TOP;chart.legend.font.size=Pt(8)
        xml=chart._chartSpace
        # Automatic plot layout avoids the template's overlapping secondary
        # category axes and cropped dates after splitting the panels.
        for layout in xml.xpath('.//c:plotArea/c:layout'):
            for child in list(layout):layout.remove(child)
        for axis in xml.xpath('.//c:catAx | .//c:dateAx'):
            delete=axis.find(C('delete'))
            if delete is not None:delete.set('val','0')
            for props in axis.xpath('./c:txPr//a:bodyPr'):
                props.set('rot','-5400000')
        for tick in xml.xpath('.//c:valAx/c:tickLblPos'):tick.set('val','nextTo')
    put(21,7,'PPI及分项价格：同比与环比（%）','caption')
    put(25,9,'M1、M2同比增速（%）','caption')
    put(23,18,'','removed');put(23,8,'各部门新增贷款（亿元；累计差为约数）','caption')
    deck.slides[7].notes_slide.notes_text_frame.text=json.dumps(high,ensure_ascii=False)
    # Summary is produced from the already-completed page facts, not another dataset.
    for sid,topics in [(23,['经济增长','工业生产','投资']),(45,['消费','消费结构']),
                       (54,['CPI','PPI']),(55,['新增人民币贷款','社融','货币']),
                       (43,['汽车市场']),(36,['财政货币']),(37,['财政货币'])]:
        selected=[n for n in narratives if n['topic'] in topics]
        put(2,sid,'；'.join(n['headline'] for n in selected)[:145]+'。')
    put(2,2,f'{m}月经济运行：从生产、需求、价格与信用观察结构变化','title')
    industrial=next((n for n in narratives if n['topic']=='工业生产'),None)
    retail=next((n for n in narratives if n['topic']=='消费'),None)
    if industrial and retail:
        get=lambda n,word:next((f for f in n['facts'] if word in f['label'] and f['period'][:7]==month),None)
        a=get(industrial,'当月同比');b=get(retail,'当月同比')
        if a and b and a['delta'] is not None and b['delta'] is not None:
            primary='工业与社零增速同步回升' if a['delta']>0 and b['delta']>0 else '工业与社零增速双双放缓' if a['delta']<0 and b['delta']<0 else '工业增速回升，消费增速有所放缓' if a['delta']>0 else '消费增速回升，工业增速有所放缓'
            title=f'{m}月{primary}'
            put(2,2,title,'title')
            growth=next((n for n in narratives if n['topic']=='经济增长'),None)
            if growth:put(growth['slide'],slots[growth['slide']]['title'],title,'title')
    fiscal=facts_for_page(contract['slides'][28]['charts'],contract,facts,month)
    for sid,words in [(36,['财政收入','财政支出']),(37,['贷款市场报价利率'])]:
        selected=[f for f in fiscal if any(w in f['label'] for w in words) and f['period'][:7]==month][:2]
        def summary_value(f):
            label=re.sub(r'\([^)]*\)|（[^）]*）','',f['label']);v=f['value'];unit=f['unit']
            if unit=='亿元' and abs(v)>=10000:v/=10000;unit='万亿元';label=f'1—{m}月累计'+label
            return label+'为'+fmt(v)+unit
        put(2,sid,'；'.join(summary_value(f) for f in selected) or '本期相关数据待核验。')
    # Policy pages are kept distinct from fact-derived economic assertions.
    # Existing August policy statements are historical evidence, not new research.
    policy=[]
    for n in (30,31):
        for t in contract['slides'][n-1]['texts']:
            if len(t['text'])>70:policy.append({'slide':n,'shape_id':t['shape_id'],'text':t['text'],'provenance':'user_provided_history'})
        put(n,2,f'{m}月政策观察：历史政策信号与当月经济事实','title')
        deck.slides[n-1].notes_slide.notes_text_frame.text=json.dumps({'policy_boundary':'沿用用户历史报告所载政策解读；非本次新采原文，不能升级为已核验因果','historical_policy':policy},ensure_ascii=False)
    put(30,5,'7月底会议和8月国务院部署作为既有政策背景。8月生产、消费与投资按已核验数据分列观察，政策效果需结合后续数据检验。')
    pmi=facts.get(('s11-shape12::PMI',period_end(month).isoformat()))
    put(30,17,f'{m}月制造业PMI为{fmt(pmi["value"])}，以当期产需与分规模指数观察景气变化。' if pmi else '本期制造业景气取证未完成，不能沿用上月走势推断政策效果。')
    put(30,21,'财政政策结合本期累计收支观察；资金投放与实际需求变化需分别核对，不能直接等同。')
    put(30,25,'车市政策影响需结合补贴执行、居民预期与终端数据核验，不仅凭政策措辞判断消费改善。')
    put(31,42,'历史会议解读仅作背景；政策传导仍需以原文与后续数据复核。')
    for n,sid in [(31,35),(31,39)]:
        frame=_shape(deck.slides[n-1],sid).text_frame;frame.margin_left=frame.margin_right=Pt(0)
    unique={}
    for update in updates:
        key=(update['slide'],update['shape_id'])
        if key in unique:update={**update,'old':unique[key]['old']}
        unique[key]=update
    return list(unique.values()),narratives,policy


def update_tables(deck,contract,facts,month,actuals):
    tables=[];overview=_shape(deck.slides[6],17).table
    # Keep the 15-month window and native table formatting.
    dates=[overview.cell(0,c).text for c in range(3,len(overview.columns)-1)]
    target=period_end(month);header=f'{target.year%100:02d}-{target.month:02d}'
    if header not in dates:
        for row in range(len(overview.rows)):
            for col in range(3,len(overview.columns)-2):set_text(overview.cell(row,col),overview.cell(row,col+1).text)
        set_text(overview.cell(0,len(overview.columns)-2),header)
    columns={c:'20'+overview.cell(0,c).text for c in range(3,len(overview.columns)-1)}
    byid={c['id']:c for s in contract['slides'] for c in s['charts']}
    for row,(label,cid,key,unit) in enumerate(OVERVIEW_ROWS,1):
        if cid.startswith('overview::'):rid=cid
        elif cid=='s24-shape9':rid='s24-shape9::current'
        else:rid=next(r['id'] for r in byid[cid]['requests'] if key in r['series_name'])
        for col,period in columns.items():
            day=period_end(period)
            if label=='GDP当季同比' and day.month%3:value=None
            else:value=facts.get((rid,day.isoformat()))
            v=value['value'] if value else None
            if label=='新增人民币贷款' and v is not None:v/=10000
            missing='--' if label=='GDP当季同比' and day.month%3 or period!=month else '缺项'
            set_text(overview.cell(row,col),missing if v is None else ('约' if value.get('approximate') else '')+fmt(v))
        a=overview.cell(row,len(overview.columns)-2).text;b=overview.cell(row,len(overview.columns)-3).text
        try:direction='↑' if float(a.removeprefix('约'))>float(b.removeprefix('约')) else '↓' if float(a.removeprefix('约'))<float(b.removeprefix('约')) else '→'
        except ValueError:direction='缺项' if a=='缺项' else '—'
        set_text(overview.cell(row,len(overview.columns)-1),direction)
    set_text(overview.cell(6,2),'指数')
    tables.append({'slide':7,'shape_id':17,'cells':[[c.text for c in r.cells] for r in overview.rows]})
    outlook=_shape(deck.slides[31],9).table
    # Explicit canonical requests avoid inferring meaning from row positions.
    rids=['s13-shape17::当月同比','s14-shape15::固定资产投资累计同比(右轴)',
          's15-shape9::社会消费品零售总额:当月同比(%,右轴)',
          's19-shape13::出口金额:当月同比(%,右轴)','s19-shape13::进口金额:当月同比(%,右轴)',
          's21-shape11::CPI:当月同比','s22-shape9::PPI同比','s24-shape9::current']
    nxt=month_end(target.year if target.month<12 else target.year+1,target.month+1 if target.month<12 else 1)
    prev=month_end(target.year if target.month>1 else target.year-1,target.month-1 if target.month>1 else 12)
    for col,text in enumerate(['经济指标','预测方向',nxt.strftime('%Y年%m月预测'),target.strftime('%Y年%m月实际'),prev.strftime('%Y年%m月实际')]):set_text(outlook.cell(0,col),text)
    forecasts=[]
    for row,rid in enumerate(rids,1):
        training=sorted([(date.fromisoformat(o['period']),o['value']) for o in actuals if o['request_id']==rid and o['period']<=target.isoformat()])
        # Duplicate publication copies are not extra training periods.
        training=sorted(dict(training).items())
        projection=None
        current=facts.get((rid,target.isoformat()));previous=facts.get((rid,prev.isoformat()))
        if len(training)>=3 and current:
            projection=forecast_series(training,nxt,'monthly','亿元' if row==8 else '%')
        prediction=projection['values'][-1]['value'] if projection and projection['values'] else None
        scale=10000 if row==8 else 1
        for col,v in [(2,prediction),(3,current['value'] if current else None),(4,previous['value'] if previous else None)]:
            approximate=(col==3 and current and current.get('approximate')) or (col==4 and previous and previous.get('approximate'))
            set_text(outlook.cell(row,col),'证据不足' if v is None and col==2 else '--' if v is None else ('约' if approximate else '')+fmt(v/scale))
        direction='待核验'
        if prediction is not None and current:direction='↑' if prediction>current['value'] else '↓' if prediction<current['value'] else '→'
        set_text(outlook.cell(row,1),direction)
        forecasts.append({'request_id':rid,'projection':projection,'target':nxt.isoformat()})
    set_text(_shape(deck.slides[31],10),f'{nxt.month}月经济前瞻：预测与本月实际分列')
    set_text(_shape(deck.slides[31],6),f'展望{nxt.month}月，结合{target.month}月已核验供需、价格与信用数据观察经济变化。定量预测仅使用原始发布证据充足的同口径历史；证据不足的指标保留空项，不以旧月预测充当实际。')
    set_text(_shape(deck.slides[31],7),'来源：本月实际与独立统计基线；预测非实际，训练证据和方法见备注')
    for sid in (6,7):
        frame=_shape(deck.slides[31],sid).text_frame;frame.margin_left=frame.margin_right=Pt(0)
    _shape(deck.slides[31],7).height=Pt(26)
    _shape(deck.slides[31],7).text_frame.margin_top=Pt(0)
    for p in _shape(deck.slides[31],7).text_frame.paragraphs:
        for run in p.runs:run.font.size=Pt(9)
    tables.append({'slide':32,'shape_id':9,'cells':[[c.text for c in r.cells] for r in outlook.rows], 'forecasts':forecasts})
    return tables


def expected_content(package,root):
    contract=inventory();history=build_history(contract);actuals=replay_sources(package,contract,root)
    facts,revisions=merged_facts(contract,history,actuals,package['report_month'])
    charts=[chart_values(c,facts,package['report_month']) for s in contract['slides'] for c in s['charts']]
    return contract,history,actuals,facts,revisions,charts


def missing_actuals(contract,actuals,month):
    missing=[]
    for r in contract['source_requests']:
        if r['frequency']=='quarterly':continue
        candidates=[o for o in actuals if o['request_id']==r['id']]
        if r['frequency'] in {'monthly','announcement_monthly'}:
            candidates=[o for o in candidates if o['period'][:7]==month]
        elif r['frequency'] in {'weekly','daily'}:
            minimum=(period_end(month)-timedelta(days=7 if r['frequency']=='weekly' else 3)).isoformat()
            candidates=[o for o in candidates if minimum<=o['period']<=period_end(month).isoformat()]
        if not candidates:missing.append(r['id'])
    return missing


def retention_annotations(deck,contract,facts,missing,month):
    records=[]
    for s in contract['slides']:
        for c in s['charts']:
            ids=[r['id'] for r in c['requests'] if r['id'] in missing]
            if not ids:continue
            periods={rid:max((p for (r,p) in facts if r==rid),default='无历史') for rid in ids}
            latest=max((p for p in periods.values() if p!='无历史'),default='无可用历史')
            title='暂无可核验数据 · 保留模板至'+latest
            shape=_shape(deck.slides[s['slide']-1],c['shape_id']);chart=shape.chart
            chart.has_title=False
            box=deck.slides[s['slide']-1].shapes.add_textbox(shape.left,shape.top,shape.width,Pt(16))
            box.text_frame.margin_top=box.text_frame.margin_bottom=Pt(0)
            box.text_frame.margin_left=box.text_frame.margin_right=Pt(0)
            frame=box.text_frame;frame.text=title
            shape.top+=Pt(18);shape.height-=Pt(18)
            for p in frame.paragraphs:
                p.alignment=2
                for run in p.runs:run.font.size=Pt(10)
            records.append({'chart_id':c['id'],'slide':s['slide'],'shape_id':c['shape_id'],
                            'annotation_shape_id':box.shape_id,'title':title,'request_ids':ids,'retained_periods':periods})
            deck.slides[s['slide']-1].notes_slide.notes_text_frame.text+='\n暂无可核验数据：'+json.dumps(records[-1],ensure_ascii=False)
    cover=_shape(deck.slides[0],8)
    pages=sorted({a['slide'] for a in records}|({7} if any(r.startswith('overview::') for r in missing) else set()))
    set_text(cover,f'部分完成 · {len(pages)}页暂无可核验数据 · 保留模板历史' if missing else '完成 · 本期数据已更新')
    return records


def build(month,package,output,ledger_path,root):
    contract,history,actuals,facts,revisions,charts=expected_content(package,root)
    deck=Presentation(str(REFERENCE))
    for c in charts:write_chart(_shape(deck.slides[c['slide']-1],c['shape_id']).chart,c)
    text,narratives,policy=render_contents(deck,contract,facts,month,package['cutoff_date'])
    tables=update_tables(deck,contract,facts,month,actuals)
    missing=missing_actuals(contract,actuals,month)
    retained=package.get('delivery_mode')=='historical_retained'
    annotations=[]
    if retained:
        invocation=json.loads((Path(root)/'inputs/request.json').read_text(encoding='utf-8'))
        if invocation.get('history_retention_confirmed') is not True:raise ValueError('historical_retention_requires_explicit_scope')
        annotations=retention_annotations(deck,contract,facts,missing,month)
        text.extend({'slide':a['slide'],'shape_id':a['annotation_shape_id'],'old':'','new':a['title'],'role':'caption'} for a in annotations)
        for update in text:
            if update['slide']==1 and update['shape_id']==8:update['new']=_shape(deck.slides[0],8).text
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);deck.save(str(output))
    required=[r for r in contract['source_requests'] if r['frequency']!='quarterly']
    ledger={'schema_version':3,'contract_version':VERSION,'mode':'historical_update','report_month':month,
            'cutoff_date':package['cutoff_date'],'issue_month':tables[-1]['forecasts'][0]['target'][:7],
            'outlook_month':tables[-1]['forecasts'][0]['target'][:7], 'source_root':str(Path(root).resolve()),
            'reference_sha256':digest(REFERENCE),'pptx_sha256':digest(output),'history':history,
            'revisions':revisions,'charts':charts,'slides':text,'table_updates':tables,'narratives':narratives,
            'policy_boundary':policy,'coverage':{'required':len(required),'verified':len(required)-len(missing)},
            'issues':[{'code':'historical_new_period_missing','request_id':rid} for rid in missing],
            'business_complete':not missing,'formal_delivery_allowed':not missing}
    if retained:
        ledger.update(delivery_mode='historical_retained',history_retention_confirmed=True,
                      formal_delivery_allowed=False,business_complete=not missing,retained_content_complete=True,
                      pending_pages=sorted({a['slide'] for a in annotations}|({7} if any(r.startswith('overview::') for r in missing) else set())),
                      retention_annotations=annotations,retained_requests=missing)
    dump(ledger_path,ledger)
    dump(Path(root)/'working/page-facts.json',narratives)
    return {'pptx':str(output),'business_complete':not missing,'issues':ledger['issues']}


def validate(pptx,ledger_path):
    ledger=json.loads(Path(ledger_path).read_text(encoding='utf-8'));root=Path(ledger['source_root']);issues=[]
    if ledger['contract_version']!=VERSION or digest(pptx)!=ledger['pptx_sha256']:issues.append({'code':'artifact_or_version_changed'})
    package=json.loads((root/'working/data-package.json').read_text(encoding='utf-8'))
    contract,history,actuals,facts,revisions,charts=expected_content(package,root)
    if history!=ledger['history'] or charts!=ledger['charts']:issues.append({'code':'historical_replay_mismatch'})
    original=Presentation(str(REFERENCE))
    texts,narratives,policy=render_contents(original,contract,facts,ledger['report_month'],ledger['cutoff_date'])
    tables=update_tables(original,contract,facts,ledger['report_month'],actuals)
    retained=ledger.get('delivery_mode')=='historical_retained'
    required=[r for r in contract['source_requests'] if r['frequency']!='quarterly']
    expected_missing=missing_actuals(contract,actuals,ledger['report_month'])
    if ledger.get('coverage')!={'required':len(required),'verified':len(required)-len(expected_missing)}:
        issues.append({'code':'coverage_replay_mismatch'})
    if retained:
        invocation=json.loads((root/'inputs/request.json').read_text(encoding='utf-8'))
        if invocation.get('history_retention_confirmed') is not True or not ledger.get('retained_content_complete') or ledger.get('formal_delivery_allowed') is not False:
            issues.append({'code':'retention_scope_missing'})
        missing=missing_actuals(contract,actuals,ledger['report_month'])
        annotations=retention_annotations(original,contract,facts,missing,ledger['report_month'])
        texts.extend({'slide':a['slide'],'shape_id':a['annotation_shape_id'],'old':'','new':a['title'],'role':'caption'} for a in annotations)
        if annotations!=ledger.get('retention_annotations') or missing!=ledger.get('retained_requests'):
            issues.append({'code':'retention_replay_mismatch'})
        expected_pages=sorted({a['slide'] for a in annotations}|({7} if any(r.startswith('overview::') for r in missing) else set()))
        if ledger.get('pending_pages')!=expected_pages or ledger.get('business_complete')!= (not missing):
            issues.append({'code':'page_completion_replay_mismatch'})
        for update in texts:
            if update['slide']==1 and update['shape_id']==8:update['new']=_shape(original.slides[0],8).text
    if texts!=ledger['slides'] or narratives!=ledger['narratives'] or tables!=ledger['table_updates']:issues.append({'code':'narrative_or_table_replay_mismatch'})
    deck=Presentation(str(pptx))
    for c in charts:
        chart=_shape(deck.slides[c['slide']-1],c['shape_id']).chart;xml=chart._chartSpace
        ext=xml.find(C('externalData'));part=chart.part.related_part(ext.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
        wb=load_workbook(io.BytesIO(part.blob),data_only=True)
        for node,expected in zip(xml.xpath('.//c:ser'),c['series']):
            if _series_name(node)!=expected['name']:issues.append({'code':'series_name_changed','chart_id':c['chart_id']})
            for path,wanted in [('./c:cat/*',c['categories']),('./c:val/c:numRef',expected['values'])]:
                ref=node.xpath(path)[0];cache=ref.find(C('numCache'))
                numeric=cache is not None
                if cache is None:cache=ref.find(C('strCache'))
                count=int(cache.find(C('ptCount')).get('val'));values=[None]*count
                for pt in cache.findall(C('pt')):
                    raw=pt.find(C('v')).text;values[int(pt.get('idx'))]=float(raw) if numeric else raw
                desired=[float(v) if numeric and v is not None else str(v) if not numeric and v is not None else v for v in wanted]
                if values!=desired:issues.append({'code':'native_chart_data_changed','chart_id':c['chart_id']})
                sheet,col,start,end,_=_range(ref.find(C('f')).text,count)
                cells=[wb[sheet].cell(row,col).value for row in range(start,end+1)]
                equal=len(cells)==len(desired) and all(
                    (a is None and b is None) or (a==b) or
                    (isinstance(a,(int,float)) and isinstance(b,(int,float)) and math.isclose(a,b,rel_tol=1e-10,abs_tol=1e-8))
                    for a,b in zip(cells,desired))
                if not equal:issues.append({'code':'cache_workbook_mismatch','chart_id':c['chart_id']})
    for update in texts:
        if _shape(deck.slides[update['slide']-1],update['shape_id']).text!=update['new']:issues.append({'code':'text_changed','slide':update['slide']})
    for update in tables:
        table=_shape(deck.slides[update['slide']-1],update['shape_id']).table
        if [[c.text for c in r.cells] for r in table.rows]!=update['cells']:issues.append({'code':'table_changed','slide':update['slide']})
    if retained:
        for a in ledger['retention_annotations']:
            box=_shape(deck.slides[a['slide']-1],a['annotation_shape_id'])
            if box.text!=a['title']:issues.append({'code':'retention_caption_missing','chart_id':a['chart_id']})
    for sid in (6,7,10):
        if _shape(deck.slides[31],sid).text!=_shape(original.slides[31],sid).text:issues.append({'code':'outlook_text_changed','shape_id':sid})
    if not retained:issues.extend(ledger['issues'])
    return {'passed':not issues,'gate':'macro_retained_business' if retained else 'macro_business','issues':issues,'artifact_sha256':digest(pptx),
            'formal_delivery_allowed':False if retained else not issues,'retained_delivery_allowed':retained and not issues,
            'report_month':ledger['report_month'],'cutoff_date':ledger['cutoff_date']}
