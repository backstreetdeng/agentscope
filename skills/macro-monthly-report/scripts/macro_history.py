"""Historical monthly profile: semantic mappings and traceable archived history.

The legacy forecast contract stays immutable. This profile resolves physical
objects from series meaning, while retaining canonical source-reader IDs.
"""
from __future__ import annotations
import hashlib
import io
import json
import math
import re
from pathlib import Path
from datetime import date
from pptx import Presentation
from openpyxl import load_workbook
from macro_contract import RULES, request_contract, month_end, parse_period, static_series, attach_public_source_policy
from inspect_macro_template import all_shapes
from macro_engine import C, _series_name, _range, dump, digest

VERSION = 'macro-history-v1'
ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / 'assets/history/2026-07.pptx'

# Meaning signatures, not slide positions. A title/series change must be reviewed.
SIGNATURES = [
 ('s4-shape5', '综合PMI'), ('s4-shape10', '美国'),
 ('s8-shape20', '唐山'), ('s8-shape25', '全钢胎'), ('s8-shape11', 'WTI'),
 ('s10-shape7', 'GDP平减'), ('s10-shape28', '消费支出拉动'),
 ('s11-shape12', 'PMI:大型'), ('s11-shape15', '新出口订单'),
 ('s12-shape13', '消费者信心'), ('s12-shape16', '企业景气'),
 ('s13-shape10', '工业增加值:制造业'),
 ('s14-shape15', '固定资产投资累计值'), ('s14-shape11', '基建'),
 ('s15-shape9', '社会消费品零售总额:当月值'), ('s16-shape8', '建筑及装潢材料类'),
 ('s17-shape15', '商品房销售额'), ('s17-shape20', '第1周'),
 ('s18-shape17', '一线城市'), ('s19-shape13', '贸易差额'),
 ('s21-shape11', '核心CPI'), ('s21-shape12', '食品烟酒'),
 ('s22-shape9', 'PPI同比'), ('s22-shape10', '生产资料'),
 ('s24-shape9', '新增人民币贷款规模'), ('s24-shape17', '住户'),
 ('s25-shape9', '社会融资规模季节性'), ('s25-shape10', '非标融资'),
 ('s26-shape11', 'M1:同比'), ('s26-shape12', '社会融资规模存量:同比'),
 ('s28-shape11', '销量当月同比'), ('s28-shape17', '汽车制造:利润'),
 ('s30-shape13', '财政收入'), ('s30-shape18', '贷款市场报价利率'),
]


def semantic_id(chart, slide_text):
    names = '|'.join(s['name'] for s in chart['series'])
    cats = '|'.join(map(str, chart['categories']))
    if '土地成交' in names:
        return 'hf-land'
    if '30城商品房' in names or '30大中城市商品房' in names:
        return 's8-shape21'
    if '一线城市' in names:
        return 's18-shape17' if '环比' in slide_text and chart['x'] < chart['slide_width']/2 else 's18-shape21'
    if '同比' in names and '环比' in names and '衣着' in cats:
        return 's21-shape12'
    if names in {'累计同比|当月同比','当月同比|累计同比'} and '工业增加值' in slide_text:
        return 's13-shape17'
    if '票据' in cats:
        return 's24-shape17'
    if re.fullmatch(r'20\d\d(?:\|20\d\d)+', names):
        if '第1周' in cats:
            return 's17-shape20'
        return 's24-shape9' if '新增人民币贷款' in slide_text else 's25-shape9'
    for ident, signature in SIGNATURES:
        # Categories distinguish radar/category snapshots and finance structures.
        corpus = names + '|' + cats
        if ident in {'s24-shape9', 's25-shape9'}:
            continue
        if signature in corpus:
            return ident
    raise ValueError('unreviewed_chart_semantics:' + names + ':' + cats[:60])


def raw_inventory(path):
    deck = Presentation(str(path))
    slides = []
    for n, slide in enumerate(deck.slides, 1):
        shapes = list(all_shapes(slide.shapes))
        texts = [{'shape_id': s.shape_id, 'text': s.text, 'x':s.left, 'y':s.top,
                  'width':s.width, 'height':s.height} for s in shapes if s.has_text_frame and s.text]
        slide_text = '\n'.join(t['text'] for t in texts)
        charts = []
        for s in shapes:
            if not s.has_chart:
                continue
            chart = s.chart
            series = []
            for node in chart._chartSpace.xpath('.//c:ser'):
                val = node.find('./'+C('val')+'/'+C('numRef')+'/'+C('numCache'))
                if val is None:
                    val=node.find('./'+C('val')+'/'+C('numLit'))
                count = int(val.find(C('ptCount')).get('val'))
                values = [None]*count
                for pt in val.findall(C('pt')):
                    values[int(pt.get('idx'))] = float(pt.find(C('v')).text)
                series.append({'name':_series_name(node), 'values':values})
            item = {'shape_id':s.shape_id, 'slide':n, 'x':s.left, 'y':s.top,
                    'slide_width':deck.slide_width, 'part':str(chart.part.partname),
                    'type':str(chart.chart_type), 'categories':[c.label for c in chart.plots[0].categories],
                    'series':series, 'physical_id':f's{n}-shape{s.shape_id}'}
            item['id'] = semantic_id(item, slide_text)
            # Old combined land/property chart is split in the latest profile.
            if item['id']=='s8-shape21' and len(series)>1:
                item['series']=[v for v in series if '商品房' in v['name']]
            charts.append(item)
        slides.append({'slide':n, 'texts':texts, 'charts':charts,
                       'tables':[{'shape_id':s.shape_id,'cells':[[c.text for c in r.cells] for r in s.table.rows]}
                                 for s in shapes if s.has_table]})
    return {'slides':slides,'slide_count':len(slides),'reference_sha256':digest(path)}


def source_requests(chart):
    if chart['id'] == 'hf-land':
        return [attach_public_source_policy({'id':'hf-land::current','chart_id':'hf-land','series_name':chart['series'][0]['name'],
                 'definition':'100城土地成交规划建面','definition_aliases':['100城土地成交规划建面'],
                 'context_groups':[['100城','100个城市'],['土地','规划建面']], 'frequency':'weekly',
                 'unit':'万平方米','source_domains':['wind.com.cn'],'history_count':1,
                 'query':'100城土地成交规划建面 {year}年{month}月 原始数据 Wind'})]
    if chart['id']=='s28-shape11':
        result=[]
        for key,name,unit in [('current','汽车销量','辆'),('yoy','汽车销量同比','%')]:
            result.append({'id':'auto-sales::'+key,'chart_id':'auto-sales','series_name':name,
               'definition':'中汽协汽车总销量当月'+('同比' if key=='yoy' else '值'),
               'definition_aliases':['汽车销量','汽车销售','汽车产销'],
               'context_groups':[['汽车'],['销量','销售'],['同比']] if key=='yoy' else [['汽车'],['销量','销售']],
               'frequency':'monthly','unit':unit,'source_domains':['caam.org.cn','caam.com.cn'],
               'history_count':1,'query':'中国汽车工业协会 汽车产销 {year}年{month}月 销量 同比'} )
        return result
    requests=request_contract(chart)
    if chart['id']=='s28-shape17':
        for r in requests:
            if '利润' not in r['series_name']:
                r.update(chart_id='overview', definition='汽车产量累计同比',
                         definition_aliases=['汽车（万辆）'],context_groups=[['汽车'],['累计'],['同比']],
                         source_domains=['stats.gov.cn'],original_publisher_domains=['stats.gov.cn'],
                         query='汽车产量 累计同比 {year}年{month}月 site:stats.gov.cn')
            else:
                r['series_name']='汽车制造业利润:累计同比'
                r['definition']='汽车制造业利润累计同比'
                r['display_series_name']='汽车制造:利润总额:累计同比'
    if chart['id']=='s8-shape21':
        for r in requests:
            r.update(definition_aliases=['30城商品房成交面积','30大中城市商品房成交面积'],
                     context_groups=[['30城','30大中城市','30个大中城市'],['商品房'],['成交面积']])
    for r in requests:
        r['history_count']=1
    return requests


def inventory(path=REFERENCE):
    inv=raw_inventory(path)
    allcharts=[c for s in inv['slides'] for c in s['charts']]
    if len(allcharts)!=36 or len({c['id'] for c in allcharts})!=36:
        raise ValueError('historical_semantic_mapping_not_unique')
    requests=[]
    for c in allcharts:
        if c['id']=='hf-land':
            sem,freq,unit,domains,measure='time','weekly','万平方米',['wind.com.cn'],'100城土地成交规划建面'
        else:
            sem,freq,unit,domains,measure=RULES[c['id']]
        c.update(semantics=sem,frequency=freq,raw_unit=unit,source_domains=domains,
                 definition=measure,reference_month='2026-07')
        c['requests']=source_requests(c)
        requests.extend(c['requests'])
    for label in ('发电量','钢材产量','水泥产量'):
        requests.append({'id':'overview::'+label+'当月同比','chart_id':'overview', 'series_name':label,
            'definition':label+'当月同比','definition_aliases':[label,label.replace('产量','')],
            'context_groups':[['同比'],['当月','月份']], 'frequency':'monthly','unit':'%',
            'history_count':1,'source_domains':['stats.gov.cn'], 'query':label+' {year}年{month}月 site:stats.gov.cn'})
    inv.update(contract_version=VERSION,chart_count=36,source_requests=requests)
    return inv


def archived_points(chart, reference_month):
    """Extract time, seasonal, dated snapshots without treating forecasts as actuals."""
    sem=chart['semantics'];freq=chart['frequency'];result=[]
    for series in chart['series']:
        name=series['name']
        if static_series(name,series['values']):
            continue
        for i,value in enumerate(series['values']):
            if value is None or not math.isfinite(value):
                continue
            if sem in {'snapshot','category_metrics'}:
                period=parse_period(name if sem=='snapshot' else reference_month,'monthly').isoformat()
                rs=[r for r in chart['requests'] if r.get('category_index')==i and
                    (sem=='snapshot' or r['series_name']==name)]
            elif sem in {'seasonal','seasonal_current'}:
                if chart['id']=='s28-shape11':
                    year=int(name) if name.isdigit() else int(reference_month[:4])
                    rs=[r for r in chart['requests'] if r['id'].endswith('current' if name.isdigit() else 'yoy')]
                else:
                    year=int(name);rs=chart['requests']
                if freq=='weekly':
                    try: period=date.fromisocalendar(year,i+1,7).isoformat()
                    except ValueError:continue
                else:
                    period=month_end(year,i+1).isoformat()
            else:
                raw=chart['categories'][i]
                if freq=='quarterly' and parse_period(raw,'daily').month not in {3,6,9,12}:
                    continue
                period=parse_period(raw,freq).isoformat()
                rs=[r for r in chart['requests'] if r.get('display_series_name',r['series_name'])==name]
            if period[:7]>reference_month and freq not in {'daily','weekly'}:
                continue
            for r in rs:
                result.append({'request_id':r['id'],'period':period,'value':value,'unit':r['unit'],
                               'frequency':freq,'series_name':name,'index':i})
    return result


def build_history(contract):
    records=[];target={c['id']:c for s in contract['slides'] for c in s['charts']}
    manifest=json.loads((ROOT/'references/history-manifest.json').read_text(encoding='utf-8'))
    for entry in manifest['entries']:
        path=ROOT/entry['path']
        if digest(path)!=entry['sha256']:
            raise ValueError('historical_asset_hash_mismatch')
        raw=raw_inventory(path)
        for slide in raw['slides']:
            for c in slide['charts']:
                destination=target.get(c['id'])
                if not destination:
                    continue
                enriched={**c,**{k:destination[k] for k in ('semantics','frequency','requests')}}
                for record in archived_points(enriched,entry['report_month']):
                    records.append({**record,'source_file':entry['path'],'source_sha256':entry['sha256'],
                                    'source_month':entry['report_month'],'slide':slide['slide'],
                                    'shape_id':c['shape_id'],'provenance':'user_provided_history'})
            # Indicators present only in the native overview table also retain history.
            if slide['slide']==7:
                for table in slide['tables']:
                    cells=table['cells']
                    if len(cells)!=17 or len(cells[0])<5:continue
                    for row,label in ((14,'发电量'),(15,'钢材产量'),(16,'水泥产量')):
                        for col,header in enumerate(cells[0]):
                            if not re.fullmatch(r'\d{2}-\d{2}',header):continue
                            try:value=float(cells[row][col])
                            except ValueError:continue
                            records.append({'request_id':'overview::'+label+'当月同比','period':parse_period('20'+header).isoformat(),
                                'value':value,'unit':'%','frequency':'monthly','series_name':label,'index':col,
                                'source_file':entry['path'],'source_sha256':entry['sha256'],'source_month':entry['report_month'],
                                'slide':7,'shape_id':table['shape_id'],'provenance':'user_provided_history'})
    groups={}
    for r in records:
        groups.setdefault((r['request_id'],r['period']),[]).append(r)
    accepted=[];conflicts=[]
    for (rid,period),variants in sorted(groups.items()):
        if any(not math.isclose(v['value'],variants[0]['value'],rel_tol=1e-10,abs_tol=1e-8) for v in variants):
            conflicts.append({'request_id':rid,'period':period,'variants':variants})
        else:
            accepted.append(variants[-1])
    return {'schema_version':1,'contract_version':VERSION,'observations':accepted,
            'conflicts':conflicts,'source_records':records,'assets':manifest['entries'],
            'boundary':'历史样例原值；有冲突的点不作为比较/预测依据，不能证明新增月份实际'}


def main():
    import argparse,os
    p=argparse.ArgumentParser();p.add_argument('output');p.add_argument('--history-output',default='working/history.json')
    args=p.parse_args();root=Path(os.environ.get('SKILL_RUN_ROOT',Path.cwd())).resolve()
    for rel in (args.output,args.history_output):
        if root not in (root/rel).resolve().parents:raise ValueError('path_outside_run')
    contract=inventory();history=build_history(contract)
    dump(root/args.output,contract);dump(root/args.history_output,history)
    print(json.dumps({'slides':33,'charts':36,'requests':len(contract['source_requests']),
                      'history_points':len(history['observations']),'conflicts':len(history['conflicts'])}))


if __name__=='__main__':main()
