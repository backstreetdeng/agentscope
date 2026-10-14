"""Generate a forecast-labelled monthly macro report from the retained template.

The script is part of the Skill package rather than the web app.  It consumes
only the template's embedded editable chart data, keeps the source template
read-only, and records every generated value in a run-local ledger.  It is a
safe baseline when the requested target-month actuals are not yet published.
"""

from __future__ import annotations

import json
import math
import os
import re
import shutil
import sys
from datetime import date
from pathlib import Path
from statistics import fmean

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


EXCEL_EPOCH = date(1899, 12, 30)
CONSTANT_SERIES_HINTS = ("分界", "临界", "基准线", "阈值")


def parse_month(value: str) -> tuple[int, int]:
    match = re.fullmatch(r"(20\d{2})-(0[1-9]|1[0-2])", value.strip())
    if not match:
        raise ValueError("目标月份必须是 YYYY-MM，例如 2026-08")
    return int(match.group(1)), int(match.group(2))


def excel_month_serial(year: int, month: int) -> int:
    return (date(year, month, 1) - EXCEL_EPOCH).days


def target_category(categories: list[object], year: int, month: int) -> object:
    last = categories[-1] if categories else ""
    if isinstance(last, (int, float)) or (isinstance(last, str) and re.fullmatch(r"\d{5}", last)):
        return excel_month_serial(year, month)
    if isinstance(last, str) and "月" in last:
        return f"{month}月"
    return f"{year}-{month:02d}"


def numeric_history(values: list[object]) -> list[float]:
    return [float(value) for value in values if isinstance(value, (int, float)) and math.isfinite(value)]


def forecast_value(name: str, values: list[object]) -> tuple[object, str]:
    history = numeric_history(values)
    if not history:
        return None, "no_history"
    if any(hint in name for hint in CONSTANT_SERIES_HINTS) or max(history) - min(history) < 1e-9:
        return round(history[-1], 6), "constant"
    if len(history) >= 24:
        return round(history[-12], 6), "seasonal_naive_12"
    window = history[-min(6, len(history)):]
    if len(window) < 2:
        return round(window[-1], 6), "last_value"
    slope = (window[-1] - window[0]) / (len(window) - 1)
    return round(window[-1] + slope, 6), "drift_6"


def replace_month_text(text: str, year: int, month: int) -> str:
    target = f"{year}年{month}月"
    # Keep the retained cover typography on one line. The forecast boundary is
    # carried by the per-slide footer and run ledger, not appended to a title
    # box that was designed for the shorter original title.
    text = re.sub(r"20\d{2}年\d{1,2}月宏观经济分析报告", f"{target}宏观经济分析报告", text)
    # The retained cover also carries a separate reporting-period label.  It
    # is metadata, not a historical observation, so roll that standalone
    # label with the report title.  Historical dates embedded in other prose
    # and charts are deliberately left untouched.
    if re.fullmatch(r"\s*20\d{2}年\d{1,2}月\s*", text):
        text = target
    text = re.sub(r"展望\d{1,2}月", f"展望{month}月", text)
    return text


def update_text(slide, year: int, month: int, note: str) -> int:
    changed = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            old = paragraph.text
            new = replace_month_text(old, year, month)
            if new != old:
                paragraph.text = new
                changed += 1
    # A small, uniform footer makes the forecasting boundary explicit on every
    # page without replacing the source template's layout or chart objects.
    footer = slide.shapes.add_textbox(Inches(0.35), Inches(7.15), Inches(12.55), Inches(0.18))
    paragraph = footer.text_frame.paragraphs[0]
    paragraph.text = note
    paragraph.font.size = Pt(6.5)
    paragraph.font.color.rgb = RGBColor(112, 128, 144)
    return changed


def update_chart(shape, year: int, month: int) -> dict:
    chart = shape.chart
    categories = list(chart.plots[0].categories) if chart.plots else []
    categories = [item.label for item in categories]
    if not categories:
        raise ValueError("图表没有类别轴，无法生成滚动预测")
    next_category = target_category(categories, year, month)
    updated_categories = categories[1:] + [next_category]
    data = CategoryChartData()
    data.categories = updated_categories
    series_receipts = []
    for series in chart.series:
        values = list(series.values)
        prediction, method = forecast_value(series.name, values)
        updated_values = values[1:] + [prediction]
        data.add_series(series.name, tuple(updated_values))
        series_receipts.append(
            {"series": series.name, "forecast": prediction, "method": method, "history_points": len(numeric_history(values))}
        )
    chart.replace_data(data)
    return {"category": next_category, "series": series_receipts}


def update_tables(slide, year: int, month: int) -> int:
    changes = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_table", False):
            continue
        table = shape.table
        for row in table.rows:
            for cell in row.cells:
                old = cell.text
                new = replace_month_text(old, year, month)
                if new != old:
                    cell.text = new
                    changes += 1
        if len(table.rows) and len(table.rows[0].cells) >= 3:
            table.cell(0, 2).text = f"{year}年{month}月预测值"
            changes += 1
    return changes


def main() -> int:
    if len(sys.argv) != 4:
        raise SystemExit("usage: generate_forecast_baseline.py <YYYY-MM> <output.pptx> <ledger.json>")
    year, month = parse_month(sys.argv[1])
    run_root = Path(os.environ.get("SKILL_RUN_ROOT", Path.cwd())).resolve()
    output = Path(sys.argv[2])
    ledger_path = Path(sys.argv[3])
    output = (run_root / output).resolve() if not output.is_absolute() else output.resolve()
    ledger_path = (run_root / ledger_path).resolve() if not ledger_path.is_absolute() else ledger_path.resolve()
    if run_root not in output.parents or run_root not in ledger_path.parents:
        raise ValueError("输出和台账必须位于本次 Skill 运行工作区内")
    package_root = Path(__file__).resolve().parents[1]
    reference = package_root / "assets" / "reference.pptx"
    if not reference.is_file():
        raise SystemExit("Skill reference.pptx 不存在")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(reference, output)
    presentation = Presentation(str(output))
    note = f"预测版｜目标月：{year}年{month}月｜历史序列来自模板内嵌可编辑数据｜目标期值为透明基准预测，非已发布实际值"
    ledger: dict[str, object] = {
        "schema_version": 1,
        "mode": "forecast",
        "report_month": f"{year}-{month:02d}",
        "data_boundary": "模板内嵌历史可编辑数据；所有目标月新增点均标为预测",
        "source_template": "assets/reference.pptx",
        "slides": [],
        "summary": {"slides": len(presentation.slides), "charts": 0, "tables": 0, "text_changes": 0},
    }
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_changes = update_text(slide, year, month, note)
        table_changes = update_tables(slide, year, month)
        charts = []
        chart_number = 0
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart_number += 1
            charts.append({"chart": chart_number, **update_chart(shape, year, month)})
        ledger["slides"].append(
            {"slide": slide_number, "text_changes": text_changes, "table_changes": table_changes, "charts": charts}
        )
        ledger["summary"]["charts"] += len(charts)
        ledger["summary"]["tables"] += sum(1 for shape in slide.shapes if getattr(shape, "has_table", False))
        ledger["summary"]["text_changes"] += text_changes + table_changes
    presentation.save(str(output))
    ledger_path.parent.mkdir(parents=True, exist_ok=True)
    ledger_path.write_text(json.dumps(ledger, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "ledger": str(ledger_path), **ledger["summary"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
