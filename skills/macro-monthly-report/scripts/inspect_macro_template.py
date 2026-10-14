"""Read-only inventory of the retained macro deck; no generic axis mutation."""
from __future__ import annotations

import json
import os
from pathlib import Path
import sys

from pptx import Presentation
from macro_contract import attach_contract


def all_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from all_shapes(shape.shapes)


def inventory(reference: Path) -> dict:
    deck = Presentation(str(reference))
    slides = []
    for number, slide in enumerate(deck.slides, 1):
        texts, charts, tables = [], [], []
        for shape in all_shapes(slide.shapes):
            if getattr(shape, "has_text_frame", False):
                texts.append({"shape_id": shape.shape_id, "text": shape.text, "x": shape.left, "y": shape.top, "width": shape.width, "height": shape.height})
            if getattr(shape, "has_chart", False):
                chart = shape.chart
                series = [{"name": s.name, "values": list(s.values)} for s in chart.series]
                charts.append({"id": f"s{number}-shape{shape.shape_id}", "shape_id": shape.shape_id,
                    "type": str(chart.chart_type), "part": str(chart.part.partname),
                    "categories": [c.label for c in chart.plots[0].categories], "series": series,
                    "plots": [str(p.__class__.__name__) for p in chart.plots]})
            if getattr(shape, "has_table", False):
                tables.append({"shape_id": shape.shape_id, "cells": [[c.text for c in r.cells] for r in shape.table.rows]})
        slides.append({"slide": number, "texts": texts, "charts": charts, "tables": tables})
    return attach_contract({"schema_version": 2, "slides": slides, "slide_count": len(slides),
        "chart_count": sum(len(s["charts"]) for s in slides), "table_count": sum(len(s["tables"]) for s in slides)})


def main():
    if len(sys.argv) != 2:
        raise SystemExit("usage: inspect_macro_template.py <inventory.json>")
    root = Path(os.environ.get("SKILL_RUN_ROOT", Path.cwd())).resolve()
    output = (root / sys.argv[1]).resolve()
    if root not in output.parents:
        raise ValueError("Inventory must stay inside run workspace")
    reference = Path(__file__).resolve().parents[1] / "assets/reference.pptx"
    result = inventory(reference)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"inventory": str(output), "slides": result["slide_count"], "charts": result["chart_count"], "tables": result["table_count"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
