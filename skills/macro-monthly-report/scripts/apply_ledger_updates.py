"""Apply a macro-monthly report working ledger to a copied presentation.

This is a Skill-owned execution asset: the agent supplies the ledger and invokes
this script inside its isolated run workspace.  It never touches the reference
template and returns a machine-readable update receipt for the validation gate.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation


def replacement_pairs(ledger: dict) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for slide in ledger.get("slides", []):
        for change in slide.get("changes", []):
            old, new = change.get("old"), change.get("new")
            if isinstance(old, str) and isinstance(new, str) and old and old != new:
                pairs.append((old, new))
    for change in ledger.get("text_replacements", []):
        old, new = change.get("old"), change.get("new")
        # A bare global month substitution can change the "latest actual"
        # period in a disclosure sentence.  Use the explicit, per-slide mapping
        # instead; unpaired residual month text must be raised by the content
        # gate rather than silently rewritten.
        if change.get("scope") == "全局月份替换":
            continue
        if isinstance(old, str) and isinstance(new, str) and old and old != new:
            pairs.append((old, new))
    # Long strings must be considered before their shorter fragments.
    return sorted(set(pairs), key=lambda item: len(item[0]), reverse=True)


def replace_text_in_shape(shape, pairs: list[tuple[str, str]]) -> int:
    if not getattr(shape, "has_text_frame", False):
        return 0
    changes = 0
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            updated = run.text
            for old, new in pairs:
                if old in updated:
                    updated = updated.replace(old, new)
            if updated != run.text:
                run.text = updated
                changes += 1
        # PowerPoint often splits a visually continuous sentence into several
        # runs.  A ledger maps the whole business sentence, so fall back to a
        # paragraph-level replacement when no individual run owns the text.
        paragraph_text = paragraph.text
        updated_paragraph = paragraph_text
        for old, new in pairs:
            if old in updated_paragraph:
                updated_paragraph = updated_paragraph.replace(old, new)
        if updated_paragraph != paragraph_text:
            paragraph.text = updated_paragraph
            changes += 1
    return changes


def main() -> int:
    if len(sys.argv) != 3:
        raise SystemExit("usage: apply_ledger_updates.py <pptx> <ledger.json>")
    pptx_path, ledger_path = map(Path, sys.argv[1:])
    if not pptx_path.is_file() or not ledger_path.is_file():
        raise SystemExit("PPTX 或内部台账不存在")
    ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
    pairs = replacement_pairs(ledger)
    if not pairs:
        raise SystemExit("台账中没有可执行的 old/new 文本映射")
    presentation = Presentation(str(pptx_path))
    changed_runs = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            changed_runs += replace_text_in_shape(shape, pairs)
    presentation.save(str(pptx_path))
    print(json.dumps({"pptx": str(pptx_path), "pairs": len(pairs), "changed_runs": changed_runs}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
