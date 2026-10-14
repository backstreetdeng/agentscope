"""Validate a generated 2.0 PPTX against the accepted 23-slide business baseline."""

from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
PROFILE = ROOT / "profiles" / "passenger_vehicle_monthly"
DEFAULT_TEMPLATE = PROFILE / "assets" / "monthly_report_template.pptx"
DEFAULT_OUTPUT = ROOT / "dist" / "month-skill-v2-preview-2026-06.pptx"


def charts(slide: object) -> list[object]:
    return [shape.chart for shape in slide.shapes if shape.has_chart]


def tables(slide: object) -> list[object]:
    return [shape.table for shape in slide.shapes if shape.has_table]


def categories(chart: object) -> tuple[int, tuple[tuple[str, ...], ...]]:
    try:
        chart_categories = chart.plots[0].categories
        return (
            chart_categories.depth,
            tuple(tuple(str(value) for value in label) for label in chart_categories.flattened_labels),
        )
    except (AttributeError, IndexError, TypeError):
        return (0, ())


def series_values(chart: object) -> list[list[float | None]]:
    return [list(series.values) for series in chart.series]


def is_close(baseline: float | None, actual: float | None) -> bool:
    if baseline is None or actual is None:
        return baseline is None and actual is None
    baseline_number = float(baseline)
    actual_number = float(actual)
    tolerance = max(0.15, abs(baseline_number) * 0.01)
    return math.isclose(baseline_number, actual_number, abs_tol=tolerance)


def validate(output: Path, template: Path) -> dict[str, object]:
    presentation = Presentation(output)
    baseline = Presentation(template)
    errors: list[str] = []
    warnings: list[str] = []

    chart_count = sum(len(charts(slide)) for slide in presentation.slides)
    table_count = sum(len(tables(slide)) for slide in presentation.slides)
    if len(presentation.slides) != 23:
        errors.append(f"成品页数为{len(presentation.slides)}，预期23。")
    if chart_count != 47:
        errors.append(f"成品图表数为{chart_count}，预期47。")
    if table_count != 37:
        errors.append(f"成品表格数为{table_count}，预期37。")

    chart_shape_mismatches: list[str] = []
    chart_value_mismatches: list[str] = []
    category_mismatches: list[str] = []
    table_shape_mismatches: list[str] = []
    blank_tables: list[str] = []
    out_of_bounds: list[str] = []
    forbidden_text: list[str] = []

    for page, (base_slide, slide) in enumerate(zip(baseline.slides, presentation.slides), 1):
        base_charts = charts(base_slide)
        actual_charts = charts(slide)
        if len(base_charts) != len(actual_charts):
            chart_shape_mismatches.append(f"P{page}: {len(base_charts)}->{len(actual_charts)} charts")
        for index, (base_chart, actual_chart) in enumerate(zip(base_charts, actual_charts), 1):
            base_series = series_values(base_chart)
            actual_series = series_values(actual_chart)
            if [len(values) for values in base_series] != [len(values) for values in actual_series]:
                chart_shape_mismatches.append(
                    f"P{page}C{index}: {[len(v) for v in base_series]}->{[len(v) for v in actual_series]}",
                )
                continue
            if len(base_series) != len(actual_series):
                chart_shape_mismatches.append(
                    f"P{page}C{index}: {len(base_series)}->{len(actual_series)} series",
                )
                continue
            mismatched_values = sum(
                1
                for base_values, actual_values in zip(base_series, actual_series)
                for base_value, actual_value in zip(base_values, actual_values)
                if not is_close(base_value, actual_value)
            )
            if mismatched_values:
                chart_value_mismatches.append(f"P{page}C{index}: {mismatched_values} values")
            base_categories = categories(base_chart)
            actual_categories = categories(actual_chart)
            if base_categories[0] > 0 and base_categories != actual_categories:
                category_mismatches.append(f"P{page}C{index}")

        base_tables = tables(base_slide)
        actual_tables = tables(slide)
        if len(base_tables) != len(actual_tables):
            table_shape_mismatches.append(f"P{page}: {len(base_tables)}->{len(actual_tables)} tables")
        for index, (base_table, actual_table) in enumerate(zip(base_tables, actual_tables), 1):
            base_shape = (len(base_table.rows), len(base_table.columns))
            actual_shape = (len(actual_table.rows), len(actual_table.columns))
            if base_shape != actual_shape:
                table_shape_mismatches.append(f"P{page}T{index}: {base_shape}->{actual_shape}")
            base_nonempty = sum(
                1 for row in base_table.rows for table_cell in row.cells if table_cell.text.strip()
            )
            actual_nonempty = sum(
                1 for row in actual_table.rows for table_cell in row.cells if table_cell.text.strip()
            )
            if base_nonempty and not actual_nonempty:
                blank_tables.append(f"P{page}T{index}")

        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                out_of_bounds.append(f"P{page}:{shape.name}:negative-origin")
            if shape.left + shape.width > presentation.slide_width + 9144:
                out_of_bounds.append(f"P{page}:{shape.name}:right")
            if shape.top + shape.height > presentation.slide_height + 9144:
                out_of_bounds.append(f"P{page}:{shape.name}:bottom")
            if getattr(shape, "has_text_frame", False):
                text = shape.text
                if re.search(r"#(?:DIV/0!|N/A|VALUE!|REF!)|2087年", text):
                    forbidden_text.append(f"P{page}:{shape.name}:{text[:40]}")

    if chart_shape_mismatches:
        errors.append("图表结构不一致：" + "；".join(chart_shape_mismatches))
    if table_shape_mismatches:
        errors.append("表格结构不一致：" + "；".join(table_shape_mismatches))
    if blank_tables:
        errors.append("发现整表文本被清空：" + "、".join(blank_tables))
    if chart_value_mismatches:
        errors.append("图表数值与6月验收基准存在显著差异：" + "；".join(chart_value_mismatches))
    if category_mismatches:
        warnings.append("图表分类标签重写（需视觉复核）：" + "、".join(category_mismatches))
    if out_of_bounds:
        errors.append("发现越界对象：" + "；".join(out_of_bounds))
    if forbidden_text:
        errors.append("发现公式错误或错误年份文本：" + "；".join(forbidden_text))

    return {
        "ok": not errors,
        "output": str(output),
        "slides": len(presentation.slides),
        "charts": chart_count,
        "tables": table_count,
        "errors": errors,
        "warnings": warnings,
        "chart_value_mismatch_count": len(chart_value_mismatches),
        "category_rewrite_count": len(category_mismatches),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    args = parser.parse_args()
    result = validate(args.output.resolve(), args.template.resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
