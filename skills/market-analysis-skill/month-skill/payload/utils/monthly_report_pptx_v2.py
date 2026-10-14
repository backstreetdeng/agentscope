"""乘用车月报 2.0：单 Sheet 纵向验数表到 23 页业务母版的映射与渲染。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from skills.market_analysis_skill.runtime import SheetSnapshot, read_xlsx_workbook

from .monthly_report_pptx import (
    DEFAULT_TEMPLATE_PATH,
    NEW_ENERGY_GREEN,
    TRADITIONAL_BLUE,
    charts_in,
    is_traditional_model,
    replace_chart,
    set_slide_body,
    set_slide_title,
    set_text_frame_text,
    tables_in,
)


EXPECTED_PAGE_MARKERS = (3, 4, 5, *range(7, 22))
PAGE_MARKER_RE = re.compile(r"^P(?P<page>\d+)\s*[-—]\s*(?P<title>.+)$", re.IGNORECASE)
RATE_RE = re.compile(r"^[+−-]?(?:\d+(?:\.\d+)?|>\d+)%$")
NUMBER_RE = re.compile(r"^-?\d+(?:\.\d+)?$")


@dataclass(frozen=True)
class PageBlock:
    page: int
    title: str
    start_row: int
    end_row: int
    rows: list[list[str]]

    def find(self, label: str, occurrence: int = 0) -> int:
        matches = [
            index for index, row in enumerate(self.rows)
            if normalize(cell(row, 0)) == normalize(label)
        ]
        if occurrence >= len(matches):
            raise ValueError(f"P{self.page}-{self.title} 未找到第 {occurrence + 1} 个“{label}”行。")
        return matches[occurrence]

    def find_any(self, label: str, occurrence: int = 0) -> tuple[int, int]:
        """返回标签在本页数据块中的相对行、列；用于跨列布局标签。"""

        target = normalize(label)
        matches = [
            (row_index, column_index)
            for row_index, row in enumerate(self.rows)
            for column_index, value in enumerate(row)
            if normalize(value) == target
        ]
        if occurrence >= len(matches):
            raise ValueError(f"P{self.page}-{self.title} 未找到第 {occurrence + 1} 个“{label}”单元格。")
        return matches[occurrence]


@dataclass(frozen=True)
class FullMonthlyData:
    sheet: SheetSnapshot
    blocks: dict[int, PageBlock]
    year: int
    month: int

    @property
    def last_year(self) -> int:
        return self.year - 1

    @property
    def period_label(self) -> str:
        return f"{self.year}年{self.month}月"

    @property
    def period_short(self) -> str:
        return f"{self.year}.{self.month}"

    @property
    def previous_period_short(self) -> str:
        year = self.year if self.month > 1 else self.year - 1
        month = self.month - 1 if self.month > 1 else 12
        return f"{year}.{month}"

    @property
    def release_label(self) -> str:
        year = self.year if self.month < 12 else self.year + 1
        month = self.month + 1 if self.month < 12 else 1
        return f"{year}年{month}月"


@dataclass(frozen=True)
class ComparisonData:
    labels: list[str]
    monthly_last: list[float | None]
    monthly_current: list[float | None]
    monthly_yoy: list[float | None]
    monthly_mom: list[float | None]
    monthly_share: list[float | None]
    cumulative_last: list[float | None]
    cumulative_current: list[float | None]
    cumulative_yoy: list[float | None]
    cumulative_share: list[float | None]
    share_periods: list[list[float | None]]
    share_period_labels: list[str]


@dataclass(frozen=True)
class ModelRow:
    name: str
    volume: float | None
    yoy_raw: str
    mom_raw: str
    level: str
    share_raw: str
    rank_change_raw: str
    cumulative: float | None = None


def is_v2_dataset(dataset_path: Path, sheet_name: str | None = None) -> bool:
    """识别以 P3/P4/.../P21 标记组织的 2.0 单 Sheet 数据源。"""

    try:
        data = load_full_monthly_data(dataset_path, sheet_name, require_complete=False)
    except (OSError, ValueError):
        return False
    return {3, 7, 21}.issubset(data.blocks)


def load_full_monthly_data(
    dataset_path: Path,
    sheet_name: str | None = None,
    *,
    require_complete: bool = True,
) -> FullMonthlyData:
    sheets = read_xlsx_workbook(dataset_path)
    sheet = select_sheet(sheets, sheet_name)
    blocks = split_page_blocks(sheet)
    if require_complete:
        missing = [page for page in EXPECTED_PAGE_MARKERS if page not in blocks]
        if missing:
            raise ValueError("2.0 验数表缺少页面分块：" + "、".join(f"P{page}" for page in missing))
    month_match = re.search(r"(1[0-2]|[1-9])月", sheet.name)
    if not month_match:
        raise ValueError(f"工作表“{sheet.name}”无法识别报告月份。")
    years = [
        int(match.group(1))
        for row in sheet.rows
        for value in row
        if (match := re.match(r"^(20\d{2})(?:年|$|[.\s])", normalize(value)))
    ]
    year = max(years) if years else 0
    if not year:
        raise ValueError("2.0 验数表无法识别报告年份。")
    return FullMonthlyData(sheet=sheet, blocks=blocks, year=year, month=int(month_match.group(1)))


def build_full_monthly_report_pptx(
    markdown: str,
    output_path: Path,
    dataset_path: Path,
    template_path: Path | None = None,
    sheet_name: str | None = None,
) -> dict[str, object]:
    """把 2.0 验数表和月报文字写入完整 23 页业务母版。"""

    template = (template_path or DEFAULT_TEMPLATE_PATH).resolve()
    if not template.is_file():
        raise FileNotFoundError(f"未找到月报 PPTX 母版：{template}")
    if not dataset_path.is_file():
        raise FileNotFoundError(f"未找到月报数据集：{dataset_path}")

    data = load_full_monthly_data(dataset_path, sheet_name)
    checks = validate_v2_data(data)
    presentation = Presentation(template)
    if len(presentation.slides) != 23:
        raise ValueError(f"2.0 月报母版必须为 23 页，当前为 {len(presentation.slides)} 页。")

    update_cover(presentation.slides[0], data)
    update_slide3(presentation.slides[2], data.blocks[3], data)
    update_slide4(presentation.slides[3], data.blocks[4], data)
    update_slide5(presentation.slides[4], data.blocks[5], data)

    standard_pages = {
        7: ("2.1 车型类别", 0),
        8: ("2.2 车型级别", 0),
        12: ("2.4 车型系别", 0),
        13: ("2.5 车型价格段", 0),
        15: ("2.6 区域市场", 1),
    }
    for page, (heading, occurrence) in standard_pages.items():
        update_standard_comparison_slide(
            presentation.slides[page - 1],
            data.blocks[page],
            data,
            page,
        )
        apply_markdown_or_fallback(
            presentation.slides[page - 1],
            markdown,
            heading,
            occurrence,
            data,
            data.blocks[page],
        )

    for page, heading in ((9, "2.3 轿车级别"), (10, "2.3 SUV级别"), (11, "2.3 MPV级别")):
        update_segment_slide(presentation.slides[page - 1], data.blocks[page], data)
        apply_markdown_or_fallback(
            presentation.slides[page - 1], markdown, heading, 0, data, data.blocks[page],
        )

    update_slide14(presentation.slides[13], data.blocks[14], data)
    update_slide16(presentation.slides[15], data.blocks[16], data)
    update_slide17(presentation.slides[16], data.blocks[17], data)
    update_slide18(presentation.slides[17], data.blocks[18], data)
    update_slide19(presentation.slides[18], data.blocks[19], data)
    update_slide20(presentation.slides[19], data.blocks[20], data)
    update_slide21(presentation.slides[20], data.blocks[21], data)

    for page, heading, occurrence in (
        (3, "1.2 零售量", 0),
        (4, "1.2.4 新品驱动", 0),
        (5, "1.3 库存", 0),
        (14, "2.6 区域市场", 0),
        (16, "2.7 企业竞争", 0),
        (17, "2.8 新能源", 0),
        (18, "2.8 新能源", 1),
        (19, "2.8 新能源", 2),
        (20, "2.8 新能源", 3),
    ):
        apply_markdown_or_fallback(
            presentation.slides[page - 1],
            markdown,
            heading,
            occurrence,
            data,
            data.blocks[page],
        )

    for page, slide in enumerate(presentation.slides, 1):
        replace_period_text(slide, data)
        update_year_labels(slide, data)
        if page not in {2, 6, 22, 23}:
            normalize_source_footer(slide)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    qa = verify_output(output_path)
    return {
        "success": True,
        "pptx_path": str(output_path),
        "slides_count": qa["slides"],
        "charts_count": qa["charts"],
        "tables_count": qa["tables"],
        "template": str(template),
        "period": data.period_label,
        "sheet_name": data.sheet.name,
        "file_size": output_path.stat().st_size,
        "engine": "business-monthly-pptx-template-v2",
        "data_checks": checks,
        "warnings": [
            item["message"] for item in checks if item["status"] == "warning"
        ],
    }


def split_page_blocks(sheet: SheetSnapshot) -> dict[int, PageBlock]:
    starts: list[tuple[int, int, str]] = []
    for index, row in enumerate(sheet.rows):
        match = PAGE_MARKER_RE.match(normalize(cell(row, 0)))
        if match:
            starts.append((index, int(match.group("page")), match.group("title").strip()))
    blocks: dict[int, PageBlock] = {}
    for position, (start, page, title) in enumerate(starts):
        end = starts[position + 1][0] if position + 1 < len(starts) else len(sheet.rows)
        blocks[page] = PageBlock(page, title, start + 1, end, sheet.rows[start:end])
    return blocks


def select_sheet(sheets: list[SheetSnapshot], sheet_name: str | None) -> SheetSnapshot:
    if sheet_name:
        target = normalize(sheet_name).lower()
        sheet = next((item for item in sheets if normalize(item.name).lower() == target), None)
        if sheet is None:
            raise ValueError(f"验数表未找到工作表：{sheet_name}。")
        return sheet
    return max(sheets, key=lambda item: sheet_month(item.name))


def validate_v2_data(data: FullMonthlyData) -> list[dict[str, str]]:
    checks: list[dict[str, str]] = [
        {
            "id": "page_blocks",
            "status": "pass",
            "message": f"已识别 {len(data.blocks)} 个内容页分块，覆盖 P3-P21 的 18 个数据页。",
        },
        {
            "id": "sheet_layout",
            "status": "pass",
            "message": f"已按单工作表纵向结构读取“{data.sheet.name}”，共 {data.sheet.row_count} 行。",
        },
    ]
    formula_errors = sum(
        1 for row in data.sheet.rows for value in row if normalize(value).startswith("#")
    )
    checks.append(
        {
            "id": "formula_error_cells",
            "status": "warning" if formula_errors else "pass",
            "message": (
                f"输入包含 {formula_errors} 个公式错误缓存；均按新车无同期值“-”处理。"
                if formula_errors else "未发现公式错误缓存。"
            ),
        },
    )
    p4 = data.blocks[4]
    has_typical_models = any("典型新车" in " ".join(row) for row in p4.rows)
    checks.append(
        {
            "id": "p4_typical_models",
            "status": "pass" if has_typical_models else "warning",
            "message": (
                "P4 已提供典型新车子图数据。"
                if has_typical_models
                else "P4 未提供右侧典型新车子图明细；本期保留业务母版中的已验收 6 月数据，后续月份需补充该区块。"
            ),
        },
    )
    for page in (7, 8, 12, 13, 15):
        comp = parse_comparison(data.blocks[page])
        deviations = [
            abs(sum(value or 0 for value in period) - 1)
            for period in comp.share_periods
        ]
        if max(deviations, default=0) > 0.002:
            checks.append(
                {
                    "id": f"share_sum_p{page}",
                    "status": "warning",
                    "message": f"P{page} 份额加总偏离 100%，最大偏差 {max(deviations) * 100:.2f} 个百分点。",
                },
            )
    return checks


def update_cover(slide: object, data: FullMonthlyData) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = normalize(shape.text)
        if "中国乘用车市场分析报告" in text:
            set_text_frame_text(shape.text_frame, f"{data.period_label}中国乘用车市场分析报告")
        elif re.fullmatch(r"20\d{2}年\d{1,2}月", text):
            set_text_frame_text(shape.text_frame, data.release_label)


def update_slide3(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    middle = block.find("中间图表")
    rows = block.rows[middle + 2:middle + 14]
    categories = [normalize(cell(row, 0)) for row in rows]
    series = [
        (f"{data.last_year}年零售量", [number(cell(row, 1)) for row in rows]),
        (f"{data.year}年零售量", [number(cell(row, 2)) for row in rows]),
        (f"{data.last_year}年增速", [number(cell(row, 3)) for row in rows]),
        (f"{data.year}年增速", [number(cell(row, 4)) for row in rows]),
    ]
    chart = charts_in(slide)[0]
    replace_chart(chart, categories, series)
    set_plot_label_formats(chart, ("0", "0.0%"))
    table = tables_in(slide)[0]
    table_start = block.find("下方表格")
    for index, source_row in enumerate(block.rows[table_start + 2:table_start + 6]):
        values = [normalize(cell(source_row, 0))]
        for value in source_row[1:13]:
            values.append(fmt_pct(value) if index == 3 else fmt_number(value, 1))
        write_table_row(table, index, values)


def update_slide4(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    start = block.find("左侧图表")
    rows = block.rows[start + 2:start + 14]
    categories = row_period_categories(rows)
    main_chart = next(chart for chart in charts_in(slide) if len(chart.plots) == 2)
    replace_chart(
        main_chart,
        categories,
        [
            ("新车增量", [number(cell(row, 2)) for row in rows]),
            ("新车增量贡献率", [number(cell(row, 3)) for row in rows]),
        ],
    )
    set_plot_label_formats(main_chart, ("0.0", "0.0%"))
    table = tables_in(slide)[0]
    table_start = block.find("左下方表格")
    for index, source_row in enumerate(block.rows[table_start + 4:table_start + 6]):
        write_table_row(
            table,
            index,
            [normalize(cell(source_row, 0)), *[
                fmt_number(value, 2 if index else 0) for value in source_row[1:13]
            ]],
        )


def update_slide5(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    left = block.find("左侧图表")
    rows = block.rows[left + 2:left + 20]
    left_chart = charts_in(slide)[0]
    replace_chart(
        left_chart,
        row_period_categories(rows),
        [
            ("乘用车库存量", [number(cell(row, 2)) for row in rows]),
            ("乘用车存销比", [number(cell(row, 3)) for row in rows]),
        ],
    )
    set_plot_label_formats(left_chart, ("0", "0.0"))
    right = block.find("右侧图表")
    rows = block.rows[right + 3:right + 7]
    right_chart = charts_in(slide)[1]
    replace_chart(
        right_chart,
        [(data.period_short, normalize(cell(row, 1))) for row in rows],
        [
            ("库存量", [number(cell(row, 2)) for row in rows]),
            ("存销比", [number(cell(row, 3)) for row in rows]),
        ],
    )
    set_plot_label_formats(right_chart, ("0", "0.0"))


def update_standard_comparison_slide(
    slide: object,
    block: PageBlock,
    data: FullMonthlyData,
    page: int,
) -> None:
    comp = parse_comparison(block)
    charts = charts_in(slide)
    blanks = [""] * len(comp.labels)
    replace_chart(
        charts[0],
        blanks,
        [
            ("", to_wan_list(comp.monthly_last)),
            ("", to_wan_list(comp.monthly_current)),
        ],
    )
    replace_chart(
        charts[1],
        blanks,
        [
            ("", to_wan_list(comp.cumulative_last)),
            ("", to_wan_list(comp.cumulative_current)),
        ],
    )
    replace_chart(
        charts[2],
        [""] * 4,
        [
            (
                "",
                [
                    round((period[index] or 0) * 100, 1)
                    for period in comp.share_periods
                ],
            )
            for index in reversed(range(len(comp.labels)))
        ],
        fixed_value_axis_maximum=100.0,
    )
    write_table_row(
        tables_in(slide)[0],
        0,
        ["环比", *[fmt_pct(value) for value in comp.monthly_mom]],
    )
    update_rate_callouts(
        slide,
        [*comp.monthly_yoy, *comp.cumulative_yoy],
        plus=True,
    )
    update_small_share_labels(slide, comp.share_periods)


def parse_comparison(block: PageBlock) -> ComparisonData:
    upper = block.find("左图上")
    labels = nonempty(block.rows[upper][1:])
    lower = block.find("左图下")
    right = block.find("右图")
    monthly_share = (
        numeric_slice(block.rows[upper + 4], 1, len(labels))
        if normalize(cell(block.rows[upper + 4], 0)) == "份额"
        else shares(numeric_slice(block.rows[upper + 2], 1, len(labels)))
    )
    cumulative_share = (
        numeric_slice(block.rows[lower + 4], 1, len(labels))
        if normalize(cell(block.rows[lower + 4], 0)) == "份额"
        else shares(numeric_slice(block.rows[lower + 2], 1, len(labels)))
    )
    right_header = right
    if not nonempty(block.rows[right_header][1:5]):
        candidate = block.rows[right_header + 1]
        if normalize(cell(candidate, 0)) in {"行标签", "分类"}:
            right_header += 1
    period_labels = nonempty(block.rows[right_header][1:5])
    share_periods = [[None] * len(labels) for _ in range(4)]
    for label_index in range(len(labels)):
        source = block.rows[right_header + 1 + label_index]
        values = numeric_slice(source, 1, 4)
        for period_index, value in enumerate(values):
            share_periods[period_index][label_index] = value
    return ComparisonData(
        labels=labels,
        monthly_last=numeric_slice(block.rows[upper + 1], 1, len(labels)),
        monthly_current=numeric_slice(block.rows[upper + 2], 1, len(labels)),
        monthly_yoy=numeric_slice(block.rows[upper + 3], 1, len(labels)),
        monthly_mom=numeric_slice(block.rows[upper + 4], 1, len(labels))
        if normalize(cell(block.rows[upper + 4], 0)) == "环比"
        else numeric_slice(block.rows[upper + 5], 1, len(labels)),
        monthly_share=monthly_share,
        cumulative_last=numeric_slice(block.rows[lower + 1], 1, len(labels)),
        cumulative_current=numeric_slice(block.rows[lower + 2], 1, len(labels)),
        cumulative_yoy=numeric_slice(block.rows[lower + 3], 1, len(labels)),
        cumulative_share=cumulative_share,
        share_periods=share_periods,
        share_period_labels=period_labels,
    )


def update_segment_slide(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    comp = parse_comparison(block)
    charts = charts_in(slide)
    blanks = [""] * len(comp.labels)
    replace_chart(charts[0], blanks, [("", to_wan_list(comp.monthly_last)), ("", to_wan_list(comp.monthly_current))])
    replace_chart(charts[1], blanks, [("", to_wan_list(comp.cumulative_last)), ("", to_wan_list(comp.cumulative_current))])
    models = parse_model_rows(block, limit=15)
    replace_chart(
        charts[2],
        [item.name for item in models],
        [(str(data.year), [to_wan(item.volume) or 0 for item in models])],
    )
    color_model_chart(charts[2], [item.name for item in models])
    tables = tables_in(slide)
    write_table_row(tables[0], 0, ["环比", *[fmt_pct(value) for value in comp.monthly_mom]])
    write_table_row(tables[0], 1, ["份额", *[fmt_pct(value) for value in comp.monthly_share]])
    write_table_row(tables[1], 0, ["份额", *[fmt_pct(value) for value in comp.cumulative_share]])
    write_model_table(tables[2], models, cumulative=True)
    total = sum(value or 0 for value in comp.cumulative_current)
    top_total = sum(item.cumulative or 0 for item in models)
    write_table_row(
        tables[3],
        0,
        [normalize(tables[3].cell(0, 0).text), fmt_pct(top_total / total if total else None)],
    )
    update_rate_callouts(slide, [*comp.monthly_yoy, *comp.cumulative_yoy], plus=True)


def parse_model_rows(block: PageBlock, limit: int) -> list[ModelRow]:
    header = next(
        index for index, row in enumerate(block.rows)
        if normalize(cell(row, 1)) == "车型" and normalize(cell(row, 2)) in {"销量", "累计销量"}
    )
    rows: list[ModelRow] = []
    for source in block.rows[header + 1:header + 1 + limit]:
        if not normalize(cell(source, 0)).isdigit():
            break
        rows.append(
            ModelRow(
                name=normalize(cell(source, 1)),
                volume=number(cell(source, 2)),
                yoy_raw=normalize(cell(source, 3)),
                mom_raw=normalize(cell(source, 4)),
                level=normalize(cell(source, 5)),
                share_raw=normalize(cell(source, 6)),
                rank_change_raw=normalize(cell(source, 7)),
                cumulative=number(cell(source, 8)),
            ),
        )
    return rows


def update_slide14(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    start = block.find("中间图表")
    rows = block.rows[start + 2:start + 33]
    replace_chart(
        charts_in(slide)[0],
        [normalize(cell(row, 0)) for row in rows],
        [
            ("销量", [number(cell(row, 1)) for row in rows]),
            ("同比增速", [number(cell(row, 2)) for row in rows]),
            ("整体市场增速", [number(cell(row, 3)) for row in rows]),
        ],
    )
    table_start = block.find("下方表格")
    table = tables_in(slide)[0]
    header = block.rows[table_start + 1]
    write_table_row(table, 0, ["", *nonempty(header[1:5])])
    for index, source in enumerate(block.rows[table_start + 2:table_start + 5], 1):
        write_table_row(table, index, [normalize(cell(source, 0)), *[fmt_pct(value, plus=index == 3) for value in source[1:5]]])


def update_slide16(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    update_company_slide(slide, block, data, top_n=15, cumulative_yoy_row=False)


def update_company_slide(
    slide: object,
    block: PageBlock,
    data: FullMonthlyData,
    *,
    top_n: int,
    cumulative_yoy_row: bool,
) -> None:
    upper = block.find("上方图表")
    upper_labels = nonempty(block.rows[upper + 1][1:1 + top_n])
    upper_last = numeric_slice(block.rows[upper + 2], 1, top_n)
    upper_current = numeric_slice(block.rows[upper + 3], 1, top_n)
    upper_yoy = numeric_slice(block.rows[upper + 4], 1, top_n)
    lower = block.find("下方图表")
    lower_label_row = lower if nonempty(block.rows[lower][1:1 + top_n]) else lower + 1
    lower_labels = nonempty(block.rows[lower_label_row][1:1 + top_n])
    lower_last = numeric_slice(block.rows[lower_label_row + 1], 1, top_n)
    lower_current = numeric_slice(block.rows[lower_label_row + 2], 1, top_n)
    if (
        cumulative_yoy_row
        and lower_label_row + 3 < len(block.rows)
        and normalize(cell(block.rows[lower_label_row + 3], 0)) == "同比"
    ):
        lower_yoy = numeric_slice(block.rows[lower_label_row + 3], 1, top_n)
    else:
        lower_yoy = [
            current / last - 1 if last not in {None, 0} and current is not None else None
            for last, current in zip(lower_last, lower_current)
        ]
    charts = charts_in(slide)
    replace_chart(charts[0], [""] * top_n, [("", upper_last), ("", upper_current)])
    replace_chart(charts[1], [""] * top_n, [("", lower_last), ("", lower_current)])
    table_start = block.find("中间表格")
    table = tables_in(slide)[0]
    for index, source in enumerate(block.rows[table_start + 1:table_start + 5]):
        label = normalize(cell(source, 0))
        write_table_row(
            table,
            index,
            [label, *[
                fmt_rank(value) if "排名" in label else fmt_pct(value, plus="Δ" in label)
                for value in source[1:1 + top_n]
            ]],
        )
    update_company_labels(slide, upper_labels, lower_labels)
    update_rate_callouts(slide, [*upper_yoy, *lower_yoy], plus=True)


def update_slide17(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    top = block.find("上方图表")
    top_rows = block.rows[top + 2:top + 14]
    lower = block.find("下方图表")
    lower_rows = block.rows[lower + 2:lower + 14]
    charts = charts_in(slide)
    replace_chart(
        charts[0],
        [normalize(cell(row, 0)) for row in top_rows],
        [
            (str(data.last_year), [number(cell(row, 1)) for row in top_rows]),
            (str(data.year), [number(cell(row, 2)) for row in top_rows]),
        ],
    )
    replace_chart(
        charts[1],
        [""] * 12,
        [
            ("", [number(cell(row, 1)) for row in lower_rows]),
            ("", [number(cell(row, 2)) for row in lower_rows]),
        ],
    )
    right = block.find("右侧图表")
    cumulative = block.rows[right + 2]
    replace_chart(
        charts[2],
        ["", ""],
        [("", [number(cell(cumulative, 1)), number(cell(cumulative, 2))])],
    )
    right_table = block.find("右侧表格")
    write_table_row(
        tables_in(slide)[0],
        0,
        [
            normalize(cell(block.rows[right_table + 2], 0)),
            fmt_pct(cell(block.rows[right_table + 2], 1)),
            fmt_pct(cell(block.rows[right_table + 2], 2)),
        ],
    )
    months = nonempty(block.rows[lower + 14][1:13])
    yoy = numeric_slice(block.rows[lower + 15], 1, len(months))
    mom = numeric_slice(block.rows[lower + 16], 1, len(months))
    write_table_row(
        tables_in(slide)[1],
        0,
        ["环比", *[fmt_pct(value) for value in mom], *([""] * (12 - len(mom)))],
    )
    cumulative_yoy = number(cell(block.rows[right + 3], 2))
    update_rate_callouts(slide, [*yoy, cumulative_yoy], plus=True, include_percent_shapes=True)


def update_slide18(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    left = block.find("左侧图表")
    rows = block.rows[left + 2:left + 15]
    chart = charts_in(slide)[0]
    replace_chart(
        chart,
        row_period_categories(rows),
        [
            ("纯电", [number(cell(row, 2)) for row in rows]),
            ("插电混动", [number(cell(row, 3)) for row in rows]),
            ("纯电增速", [number(cell(row, 4)) for row in rows]),
            ("插电混动增速", [number(cell(row, 5)) for row in rows]),
        ],
    )
    set_plot_label_formats(chart, ("0.0", "0.0%"))
    table_start = block.find("左下方表格")
    table = tables_in(slide)[0]
    for row_index, source in enumerate(block.rows[table_start + 2:table_start + 6]):
        label = normalize(cell(source, 0))
        group, energy = (label.split("/", 1) + [""])[:2]
        write_table_row(
            table,
            row_index,
            [group.replace("占比", "\n占比") if row_index % 2 == 0 else "", energy, *[
                fmt_pct(value) for value in source[1:14]
            ]],
        )
    right = block.find("右侧图表")
    last_row = block.rows[right + 2]
    current_row = block.rows[right + 3]
    replace_chart(
        charts_in(slide)[1],
        ["", ""],
        [
            ("", [number(cell(last_row, 1)), number(cell(current_row, 1))]),
            ("", [number(cell(last_row, 2)), number(cell(current_row, 2))]),
        ],
    )
    right_table = block.find("右下方表格")
    for index, source in enumerate(block.rows[right_table + 2:right_table + 4]):
        write_table_row(
            tables_in(slide)[1],
            index,
            [normalize(cell(source, 0)), fmt_pct(cell(source, 1), plus=index == 1), fmt_pct(cell(source, 2), plus=index == 1)],
        )
    update_slide18_number_labels(
        slide,
        [
            to_wan(number(cell(last_row, 1))),
            to_wan(number(cell(last_row, 2))),
            to_wan(number(cell(current_row, 1))),
            to_wan(number(cell(current_row, 2))),
        ],
    )
    yoy_row = block.rows[right + 4]
    update_rate_callouts(slide, numeric_slice(yoy_row, 1, 2), plus=True)


def update_slide19(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    anchors = ("左图上", "左图下", "右图上", "右图下")
    charts = charts_in(slide)
    tables = tables_in(slide)
    all_yoy: list[float | None] = []
    for index, anchor in enumerate(anchors):
        start = block.find(anchor)
        labels = nonempty(block.rows[start + 1][1:])
        last = numeric_slice(block.rows[start + 2], 1, len(labels))
        current = numeric_slice(block.rows[start + 3], 1, len(labels))
        yoy = numeric_slice(block.rows[start + 4], 1, len(labels))
        all_yoy.extend(yoy)
        replace_chart(charts[index], [""] * len(labels), [("", last), ("", current)])
        if index in {0, 2}:
            mom = numeric_slice(block.rows[start + 5], 1, len(labels))
            share = numeric_slice(block.rows[start + 6], 1, len(labels))
            write_table_row(tables[index], 0, ["环比", *[fmt_pct(value) for value in mom]])
            write_table_row(tables[index], 1, ["份额", *[fmt_pct(value) for value in share]])
        else:
            share = numeric_slice(block.rows[start + 5], 1, len(labels))
            write_table_row(tables[index], 0, ["份额", *[fmt_pct(value) for value in share]])
    update_rate_callouts(slide, all_yoy, plus=True)
    update_phev_small_value_label(slide, number(cell(block.rows[block.find("右图上") + 3], 1)))


def update_slide20(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    update_company_slide(slide, block, data, top_n=10, cumulative_yoy_row=True)


def update_slide21(slide: object, block: PageBlock, data: FullMonthlyData) -> None:
    left = parse_model_rows(block, limit=20)
    right_header = next(
        index for index, row in enumerate(block.rows)
        if normalize(cell(row, 1)) == "车型" and normalize(cell(row, 2)) == "累计销量"
    )
    right: list[ModelRow] = []
    for source in block.rows[right_header + 1:right_header + 21]:
        right.append(
            ModelRow(
                name=normalize(cell(source, 1)),
                volume=number(cell(source, 2)),
                yoy_raw=normalize(cell(source, 3)),
                mom_raw="",
                level=normalize(cell(source, 5)),
                share_raw=normalize(cell(source, 4)),
                rank_change_raw=normalize(cell(source, 6)),
            ),
        )
    charts = charts_in(slide)
    replace_chart(charts[0], [item.name for item in left], [(str(data.year), [item.volume or 0 for item in left])])
    replace_chart(charts[1], [item.name for item in right], [(str(data.year), [item.volume or 0 for item in right])])
    color_model_chart(charts[0], [item.name for item in left])
    color_model_chart(charts[1], [item.name for item in right])
    tables = tables_in(slide)
    write_model_table(tables[0], left, cumulative=False, include_name=False)
    left_share_row, left_share_col = block.find_any("累计份额", 0)
    right_share_row, right_share_col = block.find_any("累计份额", 1)
    write_table_row(
        tables[1],
        0,
        [
            normalize(tables[1].cell(0, 0).text),
            fmt_pct(cell(block.rows[left_share_row + 1], left_share_col)),
        ],
    )
    write_table_row(
        tables[2],
        0,
        [
            normalize(tables[2].cell(0, 0).text),
            fmt_pct(cell(block.rows[right_share_row + 1], right_share_col)),
        ],
    )
    write_model_table(tables[3], right, cumulative=False, include_name=False, cumulative_ranking=True)


def write_model_table(
    table: object,
    models: Sequence[ModelRow],
    *,
    cumulative: bool,
    include_name: bool = False,
    cumulative_ranking: bool = False,
) -> None:
    if cumulative_ranking:
        headers = ["同比", "份额", "级别", "排名变化"]
    else:
        headers = ["同比", "环比", "级别", "份额", "排名变化"]
        if cumulative:
            headers.append("累计")
    if include_name:
        headers.insert(0, "车型")
    write_table_row(table, 0, headers)
    for index in range(1, len(table.rows)):
        model = models[index - 1] if index - 1 < len(models) else None
        if model is None:
            write_table_row(table, index, [""] * len(headers))
            continue
        if cumulative_ranking:
            values = [
                fmt_pct(model.yoy_raw),
                fmt_pct(model.share_raw),
                model.level,
                fmt_rank(model.rank_change_raw),
            ]
        else:
            values = [
                fmt_pct(model.yoy_raw),
                fmt_pct(model.mom_raw),
                model.level,
                fmt_pct(model.share_raw),
                fmt_rank(model.rank_change_raw),
            ]
            if cumulative:
                values.append(fmt_number(model.cumulative, 1))
        if include_name:
            values.insert(0, model.name)
        write_table_row(table, index, values)


def apply_markdown_or_fallback(
    slide: object,
    markdown: str,
    heading_pattern: str,
    occurrence: int,
    data: FullMonthlyData,
    block: PageBlock,
) -> None:
    heading, body = extract_markdown_page(markdown, heading_pattern, occurrence)
    if heading:
        set_slide_title(slide, heading)
    if body:
        set_slide_body(slide, body, max_chars=240)
    if heading or body or (data.year == 2026 and data.month == 6):
        return
    fallback_title, fallback_body = build_fallback_text(block, data)
    if fallback_title:
        set_slide_title(slide, fallback_title)
    if fallback_body:
        set_slide_body(slide, fallback_body, max_chars=220)


def extract_markdown_page(markdown: str, pattern: str, occurrence: int) -> tuple[str, str]:
    heading_re = re.compile(rf"(?mi)^###?\s*(?P<heading>{pattern}[^\n]*)$")
    matches = list(heading_re.finditer(markdown or ""))
    if occurrence >= len(matches):
        return "", ""
    match = matches[occurrence]
    next_heading = re.search(r"(?m)^#{2,4}\s+", (markdown or "")[match.end():])
    end = match.end() + next_heading.start() if next_heading else len(markdown or "")
    body = (markdown or "")[match.end():end]
    body = re.sub(r"〔[^〕]*[ER]\d+[^〕]*〕", "", body)
    body = re.sub(r"(?m)^\s*[-*◼]\s*", "", body)
    body = re.sub(r"(?m)^\|.*\|\s*$", "", body)
    body = re.sub(r"\s+", " ", body).strip()
    return normalize(match.group("heading"))[:90], body


def build_fallback_text(block: PageBlock, data: FullMonthlyData) -> tuple[str, str]:
    try:
        comp = parse_comparison(block)
    except (ValueError, IndexError):
        return "", ""
    valid = [(label, rate) for label, rate in zip(comp.labels, comp.monthly_yoy) if rate is not None]
    if not valid:
        return "", ""
    best_label, best = max(valid, key=lambda item: item[1])
    worst_label, worst = min(valid, key=lambda item: item[1])
    prefix = {
        7: "2.1 车型类别",
        8: "2.2 车型级别",
        9: "2.3 轿车级别",
        10: "2.3 SUV级别",
        11: "2.3 MPV级别",
        12: "2.4 车型系别",
        13: "2.5 车型价格段",
        15: "2.6 区域市场",
    }.get(block.page, f"P{block.page}")
    return (
        f"{prefix}：{best_label}表现相对较好，{worst_label}同比承压",
        f"{data.month}月，{best_label}同比{fmt_pct(best)}，在本页维度中表现相对较好；"
        f"{worst_label}同比{fmt_pct(worst)}，是主要结构性拖累。具体产品、政策及供给机制需结合当月证据核验。",
    )


def replace_period_text(slide: object, data: FullMonthlyData) -> None:
    replacements = (
        (r"2026年6月", data.period_label),
        (r"2026\.5(?!\d)", data.previous_period_short),
        (r"2026\.6(?!\d)", data.period_short),
        (r"2025\.1-12", f"{data.last_year}.1-12"),
        (r"2025\.1-6", f"{data.last_year}.1-{data.month}"),
        (r"2026\.1-6", f"{data.year}.1-{data.month}"),
    )
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        original = shape.text
        updated = original
        for pattern, replacement in replacements:
            updated = re.sub(pattern, replacement, updated)
        if updated != original:
            set_text_frame_text(shape.text_frame, updated)


def update_year_labels(slide: object, data: FullMonthlyData) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = normalize(shape.text)
        if text == "2025" and text != str(data.last_year):
            set_text_frame_text(shape.text_frame, str(data.last_year))
        elif text == "2026" and text != str(data.year):
            set_text_frame_text(shape.text_frame, str(data.year))


def normalize_source_footer(slide: object) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if "数据来源：" in shape.text and "中汽数据有限公司" not in shape.text:
            set_text_frame_text(shape.text_frame, "数据来源：中汽数据有限公司终端零售数据")


def update_rate_callouts(
    slide: object,
    values: Sequence[float | None],
    *,
    plus: bool,
    include_percent_shapes: bool = False,
) -> None:
    shapes = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or getattr(shape, "has_table", False):
            continue
        text = normalize(shape.text)
        if "椭圆" in (shape.name or "") or (
            include_percent_shapes and RATE_RE.fullmatch(text)
        ):
            shapes.append(shape)
    for shape, value in zip(shapes, values):
        set_text_frame_text(shape.text_frame, fmt_pct(value, plus=plus))
        paragraph = shape.text_frame.paragraphs[0]
        if paragraph.runs:
            run = paragraph.runs[0]
            run.font.name = "微软雅黑"
            run.font.size = run.font.size
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 112)


def update_small_share_labels(slide: object, periods: Sequence[Sequence[float | None]]) -> None:
    chart_shapes = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
    if len(chart_shapes) < 3:
        return
    chart = chart_shapes[2]
    centers = [chart.left + int(chart.width * (0.14 + index * 0.24)) for index in range(4)]
    candidates = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and "椭圆" not in (shape.name or "")
        and RATE_RE.fullmatch(normalize(shape.text))
        and shape.left > chart.left
        and shape.top > chart.top + int(chart.height * 0.55)
    ]
    for shape in candidates:
        center = shape.left + shape.width // 2
        period_index = min(range(4), key=lambda index: abs(centers[index] - center))
        period = [value for value in periods[period_index] if value is not None]
        if period:
            set_text_frame_text(shape.text_frame, fmt_pct(min(period)))


def update_company_labels(
    slide: object,
    monthly: Sequence[str],
    cumulative: Sequence[str],
) -> None:
    rectangles = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.shape_type.name == "AUTO_SHAPE"
        and "椭圆" not in (shape.name or "")
        and normalize(shape.text) not in {"2025", "2026"}
        and not RATE_RE.fullmatch(normalize(shape.text))
    ]
    upper = sorted(
        [shape for shape in rectangles if 3.7 <= shape.top / 914400 <= 4.4],
        key=lambda shape: shape.left,
    )
    lower = sorted(
        [shape for shape in rectangles if 6.35 <= shape.top / 914400 <= 7.0],
        key=lambda shape: shape.left,
    )
    for shape, value in zip(upper, monthly):
        set_text_frame_text(shape.text_frame, value)
    for shape, value in zip(lower, cumulative):
        set_text_frame_text(shape.text_frame, value)


def update_slide18_number_labels(slide: object, values: Sequence[float | None]) -> None:
    labels = [
        shape for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
        and shape.shape_type.name == "AUTO_SHAPE"
        and NUMBER_RE.fullmatch(normalize(shape.text))
        and 10 <= shape.left / 914400 <= 13
        and 3 <= shape.top / 914400 <= 5
    ]
    for shape, value in zip(labels, values):
        set_text_frame_text(shape.text_frame, fmt_number(value, 1))


def update_phev_small_value_label(slide: object, value: float | None) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if (
            shape.shape_type.name == "AUTO_SHAPE"
            and "椭圆" not in (shape.name or "")
            and NUMBER_RE.fullmatch(normalize(shape.text))
            and shape.left / 914400 > 7
            and 3.8 < shape.top / 914400 < 4.4
        ):
            set_text_frame_text(shape.text_frame, fmt_number(value, 1))
            return


def color_model_chart(chart: object, names: Sequence[str]) -> None:
    if not chart.series:
        return
    points = chart.series[0].points
    for index, name in enumerate(names):
        if index >= len(points):
            break
        fill = points[index].format.fill
        fill.solid()
        fill.fore_color.rgb = TRADITIONAL_BLUE if is_traditional_model(name) else NEW_ENERGY_GREEN


def set_plot_label_formats(chart: object, formats: Sequence[str]) -> None:
    """固定组合图各 plot 的标签格式，避免数据替换后显示长小数。"""

    for plot, number_format in zip(chart.plots, formats):
        if not plot.has_data_labels:
            continue
        plot.data_labels.number_format = number_format
        plot.data_labels.number_format_is_linked = False
        # 业务母版对部分系列保存了独立 dLbls；它会覆盖 plot 级格式。
        # python-pptx 尚未公开 Series.data_labels，因此在保留布局/字体的同时，
        # 只更新系列 dLbls 中的 numFmt。
        for series in plot.series:
            data_labels = series._ser.find(qn("c:dLbls"))
            if data_labels is None:
                continue
            number_format_element = data_labels.find(qn("c:numFmt"))
            if number_format_element is None:
                number_format_element = OxmlElement("c:numFmt")
                insert_at = 0
                for child in data_labels:
                    if child.tag == qn("c:dLbl"):
                        insert_at += 1
                    else:
                        break
                data_labels.insert(insert_at, number_format_element)
            number_format_element.set("formatCode", number_format)
            number_format_element.set("sourceLinked", "0")


def write_table_row(table: object, row_index: int, values: Sequence[str]) -> None:
    if row_index >= len(table.rows):
        return
    for column_index, value in enumerate(values):
        if column_index >= len(table.columns):
            break
        set_text_frame_text(table.cell(row_index, column_index).text_frame, str(value))


def verify_output(path: Path) -> dict[str, int]:
    presentation = Presentation(path)
    charts = sum(
        1 for slide in presentation.slides for shape in slide.shapes
        if getattr(shape, "has_chart", False)
    )
    tables = sum(
        1 for slide in presentation.slides for shape in slide.shapes
        if getattr(shape, "has_table", False)
    )
    if len(presentation.slides) != 23 or charts < 35 or tables < 20:
        raise ValueError(
            f"生成文件结构异常：slides={len(presentation.slides)}, charts={charts}, tables={tables}",
        )
    return {"slides": len(presentation.slides), "charts": charts, "tables": tables}


def period_categories(labels: Sequence[str]) -> list[str | tuple[str, str]]:
    """把 `2026.6` 等期间转为PowerPoint双层“年份/月”分类轴。"""

    pairs: list[tuple[str, str]] = []
    for label in labels:
        match = re.fullmatch(r"(20\d{2})(?:年|[./-])(\d{1,2})(?:月)?", normalize(label))
        if not match:
            return list(labels)
        pairs.append((match.group(1), f"{int(match.group(2))}月"))
    return pairs


def row_period_categories(rows: Sequence[Sequence[str]]) -> list[tuple[str, str]]:
    """读取“年份仅在分组首行出现、月份逐行出现”的验数表分类轴。"""

    categories: list[tuple[str, str]] = []
    current_year = ""
    for row in rows:
        year = normalize(cell(row, 0))
        if year:
            current_year = year
        month = normalize(cell(row, 1))
        if not current_year or not month:
            raise ValueError("期间分类轴缺少年份或月份。")
        categories.append((current_year, month))
    return categories


def numeric_slice(row: Sequence[str], start: int, count: int) -> list[float | None]:
    return [number(cell(row, start + index)) for index in range(count)]


def shares(values: Sequence[float | None]) -> list[float]:
    total = sum(value or 0 for value in values)
    return [(value or 0) / total if total else 0 for value in values]


def to_wan_list(values: Sequence[float | None]) -> list[float | None]:
    return [to_wan(value) for value in values]


def to_wan(value: float | None) -> float | None:
    if value is None:
        return None
    return value / 10000 if abs(value) >= 1000 else value


def fmt_number(value: object, decimals: int = 1) -> str:
    parsed = number(value)
    if parsed is None:
        return ""
    if decimals == 0:
        return f"{parsed:.0f}"
    return f"{parsed:.{decimals}f}"


def fmt_pct(value: object, *, plus: bool = False) -> str:
    text = normalize(value)
    if not text or text.startswith("#") or text in {"-", "--"}:
        return "-"
    if text.startswith(">") or text.startswith("<"):
        return text
    parsed = number(text)
    if parsed is None:
        return text
    if abs(parsed) >= 10:
        return ">1000%" if parsed > 0 else "<-1000%"
    prefix = "+" if plus and parsed > 0 else ""
    return f"{prefix}{parsed * 100:.1f}%"


def fmt_rank(value: object) -> str:
    text = normalize(value)
    if not text or text == "0":
        return "-"
    if text == "--":
        return "--"
    try:
        parsed = int(float(text))
    except ValueError:
        return text
    return f"+{parsed}" if parsed > 0 else str(parsed) if parsed < 0 else "-"


def number(value: object) -> float | None:
    text = normalize(value).replace(",", "")
    if not text or text.startswith("#") or text in {"-", "--"} or text.startswith(">") or text.startswith("<"):
        return None
    if text.endswith("%"):
        try:
            return float(text[:-1]) / 100
        except ValueError:
            return None
    try:
        return float(text)
    except ValueError:
        return None


def nonempty(values: Iterable[object]) -> list[str]:
    return [normalize(value) for value in values if normalize(value) and normalize(value) != "`"]


def cell(row: Sequence[object], index: int) -> object:
    return row[index] if 0 <= index < len(row) else ""


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def sheet_month(name: str) -> int:
    match = re.search(r"(1[0-2]|[1-9])月", name)
    return int(match.group(1)) if match else 0
