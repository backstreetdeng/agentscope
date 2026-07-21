"""乘用车市场月报可编辑 PPTX 生成器。

该模块以业务 Profile 中的 ``monthly_report_template.pptx`` 为母版，而不是从零绘制
HTML 页面。母版中的图表、表格和文本框均为 PowerPoint 原生对象；本模块只替换
业务数据与文字内容，生成的 ``.pptx`` 可以继续在 Microsoft PowerPoint 中编辑。

页面映射严格沿用业务近三个月样例：
1. 车型类别；2. 车型级别；3. 轿车分级别；4. SUV 分级别；5. MPV 分级别。
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

from skills.market_analysis_skill.runtime import SheetSnapshot, read_xlsx_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
BUSINESS_MONTHLY_SKILL_DIR = (
    PROJECT_DIR / "skills" / "market_analysis_skill" / "profiles" / "passenger_vehicle_monthly"
)
DEFAULT_TEMPLATE_PATH = BUSINESS_MONTHLY_SKILL_DIR / "assets" / "monthly_report_template.pptx"

# 业务样例的原生表格用蓝/绿区分传统和新能源车型。验数表尚未提供能源字段，下面的
# 名单只用于已知传统燃油车型的保守识别；未知车型按新能源绿色处理，避免把新能源误标为传统。
TRADITIONAL_MODEL_TOKENS = (
    "轩逸",
    "凯美瑞",
    "迈腾",
    "奔驰",
    "宝马",
    "星越L",
    "缤越",
    "卡罗拉锐放",
    "锋兰达",
    "别克GL8",
    "赛那",
    "格瑞维亚",
    "五菱宏光",
    "传祺M8",
    "五菱宏光V",
    "五菱宏光S",
    "传祺M6",
)
TRADITIONAL_BLUE = RGBColor(47, 85, 151)
NEW_ENERGY_GREEN = RGBColor(0, 176, 80)


@dataclass(frozen=True)
class TopModel:
    """TOP 车型明细，所有销量单位沿用业务验数表的“万辆”口径。"""

    name: str
    volume: float | None
    yoy: float | None
    mom: float | None
    level: str
    share: float | None
    rank_change: str
    cumulative_volume: float | None


@dataclass(frozen=True)
class MarketBlock:
    """一类市场（车型类别、级别或细分市场）的结构化数据。"""

    title: str
    labels: list[str]
    monthly_last_year: list[float | None]
    monthly_current: list[float | None]
    yoy: list[float | None]
    mom: list[float | None]
    cumulative_last_year: list[float | None]
    cumulative_current: list[float | None]
    cumulative_yoy: list[float | None]
    share_periods: list[list[float | None]]
    top_models: list[TopModel]

    @property
    def monthly_share(self) -> list[float]:
        return shares(self.monthly_current)

    @property
    def cumulative_share(self) -> list[float]:
        return shares(self.cumulative_current)


@dataclass(frozen=True)
class MonthlyPptData:
    """五页月报所需的数据快照。"""

    period_label: str
    period_short: str
    previous_period_short: str
    year: str
    month: int
    category: MarketBlock
    level: MarketBlock
    sedan: MarketBlock
    suv: MarketBlock
    mpv: MarketBlock


def build_monthly_report_pptx(
    markdown: str,
    output_path: Path,
    dataset_path: Path,
    template_path: Path | None = None,
    sheet_name: str | None = None,
) -> dict[str, object]:
    """将月报文字与验数表写入业务母版，生成 5 页可编辑 PPTX。"""

    template = (template_path or DEFAULT_TEMPLATE_PATH).resolve()
    if not template.is_file():
        raise FileNotFoundError(f"未找到月报 PPTX 母版：{template}")
    if not dataset_path.is_file():
        raise FileNotFoundError(f"未找到月报数据集：{dataset_path}")

    data = load_monthly_ppt_data(dataset_path, sheet_name)
    presentation = Presentation(template)
    if len(presentation.slides) < 5:
        raise ValueError("业务月报母版少于 5 页，无法套用固定月报页面映射。")

    sections = {
        "category": extract_report_section(markdown, r"2\.1\s*车型类别"),
        "level": extract_report_section(markdown, r"2\.2\s*车型级别"),
        "sedan": extract_report_section(markdown, r"2\.3\s*轿车"),
        "suv": extract_report_section(markdown, r"2\.4\s*SUV"),
        "mpv": extract_report_section(markdown, r"2\.5\s*MPV"),
    }
    headings = {
        "category": extract_report_heading(markdown, r"2\.1\s*车型类别"),
        "level": extract_report_heading(markdown, r"2\.2\s*车型级别"),
        "sedan": extract_report_heading(markdown, r"2\.3\s*轿车"),
        "suv": extract_report_heading(markdown, r"2\.4\s*SUV"),
        "mpv": extract_report_heading(markdown, r"2\.5\s*MPV"),
    }

    update_overview_slide(
        presentation.slides[0], data.category, sections["category"], "2.1 车型类别", data, "车型类别", headings["category"],
    )
    update_overview_slide(
        presentation.slides[1], data.level, sections["level"], "2.2 车型级别", data, "车型级别", headings["level"],
    )
    update_segment_slide(presentation.slides[2], data.sedan, sections["sedan"], "2.3 轿车级别", "轿车", data, headings["sedan"])
    update_segment_slide(presentation.slides[3], data.suv, sections["suv"], "2.4 SUV级别", "SUV", data, headings["suv"])
    update_segment_slide(presentation.slides[4], data.mpv, sections["mpv"], "2.5 MPV级别", "MPV", data, headings["mpv"])

    for slide_index, slide in enumerate(list(presentation.slides)[:5]):
        replace_period_labels(slide, data)
        update_data_source_footer(slide)
        # 业务母版的右侧份额图在 PowerPoint 中会保留历史分类轴的渲染缓存。
        # 因此在图表上方补一层当前期间标签，保证选中任意月份时可见标签与图表数据一致。
        if slide_index < 2:
            add_share_period_label_overlays(slide, data)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return {
        "success": True,
        "pptx_path": str(output_path),
        "slides_count": 5,
        "template": str(template),
        "period": data.period_label,
        "sheet_name": sheet_name or "",
        "file_size": output_path.stat().st_size,
        "engine": "business-monthly-pptx-template",
    }


def load_monthly_ppt_data(dataset_path: Path, sheet_name: str | None = None) -> MonthlyPptData:
    """从业务验数表读取本期与上期的五页制图数据。"""

    sheets = read_xlsx_workbook(dataset_path)
    ordered = sorted(sheets, key=sheet_month_number, reverse=True)
    if not ordered:
        raise ValueError(f"验数表 {dataset_path.name} 不包含可读取的工作表。")
    if sheet_name:
        target = normalize(sheet_name).lower()
        current_sheet = next((sheet for sheet in ordered if normalize(sheet.name).lower() == target), None)
        if current_sheet is None:
            raise ValueError(f"验数表 {dataset_path.name} 未找到工作表：{sheet_name}。")
    else:
        current_sheet = ordered[0]
    previous_sheet = next(
        (sheet for sheet in ordered if sheet_month_number(sheet) < sheet_month_number(current_sheet)),
        None,
    )
    year = infer_year(current_sheet) or "2026"
    current_month = sheet_month_number(current_sheet)
    previous_month = sheet_month_number(previous_sheet) if previous_sheet else max(1, current_month - 1)

    return MonthlyPptData(
        period_label=f"{year}年{current_month}月",
        period_short=f"{year}.{current_month}",
        previous_period_short=f"{year}.{previous_month}",
        year=year,
        month=current_month,
        category=parse_market_block(current_sheet, "车型类别"),
        level=parse_market_block(current_sheet, "车型级别"),
        sedan=parse_market_block(current_sheet, "轿车-级别"),
        suv=parse_market_block(current_sheet, "SUV-级别"),
        mpv=parse_market_block(current_sheet, "MPV-级别"),
    )


def parse_market_block(sheet: SheetSnapshot, title: str) -> MarketBlock:
    """解析一个展示型交叉表分区，兼容总览和轿车/SUV/MPV TOP15 分区。"""

    rows = sheet.rows
    if not rows:
        raise ValueError(f"工作表 {sheet.name} 为空。")
    start = find_section_column(rows, title)
    if start is None:
        raise ValueError(f"工作表 {sheet.name} 未找到 {title} 分区。")
    end = find_next_section_column(rows, start)
    labels = clean_labels(rows[2][start:end] if len(rows) > 2 else [])
    if not labels:
        raise ValueError(f"工作表 {sheet.name} 的 {title} 分区缺少级别标签。")

    pairs = year_pairs(rows, start)
    if not pairs:
        raise ValueError(f"工作表 {sheet.name} 的 {title} 分区缺少 2025/2026 对比数据。")
    monthly_pair = pairs[0]
    cumulative_pair = pairs[1] if len(pairs) > 1 else pairs[0]

    monthly_last_year = row_values(rows, monthly_pair[0], start, labels)
    monthly_current = row_values(rows, monthly_pair[1], start, labels)
    yoy = row_values(rows, monthly_pair[1] + 1, start, labels)
    mom = row_values(rows, monthly_pair[1] + 2, start, labels)
    cumulative_last_year = row_values(rows, cumulative_pair[0], start, labels)
    cumulative_current = row_values(rows, cumulative_pair[1], start, labels)
    cumulative_yoy = row_values(rows, cumulative_pair[1] + 1, start, labels)
    share_periods = parse_right_chart_share_periods(rows, start, labels) or [
        shares(monthly_last_year),
        shares(monthly_current),
        shares(cumulative_last_year),
        shares(cumulative_current),
    ]

    return MarketBlock(
        title=title,
        labels=labels,
        monthly_last_year=monthly_last_year,
        monthly_current=monthly_current,
        yoy=yoy,
        mom=mom,
        cumulative_last_year=cumulative_last_year,
        cumulative_current=cumulative_current,
        cumulative_yoy=cumulative_yoy,
        share_periods=share_periods,
        top_models=parse_top_models(rows, start, end),
    )


def update_overview_slide(
    slide: object,
    block: MarketBlock,
    narrative: str,
    prefix: str,
    data: MonthlyPptData,
    label: str,
    heading: str = "",
) -> None:
    """更新车型类别或车型级别总览页的结论、图表、环比与同比标签。"""

    title = build_headline(prefix, block, heading)
    body = narrative or build_data_narrative(block, data.period_label)
    last_year = str(int(data.year) - 1)
    set_slide_title(slide, title)
    # 总览页的正文区高于图表标题，控制为约四行，避免长文本侵入图表区域。
    set_slide_body(slide, body, max_chars=220)
    tables = tables_in(slide)
    if tables:
        write_row(tables[0], ["环比", *[format_rate(value) for value in block.mom]])

    charts = charts_in(slide)
    if len(charts) >= 1:
        replace_chart(
            charts[0],
            block.labels,
            [(last_year, values_or_zero(block.monthly_last_year)), (data.year, values_or_zero(block.monthly_current))],
        )
    if len(charts) >= 2:
        replace_chart(
            charts[1],
            block.labels,
            [(f"{last_year}累计", values_or_zero(block.cumulative_last_year)), (f"{data.year}累计", values_or_zero(block.cumulative_current))],
        )
    if len(charts) >= 3:
        # 母版右图是“按期间的车型类别份额”堆积柱：横轴固定为上期、当期、
        # 上年同期累计、当年累计；每根柱由轿车/SUV/MPV/交叉型乘用车构成。
        # 验数表中的份额是 0-1 小数，母版图表以 0-100 百分比点显示。
        period_shares = [percentage_points(values) for values in block.share_periods]
        replace_chart(
            charts[2],
            [
                data.previous_period_short,
                data.period_short,
                f"{last_year}.1-12",
                f"{data.year}.1-{data.month}",
            ],
            [
                (
                    label_name,
                    [period_share[index] for period_share in period_shares],
                )
                # PowerPoint 堆积柱按系列顺序自下而上绘制。业务母版要求：
                # 交叉型（深蓝）→ MPV → SUV → 轿车（浅蓝）。
                for index, label_name in reversed(list(enumerate(block.labels)))
            ],
        )
    update_oval_rates(slide, [*block.yoy, *block.cumulative_yoy])
    clear_stale_percentage_callouts(slide)


def update_segment_slide(
    slide: object,
    block: MarketBlock,
    narrative: str,
    prefix: str,
    segment_name: str,
    data: MonthlyPptData,
    heading: str = "",
) -> None:
    """更新轿车、SUV、MPV 三页的分级别图表、TOP15 车型表及排名柱图。"""

    title = build_headline(prefix, block, heading)
    body = narrative or build_data_narrative(block, data.period_label)
    last_year = str(int(data.year) - 1)
    set_slide_title(slide, title)
    # 三页细分市场模板右侧含 TOP15 表格，图表区的顶部更高。MPV 页正文框又比
    # 轿车/SUV 页更矮，因此分别采用更紧凑的单页摘要长度。
    body_limit = 145 if segment_name == "MPV" else 170
    set_slide_body(slide, body, max_chars=body_limit)

    tables = tables_in(slide)
    if len(tables) >= 1:
        write_row(tables[0], ["环比", *[format_rate(value) for value in block.mom]])
        if len(tables[0].rows) > 1:
            write_row(tables[0], ["份额", *[format_rate(value) for value in block.monthly_share]], row_index=1)
    if len(tables) >= 2:
        write_row(tables[1], ["份额", *[format_rate(value) for value in block.cumulative_share]])
    if len(tables) >= 3:
        write_top_models_table(tables[2], block.top_models)
    if len(tables) >= 4:
        cumulative_share = sum(model.cumulative_volume or 0 for model in block.top_models)
        total = sum(value or 0 for value in block.cumulative_current)
        write_row(
            tables[3],
            [f"{segment_name}TOP15车型累计份额", format_rate(cumulative_share / total if total else None)],
        )

    charts = charts_in(slide)
    if len(charts) >= 1:
        replace_chart(
            charts[0],
            block.labels,
            [(last_year, values_or_zero(block.monthly_last_year)), (data.year, values_or_zero(block.monthly_current))],
        )
    if len(charts) >= 2:
        replace_chart(
            charts[1],
            block.labels,
            [(f"{last_year}累计", values_or_zero(block.cumulative_last_year)), (f"{data.year}累计", values_or_zero(block.cumulative_current))],
        )
    if len(charts) >= 3:
        top_models = block.top_models[:15]
        replace_chart(
            charts[2],
            [model.name for model in top_models],
            [(data.year, [model.volume or 0 for model in top_models])],
        )
        color_top_model_chart(charts[2], top_models)
    update_oval_rates(slide, [*block.yoy, *block.cumulative_yoy])
    clear_stale_percentage_callouts(slide)


def replace_chart(chart: object, categories: Sequence[str], series: Sequence[tuple[str, Sequence[float]]]) -> None:
    """替换母版图表的缓存与嵌入数据表，图表仍保持 PowerPoint 原生可编辑。"""

    chart_data = CategoryChartData()
    chart_data.categories = list(categories)
    for name, values in series:
        chart_data.add_series(name, list(values))
    chart.replace_data(chart_data)


def color_top_model_chart(chart: object, models: Sequence[TopModel]) -> None:
    """按业务规范将传统车型标蓝、新能源车型标绿。"""

    if not chart.series:
        return
    series = chart.series[0]
    for index, model in enumerate(models):
        point = series.points[index]
        fill = point.format.fill
        fill.solid()
        fill.fore_color.rgb = TRADITIONAL_BLUE if is_traditional_model(model.name) else NEW_ENERGY_GREEN


def write_top_models_table(table: object, models: Sequence[TopModel]) -> None:
    """写入母版右侧 TOP15 指标表；车型名称由左侧柱图原生类别标签展示。"""

    headers = ["同比", "环比", "级别", "份额", "排名变化", "累计"]
    write_row(table, headers)
    for row_index in range(1, len(table.rows)):
        model = models[row_index - 1] if row_index - 1 < len(models) else None
        values = [
            format_rate(model.yoy) if model else "待补充",
            format_rate(model.mom) if model else "待补充",
            model.level if model else "待补充",
            format_rate(model.share) if model else "待补充",
            model.rank_change if model else "待补充",
            format_volume(model.cumulative_volume) if model else "待补充",
        ]
        write_row(table, values, row_index=row_index)


def write_row(table: object, values: Sequence[str], row_index: int = 0) -> None:
    """写入一行表格，保留母版的边框、填色、字号和对齐样式。"""

    if row_index >= len(table.rows):
        return
    for column_index, value in enumerate(values):
        if column_index >= len(table.columns):
            break
        set_text_frame_text(table.cell(row_index, column_index).text_frame, str(value))


def set_slide_title(slide: object, text: str) -> None:
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        if shape.placeholder_format.type == PP_PLACEHOLDER.TITLE:
            set_text_frame_text(shape.text_frame, truncate(text, 56))
            return


def set_slide_body(slide: object, text: str, max_chars: int = 220) -> None:
    """将报告章节压缩为模板顶部可容纳的单页摘要。

    月报完整论证保存在 Markdown 正文中；PPT 顶部只承载本页结论和关键证据。模板
    的正文框高度固定，若直接填入整段分析会遮挡图表标题，因此按页面类型限制字符数。
    """
    for shape in slide.shapes:
        if not getattr(shape, "is_placeholder", False):
            continue
        placeholder_type = shape.placeholder_format.type
        if placeholder_type in {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT} and getattr(shape, "has_text_frame", False):
            set_text_frame_text(shape.text_frame, truncate_for_slide(text, max_chars))
            return


def replace_period_labels(slide: object, data: MonthlyPptData) -> None:
    """替换母版中可见的月份标签，避免残留“2026年5月”等样例日期。"""

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        original = shape.text or ""
        updated = original
        updated = re.sub(r"20\d{2}年\d{1,2}月", data.period_label, updated)
        updated = re.sub(r"20\d{2}\.1-\d{1,2}", f"{data.year}.1-{data.month}", updated)
        updated = re.sub(r"20\d{2}\.(?!1-)\d{1,2}", data.period_short, updated)
        if updated != original:
            set_text_frame_text(shape.text_frame, updated)


def add_share_period_label_overlays(slide: object, data: MonthlyPptData) -> None:
    """覆盖母版右图的旧期间文字，保留底层原生图表的可编辑性。"""

    chart_shapes = [shape for shape in slide.shapes if getattr(shape, "has_chart", False)]
    if len(chart_shapes) < 3:
        return

    chart_shape = chart_shapes[2]
    # python-pptx 已更新图表数据源，但该业务母版在 PowerPoint 渲染时仍可能显示旧的
    # 分类轴缓存。隐藏底层轴文字，改由当前数据驱动的文本框显示，避免出现两组日期。
    chart_shape.chart.category_axis.tick_labels.font.color.rgb = RGBColor(255, 255, 255)
    labels = [
        data.previous_period_short,
        data.period_short,
        f"{int(data.year) - 1}.1-12",
        f"{data.year}.1-{data.month}",
    ]
    label_width = int(chart_shape.width * 0.19)
    label_height = int(Pt(18))
    # 轴标签位于绘图区的最下沿。与母版原标签同位覆盖，避免出现旧日期和新日期两组文字。
    label_top = chart_shape.top + int(chart_shape.height * 0.938)
    # 右图绘图区左侧为纵轴与留白，四个类别轴标签的中心点按模板比例分布。
    for index, label in enumerate(labels):
        label_center = 0.14 + index * 0.24
        label_left = chart_shape.left + int(chart_shape.width * label_center) - label_width // 2
        overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, label_left, label_top, label_width, label_height)
        overlay.name = f"当前期间标签_{index + 1}"
        overlay.fill.solid()
        overlay.fill.fore_color.rgb = RGBColor(255, 255, 255)
        overlay.line.fill.background()
        text_frame = overlay.text_frame
        text_frame.clear()
        text_frame.margin_left = 0
        text_frame.margin_right = 0
        text_frame.margin_top = 0
        text_frame.margin_bottom = 0
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = label
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0, 0, 0)


def update_data_source_footer(slide: object) -> None:
    """统一样例中的数据来源口径，保留业务既有页脚位置和字体。"""

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if "数据来源：" in (shape.text or ""):
            suffix = "，右侧份额为在对应细分市场份额；蓝色填充为传统车型，绿色填充为新能源车型"
            set_text_frame_text(
                shape.text_frame,
                "数据来源：中汽数据有限公司终端零售数据" + (suffix if "右侧份额" in shape.text else ""),
            )


def update_oval_rates(slide: object, rates: Sequence[float | None]) -> None:
    """更新母版中圆形同比标记，并统一为业务要求的强调文字样式。"""

    oval_shapes = [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and "椭圆" in (shape.name or "")
    ]
    for shape, rate in zip(oval_shapes, rates):
        set_text_frame_text(shape.text_frame, format_rate(rate))
        first_paragraph = shape.text_frame.paragraphs[0]
        if first_paragraph.runs:
            run = first_paragraph.runs[0]
            # 问题5：五页图表中的椭圆同比标签统一使用微软雅黑、10号、加粗，
            # 避免在替换数值后继承到母版示例中的 Arial 常规字。
            run.font.name = "微软雅黑"
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 112)


def clear_stale_percentage_callouts(slide: object) -> None:
    """清除模板中未与当前数据绑定的百分比提示，避免把样例数值带入新月报。"""

    percent_pattern = re.compile(r"^[+-]?\d+(?:\.\d+)?%$")
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False) or "椭圆" in (shape.name or ""):
            continue
        if getattr(shape, "has_table", False):
            continue
        if percent_pattern.fullmatch((shape.text or "").strip()):
            set_text_frame_text(shape.text_frame, "")


def set_text_frame_text(text_frame: object, value: str) -> None:
    """替换文本但保持模板首个 Run 的字号、字体、颜色和段落格式。"""

    paragraphs = text_frame.paragraphs
    if not paragraphs:
        return
    first = paragraphs[0]
    if first.runs:
        first.runs[0].text = value
        for run in first.runs[1:]:
            run.text = ""
    else:
        first.text = value
    for paragraph in paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


def charts_in(slide: object) -> list[object]:
    return [shape.chart for shape in slide.shapes if getattr(shape, "has_chart", False)]


def tables_in(slide: object) -> list[object]:
    return [shape.table for shape in slide.shapes if getattr(shape, "has_table", False)]


def find_section_column(rows: list[list[str]], title: str) -> int | None:
    for index, value in enumerate(rows[0] if rows else []):
        if normalize(value) == title:
            return index
    return None


def find_next_section_column(rows: list[list[str]], start: int) -> int:
    titles = {"车型类别", "车型级别", "轿车-级别", "SUV-级别", "MPV-级别"}
    header = rows[0] if rows else []
    for index in range(start + 1, len(header)):
        if normalize(header[index]) in titles:
            return index
    return len(header)


def clean_labels(values: Iterable[str]) -> list[str]:
    labels: list[str] = []
    for value in values:
        label = normalize(value)
        if not label or label.startswith("左图"):
            continue
        labels.append(label)
    return labels


def year_pairs(rows: list[list[str]], start: int) -> list[tuple[int, int]]:
    pairs: list[tuple[int, int]] = []
    for row_index in range(len(rows) - 1):
        previous_year = cell(rows, row_index, start)
        current_year = cell(rows, row_index + 1, start)
        if (
            re.fullmatch(r"20\d{2}", previous_year)
            and re.fullmatch(r"20\d{2}", current_year)
            and int(current_year) == int(previous_year) + 1
        ):
            pairs.append((row_index, row_index + 1))
    return pairs


def row_values(rows: list[list[str]], row_index: int, start: int, labels: Sequence[str]) -> list[float | None]:
    return [normalize_sales_volume(to_float(cell(rows, row_index, start + offset + 1))) for offset in range(len(labels))]


def parse_top_models(rows: list[list[str]], start: int, end: int) -> list[TopModel]:
    """解析样例验数表中排列在分级别数据右侧的 TOP15 车型。"""

    header_row = None
    for row_index, row in enumerate(rows):
        if cell(rows, row_index, start + 1) == "车型" and cell(rows, row_index, start + 2) == "销量":
            header_row = row_index
            break
    if header_row is None:
        return []

    models: list[TopModel] = []
    for row_index in range(header_row + 1, len(rows)):
        rank = normalize(cell(rows, row_index, start))
        name = normalize(cell(rows, row_index, start + 1))
        if not rank.isdigit() or not name:
            if models:
                break
            continue
        models.append(
            TopModel(
                name=name,
                volume=normalize_sales_volume(to_float(cell(rows, row_index, start + 2))),
                yoy=to_float(cell(rows, row_index, start + 3)),
                mom=to_float(cell(rows, row_index, start + 4)),
                level=normalize(cell(rows, row_index, start + 5)),
                share=to_float(cell(rows, row_index, start + 6)),
                rank_change=format_rank_change(cell(rows, row_index, start + 7)),
                cumulative_volume=normalize_sales_volume(to_float(cell(rows, row_index, start + 8))),
            ),
        )
    return models[:15]


def parse_right_chart_share_periods(
    rows: list[list[str]],
    start: int,
    labels: Sequence[str],
) -> list[list[float | None]] | None:
    """读取验数表“右图”区域的四期标准份额，避免由左图销量重新推算。"""

    header_row = next(
        (index for index, row in enumerate(rows) if normalize(cell(rows, index, start)) == "右图"),
        None,
    )
    if header_row is None:
        return None

    result: list[list[float | None]] = [[None] * len(labels) for _ in range(4)]
    matched = 0
    for row_index in range(header_row + 1, len(rows)):
        label = normalize(cell(rows, row_index, start))
        if label not in labels:
            if matched:
                break
            continue
        label_index = labels.index(label)
        values = [to_float(cell(rows, row_index, start + offset)) for offset in range(1, 5)]
        if any(value is None for value in values):
            return None
        for period_index, value in enumerate(values):
            result[period_index][label_index] = value
        matched += 1

    return result if matched == len(labels) else None


def extract_report_section(markdown: str, marker: str) -> str:
    """提取当前文字月报对应章节，作为 PPT 页面的因果分析正文。"""

    pattern = re.compile(
        rf"(?ms)^###?\s*{marker}[^\n]*\n(?P<body>.*?)(?=^###?\s+|\Z)",
        re.IGNORECASE,
    )
    match = pattern.search(markdown or "")
    if not match:
        return ""
    body = match.group("body")
    # 月报正文的第一个“◼”段被业务 Skill 定义为可直接进入当前页顶部的市场深度分析。
    # 其后的数据表与补充说明保留在文字报告，不再被固定高度的 PPT 正文框截断。
    summary_match = re.search(r"(?ms)^\s*◼\s*(?P<summary>.*?)(?=\n\s*\n|\Z)", body)
    if summary_match:
        body = summary_match.group("summary")
    body = re.sub(r"^\s*[-*◼]\s*", "", body, flags=re.MULTILINE)
    body = re.sub(r"\|[^\n]+\|", "", body)
    # 审计草稿中的〔E/R〕编号只供服务端门禁使用；PPT 只呈现业务结论和关键事实。
    body = re.sub(r"〔[^〕]*[ER]\d+[^〕]*〕", "", body)
    # 车型明细用于后台筛选，不应作为咨询交付语言出现在正文中。
    body = re.sub(r"(?:TOP\s*15\s*样本(?:内|中)?|样本(?:内|中))", "", body, flags=re.IGNORECASE)
    body = re.sub(r"\n{2,}", "\n", body)
    return re.sub(r"\s+", " ", body).strip()


def extract_report_heading(markdown: str, marker: str) -> str:
    """取模型给出的结论式章节标题，供 PPT 保留专家的“原因 + 结果”表达。"""

    pattern = re.compile(rf"(?mi)^###\s*(?P<heading>{marker}[^\n]*)$")
    match = pattern.search(markdown or "")
    if not match:
        return ""
    heading = re.sub(r"〔[^〕]*[ER]\d+[^〕]*〕", "", match.group("heading"))
    return re.sub(r"\s+", " ", heading).strip()


def build_headline(prefix: str, block: MarketBlock, report_heading: str = "") -> str:
    """按业务样例的“结论式标题”风格生成 1 句可落在 PPT 顶部的标题。"""

    # 业务专家写出的“原因 + 结果”标题优先于数据规则的兜底标题，既保留月报的
    # 叙事逻辑，也不把模板固定为“最好/最差类别”的机械句型。
    if report_heading and "：" in report_heading:
        return report_heading[:74].rstrip("，；; ")

    valid = [(label, rate) for label, rate in zip(block.labels, block.yoy) if rate is not None]
    if not valid:
        return f"{prefix}：当前数据待补充，暂不形成结构性判断"
    best_label, best_rate = max(valid, key=lambda item: item[1])
    worst_label, worst_rate = min(valid, key=lambda item: item[1])
    best_phrase = "同比正增长" if best_rate >= 0 else "同比降幅相对较小"
    worst_phrase = "同比承压" if worst_rate < 0 else "增速相对偏弱"
    return f"{prefix}：{best_label}市场{best_phrase}，{worst_label}市场{worst_phrase}"


def build_data_narrative(block: MarketBlock, period_label: str) -> str:
    """当模型章节缺失时，根据验数表输出保守、可验证的兜底分析段。"""

    valid = [(label, rate) for label, rate in zip(block.labels, block.yoy) if rate is not None]
    if not valid:
        return f"{period_label}当前数据仅覆盖部分结构指标，建议补充车型级明细后再形成原因判断。"
    best_label, best_rate = max(valid, key=lambda item: item[1])
    worst_label, worst_rate = min(valid, key=lambda item: item[1])
    return (
        f"{period_label}，{best_label}同比{format_rate(best_rate)}，在本层级中表现相对较好；"
        f"{worst_label}同比{format_rate(worst_rate)}，是当前主要拖累项。"
        "当前验数表可支撑结构变化判断；具体产品、政策和品牌驱动需结合车型明细与公开证据进一步核验。"
    )


def shares(values: Sequence[float | None]) -> list[float]:
    total = sum(value or 0 for value in values)
    if not total:
        return [0.0 for _ in values]
    return [(value or 0) / total for value in values]


def percentage_points(values: Sequence[float | None]) -> list[float]:
    """将已计算的 0-1 份额转换为母版堆积图使用的百分比点。"""

    return [round((value or 0) * 100, 1) for value in values]


def values_or_zero(values: Sequence[float | None]) -> list[float]:
    return [value or 0.0 for value in values]


def format_rate(value: float | None) -> str:
    """按业务月报规范显示比例：正数不加 ``+``，负数保留 ``-``。"""

    if value is None:
        return "-"
    if abs(value) >= 10:
        return ">1000%" if value > 10 else "<-1000%"
    return f"{value * 100:.1f}%" if value else "0.0%"


def format_volume(value: float | None) -> str:
    return "-" if value is None else f"{value:.1f}"


def format_rank_change(value: str) -> str:
    text = normalize(value)
    if not text or text == "0":
        return "-"
    if text.startswith("+"):
        return text[1:]
    return text


def is_traditional_model(name: str) -> bool:
    return any(token in name for token in TRADITIONAL_MODEL_TOKENS)


def infer_year(sheet: SheetSnapshot) -> str:
    for row_index in range(len(sheet.rows) - 1):
        previous_year = cell(sheet.rows, row_index, 0)
        current_year = cell(sheet.rows, row_index + 1, 0)
        if (
            re.fullmatch(r"20\d{2}", previous_year)
            and re.fullmatch(r"20\d{2}", current_year)
            and int(current_year) == int(previous_year) + 1
        ):
            return current_year
    return ""


def sheet_month_number(sheet: SheetSnapshot | None) -> int:
    if sheet is None:
        return 0
    match = re.search(r"(\d{1,2})月", sheet.name)
    return int(match.group(1)) if match else 0


def cell(rows: list[list[str]], row_index: int, column_index: int) -> str:
    if row_index < 0 or column_index < 0 or row_index >= len(rows):
        return ""
    row = rows[row_index]
    return row[column_index] if column_index < len(row) else ""


def to_float(value: str) -> float | None:
    text = normalize(value)
    if not text or text.startswith("#"):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def normalize_sales_volume(value: float | None) -> float | None:
    """统一到业务样例图表使用的“万辆”口径；级别和 TOP15 已是万辆则保持原值。"""

    if value is None:
        return None
    return value / 10000 if abs(value) >= 1000 else value


def normalize(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def truncate(text: str, limit: int) -> str:
    normalized = normalize(text)
    return normalized if len(normalized) <= limit else normalized[: limit - 1].rstrip("，。；；") + "。"


def truncate_for_slide(text: str, limit: int) -> str:
    """优先在完整句/分句结束处截断 PPT 顶部摘要，避免出现半个级别名称。"""

    normalized = normalize(text)
    if len(normalized) <= limit:
        return normalized

    candidate = normalized[:limit]
    sentence_boundary = max(candidate.rfind(separator) for separator in ("。", "；", "！", "？"))
    if sentence_boundary >= 0:
        return candidate[: sentence_boundary + 1]

    # 当输入本身没有完整句号时，至少在逗号处收束为一句，避免在数字、车型级别或
    # 英文名称中间截断。找不到可靠边界时才回退到通用截断。
    comma_position = max(candidate.rfind("，"), candidate.rfind("、"))
    if comma_position >= max(36, int(limit * 0.45)):
        return candidate[:comma_position].rstrip("，、") + "。"
    return truncate(normalized, limit)
