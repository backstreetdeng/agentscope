"""月报 2.0 的 23 页离线回归测试。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from utils.monthly_report_pptx import build_monthly_report_pptx
from utils.monthly_report_pptx_v2 import is_v2_dataset, load_full_monthly_data


PROJECT_DIR = Path(__file__).resolve().parents[2]
PROFILE_PATH = PROJECT_DIR / "profiles" / "passenger_vehicle_monthly"
DATASET_PATH = PROFILE_PATH / "assets" / "demo_monthly_data.xlsx"
TEMPLATE_PATH = PROFILE_PATH / "assets" / "monthly_report_template.pptx"


@unittest.skipUnless(DATASET_PATH.is_file() and TEMPLATE_PATH.is_file(), "缺少2.0验数表或PPTX母版")
class MonthlyReportPptxV2Test(unittest.TestCase):
    def test_detects_complete_page_block_dataset(self) -> None:
        self.assertTrue(is_v2_dataset(DATASET_PATH, "6月"))
        data = load_full_monthly_data(DATASET_PATH, "6月")
        self.assertEqual((data.year, data.month), (2026, 6))
        self.assertEqual(sorted(data.blocks), [3, 4, 5, *range(7, 22)])

    def test_builds_exact_23_slide_editable_report(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "month_skill_v2.pptx"
            result = build_monthly_report_pptx(
                "",
                output,
                DATASET_PATH,
                TEMPLATE_PATH,
                sheet_name="6月",
            )
            self.assertEqual(result["slides_count"], 23)
            self.assertEqual(result["charts_count"], 47)
            self.assertEqual(result["tables_count"], 37)
            self.assertEqual(result["period"], "2026年6月")
            self.assertIn("P4 未提供右侧典型新车子图明细", " ".join(result["warnings"]))

            presentation = Presentation(output)
            self.assertEqual(len(presentation.slides), 23)
            chart_count = sum(
                1 for slide in presentation.slides for shape in slide.shapes if shape.has_chart
            )
            table_count = sum(
                1 for slide in presentation.slides for shape in slide.shapes if shape.has_table
            )
            self.assertEqual((chart_count, table_count), (47, 37))

            p21_tables = [
                shape.table for shape in presentation.slides[20].shapes if shape.has_table
            ]
            self.assertEqual(p21_tables[1].cell(0, 1).text, "21.9%")
            self.assertEqual(p21_tables[2].cell(0, 1).text, "21.1%")

            all_text = " ".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertNotIn("#DIV/0!", all_text)
            self.assertIn("2026年6月", all_text)

    def test_p8_share_chart_periods_each_sum_to_100(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "share_check.pptx"
            build_monthly_report_pptx("", output, DATASET_PATH, TEMPLATE_PATH, "6月")
            presentation = Presentation(output)
            charts = [
                shape.chart for shape in presentation.slides[7].shapes if shape.has_chart
            ]
            share_chart = charts[2]
            period_sums = [
                sum(series.values[period_index] for series in share_chart.series)
                for period_index in range(4)
            ]
            for value in period_sums:
                self.assertAlmostEqual(value, 100.0, delta=0.2)


if __name__ == "__main__":
    unittest.main()
