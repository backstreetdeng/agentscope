"""月报可编辑 PPTX 的离线回归测试，不调用 MiniMax API。"""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from utils.monthly_report_pptx import (
    build_monthly_report_pptx,
    format_rank_change,
    extract_report_heading,
    extract_report_section,
    format_rate,
    truncate,
    truncate_for_slide,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
PROFILE_PATH = PROJECT_DIR / "skills" / "market_analysis_skill" / "profiles" / "passenger_vehicle_monthly"
DATASET_PATH = PROFILE_PATH / "assets" / "demo_monthly_data.xlsx"
TEMPLATE_PATH = PROFILE_PATH / "assets" / "monthly_report_template.pptx"

SAMPLE_REPORT = """# 2026年5月乘用车市场月报

### 2.1 车型类别
轿车和SUV仍是主体，SUV降幅相对较小，建议继续跟踪结构份额变化。

### 2.2 车型级别
B级市场体量最大，A00级承压明显，需观察新品补位节奏。

### 2.3 轿车级别分析
轿车B级和A级承压，TOP车型分化加快。

### 2.4 SUV级别分析
SUV A0与C+级表现韧性较强，新能源产品贡献增量。

### 2.5 MPV级别分析
MPV市场规模较小，A级同比正增长但需持续验证。
"""


@unittest.skipUnless(DATASET_PATH.is_file() and TEMPLATE_PATH.is_file(), "缺少业务验数表或 PPTX 模板")
class MonthlyReportPptxTest(unittest.TestCase):
    def test_six_month_sales_charts_clear_template_axis_ceiling(self) -> None:
        """问题6：销量超过5月母版旧上限时，柱高必须按真实数值自动缩放。"""

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "六月纵轴校验.pptx"
            build_monthly_report_pptx(
                SAMPLE_REPORT,
                output_path,
                DATASET_PATH,
                sheet_name="6月",
            )
            presentation = Presentation(output_path)
            charts = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart]

            monthly_chart, cumulative_chart, share_chart = charts[:3]
            self.assertEqual(
                [round(value, 4) for value in monthly_chart.series[0].values],
                [97.3963, 106.6977, 10.8309, 1.5728],
            )
            self.assertEqual(
                [round(value, 4) for value in monthly_chart.series[1].values],
                [68.3574, 93.2123, 7.7313, 0.9169],
            )
            self.assertIsNone(monthly_chart.value_axis.maximum_scale)
            self.assertIsNone(cumulative_chart.value_axis.maximum_scale)
            self.assertEqual(share_chart.value_axis.maximum_scale, 100.0)

    def test_rank_change_positive_sign_is_removed(self) -> None:
        self.assertEqual(format_rank_change("+4"), "4")
        self.assertEqual(format_rank_change("4"), "4")
        self.assertEqual(format_rank_change("-2"), "-2")

    def test_positive_rates_do_not_show_plus_sign_in_ppt(self) -> None:
        """问题4：同比、环比、份额的正值必须不显示加号。"""

        self.assertEqual(format_rate(0.049), "4.9%")
        self.assertEqual(format_rate(0.087), "8.7%")
        self.assertEqual(format_rate(0.0), "0.0%")
        self.assertEqual(format_rate(-0.113), "-11.3%")

    def test_all_ppt_tables_and_rate_callouts_hide_positive_plus_sign(self) -> None:
        """问题4：业务母版生成后的五页表格和同比圆形标记都遵循同一格式。"""

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "正数格式校验.pptx"
            build_monthly_report_pptx(SAMPLE_REPORT, output_path, DATASET_PATH)
            presentation = Presentation(output_path)

            rate_values: list[str] = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            rate_values.extend(cell.text.strip() for cell in row.cells)
                    elif "椭圆" in (getattr(shape, "name", "") or "") and getattr(shape, "has_text_frame", False):
                        rate_values.append(shape.text.strip())

            self.assertTrue(any(value == "4.9%" for value in rate_values))
            self.assertFalse(any(value.startswith("+") and value.endswith("%") for value in rate_values))

    def test_all_oval_rate_callouts_use_business_font_style(self) -> None:
        """问题5：五页图表的椭圆同比标签必须统一为微软雅黑、10号、加粗。"""

        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "椭圆标签样式校验.pptx"
            build_monthly_report_pptx(SAMPLE_REPORT, output_path, DATASET_PATH)
            presentation = Presentation(output_path)

            oval_runs = []
            for slide in presentation.slides:
                for shape in slide.shapes:
                    if "椭圆" not in (getattr(shape, "name", "") or ""):
                        continue
                    if not getattr(shape, "has_text_frame", False):
                        continue
                    paragraph = shape.text_frame.paragraphs[0]
                    if paragraph.runs:
                        oval_runs.append(paragraph.runs[0])

            self.assertGreaterEqual(len(oval_runs), 1)
            for run in oval_runs:
                self.assertEqual(run.font.name, "微软雅黑")
                self.assertEqual(run.font.size.pt, 10)
                self.assertTrue(run.font.bold)

    def test_ppt_section_strips_external_evidence_markers_but_keeps_conclusion(self) -> None:
        markdown = "### 2.1 车型类别\n5月传统车表现承压〔E1〕，SUV跌幅相对较小〔R3〕。\n"
        section = extract_report_section(markdown, r"2\.1\s*车型类别")

        self.assertEqual(section, "5月传统车表现承压，SUV跌幅相对较小。")

    def test_ppt_uses_first_business_summary_not_later_data_appendix(self) -> None:
        summary = "◼ 5月，SUV表现相对稳定；星愿与海鸥变化带来结构分化；行业高基数影响仍待结合证据核验；建议跟踪新品承接。〔E1〕"
        markdown = f"### 2.1 车型类别\n{summary}\n\n当月销量、同比、环比等明细数据另附。"

        section = extract_report_section(markdown, r"2\.1\s*车型类别")

        self.assertEqual(section, summary.replace("◼ ", "").replace("〔E1〕", ""))

    def test_ppt_uses_conclusion_style_section_heading(self) -> None:
        markdown = "### 2.4 SUV级别分析：新品承接与传统车承压并存，A级市场同比走弱\n正文。"

        heading = extract_report_heading(markdown, r"2\.4\s*SUV")

        self.assertEqual(heading, "2.4 SUV级别分析：新品承接与传统车承压并存，A级市场同比走弱")

    def test_ppt_hides_internal_audit_and_sample_language(self) -> None:
        report = SAMPLE_REPORT.replace(
            "### 2.1 车型类别\n轿车和SUV仍是主体，SUV降幅相对较小，建议继续跟踪结构份额变化。",
            "### 2.1 车型类别：同期高基数下SUV相对抗跌，MPV市场承压\n"
            "◼ 5月，SUV市场同比降幅相对较小；星愿与海鸥的变化体现入门级产品分化；"
            "高油价影响传统车表现〔R3〕；整体看，应跟踪新品承接。TOP15样本中仅作内部筛选。",
        )
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "交付表达校验.pptx"
            build_monthly_report_pptx(report, output_path, DATASET_PATH)
            presentation = Presentation(output_path)
            slide_text = " ".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False)
            )

        self.assertIn("同期高基数下SUV相对抗跌", slide_text)
        self.assertNotIn("〔R3〕", slide_text)
        self.assertNotIn("TOP15样本", slide_text)

    def test_builds_five_editable_slides_from_business_template(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "月报测试.pptx"
            result = build_monthly_report_pptx(SAMPLE_REPORT, output_path, DATASET_PATH, sheet_name="5月")

            self.assertTrue(result["success"])
            self.assertEqual(result["slides_count"], 5)
            self.assertEqual(result["period"], "2026年5月")
            self.assertTrue(output_path.is_file())

            presentation = Presentation(output_path)
            self.assertEqual(len(presentation.slides), 5)
            self.assertEqual(
                len([shape for shape in presentation.slides[0].shapes if shape.has_chart]),
                3,
            )
            self.assertEqual(
                len([shape for shape in presentation.slides[3].shapes if shape.has_table]),
                4,
            )

            first_chart = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart][0]
            first_series = list(first_chart.series[0].values)
            self.assertEqual([round(value, 4) for value in first_series], [85.0076, 92.7674, 9.6054, 1.4636])

            share_chart = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart][2]
            self.assertEqual(
                [[round(value, 1) for value in series.values] for series in share_chart.series],
                [
                    [0.8, 0.7, 0.7, 0.6],
                    [5.0, 4.4, 4.8, 4.9],
                    [52.6, 54.4, 48.9, 53.7],
                    [41.6, 40.5, 45.6, 40.8],
                ],
            )

            self.assertEqual(
                [category.label for category in share_chart.plots[0].categories],
                ["2026.4", "2026.5", "2025.1-12", "2026.1-5"],
            )

            slide_text = " ".join(
                shape.text
                for shape in presentation.slides[3].shapes
                if getattr(shape, "has_text_frame", False)
            )
            self.assertIn("2.4 SUV级别", slide_text)
            self.assertIn("SUV A0与C+级表现韧性较强", slide_text)

    def test_can_render_specified_month_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output_path = Path(directory) / "四月月报测试.pptx"
            result = build_monthly_report_pptx(
                SAMPLE_REPORT,
                output_path,
                DATASET_PATH,
                sheet_name="4月",
            )

            self.assertTrue(result["success"])
            self.assertEqual(result["period"], "2026年4月")
            self.assertEqual(result["sheet_name"], "4月")
            presentation = Presentation(output_path)
            share_chart = [shape.chart for shape in presentation.slides[0].shapes if shape.has_chart][2]
            self.assertEqual(
                [category.label for category in share_chart.plots[0].categories],
                ["2026.3", "2026.4", "2025.1-12", "2026.1-4"],
            )
            overlay_labels = [
                shape.text
                for shape in presentation.slides[0].shapes
                if getattr(shape, "name", "").startswith("当前期间标签_")
            ]
            self.assertEqual(overlay_labels, ["2026.3", "2026.4", "2025.1-12", "2026.1-4"])

    def test_slide_summary_is_limited_to_template_capacity(self) -> None:
        long_text = "市场结构分析" * 80

        self.assertLessEqual(len(truncate(long_text, 220)), 220)
        self.assertLessEqual(len(truncate(long_text, 170)), 170)
        self.assertLessEqual(len(truncate(long_text, 145)), 145)

    def test_slide_summary_ends_at_a_complete_sentence(self) -> None:
        text = "第一句已完整。第二句很长但不能被截成半句话，仍然继续补充更多文字。"

        self.assertEqual(truncate_for_slide(text, 18), "第一句已完整。")
