#!/usr/bin/env python
"""Build a deterministic vehicle reputation POC from DCD + Autohome Excel exports."""
from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import psycopg2
from dotenv import load_dotenv
from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parents[1]
load_dotenv(PROJECT_DIR / ".env.new-car", override=False)

DCD_XLSX = Path(r"D:\2024年度工作日志和备忘录\数字化转型产品\4.0 同事组\5.4 晓亚\乘用车新车监测板块-0921\懂车帝口碑明细导出.xlsx")
AH_XLSX = Path(r"D:\2024年度工作日志和备忘录\数字化转型产品\4.0 同事组\5.4 晓亚\乘用车新车监测板块-0921\汽车之家口碑明细_20260921.xlsx")

PG = {
    "host": os.getenv("NEW_CAR_PG_HOST", "192.168.3.146"),
    "port": int(os.getenv("NEW_CAR_PG_PORT", "5432")),
    "dbname": os.getenv("NEW_CAR_PG_DATABASE", "vectordb"),
    "user": os.getenv("NEW_CAR_PG_USER", "vectordb"),
    "password": os.getenv("NEW_CAR_PG_PASSWORD", ""),
}

PRIMARY_KEYWORDS = {
    "外观": ["外观", "颜值", "造型", "设计", "线条", "车灯", "大灯", "轮毂", "回头率"],
    "内饰": ["内饰", "做工", "用料", "材质", "异味", "塑料感", "氛围灯"],
    "空间": ["空间", "后排", "前排", "头部", "腿部", "储物", "后备", "座舱"],
    "动力": ["动力", "加速", "提速", "起步", "超车", "发动机", "变速箱"],
    "操控": ["操控", "转向", "方向盘", "底盘", "刹车", "悬架", "车身稳定"],
    "舒适性": ["舒适", "座椅", "隔音", "噪音", "胎噪", "风噪", "减震", "颠簸", "空调"],
    "油耗能耗": ["油耗", "能耗", "电耗", "续航", "充电", "省油", "费油"],
    "配置智能": ["配置", "车机", "导航", "语音", "智驾", "辅助驾驶", "智能座舱", "音响", "功能"],
    "性价比": ["性价比", "价格", "优惠", "保值", "划算", "成本", "贵", "便宜"],
}

SECONDARY = {
    "外观": [("造型设计", ["造型", "设计", "颜值", "好看", "帅"]), ("车身线条", ["线条", "侧面", "溜背"]), ("灯组设计", ["车灯", "大灯", "尾灯"]), ("轮毂造型", ["轮毂"])],
    "内饰": [("用料质感", ["用料", "材质", "软包", "质感"]), ("做工细节", ["做工", "缝线", "接缝"]), ("内饰设计", ["内饰", "氛围灯"]), ("车内异味", ["异味", "味道"])],
    "空间": [("乘坐空间", ["乘坐", "前排", "腿部", "宽敞"]), ("后排空间", ["后排"]), ("头部空间", ["头部"]), ("储物能力", ["储物"]), ("后备厢空间", ["后备"]), ("空间实用性", ["空间"])],
    "动力": [("加速响应", ["加速", "提速", "起步", "响应"]), ("高速超车", ["超车", "高速"]), ("动力储备", ["动力", "有劲", "强劲"]), ("动力平顺", ["平顺", "变速箱"]), ("动力不足", ["肉", "不足", "乏力"])],
    "操控": [("转向手感", ["转向", "方向盘"]), ("底盘质感", ["底盘"]), ("制动表现", ["刹车", "制动"]), ("悬架支撑", ["悬架", "支撑"]), ("车身稳定", ["稳定", "侧倾"])],
    "舒适性": [("座椅舒适", ["座椅"]), ("隔音表现", ["隔音", "噪音", "胎噪", "风噪"]), ("滤振表现", ["减震", "颠簸", "滤振"]), ("空调体验", ["空调"]), ("乘坐舒适", ["舒适"])],
    "油耗能耗": [("油耗表现", ["油耗", "省油", "费油"]), ("能耗表现", ["能耗", "电耗"]), ("续航表现", ["续航"]), ("补能便利", ["充电", "补能"])],
    "配置智能": [("车机体验", ["车机", "导航", "语音"]), ("辅助驾驶", ["智驾", "辅助驾驶"]), ("智能座舱", ["智能座舱"]), ("音响体验", ["音响"]), ("配置丰富度", ["配置", "功能"])],
    "性价比": [("价格竞争力", ["价格", "贵", "便宜"]), ("优惠力度", ["优惠"]), ("配置价值", ["性价比", "划算"]), ("用车成本", ["成本", "保值"])],
}

DCD_SCORE_MAP = {
    "外观评分": "外观", "内饰评分": "内饰", "配置评分": "配置智能",
    "空间评分": "空间", "舒适性评分": "舒适性", "油耗评分": "油耗能耗",
    "操控评分": "操控", "动力评分": "动力", "辅助驾驶评分": "配置智能",
    "智能座舱评分": "配置智能", "动力操控评分": "操控", "续航评分": "油耗能耗",
    "驾驶感受评分": "舒适性",
}

def classify_text(text: str) -> set[str]:
    return {label for label, words in PRIMARY_KEYWORDS.items() if any(word in text for word in words)}

def secondary_labels(label: str, texts: list[str], direction: str) -> list[dict]:
    counts = Counter()
    for text in texts:
        for name, words in SECONDARY[label]:
            if any(word in text for word in words):
                counts[name] += 1
    if not counts:
        fallback = "表现获得认可" if direction == "positive" else "体验仍需优化"
        return [{"name": fallback, "count": len(texts)}]
    return [{"name": name, "count": count} for name, count in counts.most_common(3)]

def add_mention(store, model, direction, label, text, source):
    item = store[model][direction][label]
    item["count"] += 1
    item["texts"].append(text)
    item["sources"][source] += 1

def read_autohome(store, allowed):
    raw_counts = Counter()
    wb = load_workbook(AH_XLSX, read_only=True, data_only=True)
    ws = wb["口碑明细"]
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows)); pos = {name: headers.index(name) for name in ["车型名称", "满意", "不满意"]}
    for row in rows:
        model = str(row[pos["车型名称"]] or "").strip()
        if model not in allowed: continue
        raw_counts[model] += 1
        for direction, col in (("positive", "满意"), ("negative", "不满意")):
            text = str(row[pos[col]] or "").strip()
            if not text: continue
            for label in classify_text(text):
                add_mention(store, model, direction, label, text, "autohome")
    wb.close()
    return raw_counts

def read_dongchedi(store, allowed):
    raw_counts = Counter()
    wb = load_workbook(DCD_XLSX, read_only=True, data_only=True)
    ws = wb["口碑明细"]
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows)); model_idx = headers.index("车系名称"); text_idx = headers.index("正文")
    score_cols = [(headers.index(col), label) for col, label in DCD_SCORE_MAP.items() if col in headers]
    for row in rows:
        model = str(row[model_idx] or "").strip()
        if model not in allowed: continue
        raw_counts[model] += 1
        text = str(row[text_idx] or "").strip()
        for idx, label in score_cols:
            try: score = float(row[idx])
            except (TypeError, ValueError): continue
            if score >= 4.5:
                add_mention(store, model, "positive", label, text, "dongchedi")
            elif score <= 3.0:
                add_mention(store, model, "negative", label, text, "dongchedi")
    wb.close()
    return raw_counts

def main():
    conn = psycopg2.connect(**PG)
    cur = conn.cursor()
    cur.execute("SELECT DISTINCT model FROM new_car_model WHERE model IS NOT NULL")
    allowed = {r[0] for r in cur.fetchall()}
    store = defaultdict(lambda: defaultdict(lambda: defaultdict(lambda: {"count": 0, "texts": [], "sources": Counter()})))
    ah_counts = read_autohome(store, allowed)
    dcd_counts = read_dongchedi(store, allowed)
    batch = datetime.now().strftime("%Y%m%d")
    cur.execute("""
        CREATE TABLE IF NOT EXISTS new_car_reputation_summary (
            model TEXT NOT NULL, direction TEXT NOT NULL, rank SMALLINT NOT NULL,
            primary_label TEXT NOT NULL, secondary_labels JSONB NOT NULL,
            mention_count INT NOT NULL, mention_share NUMERIC(6,4) NOT NULL,
            dongchedi_count INT NOT NULL, autohome_count INT NOT NULL,
            source_mentions JSONB NOT NULL, import_batch TEXT NOT NULL,
            PRIMARY KEY (model, direction, rank)
        )
    """)
    cur.execute("DELETE FROM new_car_reputation_summary")
    inserted = 0
    for model, directions in store.items():
        for direction in ("positive", "negative"):
            labels = directions[direction]
            total_mentions = sum(item["count"] for item in labels.values())
            if not total_mentions: continue
            ranked = sorted(labels.items(), key=lambda pair: (-pair[1]["count"], pair[0]))[:3]
            for rank, (label, item) in enumerate(ranked, 1):
                secondaries = secondary_labels(label, item["texts"], direction)
                cur.execute("""
                    INSERT INTO new_car_reputation_summary
                    (model,direction,rank,primary_label,secondary_labels,mention_count,mention_share,
                     dongchedi_count,autohome_count,source_mentions,import_batch)
                    VALUES (%s,%s,%s,%s,%s::jsonb,%s,%s,%s,%s,%s::jsonb,%s)
                """, (model, direction, rank, label, json.dumps(secondaries, ensure_ascii=False),
                      item["count"], item["count"] / total_mentions, dcd_counts[model], ah_counts[model],
                      json.dumps(dict(item["sources"]), ensure_ascii=False), batch))
                inserted += 1
    conn.commit()
    cur.execute("SELECT count(DISTINCT model), count(*) FROM new_car_reputation_summary")
    models, rows = cur.fetchone()
    cur.close(); conn.close()
    print(f"models={models} rows={rows} inserted={inserted} batch={batch}")

if __name__ == "__main__":
    main()
