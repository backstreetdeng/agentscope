#!/usr/bin/env python
"""Import monthly sales and build taxonomy-aligned reputation summaries for eight vehicles."""
from __future__ import annotations

import json
import os
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import psycopg2
from dotenv import load_dotenv
from openpyxl import load_workbook

PROJECT_DIR = Path(__file__).resolve().parents[1]
load_dotenv(PROJECT_DIR / ".env.new-car", override=False)

SOURCE_DIR = Path(r"E:/AI/migration-backup/5.4 晓亚\乘用车-新车明细板块-0927")
SALES_XLSX = SOURCE_DIR / "典型车型上市后销量.xlsx"
AH_XLSX = SOURCE_DIR / "汽车之家口碑明细_20260924.xlsx"
DCD_XLSX = SOURCE_DIR / "懂车帝口碑明细导出.xlsx"
TAXONOMY_XLSX = SOURCE_DIR / "汽车口碑标签字典_V1.0.xlsx"
TARGET_MODELS = ("问界M8", "领克900", "理想i8", "蔚来ES8", "理想L9", "魏牌V9X", "极氪9X", "小鹏GX")

PG = {
    "host": os.getenv("NEW_CAR_PG_HOST", "192.168.3.146"),
    "port": int(os.getenv("NEW_CAR_PG_PORT", "5432")),
    "dbname": os.getenv("NEW_CAR_PG_DATABASE", "vectordb"),
    "user": os.getenv("NEW_CAR_PG_USER", "vectordb"),
    "password": os.getenv("NEW_CAR_PG_PASSWORD", ""),
}


def norm(value: object) -> str:
    return re.sub(r"[\s·._-]+", "", str(value or "")).upper()


TARGET_BY_NORM = {norm(model): model for model in TARGET_MODELS}

# The rules are aligned to V1.0 neutral L3 names. Each hit is deduplicated per review.
# Phrases are deliberately specific so that the aggregation remains auditable and rerunnable.
TAG_PATTERNS: dict[str, tuple[str, ...]] = {
    "整体造型": ("外观", "造型", "颜值", "大气", "霸气", "好看", "耐看", "气场"),
    "车身比例": ("车身比例", "尺寸协调", "车长", "车宽"),
    "前脸设计": ("前脸", "车头"), "尾部设计": ("尾部", "车尾", "尾灯"),
    "轮毂造型": ("轮毂",), "车身配色": ("配色", "颜色", "车漆"),
    "外部装配工艺": ("钣金", "接缝", "外部做工", "装配"),
    "整体风格": ("内饰", "座舱设计"), "内饰配色": ("内饰颜色", "内饰配色"),
    "软性材质覆盖": ("软包", "真皮", "硬塑料"), "材质触感": ("用料", "材质", "质感"),
    "内饰装配工艺": ("内饰做工", "饰板", "音响罩"), "饰件耐用性": ("磨损", "脱落", "掉皮"),
    "前排空间": ("前排空间",), "第二排空间": ("二排空间", "第二排空间", "中排空间"),
    "第三排空间": ("三排空间", "第三排空间", "第三排能坐", "三排能坐"),
    "后备箱容量": ("后备箱", "后备厢"), "满员行李空间": ("满员", "六个人出门", "六人出门"),
    "前备箱实用性": ("前备箱",), "中控区域储物": ("中控储物",),
    "杯架实用性": ("杯架",), "其他储物空间": ("储物", "收纳"),
    "座椅放倒与组合": ("座椅放倒", "座椅组合", "大四座"),
    "上下车便利性": ("上下车", "进出第三排", "进出三排", "中间过道"),
    "车门尾门操作": ("车门", "尾门"), "车窗天窗操作": ("车窗", "天窗"),
    "座椅软硬与支撑": ("座椅", "零重力", "坐着", "包裹", "腰部支撑"),
    "座椅调节": ("座椅调节",), "座椅通风": ("座椅通风", "通风"),
    "座椅加热": ("座椅加热", "加热"), "座椅按摩": ("座椅按摩", "按摩"),
    "路面滤震": ("滤震", "减震", "坑洼", "减速带", "颠簸", "悬架舒适"),
    "车身晃动与晕车": ("晃动", "晕车", "船感"), "风噪": ("风噪",),
    "胎噪路噪": ("胎噪", "路噪"), "动力系统噪声振动": ("发动机声音", "发动机噪音", "震动", "抖动"),
    "异常响声": ("异响",), "空调制冷": ("制冷", "空调"), "空调制热": ("制热",),
    "车内气味": ("异味", "味道", "刺鼻"), "遮阳隔热": ("隔热", "晒"),
    "车载冰箱": ("冰箱",), "对外放电": ("对外放电", "外放电"),
    "起步响应": ("起步", "动力响应"), "加速能力": ("加速", "提速", "超车", "推背"),
    "满载动力": ("满载动力",), "换挡平顺性": ("换挡", "油电切换", "顿挫"),
    "转向手感": ("转向", "方向盘"), "制动脚感": ("刹车", "制动"),
    "动能回收体验": ("动能回收",), "车身稳定性": ("稳定", "侧倾", "底盘扎实"),
    "转弯灵活性": ("转弯", "掉头", "转向半径", "后轮转向", "灵活"),
    "通过性": ("通过性", "离地间隙", "蹭底"), "脱困能力": ("脱困", "蟹行", "原地掉头"),
    "实际续航": ("实际续航", "续航", "能跑"), "续航达成率": ("达成率", "官宣", "标称"),
    "高速续航": ("高速续航", "全程高速"), "低温续航": ("低温", "冬天", "冬季续航"),
    "燃油消耗": ("油耗",), "电能消耗": ("电耗",), "馈电油耗": ("馈电", "亏电油耗"),
    "驻车能耗": ("驻车掉电",), "里程显示准确性": ("续航显示", "表显"),
    "电池容量保持": ("电池衰减", "容量衰减"), "快充速度": ("快充", "超充"),
    "交流充电速度": ("家充速度", "慢充"), "充电兼容与稳定": ("充电失败", "充电中断"),
    "充电口便利性": ("充电口",), "家充安装": ("家充安装",),
    "公共充电可达性": ("充电桩", "公共充电"), "充电排队": ("充电排队",),
    "换电便利性": ("换电",), "加油便利性": ("加油方便",),
    "响应速度": ("车机", "卡顿", "反应速度", "响应速度", "流畅"),
    "系统稳定性": ("死机", "黑屏", "系统异常", "车机故障", "小毛病"),
    "界面操作": ("操作逻辑", "车机逻辑", "界面", "实体按键"), "屏幕显示": ("屏幕", "反光", "显示清晰"),
    "语音识别": ("语音", "小P", "理想同学", "小魏同学"), "语音控制范围": ("语音控制",),
    "导航准确性": ("导航",), "手机互联": ("手机互联", "CARPLAY", "HICAR"),
    "应用丰富度": ("APP", "应用"), "远程控车": ("远程控制", "远程控车"),
    "钥匙与解锁": ("数字钥匙", "蓝牙钥匙", "UWB"), "手机充电": ("无线充电", "USB", "手机充电"),
    "音响音质": ("音响", "音质"), "影音娱乐": ("娱乐屏", "追剧", "影音"), "车机升级体验": ("OTA", "升级"),
    "巡航跟车": ("跟车", "巡航"), "车道保持": ("车道保持",), "高速领航": ("高速NOA", "高速领航", "高速智驾"),
    "城区领航": ("城区NOA", "城区领航", "城市辅助驾驶", "城区智驾"),
    "接管与退出": ("接管", "紧急避让"), "辅助功能可用范围": ("智驾", "辅助驾驶"),
    "自动泊车": ("自动泊车",), "遥控泊车": ("遥控泊车", "召唤"),
    "泊车影像": ("全景影像", "倒车影像", "泊车影像", "倒车画面"), "泊车雷达": ("泊车雷达", "雷达"),
    "紧急制动表现": ("主动刹车", "紧急制动", "AEB"), "碰撞预警": ("碰撞预警",),
    "驾驶视野": ("视野", "盲区"), "照明效果": ("照明", "大灯"),
    "碰撞保护感受": ("安全性", "碰撞", "车身用料", "安全感"),
    "约束系统": ("安全带", "气囊"), "儿童安全设施": ("儿童座椅",),
    "电池消防安全感受": ("电池安全", "自燃"), "隐私与数据安全": ("隐私", "数据安全"),
    "成交价格": ("价格", "车价", "落地价"), "优惠力度": ("优惠", "权益"), "版本价差": ("版本价差", "贵了"),
    "整体性价比": ("性价比", "划算", "值得入手", "对得起"), "配置价值": ("配置丰富", "配置齐", "标配", "选装"),
    "降价感受": ("降价", "背刺"), "保值表现": ("保值", "二手车", "亏掉"),
    "能源费用": ("电费", "油费", "用电成本"), "保险费用": ("保险",), "保养费用": ("保养",),
    "维修费用": ("维修费",), "耗材更换费用": ("轮胎贵", "耗材"), "电池更换费用": ("换电池",),
    "订阅及租赁费用": ("订阅", "租赁"), "综合用车费用": ("用车成本", "养车"),
    "销售服务": ("销售", "门店服务"), "试驾体验": ("试驾",), "报价透明度": ("报价", "附加费"),
    "金融办理": ("贷款", "金融"), "置换办理": ("置换",), "合同与承诺": ("承诺", "合同"),
    "等待时长": ("等车", "提车得", "交付等待"), "验车与交付车况": ("提车时", "交付车况"), "交付流程": ("交付流程",),
    "网点便利性": ("售后网点", "服务网点"), "预约便利性": ("售后预约",),
    "服务态度": ("售后态度",), "处理时效": ("售后", "处理慢", "迟迟"), "维修效果": ("维修",),
    "配件供应": ("配件",), "质保执行": ("质保",), "客服投诉处理": ("客服", "投诉"), "道路救援": ("救援",),
    "品牌认同": ("品牌形象", "品牌认可"), "身份认同": ("情绪价值", "有面子", "装逼"),
    "技术信任": ("技术", "技术实力"), "诚信信任": ("诚信", "承诺"), "经营稳定性": ("倒闭", "经营"),
    "整体满意度": ("整体满意", "非常满意", "很满意"), "预期符合度": ("超出预期", "不及预期"),
    "整体可靠性感受": ("省心", "故障率", "毛病多", "耐用性", "可靠性"),
    "购买后悔程度": ("后悔",), "推荐意愿": ("推荐", "不建议购买"), "再购意愿": ("再次购买", "再买"),
}

CUSTOM_TAGS = {
    "六座乘员平权": ("空间与实用性", "乘坐空间", ("六座平权", "座位平权", "每一排都", "三排都能坐", "六个座位")),
    "大车停车便利性": ("空间与实用性", "布局与便利", ("停车费劲", "停车不方便", "车位窄", "车身太宽", "车身尺寸偏大", "机械车位", "大车笨重")),
    "辅助驾驶拟人度": ("辅助驾驶与泊车", "辅助行车", ("拟人", "老司机", "驾驶策略")),
    "软件功能可靠性": ("智能座舱", "车机系统", ("功能失效", "设置不成功", "保存不生效", "等待升级修复")),
}

POSITIVE_CUES = ("满意", "好看", "漂亮", "大气", "霸气", "舒服", "舒适", "宽敞", "充足", "够用", "方便", "顺畅", "流畅", "精准", "轻快", "有力", "强劲", "平顺", "安静", "稳", "省", "快", "灵活", "实用", "丰富", "扎实", "细腻", "喜欢", "不错", "很好", "优秀", "清晰", "高级", "惊喜", "认可", "放心", "值得")
NEGATIVE_CUES = ("不满意", "不好", "不便", "不足", "局促", "拥挤", "偏小", "太小", "太大", "偏大", "慢", "卡顿", "死机", "黑屏", "异响", "噪音大", "噪声大", "贵", "偏高", "费劲", "笨重", "缩水", "衰减", "顿挫", "后悔", "担忧", "不准", "失败", "反光", "复杂", "困难", "抖动", "颠簸", "故障", "毛病", "差", "一般", "吐槽", "遗憾", "焦虑", "不聪明", "不流畅", "不方便", "不够", "没有", "缺少")


def load_taxonomy() -> dict[str, dict[str, str]]:
    wb = load_workbook(TAXONOMY_XLSX, read_only=True, data_only=True)
    ws = wb["三级标签字典"]
    rows = ws.iter_rows(values_only=True)
    for row in rows:
        if row and row[0] == "三级编码":
            break
    taxonomy = {}
    for row in rows:
        if not row or not row[0]:
            continue
        taxonomy[str(row[3])] = {"code": str(row[0]), "l1": str(row[1]), "l2": str(row[2])}
    wb.close()
    for name, (l1, l2, _patterns) in CUSTOM_TAGS.items():
        taxonomy[name] = {"code": f"CUSTOM-{len(taxonomy)+1:03d}", "l1": l1, "l2": l2, "custom": True}
    return taxonomy


def split_clauses(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"[\r\n。！？；;]+", text or "") if len(part.strip()) >= 2]


def matched_tags(text: str) -> set[str]:
    upper = text.upper()
    matches = {tag for tag, patterns in TAG_PATTERNS.items() if any(pattern.upper() in upper for pattern in patterns)}
    matches.update(tag for tag, (_l1, _l2, patterns) in CUSTOM_TAGS.items() if any(pattern.upper() in upper for pattern in patterns))
    return matches


def infer_direction(clause: str) -> str | None:
    pos = sum(cue in clause for cue in POSITIVE_CUES)
    neg = sum(cue in clause for cue in NEGATIVE_CUES)
    if pos == neg:
        return None
    return "positive" if pos > neg else "negative"


def add_review(
    store, details, model: str, direction: str, review_key: str, text: str, source: str,
    taxonomy, forced: bool = False, review_text: str = "",
):
    matched_clauses: dict[str, str] = {}
    for clause in split_clauses(text):
        if forced or infer_direction(clause) == direction:
            for tag in matched_tags(clause):
                matched_clauses.setdefault(tag, clause)
    if forced and not matched_clauses:
        for tag in matched_tags(text):
            matched_clauses.setdefault(tag, text.strip())
    for tag, excerpt in matched_clauses.items():
        meta = taxonomy[tag]
        l1 = meta["l1"]
        store[model][direction]["l1"][l1]["reviews"].add(review_key)
        store[model][direction]["l1"][l1]["sources"][source].add(review_key)
        store[model][direction]["l3"][tag]["reviews"].add(review_key)
        store[model][direction]["l3"][tag]["sources"][source].add(review_key)
        details.setdefault((model, direction, source, review_key, tag), {
            "model": model, "direction": direction, "source": source,
            "review_key": review_key, "primary_label": l1,
            "secondary_label": meta["l2"], "tertiary_label": tag,
            "taxonomy_code": meta["code"], "is_custom": bool(meta.get("custom")),
            "excerpt": excerpt, "review_text": review_text or text,
        })


def read_autohome(store, details, taxonomy):
    counts = Counter()
    wb = load_workbook(AH_XLSX, read_only=True, data_only=True)
    ws = wb["口碑明细"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    idx = {name: headers.index(name) for name in ("koubei_id", "车系", "车型名称", "口碑全文", "满意", "不满意")}
    for row in rows:
        source_name = row[idx["车系"]] or row[idx["车型名称"]]
        model = TARGET_BY_NORM.get(norm(source_name))
        if not model:
            continue
        review_key = f"autohome:{row[idx['koubei_id']]}"
        counts[model] += 1
        positive = str(row[idx["满意"]] or "").strip()
        negative = str(row[idx["不满意"]] or "").strip()
        full = str(row[idx["口碑全文"]] or "").strip()
        if positive:
            add_review(store, details, model, "positive", review_key, positive, "autohome", taxonomy, forced=True, review_text=full or positive)
        if negative:
            add_review(store, details, model, "negative", review_key, negative, "autohome", taxonomy, forced=True, review_text=full or negative)
        # The full review contains richer aspect paragraphs. Only clauses with explicit sentiment enter the ranking.
        if full:
            add_review(store, details, model, "positive", review_key, full, "autohome", taxonomy, review_text=full)
            add_review(store, details, model, "negative", review_key, full, "autohome", taxonomy, review_text=full)
    wb.close()
    return counts


DCD_L1_SCORES = {
    "外观设计": "外观评分", "内饰设计与做工": "内饰评分", "价格与购买价值": "配置评分",
    "空间与实用性": "空间评分", "乘坐舒适性": "舒适性评分", "续航与能耗": "续航评分",
    "动力与操控": "动力操控评分", "辅助驾驶与泊车": "辅助驾驶评分", "智能座舱": "智能座舱评分",
}


def read_dongchedi(store, details, taxonomy):
    counts = Counter()
    wb = load_workbook(DCD_XLSX, read_only=True, data_only=True)
    ws = wb["口碑明细"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    pos = {name: i for i, name in enumerate(headers)}
    for row in rows:
        source_name = row[pos["车系名称"]] or row[pos["车名"]]
        model = TARGET_BY_NORM.get(norm(source_name))
        if not model:
            continue
        review_key = f"dongchedi:{row[pos['口碑ID']]}"
        text = str(row[pos["正文"]] or "").strip()
        counts[model] += 1
        if not text:
            continue
        before = {
            direction: {tag: len(item["reviews"]) for tag, item in store[model][direction]["l3"].items()}
            for direction in ("positive", "negative")
        }
        add_review(store, details, model, "positive", review_key, text, "dongchedi", taxonomy, review_text=text)
        add_review(store, details, model, "negative", review_key, text, "dongchedi", taxonomy, review_text=text)
        # Score fallback is applied only to matched text topics that lacked an explicit sentiment phrase.
        for clause in split_clauses(text):
            if infer_direction(clause) is not None:
                continue
            for tag in matched_tags(clause):
                l1 = taxonomy[tag]["l1"]
                score_col = DCD_L1_SCORES.get(l1)
                if not score_col or score_col not in pos:
                    continue
                try:
                    score = float(row[pos[score_col]])
                except (TypeError, ValueError):
                    continue
                direction = "positive" if score >= 4.5 else "negative" if score <= 3.0 else None
                if direction:
                    add_review(store, details, model, direction, review_key, clause, "dongchedi", taxonomy, forced=True, review_text=text)
    wb.close()
    return counts


def import_sales(cursor, batch: str) -> int:
    wb = load_workbook(SALES_XLSX, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS new_car_monthly_sales (
            model TEXT NOT NULL, sales_month DATE NOT NULL, month_index SMALLINT NOT NULL,
            sales_volume INT NOT NULL, import_batch TEXT NOT NULL,
            PRIMARY KEY (model, sales_month)
        )
    """)
    cursor.execute("DELETE FROM new_car_monthly_sales WHERE model = ANY(%s)", (list(TARGET_MODELS),))
    records = []
    first_seen = {model: 0 for model in TARGET_MODELS}
    for row in rows:
        month = row[0]
        for col, raw_model in enumerate(headers[1:9], 1):
            model = TARGET_BY_NORM.get(norm(raw_model))
            value = row[col]
            if not model or value in (None, ""):
                continue
            first_seen[model] += 1
            records.append((model, month.date(), first_seen[model], int(value), batch))
    cursor.executemany(
        "INSERT INTO new_car_monthly_sales (model,sales_month,month_index,sales_volume,import_batch) VALUES (%s,%s,%s,%s,%s)",
        records,
    )
    wb.close()
    return len(records)


def import_reputation(cursor, batch: str, taxonomy) -> tuple[int, dict[str, dict[str, int]]]:
    store = defaultdict(lambda: defaultdict(lambda: {
        "l1": defaultdict(lambda: {"reviews": set(), "sources": defaultdict(set)}),
        "l3": defaultdict(lambda: {"reviews": set(), "sources": defaultdict(set)}),
    }))
    details: dict[tuple[str, str, str, str, str], dict[str, object]] = {}
    ah_counts = read_autohome(store, details, taxonomy)
    dcd_counts = read_dongchedi(store, details, taxonomy)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS new_car_reputation_taxonomy_summary (
            model TEXT NOT NULL, direction TEXT NOT NULL, rank SMALLINT NOT NULL,
            primary_label TEXT NOT NULL, tertiary_labels JSONB NOT NULL,
            mention_count INT NOT NULL, mention_share NUMERIC(7,4) NOT NULL,
            dongchedi_count INT NOT NULL, autohome_count INT NOT NULL,
            source_mentions JSONB NOT NULL, taxonomy_version TEXT NOT NULL,
            import_batch TEXT NOT NULL, PRIMARY KEY (model, direction, rank)
        )
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS new_car_reputation_review_tag (
            model TEXT NOT NULL, direction TEXT NOT NULL, source TEXT NOT NULL,
            review_key TEXT NOT NULL, primary_label TEXT NOT NULL,
            secondary_label TEXT NOT NULL, tertiary_label TEXT NOT NULL,
            taxonomy_code TEXT NOT NULL, is_custom BOOLEAN NOT NULL DEFAULT FALSE,
            excerpt TEXT NOT NULL, review_text TEXT NOT NULL, import_batch TEXT NOT NULL,
            PRIMARY KEY (model, direction, source, review_key, tertiary_label)
        )
    """)
    cursor.execute("DELETE FROM new_car_reputation_taxonomy_summary WHERE model = ANY(%s)", (list(TARGET_MODELS),))
    cursor.execute("DELETE FROM new_car_reputation_review_tag WHERE model = ANY(%s)", (list(TARGET_MODELS),))
    inserted = 0
    coverage = {}
    for model in TARGET_MODELS:
        total_reviews = ah_counts[model] + dcd_counts[model]
        coverage[model] = {"autohome": ah_counts[model], "dongchedi": dcd_counts[model]}
        if not total_reviews:
            continue
        for direction in ("positive", "negative"):
            l1_items = store[model][direction]["l1"]
            ranked = sorted(l1_items.items(), key=lambda item: (-len(item[1]["reviews"]), item[0]))[:5]
            for rank, (l1, l1_data) in enumerate(ranked, 1):
                tertiary = []
                for tag, tag_data in store[model][direction]["l3"].items():
                    meta = taxonomy[tag]
                    if meta["l1"] != l1:
                        continue
                    tertiary.append({
                        "name": tag, "level2": meta["l2"], "code": meta["code"],
                        "count": len(tag_data["reviews"]), "custom": bool(meta.get("custom")),
                    })
                tertiary.sort(key=lambda item: (-item["count"], item["code"]))
                source_mentions = {source: len(keys) for source, keys in l1_data["sources"].items()}
                cursor.execute("""
                    INSERT INTO new_car_reputation_taxonomy_summary
                    (model,direction,rank,primary_label,tertiary_labels,mention_count,mention_share,
                     dongchedi_count,autohome_count,source_mentions,taxonomy_version,import_batch)
                    VALUES (%s,%s,%s,%s,%s::jsonb,%s,%s,%s,%s,%s::jsonb,%s,%s)
                """, (
                    model, direction, rank, l1, json.dumps(tertiary[:4], ensure_ascii=False),
                    len(l1_data["reviews"]), len(l1_data["reviews"]) / total_reviews,
                    dcd_counts[model], ah_counts[model], json.dumps(source_mentions, ensure_ascii=False),
                    "V1.0+POC", batch,
                ))
                inserted += 1
    detail_rows = [
        (
            item["model"], item["direction"], item["source"], item["review_key"],
            item["primary_label"], item["secondary_label"], item["tertiary_label"],
            item["taxonomy_code"], item["is_custom"], item["excerpt"], item["review_text"], batch,
        )
        for item in details.values()
    ]
    cursor.executemany("""
        INSERT INTO new_car_reputation_review_tag
        (model,direction,source,review_key,primary_label,secondary_label,tertiary_label,
         taxonomy_code,is_custom,excerpt,review_text,import_batch)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    """, detail_rows)
    return inserted, coverage


def main():
    batch = datetime.now().strftime("%Y%m%d")
    taxonomy = load_taxonomy()
    connection = psycopg2.connect(**PG)
    try:
        with connection, connection.cursor() as cursor:
            sales_rows = import_sales(cursor, batch)
            reputation_rows, coverage = import_reputation(cursor, batch, taxonomy)
        print(json.dumps({
            "batch": batch, "sales_rows": sales_rows, "reputation_rows": reputation_rows,
            "coverage": coverage,
        }, ensure_ascii=False, indent=2))
    finally:
        connection.close()


if __name__ == "__main__":
    main()
