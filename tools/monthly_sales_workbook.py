"""Validate business Excel and aggregate actual quantities by month/model/energy."""
import hashlib
import json
import unicodedata
from collections import defaultdict
from datetime import date
from pathlib import Path
from openpyxl import load_workbook

ENERGY_ALIASES={'纯电动':'纯电动','插电式混合动力':'插电式混合动力','插电式':'插电式混合动力',
    '增程式':'插电式混合动力','常规混合动力':'常规混合动力','传统燃料':'传统燃料',
    '燃料电池':'燃料电池','氢燃料':'燃料电池'}

def model_key(value):
    return ''.join(unicodedata.normalize('NFKC',str(value)).split()).casefold()

def read_workbook(path):
    path=Path(path);digest=hashlib.sha256(path.read_bytes()).hexdigest()
    workbook=load_workbook(path,read_only=True,data_only=True)
    grouped=defaultdict(lambda:{'sales':0,'source_rows':[],'source_models':set()})
    raw_count=total=0
    try:
        sheet=workbook.active;records=sheet.iter_rows(values_only=True)
        headers=list(next(records))
        if headers!=['年','月','车型','新能源类型','数量']:raise ValueError('销量Excel字段与约定不一致')
        for row_number,row in enumerate(records,2):
            if all(v is None for v in row):continue
            year,month,model,energy,quantity=row
            # Excel stores numeric vehicle names such as Porsche 718 as numbers.
            if isinstance(model,int) and not isinstance(model,bool):model=str(model)
            if (not isinstance(year,int) or not isinstance(month,int) or not 1<=month<=12
                or not isinstance(model,str) or not model.strip() or energy not in ENERGY_ALIASES
                or isinstance(quantity,bool) or not isinstance(quantity,(int,float))
                or quantity<0 or int(quantity)!=quantity):raise ValueError(f'销量Excel第{row_number}行无效')
            day=date(year,month,1).isoformat();quantity=int(quantity)
            key=(model_key(model),ENERGY_ALIASES[energy],day)
            item=grouped[key];item['sales']+=quantity;item['source_rows'].append(row_number);item['source_models'].add(model.strip())
            raw_count+=1;total+=quantity
    finally:workbook.close()
    rows=[dict(model_key=key[0],energy_type=key[1],month=key[2],model=sorted(item['source_models'])[0],
               sales=item['sales'],source_rows=item['source_rows'],source_models=sorted(item['source_models']))
          for key,item in sorted(grouped.items())]
    assert sum(r['sales'] for r in rows)==total
    return {'source_name':path.name,'source_sha256':digest,'raw_rows':raw_count,'records':len(rows),
        'models':len({r['model_key'] for r in rows}),'total_sales':total,
        'first_month':min(r['month'] for r in rows),'last_month':max(r['month'] for r in rows),'rows':rows}
